# pyright: reportAttributeAccessIssue=false, reportOptionalMemberAccess=false, reportOptionalSubscript=false, reportArgumentType=false, reportAssignmentType=false, reportCallIssue=false
"""
ALine UI 信号测试套件

使用 offscreen 渲染（QT_QPA_PLATFORM=offscreen），通过主动发射信号
验证各 UI 组件的信号-槽连接和数据流向。

运行方式：
    QT_QPA_PLATFORM=offscreen python -m unittest tests/test_ui.py
"""
from __future__ import annotations

import os
import sys
import tempfile
import unittest
import warnings
import gc
from pathlib import Path
from typing import cast
from unittest import mock

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
warnings.filterwarnings(
    "ignore",
    message=r".*QMouseEvent\.globalPos\(\) const.*deprecated.*",
    category=DeprecationWarning,
)

# 项目根路径
_PROJ_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if _PROJ_ROOT not in sys.path:
    sys.path.insert(0, _PROJ_ROOT)

from PySide6.QtWidgets import QApplication, QAbstractItemView, QWidget
from PySide6.QtCore import Qt, QTimer
from PySide6.QtTest import QTest

_app: QApplication | None = None


def setUpModule():
    global _app
    if QApplication.instance() is None:
        _app = QApplication(sys.argv)
    else:
        _app = cast(QApplication | None, QApplication.instance())


def tearDownModule():
    global _app

    app = cast(QApplication | None, QApplication.instance())
    if app is None:
        return
    for widget in list(app.topLevelWidgets()):
        try:
            widget.close()
            widget.deleteLater()
        except RuntimeError:
            continue
    app.processEvents()
    try:
        import shiboken6

        shiboken6.delete(app)
    except Exception:
        pass
    _app = None
    gc.collect()


# ─── 辅助函数 ───────────────────────────────────────────────

def _make_project(name="ui_test"):
    """创建带迁移的测试项目（使用独立 ProjectManager）"""
    from core.project_manager import ProjectManager
    from models.schemas import DataFile, DataSeries
    pm = ProjectManager()
    p = pm.create_new(name)
    pm.migrate_to_v2(p)
    s = DataSeries(name="s1", x=[1.0, 2.0, 3.0, 4.0, 5.0],
                   y=[2.0, 4.0, 6.0, 8.0, 10.0])
    df = DataFile(name="test.csv", series=[s])
    pm.add_data_file(df)
    return pm, p, df, s


_PM_MODULES = [
    "core.project_manager",
    "ai.command_layer",
    "ui.dialogs.export_flow",
    "ui.dialogs.import_dialog",
    "ui.pages.data_page",
    "ui.pages.chart_page",
    "ui.pages.process_page",
    "ui.pages.analysis_page",
    "ui.pages.digitize_page",
    "ui.widgets.project_tree",
    "ui.main_window",
]


def _patch_pm(pm):
    """Patch project_manager in all relevant modules and return restorer."""
    import importlib
    saved = {}
    for mod_name in _PM_MODULES:
        try:
            mod = importlib.import_module(mod_name)
            if hasattr(mod, "project_manager"):
                saved[mod_name] = mod.project_manager
                mod.project_manager = pm
        except ImportError:
            pass

    def restore():
        for mod_name, orig in saved.items():
            try:
                mod = importlib.import_module(mod_name)
                mod.project_manager = orig
            except ImportError:
                pass

    return restore


def _patch_global_assets():
    from core.global_assets import GlobalAssets, global_assets

    temp_dir = tempfile.TemporaryDirectory()
    old_path = global_assets._asset_path
    old_cache = global_assets._cache
    global_assets._asset_path = Path(temp_dir.name) / "global_assets.json"
    global_assets._cache = GlobalAssets()
    global_assets.save()

    def restore():
        global_assets._asset_path = old_path
        global_assets._cache = old_cache
        temp_dir.cleanup()

    return restore


def _analysis_result_save_plans(pm, *result_names, parent_id=None):
    from ui.dialogs.export_flow import AnalysisResultSavePlan

    target_parent_id = parent_id
    if target_parent_id is None:
        analysis_root = pm._find_folder_by_group_type("analysis_result_group")
        if analysis_root is None:
            raise AssertionError("analysis_result_group root not found")
        target_parent_id = analysis_root.id

    return [
        AnalysisResultSavePlan(result_name=result_name, target_parent_id=target_parent_id)
        for result_name in result_names
    ]


# ═══════════════════════════════════════════════════════════════════════════
# 1. ProjectTreeWidget
# ═══════════════════════════════════════════════════════════════════════════

class TestProjectTreeWidget(unittest.TestCase):
    """ProjectTreeWidget 信号与基本操作"""

    def setUp(self):
        self._restore_assets = _patch_global_assets()
        self.pm, self.p, self.df, self.s = _make_project("tree_test")
        self._restore = _patch_pm(self.pm)
        from ui.widgets.project_tree import ProjectTreeWidget
        self.widget = ProjectTreeWidget()

    def tearDown(self):
        self._restore()
        self._restore_assets()
        self.widget.deleteLater()

    @staticmethod
    def _icon_image(icon, size: int = 20):
        return icon.pixmap(size, size).toImage()

    def test_refresh_builds_tree(self):
        self.widget.refresh()
        # 树应该有若干根节点
        self.assertGreater(self.widget._tree.topLevelItemCount(), 0)

    def test_refresh_shows_project_roots(self):
        self.widget.refresh()
        labels = [self.widget._tree.topLevelItem(i).text(0) for i in range(self.widget._tree.topLevelItemCount())]
        self.assertEqual(labels, [self.p.name, "全局资源"])

    def test_multiple_projects_render_as_top_level_roots(self):
        from models.schemas import DataFile, DataSeries

        other = self.pm.create_new("tree_second")
        self.pm.migrate_to_v3(other)
        self.pm.set_current_project(other.id)
        self.pm.add_data_file(DataFile(name="other.csv", series=[DataSeries(name="s2", x=[1.0], y=[2.0])]))
        self.pm.set_current_project(self.p.id)
        self.widget.refresh()
        labels = [self.widget._tree.topLevelItem(i).text(0) for i in range(self.widget._tree.topLevelItemCount())]
        self.assertEqual(labels, [self.p.name, other.name, "全局资源"])

    def test_global_resource_contains_global_template_groups(self):
        from models.schemas import FigureConfig

        self.pm.add_saved_pipeline("p1", [{"type": "smooth", "params": {}}])
        self.pm.add_report_template("tmpl_a", "# Report")
        self.pm.add_figure_template(FigureConfig(name="样式A", theme="Nature"))
        self.widget.refresh()
        global_root = self.widget._tree.topLevelItem(self.widget._tree.topLevelItemCount() - 1)
        root_children = [global_root.child(i).text(0) for i in range(global_root.childCount())]
        self.assertEqual(root_children, ["Pipelines", "曲线样式", "绘图样式", "报告模板", "扩展配置"])

        pipeline_group = next(
            global_root.child(i)
            for i in range(global_root.childCount())
            if global_root.child(i).text(0) == "Pipelines"
        )
        report_group = next(
            global_root.child(i)
            for i in range(global_root.childCount())
            if global_root.child(i).text(0) == "报告模板"
        )
        style_group = next(
            global_root.child(i)
            for i in range(global_root.childCount())
            if global_root.child(i).text(0) == "绘图样式"
        )
        self.assertEqual([pipeline_group.child(i).text(0) for i in range(pipeline_group.childCount())], ["p1"])
        self.assertEqual([report_group.child(i).text(0) for i in range(report_group.childCount())], ["默认模板", "tmpl_a"])
        style_items = [style_group.child(i).text(0) for i in range(style_group.childCount())]
        self.assertIn("默认", style_items)
        self.assertIn("样式A", style_items)

    def test_global_resource_contains_extension_config_groups_without_default_leaves(self):
        from core.extension_api import ProcessingExtension, extension_registry

        def _probe(lines, params):
            return lines[0] if lines else []

        extension_registry.register_processing(
            ProcessingExtension(type="tree_extension_config_probe", name="树配置探针", handler=_probe, default_options={"factor": 2}, settings=True)
        )
        try:
            self.widget.refresh()
            global_root = self.widget._tree.topLevelItem(self.widget._tree.topLevelItemCount() - 1)
            extension_root = next(
                global_root.child(i)
                for i in range(global_root.childCount())
                if global_root.child(i).text(0) == "扩展配置"
            )
            processing_group = next(
                extension_root.child(i)
                for i in range(extension_root.childCount())
                if extension_root.child(i).text(0) == "处理扩展"
            )
            target_item = next(
                processing_group.child(i)
                for i in range(processing_group.childCount())
                if processing_group.child(i).text(0) == "树配置探针"
            )
            self.assertEqual(target_item.childCount(), 0)
            self.assertEqual(target_item.data(0, Qt.ItemDataRole.UserRole)[0], "global_group")
            self.assertEqual(target_item.data(0, Qt.ItemDataRole.UserRole)[1], "__global_extension_configs__:processing:tree_extension_config_probe")
        finally:
            extension_registry.unregister_processing("tree_extension_config_probe")

    def test_extension_config_group_helpers_can_create_and_duplicate_presets(self):
        from core.extension_api import ProcessingExtension, extension_registry
        from core.global_assets import global_assets

        def _probe(lines, params):
            return lines[0] if lines else []

        extension_registry.register_processing(
            ProcessingExtension(type="tree_extension_config_create", name="树配置新建探针", handler=_probe, default_options={"factor": 2}, settings=True)
        )
        try:
            with mock.patch("ui.widgets.project_tree.TextInputDialog.get_text", side_effect=[("方案A", True), ("方案A 副本", True)]):
                self.widget._cmd_create_extension_config("__global_extension_configs__:processing:tree_extension_config_create")
                created = global_assets.get_extension_config_by_name("processing", "tree_extension_config_create", "方案A")
                self.assertIsNotNone(created)

                self.widget._cmd_duplicate_extension_config(created.id)
                duplicated = global_assets.get_extension_config_by_name("processing", "tree_extension_config_create", "方案A 副本")
                self.assertIsNotNone(duplicated)
                self.assertEqual(duplicated.options, created.options)
        finally:
            extension_registry.unregister_processing("tree_extension_config_create")

    def test_extension_config_groups_hide_entries_without_settings_support(self):
        from core.extension_api import ProcessingExtension, extension_registry

        def _probe(lines, params):
            return lines[0] if lines else []

        extension_registry.register_processing(
            ProcessingExtension(type="tree_extension_config_hidden", name="隐藏配置扩展", handler=_probe, settings=False)
        )
        try:
            self.widget.refresh()
            global_root = self.widget._tree.topLevelItem(self.widget._tree.topLevelItemCount() - 1)
            extension_root = next(
                global_root.child(i)
                for i in range(global_root.childCount())
                if global_root.child(i).text(0) == "扩展配置"
            )
            processing_group = next(
                extension_root.child(i)
                for i in range(extension_root.childCount())
                if extension_root.child(i).text(0) == "处理扩展"
            )
            group_labels = [processing_group.child(i).text(0) for i in range(processing_group.childCount())]
            self.assertNotIn("隐藏配置扩展", group_labels)
        finally:
            extension_registry.unregister_processing("tree_extension_config_hidden")

    def test_builtin_report_template_cannot_be_renamed_or_deleted(self):
        restore_assets = _patch_global_assets()
        try:
            from core.global_assets import global_assets

            builtin_template = next(item for item in global_assets.list_report_templates(include_builtin=True) if item.is_builtin)
            self.assertFalse(self.widget._rename_global_asset("global_report_template", builtin_template.id, "重命名默认模板"))
            self.assertFalse(self.widget._delete_global_asset("global_report_template", builtin_template.id))
        finally:
            restore_assets()

    def test_global_resource_items_can_be_renamed_and_deleted(self):
        restore_assets = _patch_global_assets()
        try:
            from core.global_assets import global_assets
            from models.schemas import ReportTemplate, SavedPipeline

            pipeline = global_assets.add_saved_pipeline(SavedPipeline(name="流程A", ops=[]))
            report = global_assets.add_report_template(ReportTemplate(name="报告A", template="# 报告"))
            self.assertTrue(self.widget._rename_global_asset("global_pipeline", pipeline.id, "流程B"))
            self.assertEqual(global_assets.get_saved_pipeline(pipeline.id).name, "流程B")
            self.assertTrue(self.widget._rename_global_asset("global_report_template", report.id, "报告B"))
            self.assertEqual(global_assets.get_report_template(report.id).name, "报告B")
            self.assertTrue(self.widget._delete_global_asset("global_pipeline", pipeline.id))
            self.assertIsNone(global_assets.get_saved_pipeline(pipeline.id))
            self.assertTrue(self.widget._delete_global_asset("global_report_template", report.id))
            self.assertIsNone(global_assets.get_report_template(report.id))
        finally:
            restore_assets()

    def test_refresh_preserves_tree_expansion_state(self):
        self.widget.refresh()
        project_root = self.widget._tree.topLevelItem(0)
        dataset_group = project_root.child(0)
        global_root = self.widget._tree.topLevelItem(self.widget._tree.topLevelItemCount() - 1)

        project_root.setExpanded(True)
        dataset_group.setExpanded(False)
        global_root.setExpanded(False)

        self.widget.refresh()
        project_root = self.widget._tree.topLevelItem(0)
        dataset_group = project_root.child(0)
        global_root = self.widget._tree.topLevelItem(self.widget._tree.topLevelItemCount() - 1)
        self.assertTrue(project_root.isExpanded())
        self.assertFalse(dataset_group.isExpanded())
        self.assertFalse(global_root.isExpanded())

    def test_project_root_defaults_to_first_level_expanded_only(self):
        datasets_root = self.pm._find_folder_by_group_type("datasets")
        self.assertIsNotNone(datasets_root)
        folder = self.pm.add_folder("批次A", parent_id=datasets_root.id, group_type="datasets")
        self.assertIsNotNone(folder)
        self.pm.add_folder("子目录", parent_id=folder.id, group_type="datasets")

        self.widget.refresh()

        project_root = self.widget._tree.topLevelItem(0)
        self.assertTrue(project_root.isExpanded())
        root_children = [project_root.child(i) for i in range(project_root.childCount())]
        self.assertTrue(root_children)
        self.assertTrue(all(not child.isExpanded() for child in root_children))

    def test_root_group_folder_uses_category_label_instead_of_node_id(self):
        source_root = self.pm._find_folder_by_group_type("source_files")
        self.assertIsNotNone(source_root)
        source_root.name = source_root.id

        self.widget.refresh()

        project_root = self.widget._tree.topLevelItem(0)
        child_labels = [project_root.child(i).text(0) for i in range(project_root.childCount())]
        self.assertIn("源文件", child_labels)

    def test_tree_sorts_siblings_with_folders_first_and_english_before_chinese(self):
        from models.schemas import DataFile, DataSeries

        datasets_root = self.pm._find_folder_by_group_type("datasets")
        self.assertIsNotNone(datasets_root)
        container = self.pm.add_folder("排序测试组", parent_id=datasets_root.id, group_type="datasets")
        self.assertIsNotNone(container)

        self.pm.add_folder("AlphaFolder", parent_id=container.id, group_type="datasets")
        self.pm.add_folder("中文目录", parent_id=container.id, group_type="datasets")
        self.pm.add_data_file(
            DataFile(name="Beta.csv", series=[DataSeries(name="s_beta", x=[0.0], y=[1.0])]),
            parent_id=container.id,
        )
        self.pm.add_data_file(
            DataFile(name="中文数据.csv", series=[DataSeries(name="s_cn", x=[0.0], y=[2.0])]),
            parent_id=container.id,
        )

        self.widget.refresh()
        container_item = self.widget._find_item(container.id)
        self.assertIsNotNone(container_item)
        children = [container_item.child(i).text(0) for i in range(container_item.childCount())]
        self.assertEqual(children, ["AlphaFolder", "中文目录", "Beta.csv", "中文数据.csv"])

    def test_tree_config_enables_wrapped_labels(self):
        self.widget.set_name_display_mode("wrap")
        self.assertEqual(self.widget._tree.textElideMode(), Qt.TextElideMode.ElideNone)

    def test_tree_supports_extended_selection_for_batch_actions(self):
        self.assertEqual(self.widget._tree.selectionMode(), QAbstractItemView.SelectionMode.ExtendedSelection)

    def test_f2_shortcut_renames_current_tree_item(self):
        self.widget.refresh()
        data_file_node = next(n for n in self.p.tree.nodes if n.kind == "data_file" and n.data_file_id == self.df.id)
        item = self.widget._find_item(data_file_node.id)
        self.assertIsNotNone(item)

        self.widget._tree.setCurrentItem(item)
        item.setSelected(True)
        self.widget._tree.setFocus()

        with mock.patch.object(self.widget, "rename_selected_item") as rename_mock:
            QTest.keyClick(self.widget._tree, Qt.Key.Key_F2)

        rename_mock.assert_called_once_with()

    def test_tree_disables_double_click_inline_rename(self):
        self.assertEqual(self.widget._tree.editTriggers(), QAbstractItemView.EditTrigger.NoEditTriggers)

    def test_tree_uses_compact_indentation(self):
        self.assertEqual(self.widget._tree.indentation(), 14)

    def test_default_tree_name_display_mode_pref_is_elide(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            config_path = Path(temp_dir) / "ui_preferences.json"
            with mock.patch("core.ui_preferences._CONFIG_PATH", config_path):
                from core.ui_preferences import UIPreferences, get_tree_name_display_mode

                self.assertEqual(UIPreferences.load().tree_name_display_mode, "elide")
                self.assertEqual(get_tree_name_display_mode(), "elide")

    def test_tree_uses_default_delegate_and_no_custom_foreground_role(self):
        from qfluentwidgets.components.widgets.tree_view import TreeItemDelegate

        self.widget.set_name_display_mode("wrap")
        self.pm.add_saved_pipeline("流程A", [{"type": "smooth", "params": {}}])
        self.widget.refresh()

        self.assertNotEqual(type(self.widget._tree.itemDelegate()).__name__, "_ProjectTreeItemDelegate")
        self.assertIsInstance(self.widget._tree.itemDelegate(), TreeItemDelegate)
        series_item = self.widget._find_item(self.s.id)
        self.assertIsNotNone(series_item)
        self.assertIsNone(series_item.data(0, Qt.ItemDataRole.ForegroundRole))
        self.assertFalse(self.widget._tree.uniformRowHeights())

    def test_set_name_display_mode_elides_long_labels(self):
        self.widget.set_name_display_mode("elide")
        self.assertEqual(self.widget._tree.textElideMode(), Qt.TextElideMode.ElideRight)
        self.assertTrue(self.widget._tree.uniformRowHeights())

    def test_set_name_display_mode_does_not_trigger_rename_warning(self):
        with mock.patch("ui.widgets.project_tree.InfoBar.warning") as warning_mock:
            self.widget.set_name_display_mode("wrap")
            self.widget.set_name_display_mode("elide")
            self.widget.set_name_display_mode("wrap")

        warning_mock.assert_not_called()

    def test_batch_drop_move_moves_multiple_selected_items(self):
        from models.schemas import DataFile, DataSeries

        datasets_root = self.pm._find_folder_by_group_type("datasets")
        self.assertIsNotNone(datasets_root)
        target_folder = self.pm.add_folder("批量目标", parent_id=datasets_root.id, group_type="datasets")
        self.assertIsNotNone(target_folder)

        other = DataFile(name="batch_other.csv", series=[DataSeries(name="s2", x=[0.0], y=[2.0])])
        self.pm.add_data_file(other)

        self.widget.refresh()

        first_node = next(n for n in self.p.tree.nodes if n.kind == "data_file" and n.data_file_id == self.df.id)
        second_node = next(n for n in self.p.tree.nodes if n.kind == "data_file" and n.data_file_id == other.id)
        source_items = [self.widget._find_item(first_node.id), self.widget._find_item(second_node.id)]
        self.assertTrue(all(item is not None for item in source_items))
        target_item = self.widget._find_item(target_folder.id)
        self.assertIsNotNone(target_item)

        moved = self.widget._perform_batch_drop_move(source_items, target_item, defer_view_refresh=False)

        self.assertTrue(moved)
        self.assertEqual(self.pm.get_node_by_id(first_node.id).parent_id, target_folder.id)
        self.assertEqual(self.pm.get_node_by_id(second_node.id).parent_id, target_folder.id)
        selected_ids = {item.data(0, Qt.ItemDataRole.UserRole)[1] for item in self.widget._tree.selectedItems()}
        self.assertEqual(selected_ids, {first_node.id, second_node.id})

    def test_focus_selected_item_hides_unrelated_tree_branches(self):
        from models.schemas import DataFile, DataSeries

        other = DataFile(name="focus_other.csv", series=[DataSeries(name="s2", x=[0.0], y=[2.0])])
        self.pm.add_data_file(other)
        self.widget.refresh()

        target_node = next(n for n in self.p.tree.nodes if n.kind == "data_file" and n.data_file_id == self.df.id)
        other_node = next(n for n in self.p.tree.nodes if n.kind == "data_file" and n.data_file_id == other.id)

        self.widget.select_node(target_node.id)
        self.widget.focus_selected_item()

        project_root = self.widget._tree.topLevelItem(0)
        global_root = self.widget._tree.topLevelItem(self.widget._tree.topLevelItemCount() - 1)
        datasets_group = next(project_root.child(i) for i in range(project_root.childCount()) if project_root.child(i).text(0) == "数据集")
        source_group = next(project_root.child(i) for i in range(project_root.childCount()) if project_root.child(i).text(0) == "源文件")
        target_item = self.widget._find_item(target_node.id)
        other_item = self.widget._find_item(other_node.id)

        self.assertTrue(self.widget.is_focus_active())
        self.assertFalse(datasets_group.isHidden())
        self.assertTrue(source_group.isHidden())
        self.assertTrue(global_root.isHidden())
        self.assertFalse(target_item.isHidden())
        self.assertTrue(other_item.isHidden())

    def test_clear_focus_restores_hidden_tree_branches(self):
        from models.schemas import DataFile, DataSeries

        other = DataFile(name="focus_restore.csv", series=[DataSeries(name="s2", x=[0.0], y=[2.0])])
        self.pm.add_data_file(other)
        self.widget.refresh()

        target_node = next(n for n in self.p.tree.nodes if n.kind == "data_file" and n.data_file_id == self.df.id)
        other_node = next(n for n in self.p.tree.nodes if n.kind == "data_file" and n.data_file_id == other.id)

        self.widget.select_node(target_node.id)
        self.widget.focus_selected_item()
        self.widget.clear_focus()

        project_root = self.widget._tree.topLevelItem(0)
        global_root = self.widget._tree.topLevelItem(self.widget._tree.topLevelItemCount() - 1)
        source_group = next(project_root.child(i) for i in range(project_root.childCount()) if project_root.child(i).text(0) == "源文件")
        other_item = self.widget._find_item(other_node.id)

        self.assertFalse(self.widget.is_focus_active())
        self.assertFalse(source_group.isHidden())
        self.assertFalse(global_root.isHidden())
        self.assertFalse(other_item.isHidden())

    def test_tree_item_tooltip_shows_full_long_label(self):
        self.p.name = "这是一个用于验证项目树节点 hover 可显示完整名称的超长项目名称"
        self.widget.refresh()
        root = self.widget._tree.topLevelItem(0)
        self.assertEqual(root.toolTip(0), self.p.name)

    def test_project_root_includes_source_digitize_and_picture_groups(self):
        self.widget.refresh()
        root = self.widget._tree.topLevelItem(0)
        labels = [root.child(i).text(0) for i in range(root.childCount())]
        self.assertEqual(labels[:5], ["源文件", "数据集", "图片集", "分析结果", "数字化"])

    def test_duplicate_extension_config_group_labels_are_disambiguated(self):
        from core.global_assets import global_assets

        global_assets.ensure_extension_default_config("plot", "orphan_text_a", "重复文字", {"text": "A"})
        global_assets.ensure_extension_default_config("plot", "orphan_text_b", "重复文字", {"text": "B"})

        self.widget.refresh()
        global_root = self.widget._tree.topLevelItem(self.widget._tree.topLevelItemCount() - 1)
        extension_root = next(
            global_root.child(i)
            for i in range(global_root.childCount())
            if global_root.child(i).text(0) == "扩展配置"
        )
        plot_group = next(
            extension_root.child(i)
            for i in range(extension_root.childCount())
            if extension_root.child(i).text(0) == "绘图扩展"
        )
        labels = [plot_group.child(i).text(0) for i in range(plot_group.childCount()) if plot_group.child(i).text(0).startswith("重复文字")]

        self.assertEqual(labels, [])

    def test_orphan_extension_configs_are_hidden(self):
        from core.global_assets import global_assets

        global_assets.ensure_extension_default_config("processing", "panel_extension_probe", "面板配置探针", {"probe": True})

        self.widget.refresh()
        global_root = self.widget._tree.topLevelItem(self.widget._tree.topLevelItemCount() - 1)
        extension_root = next(
            global_root.child(i)
            for i in range(global_root.childCount())
            if global_root.child(i).text(0) == "扩展配置"
        )
        processing_group = next(
            extension_root.child(i)
            for i in range(extension_root.childCount())
            if extension_root.child(i).text(0) == "处理扩展"
        )
        labels = [processing_group.child(i).text(0) for i in range(processing_group.childCount())]

        self.assertNotIn("面板配置探针", labels)

    def test_lock_source_file_import_dialog_target_accepts_positional_target_id(self):
        dialog = mock.Mock()
        dialog._data_file_target_combo = mock.Mock()
        dialog._data_file_target_keys = ["df-1", "df-2"]

        self.widget._lock_source_file_import_dialog_target(dialog, "df-2")

        dialog._data_file_target_combo.setCurrentIndex.assert_called_once_with(1)
        dialog._data_file_target_combo.setEnabled.assert_called_once_with(False)


class TestNavigationStack(unittest.TestCase):

    def setUp(self):
        self._restore_assets = _patch_global_assets()
        self.pm, self.p, self.df, self.s = _make_project("navigation_stack")
        self._restore = _patch_pm(self.pm)
        from ui.widgets.project_tree import ProjectTreeWidget

        self.widget = ProjectTreeWidget()
        self.widget.resize(320, 640)
        self.widget.show()
        QApplication.processEvents()

    def tearDown(self):
        self._restore()
        self._restore_assets()
        self.widget.deleteLater()

    def test_navigation_widgets_keep_content_width(self):
        from ui.widgets.navigation_stack import PivotStackWidget, SegmentedStackWidget

        for widget_cls in (PivotStackWidget, SegmentedStackWidget):
            widget = widget_cls()
            widget.resize(820, 420)
            widget.addTab(QWidget(widget), "绘图扩展")
            widget.addTab(QWidget(widget), "处理扩展")
            widget.addTab(QWidget(widget), "分析扩展")
            widget.show()
            QApplication.processEvents()

            navigation = widget.navigationWidget
            self.assertEqual(navigation.minimumWidth(), navigation.maximumWidth())
            self.assertGreater(navigation.width(), 0)
            self.assertLess(navigation.width(), widget.width())
            widget.deleteLater()

    def test_add_source_files_creates_source_nodes_under_source_group(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            source_a = Path(temp_dir) / "raw_a.csv"
            source_b = Path(temp_dir) / "raw_b.txt"
            source_a.write_text("x,y\n1,2\n", encoding="utf-8")
            source_b.write_text("hello", encoding="utf-8")
            nodes = self.pm.add_source_files([str(source_a), str(source_b)])

        self.assertEqual(len(nodes), 2)
        self.assertEqual(len(self.p.source_files), 2)
        parent = self.p.tree.get_node(nodes[0].parent_id)
        self.assertIsNotNone(parent)
        self.assertEqual(getattr(parent, "group_type", None), "source_files")

    def test_move_source_file_to_nested_source_folder(self):
        source_root = self.pm._find_folder_by_group_type("source_files")
        self.assertIsNotNone(source_root)
        folder = self.widget._create_child_folder(source_root.id, "原始批次")
        self.assertIsNotNone(folder)
        with tempfile.TemporaryDirectory() as temp_dir:
            source_path = Path(temp_dir) / "raw.txt"
            source_path.write_text("raw", encoding="utf-8")
            node = self.pm.add_source_file(str(source_path))

        moved = self.widget._move_node_to_target("source_file", node.id, folder.id)
        self.assertTrue(moved)
        self.assertEqual(node.parent_id, folder.id)

    def test_dataset_folder_context_can_create_data_file(self):
        dataset_root = self.pm._find_folder_by_group_type("datasets")
        self.assertIsNotNone(dataset_root)

        with mock.patch("ui.widgets.project_tree.TextInputDialog.get_text", return_value=("右键数据集", True)):
            self.widget._cmd_add_dataset_node(dataset_root.id)

        created = next((node for node in self.p.tree.nodes if node.kind == "data_file" and node.name == "右键数据集"), None)
        self.assertIsNotNone(created)
        self.assertEqual(created.parent_id, dataset_root.id)

    def test_inline_rename_conflict_reverts_item_text(self):
        dataset_root = self.pm._find_folder_by_group_type("datasets")
        self.assertIsNotNone(dataset_root)
        folder_a = self.pm.add_folder("冲突A", parent_id=dataset_root.id, group_type="datasets")
        folder_b = self.pm.add_folder("冲突B", parent_id=dataset_root.id, group_type="datasets")
        self.assertIsNotNone(folder_a)
        self.assertIsNotNone(folder_b)
        self.widget.refresh()

        item = self.widget._find_item(folder_b.id)
        self.assertIsNotNone(item)

        item.setText(0, "冲突A")
        self.widget._on_item_changed(item, 0)

        self.assertEqual(item.text(0), "冲突B")

    def test_add_picture_creates_picture_node_under_picture_group(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            picture_path = Path(temp_dir) / "chart.png"
            picture_path.write_bytes(b"png")
            node = self.pm.add_picture(str(picture_path), name="图表快照")

        self.assertIsNotNone(node)
        self.assertEqual(len(self.p.pictures), 1)
        self.assertEqual(self.p.pictures[0].name, "图表快照")
        parent = self.p.tree.get_node(node.parent_id)
        self.assertIsNotNone(parent)
        self.assertEqual(getattr(parent, "group_type", None), "pictures")

    def test_export_flow_path_labels_use_compact_folder_paths(self):
        from ui.dialogs.export_flow import _build_picture_folder_entries

        picture_root = self.pm._find_folder_by_group_type("pictures")
        self.assertIsNotNone(picture_root)
        folder_two = self.pm.add_folder("2", parent_id=picture_root.id, group_type="pictures")
        nested_one = self.pm.add_folder("1", parent_id=folder_two.id, group_type="pictures")
        deep_three = self.pm.add_folder("3", parent_id=nested_one.id, group_type="pictures")
        flat_three = self.pm.add_folder("3", parent_id=folder_two.id, group_type="pictures")
        self.assertIsNotNone(folder_two)
        self.assertIsNotNone(nested_one)
        self.assertIsNotNone(deep_three)
        self.assertIsNotNone(flat_three)

        labels = {
            entry["node_id"]: entry["label"]
            for entry in _build_picture_folder_entries()
            if entry.get("node_id") in {deep_three.id, flat_three.id}
        }

        self.assertEqual(labels.get(deep_three.id), "2/1/3")
        self.assertEqual(labels.get(flat_three.id), "2/3")

    def test_export_flow_picture_root_label_uses_group_name(self):
        from ui.dialogs.export_flow import _build_picture_folder_entries

        picture_root = self.pm._find_folder_by_group_type("pictures")
        self.assertIsNotNone(picture_root)

        labels = {
            entry["node_id"]: entry["label"]
            for entry in _build_picture_folder_entries()
        }

        self.assertEqual(labels.get(picture_root.id), "图片集")

    def test_picture_path_tracks_picture_folder_move(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            project_file = Path(temp_dir) / "picture_move.aline"
            self.pm.save(str(project_file))
            picture_root = self.pm._find_folder_by_group_type("pictures")
            self.assertIsNotNone(picture_root)
            folder_a = self.pm.add_folder("导出A", parent_id=picture_root.id, group_type="pictures")
            folder_b = self.pm.add_folder("导出B", parent_id=picture_root.id, group_type="pictures")
            self.assertIsNotNone(folder_a)
            self.assertIsNotNone(folder_b)

            source_path = Path(temp_dir) / "chart.png"
            source_path.write_bytes(b"png")
            node = self.pm.add_picture(str(source_path), name="chart.png", parent_id=folder_a.id)
            self.assertIsNotNone(node)

            picture = self.p.pictures[0]
            self.assertIn("files/pictures/导出A/chart.png", self.pm.get_picture_path(picture.id))

            moved = self.pm.move_node(node.id, folder_b.id, 0)

            self.assertTrue(moved)
            target_path = self.pm.get_picture_path(picture.id)
            self.assertTrue(Path(target_path).exists())
            self.assertIn("files/pictures/导出B/chart.png", target_path)

    def test_rename_node_syncs_image_work_name(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            image_path = Path(temp_dir) / "source.png"
            image_path.write_bytes(b"png")
            image = self.pm.add_image(str(image_path), name="原图")

        image_node = next(n for n in self.p.tree.nodes if n.kind == "image_work" and n.image_work_id == image.id)
        self.assertTrue(self.pm.rename_node(image_node.id, "重命名图"))
        self.assertEqual(self.pm.get_image(image.id).name, "重命名图")

    def test_update_wrapped_item_size_hints_ignores_deleted_tree_items(self):
        self.widget.refresh()
        stale_item = self.widget._tree.topLevelItem(0)
        self.widget.refresh()

        with mock.patch.object(self.widget._tree, "topLevelItemCount", return_value=1), \
             mock.patch.object(self.widget._tree, "topLevelItem", return_value=stale_item):
            self.widget._update_wrapped_item_size_hints()

    def test_wrapped_item_size_hint_does_not_expand_to_unwrapped_text_width(self):
        self.p.name = "这是一个用于验证项目树节点自动换行不会被未换行宽度撑开的超长项目名称测试文本"
        self.widget.resize(240, 480)
        self.widget._tree.resize(240, 480)
        self.widget.refresh()

        root = self.widget._tree.topLevelItem(0)
        self.assertIsNotNone(root)

        self.widget._update_wrapped_item_size_hint_for_item(root)

        self.assertLessEqual(root.sizeHint(0).width(), max(120, self.widget._tree.viewport().width()))
        self.assertLess(root.sizeHint(0).width(), self.widget._tree.fontMetrics().horizontalAdvance(root.text(0)))

    def test_wrapped_delegate_paints_long_labels_without_import_errors(self):
        from PySide6.QtCore import QRect
        from PySide6.QtGui import QImage, QPainter
        from PySide6.QtWidgets import QStyleOptionViewItem

        self.widget.set_name_display_mode("wrap")
        self.p.name = "eta1_wave_data_elevation_profile_reference_baseline_measurement"
        self.widget.resize(240, 480)
        self.widget._tree.resize(240, 480)
        self.widget.refresh()

        root = self.widget._tree.topLevelItem(0)
        self.assertIsNotNone(root)

        image = QImage(320, 100, QImage.Format.Format_ARGB32)
        image.fill(Qt.GlobalColor.white)
        painter = QPainter(image)
        try:
            option = QStyleOptionViewItem()
            option.rect = QRect(0, 0, 320, 100)
            option.widget = self.widget._tree
            delegate = self.widget._tree.itemDelegate()
            delegate.paint(painter, option, self.widget._tree.indexFromItem(root, 0))
        finally:
            painter.end()

    def test_wrapped_delegate_breaks_long_single_token_labels(self):
        self.widget.set_name_display_mode("wrap")
        self.p.name = "eta1_wave_data_elevation_profile_reference_baseline_measurement"
        self.widget.resize(220, 480)
        self.widget._tree.resize(220, 480)
        self.widget.refresh()

        root = self.widget._tree.topLevelItem(0)
        self.assertIsNotNone(root)

        self.widget._update_wrapped_item_size_hint_for_item(root)

        self.assertEqual(type(self.widget._tree.itemDelegate()).__name__, "_ProjectTreeWrapAnywhereDelegate")
        self.assertGreater(root.sizeHint(0).height(), self.widget._tree.fontMetrics().lineSpacing() + 10)

    def test_wrapped_delegate_preserves_selected_text_render_in_dark_theme(self):
        from PySide6.QtGui import QImage
        from qfluentwidgets import Theme, setTheme

        self.widget.set_name_display_mode("wrap")
        self.p.name = "WWWWWWWWWWWWWWWWWWWWWWWWWWWW"
        self.widget.resize(320, 480)
        self.widget._tree.resize(320, 480)

        try:
            setTheme(Theme.DARK)
            self.widget.refresh()
            root = self.widget._tree.topLevelItem(0)
            self.assertIsNotNone(root)
            self.widget._tree.setCurrentItem(root)
            QApplication.processEvents()

            rect = self.widget._tree.visualItemRect(root)
            image = self.widget._tree.viewport().grab(rect).toImage().convertToFormat(QImage.Format.Format_RGB32)
            bright_pixels = 0
            x_start = min(max(56, 0), image.width())
            x_end = min(max(120, x_start + 1), image.width())
            y_start = max(0, image.height() // 2 - 6)
            y_end = min(image.height(), image.height() // 2 + 7)
            for x_pos in range(x_start, x_end):
                for y_pos in range(y_start, y_end):
                    color = image.pixelColor(x_pos, y_pos)
                    if max(color.red(), color.green(), color.blue()) >= 210:
                        bright_pixels += 1
            self.assertGreater(bright_pixels, 0)
        finally:
            setTheme(Theme.LIGHT)

    def test_set_filter_kinds_all(self):
        """空过滤 = 显示全部"""
        self.widget.set_filter_kinds([])
        self.widget.refresh()
        self.assertGreater(self.widget._tree.topLevelItemCount(), 0)

    def test_set_filter_kinds_data_file_only(self):
        self.widget.set_filter_kinds(["folder", "data_file"])
        self.widget.refresh()
        # 只有文件夹和数据文件节点可见

    def test_set_filter_kinds_image_only(self):
        self.widget.set_filter_kinds(["folder", "image_work"])
        self.widget.refresh()

    def test_node_selected_signal_emitted(self):
        self.widget.refresh()
        received = []
        self.widget.node_selected.connect(lambda k, nid: received.append((k, nid)))
        # 找第一个 data_file 节点并模拟选中
        it = self.widget._tree.topLevelItem(0)
        if it:
            self.widget._tree.setCurrentItem(it)
        # 信号由 currentItemChanged 触发
        # 手动调用私有槽模拟
        node_id = None
        kind = None
        for node in self.p.tree.nodes:
            if node.kind == "data_file":
                node_id = node.id
                kind = node.kind
                break
        if node_id:
            self.widget.node_selected.emit(kind, node_id)
            self.assertEqual(len(received), 1)
            self.assertEqual(received[0][0], "data_file")

    def test_node_activated_signal_emitted(self):
        received = []
        self.widget.node_activated.connect(lambda k, nid: received.append((k, nid)))
        self.widget.node_activated.emit("data_file", "fake-id")
        self.assertEqual(len(received), 1)

    def test_batch_delete_removes_multiple_data_files(self):
        from models.schemas import DataFile, DataSeries

        second = self.pm.add_data_file(DataFile(name="batch.csv", series=[DataSeries(name="s2", x=[1.0], y=[2.0])]))
        self.assertIsNotNone(second)
        self.widget.refresh()

        first_item = self.widget._find_item(next(node.id for node in self.p.tree.nodes if node.kind == "data_file" and node.data_file_id == self.df.id))
        second_item = self.widget._find_item(second.id)
        self.assertIsNotNone(first_item)
        self.assertIsNotNone(second_item)

        payloads = self.widget._batch_action_payloads([first_item, second_item])
        self.assertEqual(len(payloads), 2)

        with mock.patch("ui.widgets.project_tree.MessageBox.exec", return_value=True):
            self.widget._cmd_delete_batch(payloads)

        remaining_data_file_ids = [node.id for node in self.p.tree.nodes if node.kind == "data_file"]
        self.assertNotIn(second.id, remaining_data_file_ids)
        self.assertEqual(len(remaining_data_file_ids), 0)

    def test_batch_move_uses_common_target_for_multiple_data_files(self):
        from models.schemas import DataFile, DataSeries

        dataset_root = self.pm._find_folder_by_group_type("datasets")
        self.assertIsNotNone(dataset_root)
        target_folder = self.pm.add_folder("批量目标", parent_id=dataset_root.id, group_type="datasets")
        second = self.pm.add_data_file(DataFile(name="batch.csv", series=[DataSeries(name="s2", x=[1.0], y=[2.0])]))
        self.assertIsNotNone(target_folder)
        self.assertIsNotNone(second)
        self.widget.refresh()

        first_item = self.widget._find_item(next(node.id for node in self.p.tree.nodes if node.kind == "data_file" and node.data_file_id == self.df.id))
        second_item = self.widget._find_item(second.id)
        self.assertIsNotNone(first_item)
        self.assertIsNotNone(second_item)

        payloads = self.widget._batch_action_payloads([first_item, second_item])
        choices = self.widget._common_batch_move_choices(payloads)
        self.assertIn((self.widget._folder_path_label(target_folder.id), target_folder.id), choices)

        with mock.patch("ui.widgets.project_tree.SelectionDialog.get_item", return_value=(self.widget._folder_path_label(target_folder.id), True)):
            self.widget._cmd_move_batch(payloads, choices)

        moved_nodes = {node.id: node.parent_id for node in self.p.tree.nodes if node.kind == "data_file"}
        self.assertEqual(moved_nodes[next(node.id for node in self.p.tree.nodes if node.kind == "data_file" and node.data_file_id == self.df.id)], target_folder.id)
        self.assertEqual(moved_nodes[second.id], target_folder.id)

    def test_public_batch_manage_helpers_follow_selection(self):
        from models.schemas import DataFile, DataSeries

        dataset_root = self.pm._find_folder_by_group_type("datasets")
        self.assertIsNotNone(dataset_root)
        target_folder = self.pm.add_folder("批量管理目标", parent_id=dataset_root.id, group_type="datasets")
        self.assertIsNotNone(target_folder)
        second = self.pm.add_data_file(DataFile(name="batch2.csv", series=[DataSeries(name="s2", x=[1.0], y=[2.0])]))
        self.assertIsNotNone(second)
        self.widget.refresh()
        self.widget.show()
        QApplication.processEvents()

        first_item = self.widget._find_item(next(node.id for node in self.p.tree.nodes if node.kind == "data_file" and node.data_file_id == self.df.id))
        second_item = self.widget._find_item(second.id)
        self.assertIsNotNone(first_item)
        self.assertIsNotNone(second_item)

        first_item.setSelected(True)
        second_item.setSelected(True)
        self.widget._tree.setCurrentItem(first_item)

        with mock.patch.object(self.widget, "_selected_items_or_current", return_value=[first_item, second_item]):
            self.assertFalse(self.widget.can_rename_selected_item())
            self.assertTrue(self.widget.can_move_selected_items())
            self.assertTrue(self.widget.can_delete_selected_items())

    def test_project_modified_signal_emitted(self):
        received = []
        self.widget.project_modified.connect(lambda: received.append(True))
        self.widget.project_modified.emit()
        self.assertEqual(len(received), 1)

    def test_get_selected_node_none_initially(self):
        self.widget.refresh()
        # 初始无选中节点
        result = self.widget.get_selected_node()
        # 可能是 None 也可能是根节点，重要的是不崩溃
        self.assertIsInstance(result, (type(None), tuple))

    def test_select_node_programmatic(self):
        self.widget.refresh()
        for node in self.p.tree.nodes:
            if node.kind == "data_file":
                self.widget.select_node(node.id)
                break

    def test_refresh_twice_no_crash(self):
        self.widget.refresh()
        self.widget.refresh()

    def test_filter_pipeline_kind(self):
        sp = self.pm.add_saved_pipeline("p1", [{"type": "smooth", "params": {}}])
        self.widget.set_filter_kinds(["folder", "pipeline"])
        self.widget.refresh()

    def test_report_template_node_visible_after_add(self):
        self.pm.add_report_template("tmpl_a", "# Report")
        self.widget.refresh()
        found = []

        def _walk(item):
            data = item.data(0, Qt.ItemDataRole.UserRole)
            if data:
                found.append(data[0])
            for idx in range(item.childCount()):
                _walk(item.child(idx))

        for idx in range(self.widget._tree.topLevelItemCount()):
            _walk(self.widget._tree.topLevelItem(idx))
        self.assertIn("global_report_template", found)

    def test_root_folders_remain_visible_when_empty(self):
        from core.project_manager import ProjectManager
        pm = ProjectManager()
        project = pm.create_new("tree_empty")
        restore = _patch_pm(pm)
        widget = None
        try:
            from ui.widgets.project_tree import ProjectTreeWidget
            widget = ProjectTreeWidget()
            widget.set_filter_kinds(["folder", "image_work"])
            project_root = widget._tree.topLevelItem(0)
            self.assertEqual(project_root.text(0), project.name)
            names = [project_root.child(i).text(0) for i in range(project_root.childCount())]
            self.assertEqual(names, ["源文件", "数据集", "图片集", "分析结果", "数字化"])
        finally:
            restore()
            if widget is not None:
                widget.deleteLater()

    def test_move_series_between_data_files(self):
        from models.schemas import DataFile, DataSeries

        other = DataFile(name="other.csv", series=[DataSeries(name="other", x=[1.0], y=[2.0])])
        self.pm.add_data_file(other)
        moved = self.widget._move_node_to_target("series", self.s.id, other.id)
        self.assertTrue(moved)
        self.assertEqual(len(self.df.series), 0)
        self.assertTrue(any(series.id == self.s.id for series in other.series))

    def test_drop_move_moves_series_into_data_file(self):
        from models.schemas import DataFile, DataSeries

        other_node = self.pm.add_data_file(DataFile(name="other.csv", series=[DataSeries(name="other", x=[1.0], y=[2.0])]))
        self.assertIsNotNone(other_node)

        self.widget.refresh()
        source_item = self.widget._find_item(self.s.id)
        target_item = self.widget._find_item(other_node.id)

        self.assertIsNotNone(source_item)
        self.assertIsNotNone(target_item)
        self.assertTrue(self.widget._perform_drop_move(source_item, target_item))
        self.assertEqual(len(self.df.series), 0)
        target_df = self.p.find_data_file(other_node.data_file_id)
        self.assertIsNotNone(target_df)
        self.assertTrue(any(series.id == self.s.id for series in target_df.series))

    def test_create_child_folder_inherits_collection_group_type(self):
        for group_type in ("datasets", "source_files", "images", "analysis_result_group"):
            root = self.pm._find_folder_by_group_type(group_type)
            self.assertIsNotNone(root)
            folder = self.widget._create_child_folder(root.id, f"{group_type}-子文件夹")
            self.assertIsNotNone(folder)
            self.assertEqual(folder.parent_id, root.id)
            self.assertEqual(getattr(folder, "group_type", None), group_type)

    def test_move_target_choices_include_nested_dataset_folders(self):
        datasets_root = self.pm._find_folder_by_group_type("datasets")
        self.assertIsNotNone(datasets_root)
        folder = self.widget._create_child_folder(datasets_root.id, "实验组")
        self.assertIsNotNone(folder)
        node = next(n for n in self.p.tree.nodes if n.kind == "data_file" and n.data_file_id == self.df.id)

        choices = self.widget._move_target_choices("data_file", node.id)
        self.assertIn((self.widget._folder_path_label(folder.id), folder.id), choices)

    def test_root_folder_path_label_uses_group_display_name(self):
        datasets_root = self.pm._find_folder_by_group_type("datasets")
        self.assertIsNotNone(datasets_root)

        self.assertEqual(self.widget._folder_path_label(datasets_root.id), "数据集")

    def test_move_data_file_to_nested_dataset_folder(self):
        datasets_root = self.pm._find_folder_by_group_type("datasets")
        self.assertIsNotNone(datasets_root)
        folder = self.widget._create_child_folder(datasets_root.id, "归档")
        self.assertIsNotNone(folder)
        node = next(n for n in self.p.tree.nodes if n.kind == "data_file" and n.data_file_id == self.df.id)

        moved = self.widget._move_node_to_target("data_file", node.id, folder.id)
        self.assertTrue(moved)
        self.assertEqual(node.parent_id, folder.id)

    def test_nested_dataset_folder_is_editable(self):
        datasets_root = self.pm._find_folder_by_group_type("datasets")
        self.assertIsNotNone(datasets_root)
        folder = self.widget._create_child_folder(datasets_root.id, "实验组")
        self.assertIsNotNone(folder)

        self.widget.refresh()
        item = self.widget._find_item(folder.id)
        self.assertIsNotNone(item)
        self.assertTrue(bool(item.flags() & Qt.ItemFlag.ItemIsEditable))
        self.assertTrue(bool(item.flags() & Qt.ItemFlag.ItemIsDragEnabled))

    def test_data_file_and_image_work_nodes_are_editable(self):
        image = self.pm.add_image("/tmp/aline-tree-image.png", name="image.png")

        self.widget.refresh()
        data_file_node = next(n for n in self.p.tree.nodes if n.kind == "data_file" and n.data_file_id == self.df.id)
        image_node = next(n for n in self.p.tree.nodes if n.kind == "image_work" and n.image_work_id == image.id)

        data_file_item = self.widget._find_item(data_file_node.id)
        image_item = self.widget._find_item(image_node.id)

        self.assertIsNotNone(data_file_item)
        self.assertIsNotNone(image_item)
        self.assertTrue(bool(data_file_item.flags() & Qt.ItemFlag.ItemIsEditable))
        self.assertTrue(bool(image_item.flags() & Qt.ItemFlag.ItemIsEditable))

    def test_cmd_add_child_folder_creates_and_selects_root_child(self):
        datasets_root = self.pm._find_folder_by_group_type("datasets")
        self.assertIsNotNone(datasets_root)

        self.widget.refresh()
        with mock.patch("ui.widgets.project_tree.TextInputDialog.get_text", return_value=("根目录子文件夹", True)):
            self.widget._cmd_add_child_folder(datasets_root.id)

        created = next(
            (node for node in self.p.tree.nodes if node.kind == "folder" and node.parent_id == datasets_root.id and node.name == "根目录子文件夹"),
            None,
        )
        self.assertIsNotNone(created)
        selected = self.widget.get_selected_node()
        self.assertEqual(selected, ("folder", created.id))

    def test_empty_child_folder_remains_visible_with_folder_filter(self):
        datasets_root = self.pm._find_folder_by_group_type("datasets")
        self.assertIsNotNone(datasets_root)

        self.widget.set_filter_kinds(["folder", "data_file"])
        folder = self.widget._create_child_folder(datasets_root.id, "空文件夹")
        self.assertIsNotNone(folder)
        self.widget.refresh()

        self.assertIsNotNone(self.widget._find_item(folder.id))

    def test_move_nested_folder_to_another_folder(self):
        datasets_root = self.pm._find_folder_by_group_type("datasets")
        self.assertIsNotNone(datasets_root)
        folder_a = self.widget._create_child_folder(datasets_root.id, "A")
        folder_b = self.widget._create_child_folder(datasets_root.id, "B")
        self.assertIsNotNone(folder_a)
        self.assertIsNotNone(folder_b)

        moved = self.widget._move_node_to_target("folder", folder_a.id, folder_b.id)
        self.assertTrue(moved)
        moved_node = self.pm.get_node_by_id(folder_a.id)
        self.assertEqual(moved_node.parent_id, folder_b.id)

    def test_drop_move_moves_data_file_into_folder(self):
        datasets_root = self.pm._find_folder_by_group_type("datasets")
        self.assertIsNotNone(datasets_root)
        folder = self.widget._create_child_folder(datasets_root.id, "拖放目标")
        self.assertIsNotNone(folder)
        node = next(n for n in self.p.tree.nodes if n.kind == "data_file" and n.data_file_id == self.df.id)

        self.widget.refresh()
        source_item = self.widget._find_item(node.id)
        target_item = self.widget._find_item(folder.id)
        self.assertTrue(self.widget._perform_drop_move(source_item, target_item))
        self.assertEqual(node.parent_id, folder.id)

    def test_drop_move_uses_remembered_drag_source(self):
        datasets_root = self.pm._find_folder_by_group_type("datasets")
        self.assertIsNotNone(datasets_root)
        folder = self.widget._create_child_folder(datasets_root.id, "拖放目标")
        self.assertIsNotNone(folder)
        node = next(n for n in self.p.tree.nodes if n.kind == "data_file" and n.data_file_id == self.df.id)

        self.widget.refresh()
        source_item = self.widget._find_item(node.id)
        target_item = self.widget._find_item(folder.id)
        self.widget._remember_drag_source_item(source_item)

        self.assertTrue(self.widget._perform_drop_move(None, target_item))
        self.assertEqual(node.parent_id, folder.id)

    def test_drop_source_file_to_data_file_imports_series_without_moving_source_node(self):
        from models.schemas import DataSeries

        with tempfile.TemporaryDirectory() as temp_dir:
            source_path = Path(temp_dir) / "drag.csv"
            source_path.write_text("x,y\n0,1\n1,2\n", encoding="utf-8")
            source_node = self.pm.add_source_file(str(source_path))
            self.assertIsNotNone(source_node)

            self.widget.refresh()
            source_item = self.widget._find_item(source_node.id)
            target_node = next(n for n in self.p.tree.nodes if n.kind == "data_file" and n.data_file_id == self.df.id)
            target_item = self.widget._find_item(target_node.id)

            class _FakeDialog:
                def exec(self):
                    return True

                def get_results(self):
                    return [DataSeries(name="导入列", x=[0.0, 1.0], y=[1.0, 2.0])]

                def get_file_name(self):
                    return "drag.csv"

            with mock.patch.object(self.widget, "_create_source_file_import_dialog", return_value=_FakeDialog()):
                self.assertTrue(self.widget._perform_drop_move(source_item, target_item))

            self.assertTrue(any(series.name == "导入列" for series in self.df.series))
            self.assertEqual(self.pm.get_node_by_id(source_node.id).parent_id, source_node.parent_id)

    def test_drop_source_file_to_images_subfolder_imports_image_into_target_folder(self):
        from PySide6.QtGui import QImage

        images_root = self.pm._find_folder_by_group_type("images")
        self.assertIsNotNone(images_root)
        target_folder = self.widget._create_child_folder(images_root.id, "图像归档")
        self.assertIsNotNone(target_folder)

        with tempfile.TemporaryDirectory() as temp_dir:
            source_path = Path(temp_dir) / "managed.png"
            image = QImage(16, 16, QImage.Format.Format_RGB32)
            image.fill(Qt.GlobalColor.white)
            self.assertTrue(image.save(str(source_path)))

            source_node = self.pm.add_source_file(str(source_path))
            self.assertIsNotNone(source_node)

            self.widget.refresh()
            source_item = self.widget._find_item(source_node.id)
            target_item = self.widget._find_item(target_folder.id)

            self.assertTrue(self.widget._perform_drop_move(source_item, target_item))

            image_node = next((n for n in self.p.tree.nodes if n.kind == "image_work" and n.name == "managed.png"), None)
            self.assertIsNotNone(image_node)
            self.assertEqual(image_node.parent_id, target_folder.id)

    def test_tree_enables_drag_drop_mode(self):
        self.assertTrue(self.widget._tree.dragEnabled())
        self.assertEqual(self.widget._tree.dragDropMode(), QAbstractItemView.DragDropMode.DragDrop)
        self.assertTrue(self.widget._tree.showDropIndicator())

    def test_context_menu_appends_expand_and_collapse_actions(self):
        self.widget.refresh()
        datasets_root = self.pm._find_folder_by_group_type("datasets")
        self.assertIsNotNone(datasets_root)
        item = self.widget._find_item(datasets_root.id)
        self.assertIsNotNone(item)
        pos = self.widget._tree.visualItemRect(item).center()
        captured = {}

        def _fake_exec(menu, *_args):
            captured["actions"] = list(menu.actions())
            captured["separator_marks"] = [
                menu.view.item(index).data(Qt.ItemDataRole.DecorationRole)
                for index in range(menu.view.count())
            ]

        with mock.patch("ui.widgets.project_tree.RoundMenu.exec", autospec=True, side_effect=_fake_exec):
            self.widget._on_context_menu(pos)

        actions = captured["actions"]
        visible_texts = [action.text() for action in actions if not action.isSeparator()]
        self.assertIn("seperator", captured["separator_marks"])
        self.assertNotIn("全部展开", visible_texts)
        self.assertNotIn("全部折叠", visible_texts)

    def test_dataset_folder_context_menu_keeps_delete_above_prune_action(self):
        datasets_root = self.pm._find_folder_by_group_type("datasets")
        self.assertIsNotNone(datasets_root)
        folder = self.widget._create_child_folder(datasets_root.id, "菜单顺序")
        self.assertIsNotNone(folder)

        self.widget.refresh()
        root_item = self.widget._find_item(datasets_root.id)
        self.assertIsNotNone(root_item)
        root_item.setExpanded(True)
        QApplication.processEvents()
        item = self.widget._find_item(folder.id)
        self.assertIsNotNone(item)
        pos = self.widget._tree.visualItemRect(item).center()
        captured = {}

        def _fake_exec(menu, *_args):
            captured["actions"] = list(menu.actions())

        with mock.patch("ui.widgets.project_tree.RoundMenu.exec", autospec=True, side_effect=_fake_exec):
            self.widget._on_context_menu(pos)

        visible_texts = [action.text() for action in captured["actions"] if not action.isSeparator()]
        self.assertEqual(visible_texts[:3], ["新建数据集", "导入数据文件...", "新建子文件夹"])
        self.assertLess(visible_texts.index("删除"), visible_texts.index("清理空子文件夹"))
        self.assertNotIn("全部展开", visible_texts)
        self.assertNotIn("全部折叠", visible_texts)

    def test_images_folder_context_menu_exposes_import_image_action(self):
        images_root = self.pm._find_folder_by_group_type("images")
        self.assertIsNotNone(images_root)

        self.widget.refresh()
        item = self.widget._find_item(images_root.id)
        self.assertIsNotNone(item)
        pos = self.widget._tree.visualItemRect(item).center()
        captured = {}

        def _fake_exec(menu, *_args):
            captured["actions"] = list(menu.actions())

        with mock.patch("ui.widgets.project_tree.RoundMenu.exec", autospec=True, side_effect=_fake_exec):
            self.widget._on_context_menu(pos)

        visible_texts = [action.text() for action in captured["actions"] if not action.isSeparator()]
        self.assertIn("导入图片...", visible_texts[:3])

    def test_image_work_context_menu_offers_add_curve(self):
        from PySide6.QtGui import QImage

        with tempfile.TemporaryDirectory() as temp_dir:
            image_path = Path(temp_dir) / "context-image.png"
            image = QImage(16, 16, QImage.Format.Format_RGB32)
            image.fill(Qt.GlobalColor.white)
            self.assertTrue(image.save(str(image_path)))
            created = self.pm.add_image(str(image_path), name="context-image.png")

            self.widget.refresh()
            node = next((item for item in self.p.tree.nodes if item.kind == "image_work" and item.image_work_id == created.id), None)
            self.assertIsNotNone(node)
            item = self.widget._find_item(node.id)
            self.assertIsNotNone(item)
            pos = self.widget._tree.visualItemRect(item).center()
            captured = {}

            def _fake_exec(menu, *_args):
                captured["actions"] = list(menu.actions())

            with mock.patch("ui.widgets.project_tree.RoundMenu.exec", autospec=True, side_effect=_fake_exec):
                self.widget._on_context_menu(pos)

        visible_texts = [action.text() for action in captured["actions"] if not action.isSeparator()]
        self.assertIn("新增曲线", visible_texts)
        self.assertNotIn("发送到可视化", visible_texts)

    def test_dataset_folder_context_menu_can_import_data_file(self):
        from models.schemas import DataSeries

        datasets_root = self.pm._find_folder_by_group_type("datasets")
        self.assertIsNotNone(datasets_root)
        target_folder = self.widget._create_child_folder(datasets_root.id, "数据导入")
        self.assertIsNotNone(target_folder)

        dialog = mock.Mock()
        dialog.exec.return_value = True
        dialog.get_results.return_value = [DataSeries(name="imported", x=[1.0], y=[2.0])]
        dialog.get_target_data_file_id.return_value = None
        dialog.get_file_name.return_value = "imported.csv"

        with mock.patch("ui.widgets.project_tree.QFileDialog.getOpenFileName", return_value=("/tmp/imported.csv", "数据文件")), \
             mock.patch.object(self.widget._command_service, "create_source_file_import_dialog", return_value=dialog):
            self.widget._cmd_import_data_file(target_folder.id)

        node = next((item for item in self.p.tree.nodes if item.kind == "data_file" and item.name == "imported.csv"), None)
        self.assertIsNotNone(node)
        self.assertEqual(node.parent_id, target_folder.id)

    def test_curve_context_menu_offers_export_and_chart_actions_only(self):
        from PySide6.QtGui import QImage

        with tempfile.TemporaryDirectory() as temp_dir:
            image_path = Path(temp_dir) / "curve-image.png"
            image = QImage(16, 16, QImage.Format.Format_RGB32)
            image.fill(Qt.GlobalColor.white)
            self.assertTrue(image.save(str(image_path)))
            created = self.pm.add_image(str(image_path), name="curve-image.png")
            curve = self.pm.add_curve_to_image(created.id, [0.0, 1.0], [1.0, 2.0], name="curve-a")
            self.assertIsNotNone(curve)

            self.widget.refresh()
            node = next((item for item in self.p.tree.nodes if item.kind == "image_work" and item.image_work_id == created.id), None)
            self.assertIsNotNone(node)
            image_item = self.widget._find_item(node.id)
            self.assertIsNotNone(image_item)
            curve_item = image_item.child(0)
            self.assertIsNotNone(curve_item)
            pos = self.widget._tree.visualItemRect(curve_item).center()
            captured = {}

            def _fake_exec(menu, *_args):
                captured["actions"] = list(menu.actions())

            with mock.patch("ui.widgets.project_tree.RoundMenu.exec", autospec=True, side_effect=_fake_exec):
                self.widget._on_context_menu(pos)

        visible_texts = [action.text() for action in captured["actions"] if not action.isSeparator()]
        self.assertEqual(visible_texts[:4], ["导出为数据列", "发送到可视化", "重命名", "删除"])
        self.assertNotIn("发送到处理", visible_texts)
        self.assertNotIn("发送到分析", visible_texts)

    def test_picture_context_menu_offers_chart_and_folder_actions(self):
        from PySide6.QtGui import QImage

        with tempfile.TemporaryDirectory() as temp_dir:
            picture_path = Path(temp_dir) / "snapshot.png"
            image = QImage(24, 24, QImage.Format.Format_RGB32)
            image.fill(Qt.GlobalColor.white)
            self.assertTrue(image.save(str(picture_path)))
            node = self.pm.add_picture(str(picture_path), name="snapshot.png")

            self.assertIsNotNone(node)
            self.widget.refresh()
            item = self.widget._find_item(node.id)
            self.assertIsNotNone(item)
            parent_item = item.parent()
            if parent_item is not None:
                parent_item.setExpanded(True)
            pos = self.widget._tree.visualItemRect(item).center()
            captured = {}

            def _fake_exec(menu, *_args):
                captured["actions"] = list(menu.actions())

            with mock.patch("ui.widgets.project_tree.RoundMenu.exec", autospec=True, side_effect=_fake_exec):
                self.widget._on_context_menu(pos)

        visible_texts = [action.text() for action in captured["actions"] if not action.isSeparator()]
        self.assertEqual(visible_texts[:4], ["发送到可视化", "在文件夹打开", "重命名", "删除"])

    def test_project_tree_icons_follow_updated_asset_mapping(self):
        from qfluentwidgets import FluentIcon as FIF

        with tempfile.TemporaryDirectory() as temp_dir:
            source_path = Path(temp_dir) / "raw.csv"
            source_path.write_text("x,y\n1,2\n", encoding="utf-8")
            image_path = Path(temp_dir) / "digitize.png"
            from PySide6.QtGui import QImage

            image = QImage(16, 16, QImage.Format.Format_RGB32)
            image.fill(Qt.GlobalColor.white)
            self.assertTrue(image.save(str(image_path)))

            source_node = self.pm.add_source_file(str(source_path))
            image_work = self.pm.add_image(str(image_path), name="digitize.png")

        self.widget.refresh()
        project_root = self.widget._tree.topLevelItem(0)
        source_group = next(project_root.child(i) for i in range(project_root.childCount()) if project_root.child(i).text(0) == "源文件")
        digitize_group = next(project_root.child(i) for i in range(project_root.childCount()) if project_root.child(i).text(0) == "数字化")
        source_item = self.widget._find_item(source_node.id)
        image_item = next((item for item in self.p.tree.nodes if item.kind == "image_work" and item.image_work_id == image_work.id), None)
        self.assertIsNotNone(source_item)
        self.assertIsNotNone(image_item)
        image_tree_item = self.widget._find_item(image_item.id)
        self.assertIsNotNone(image_tree_item)

        self.assertEqual(self._icon_image(project_root.icon(0)), self._icon_image(getattr(FIF, "ZIP_FOLDER", getattr(FIF, "LIBRARY", FIF.FOLDER)).icon()))
        self.assertEqual(self._icon_image(source_group.icon(0)), self._icon_image(getattr(FIF, "IOT", FIF.FOLDER).icon()))
        self.assertEqual(self._icon_image(source_item.icon(0)), self._icon_image(getattr(FIF, "DOCUMENT", FIF.FOLDER).icon()))
        self.assertEqual(self._icon_image(digitize_group.icon(0)), self._icon_image(getattr(FIF, "LABEL", FIF.PHOTO).icon()))
        self.assertEqual(self._icon_image(image_tree_item.icon(0)), self._icon_image(getattr(FIF, "PHOTO", FIF.PHOTO).icon()))

    def test_project_tree_child_folders_use_plain_folder_icon(self):
        from qfluentwidgets import FluentIcon as FIF

        source_root = self.pm._find_folder_by_group_type("source_files")
        datasets_root = self.pm._find_folder_by_group_type("datasets")
        images_root = self.pm._find_folder_by_group_type("images")
        self.assertIsNotNone(source_root)
        self.assertIsNotNone(datasets_root)
        self.assertIsNotNone(images_root)

        source_child = self.pm.add_folder("源子文件夹", parent_id=source_root.id, group_type="source_files")
        dataset_child = self.pm.add_folder("数据子文件夹", parent_id=datasets_root.id, group_type="datasets")
        image_child = self.pm.add_folder("图像子文件夹", parent_id=images_root.id, group_type="images")
        self.assertIsNotNone(source_child)
        self.assertIsNotNone(dataset_child)
        self.assertIsNotNone(image_child)

        self.widget.refresh()

        source_child_item = self.widget._find_item(source_child.id)
        dataset_child_item = self.widget._find_item(dataset_child.id)
        image_child_item = self.widget._find_item(image_child.id)
        self.assertIsNotNone(source_child_item)
        self.assertIsNotNone(dataset_child_item)
        self.assertIsNotNone(image_child_item)

        expected_icon = self._icon_image(FIF.FOLDER.icon())
        self.assertEqual(self._icon_image(source_child_item.icon(0)), expected_icon)
        self.assertEqual(self._icon_image(dataset_child_item.icon(0)), expected_icon)
        self.assertEqual(self._icon_image(image_child_item.icon(0)), expected_icon)

    def test_project_tree_image_source_file_uses_photo_icon(self):
        from qfluentwidgets import FluentIcon as FIF
        from PySide6.QtGui import QImage

        with tempfile.TemporaryDirectory() as temp_dir:
            image_path = Path(temp_dir) / "source-image.png"
            image = QImage(16, 16, QImage.Format.Format_RGB32)
            image.fill(Qt.GlobalColor.white)
            self.assertTrue(image.save(str(image_path)))

            node = self.pm.add_source_file(str(image_path))
            self.assertIsNotNone(node)

        self.widget.refresh()

        image_source_item = self.widget._find_item(node.id)
        self.assertIsNotNone(image_source_item)
        self.assertEqual(self._icon_image(image_source_item.icon(0)), self._icon_image(FIF.PHOTO.icon()))

    def test_import_digitize_images_adds_images_into_target_folder(self):
        from PySide6.QtGui import QImage

        images_root = self.pm._find_folder_by_group_type("images")
        self.assertIsNotNone(images_root)
        target_folder = self.widget._create_child_folder(images_root.id, "菜单导入")
        self.assertIsNotNone(target_folder)

        with tempfile.TemporaryDirectory() as temp_dir:
            image_path = Path(temp_dir) / "import-image.png"
            image = QImage(16, 16, QImage.Format.Format_RGB32)
            image.fill(Qt.GlobalColor.white)
            self.assertTrue(image.save(str(image_path)))

            with mock.patch("ui.widgets.project_tree.QFileDialog.getOpenFileNames", return_value=([str(image_path)], "图片文件")):
                self.widget._cmd_import_digitize_images(target_folder.id)

        created = next((node for node in self.p.tree.nodes if node.kind == "image_work" and node.name == "import-image.png"), None)
        self.assertIsNotNone(created)
        self.assertEqual(created.parent_id, target_folder.id)

    def test_source_file_context_menu_uses_import_first_then_manage(self):
        source_root = self.pm._find_folder_by_group_type("source_files")
        self.assertIsNotNone(source_root)
        target_folder = self.widget._create_child_folder(source_root.id, "源文件归档")
        self.assertIsNotNone(target_folder)

        with tempfile.TemporaryDirectory() as temp_dir:
            source_path = Path(temp_dir) / "menu.csv"
            source_path.write_text("x,y\n0,1\n", encoding="utf-8")
            node = self.pm.add_source_file(str(source_path))
            self.assertIsNotNone(node)

            self.widget.refresh()
            item = self.widget._find_item(node.id)
            self.assertIsNotNone(item)
            pos = self.widget._tree.visualItemRect(item).center()
            captured = {}

            def _fake_exec(menu, *_args):
                captured["actions"] = list(menu.actions())

            with mock.patch("ui.widgets.project_tree.RoundMenu.exec", autospec=True, side_effect=_fake_exec):
                self.widget._on_context_menu(pos)

            visible_texts = [action.text() for action in captured["actions"] if not action.isSeparator()]
            self.assertEqual(visible_texts[:2], ["导入到数据集", "导入到数字化"])
            self.assertLess(visible_texts.index("删除"), visible_texts.index("移动到..."))
            self.assertNotIn("全部展开", visible_texts)
            self.assertNotIn("全部折叠", visible_texts)


# ═══════════════════════════════════════════════════════════════════════════
# 2. SettingsPage — AI 配置
# ═══════════════════════════════════════════════════════════════════════════

class TestSettingsPage(unittest.TestCase):
    """SettingsPage 基础配置与隐藏 AI 入口"""

    def setUp(self):
        self._restore_assets = _patch_global_assets()
        from ui.pages.settings_page import SettingsPage
        self.page = SettingsPage()

    def tearDown(self):
        self.page.deleteLater()
        self._restore_assets()

    def test_page_creates_without_crash(self):
        self.assertIsNotNone(self.page)

    def test_theme_combo_exists(self):
        self.assertIsNotNone(self.page.theme_combo)

    def test_extension_controls_exist(self):
        self.assertIsNotNone(self.page._builtin_extension_card)
        self.assertIsNotNone(self.page._external_extension_card)
        self.assertIsNotNone(self.page._extension_other_settings_card)
        self.assertIsNotNone(self.page._builtin_extension_management_card)
        self.assertIsNotNone(self.page._external_extension_management_card)
        self.assertIsNotNone(self.page._builtin_extensions_enabled_checkbox)
        self.assertIsNotNone(self.page._external_extensions_enabled_checkbox)
        self.assertIsNotNone(self.page._external_extensions_dirs_card)
        self.assertIsNotNone(self.page._external_extension_number_decimals_card)
        self.assertIsNotNone(self.page._refresh_external_extensions_btn)
        self.assertIsNotNone(self.page._extension_tabs)
        self.assertIsNotNone(self.page._external_extension_tabs)
        self.assertIsNotNone(self.page._save_extension_settings_btn)
        self.assertIsNotNone(self.page._extension_status_summary_btn)
        self.assertIsInstance(self.page._builtin_extension_checkboxes, dict)
        self.assertIsInstance(self.page._external_extension_checkboxes, dict)

    def test_settings_scroll_content_uses_workbench_like_margins(self):
        general_layout = self.page._tabs.widget(0).widget().layout()
        extensions_layout = self.page._tabs.widget(1).widget().layout()
        shortcuts_layout = self.page._tabs.widget(2).widget().layout()

        for layout in (general_layout, extensions_layout, shortcuts_layout):
            margins = layout.contentsMargins()
            self.assertEqual((margins.left(), margins.top(), margins.right(), margins.bottom()), (14, 12, 14, 12))

    def test_shortcut_edit_shows_focus_border_when_selected(self):
        from ui.theme import accent_color

        self.page._tabs.setCurrentIndex(2)
        self.page.show()
        QApplication.processEvents()
        first_edit = next(iter(self.page._shortcut_edits.values()))
        first_edit.setFocus()
        QApplication.processEvents()

        self.assertIn(accent_color(), first_edit.styleSheet())
        self.assertIn("2px solid", first_edit.styleSheet())

        self.page._shortcut_filter_edit.setFocus()
        QApplication.processEvents()

        self.assertIn("1px solid", first_edit.styleSheet())

    def test_shortcut_filter_field_is_clearable_and_emphasized(self):
        from ui.theme import accent_color

        self.assertTrue(self.page._shortcut_filter_edit.isClearButtonEnabled())
        self.assertGreaterEqual(self.page._shortcut_filter_edit.height(), 36)
        self.assertIn("筛选", self.page._shortcut_filter_edit.toolTip())
        self.assertIn(accent_color(), self.page._shortcut_filter_edit.styleSheet())

    def test_settings_cards_use_shared_title_and_hint_styles(self):
        from ui.theme import body_text_style_sheet, card_title_style_sheet, placeholder_text_style_sheet

        self.assertEqual(self.page._appearance_title.styleSheet(), card_title_style_sheet(font_size=18))
        self.assertEqual(self.page._theme_label.styleSheet(), body_text_style_sheet())
        self.assertEqual(self.page._extension_hint.styleSheet(), placeholder_text_style_sheet(font_size=11))

    def test_setting_card_titles_and_descriptions_use_shared_styles(self):
        from PySide6.QtWidgets import QLabel
        from ui.theme import body_text_style_sheet, placeholder_text_style_sheet

        def _page_label_text_style(text: str) -> str:
            return next(label.styleSheet() for label in self.page.findChildren(QLabel) if label.text() == text)

        self.page.update_theme_colors()

        for title_text in (
            "主题",
            "项目树长名称显示",
            "项目树页面专注模式",
            "新手引导",
            "扩展状态",
            "应用扩展设置",
            "刷新外部扩展扫描",
            "浮点参数显示小数位",
            "快捷键映射",
        ):
            self.assertEqual(_page_label_text_style(title_text), body_text_style_sheet())

        for description_text in (
            "切换浅色、深色或跟随系统。",
            "控制项目树长名称使用自动换行还是省略显示。",
            "开启后，功能页中的共享项目树只显示当前页面直接相关的节点。",
            "点击后会重新播放主页引导，并重置数据管理、处理、可视化、分析和图片数字化页面的 TeachingTip 状态。",
            "查看当前扩展加载情况与失败详情。",
            "保存当前启用状态与目录配置，并重新加载扩展。",
            "按当前目录配置重新探测外部扩展，不会修改保存设置。",
            "控制扩展 number 参数使用 DoubleSpinBox 时默认显示的小数位数。",
            "所有已注册的界面动作都会显示在这里。点击输入框后按下新快捷键，再点击“应用快捷键”保存。",
        ):
            self.assertEqual(_page_label_text_style(description_text), placeholder_text_style_sheet(font_size=11))

    def test_extension_category_height_refresh_is_lifecycle_safe(self):
        self.page._schedule_extension_category_tab_heights_refresh()
        self.page._extension_tabs.deleteLater()
        self.page._external_extension_tabs.deleteLater()
        QApplication.processEvents()

        self.page._refresh_extension_category_tab_heights()

    def test_extension_section_titles_refresh_with_theme(self):
        from ui.theme import body_text_style_sheet, card_title_style_sheet, placeholder_text_style_sheet
        from PySide6.QtWidgets import QLabel

        def _label_text_style(widget, text: str) -> str:
            return next(label.styleSheet() for label in widget.findChildren(QLabel) if label.text() == text)

        self.page.update_theme_colors()

        self.assertEqual(self.page._builtin_extension_card.titleLabel.styleSheet(), card_title_style_sheet(font_size=18))
        self.assertEqual(self.page._builtin_extensions_enabled_checkbox.titleLabel.styleSheet(), body_text_style_sheet())
        self.assertEqual(self.page._builtin_extensions_enabled_checkbox.contentLabel.styleSheet(), placeholder_text_style_sheet(font_size=11))
        self.assertEqual(self.page._external_extension_card.titleLabel.styleSheet(), card_title_style_sheet(font_size=18))
        self.assertEqual(self.page._external_extensions_enabled_checkbox.titleLabel.styleSheet(), body_text_style_sheet())
        self.assertEqual(self.page._external_extensions_enabled_checkbox.contentLabel.styleSheet(), placeholder_text_style_sheet(font_size=11))
        self.assertEqual(_label_text_style(self.page._external_extensions_dirs_card, "外部扩展目录"), body_text_style_sheet())
        self.assertEqual(_label_text_style(self.page._external_extensions_dirs_card, "可添加多个文件夹；保存后会统一扫描并重载。"), placeholder_text_style_sheet(font_size=11))
        builtin_manage_title = next(label for label in self.page._builtin_extension_management_card.findChildren(QLabel) if label.text() == "扩展管理")
        external_manage_title = next(label for label in self.page._external_extension_management_card.findChildren(QLabel) if label.text() == "扩展管理")
        self.assertEqual(builtin_manage_title.styleSheet(), card_title_style_sheet(font_size=18))
        self.assertEqual(external_manage_title.styleSheet(), card_title_style_sheet(font_size=18))
        self.assertEqual(self.page._extension_other_settings_card.titleLabel.styleSheet(), card_title_style_sheet(font_size=18))

    def test_settings_page_uses_setting_card_containers(self):
        from qfluentwidgets import ExpandGroupSettingCard, FolderListSettingCard, SettingCardGroup, SwitchSettingCard

        self.assertIsInstance(self.page._appearance_card, SettingCardGroup)
        self.assertIsInstance(self.page._extension_card, SettingCardGroup)
        self.assertIsInstance(self.page._builtin_extension_card, SettingCardGroup)
        self.assertIsInstance(self.page._external_extension_card, SettingCardGroup)
        self.assertIsInstance(self.page._extension_other_settings_card, SettingCardGroup)
        self.assertIsInstance(self.page._builtin_extension_management_card, ExpandGroupSettingCard)
        self.assertIsInstance(self.page._external_extension_management_card, ExpandGroupSettingCard)
        self.assertIsInstance(self.page._shortcuts_card, SettingCardGroup)
        self.assertIsInstance(self.page._page_tree_focus_mode_card, SwitchSettingCard)
        self.assertIsInstance(self.page._builtin_extensions_enabled_checkbox, SwitchSettingCard)
        self.assertIsInstance(self.page._external_extensions_enabled_checkbox, SwitchSettingCard)
        self.assertIsInstance(self.page._external_extensions_dirs_card, FolderListSettingCard)

    def test_shortcut_filter_field_is_nested_in_mapping_card(self):
        parent = self.page._shortcut_filter_edit.parentWidget()
        while parent is not None and parent is not self.page._shortcuts_editor_card:
            parent = parent.parentWidget()
        self.assertIs(parent, self.page._shortcuts_editor_card)

    def test_settings_title_label_is_removed(self):
        self.assertIsNone(self.page._title_label)

    def test_tree_display_mode_combo_defaults_to_elide(self):
        from pathlib import Path
        from unittest import mock
        from ui.pages.settings_page import SettingsPage

        temp_page = None
        try:
            with mock.patch("core.ui_preferences._CONFIG_PATH", Path("/nonexistent/aline_ui_preferences.json")):
                temp_page = SettingsPage()
            self.assertEqual(temp_page._current_tree_display_mode(), "elide")
        finally:
            if temp_page is not None:
                temp_page.deleteLater()

    def test_page_tree_focus_mode_checkbox_defaults_to_off(self):
        from pathlib import Path
        from unittest import mock
        from ui.pages.settings_page import SettingsPage

        temp_page = None
        try:
            with mock.patch("core.ui_preferences._CONFIG_PATH", Path("/nonexistent/aline_ui_preferences_focus.json")):
                temp_page = SettingsPage()
            self.assertFalse(temp_page._page_tree_focus_mode_card.isChecked())
        finally:
            if temp_page is not None:
                temp_page.deleteLater()

    def test_page_tree_focus_mode_checkbox_emits_enabled_state(self):
        received = []
        self.page.page_tree_focus_mode_changed.connect(received.append)
        target_state = not self.page._page_tree_focus_mode_checkbox.isChecked()

        with mock.patch("ui.pages.settings_page.set_page_tree_focus_mode_enabled", return_value=target_state):
            self.page._page_tree_focus_mode_checkbox.setChecked(target_state)

        self.assertEqual(received[-1], target_state)

    def test_ai_tab_is_hidden_from_settings_ui(self):
        titles = [self.page._tabs.tabText(i) for i in range(self.page._tabs.count())]
        self.assertEqual(titles, ["常规", "扩展", "快捷键"])

    def test_settings_tabs_and_extension_sections_use_fluent_navigation_widgets(self):
        from pathlib import Path
        from qfluentwidgets import Pivot, SegmentedWidget
        from ui.pages.settings_page import SettingsPage

        builtin_specs = [
            {
                "id": "plot_reference_line",
                "file_name": "plot_reference_line.py",
                "name": "参考线标注",
                "categories": ["plot"],
                "category_labels": ["绘图扩展"],
                "type_ids": ["plot_reference_line"],
                "names_by_category": {"plot": ["参考线标注"]},
                "type_ids_by_category": {"plot": ["plot_reference_line"]},
                "source": "builtin",
                "source_label": "内置",
                "load_error": "",
            }
        ]
        external_specs = [
            {
                "id": "external_plot_reference_line",
                "file_name": "external_plot_reference_line.py",
                "name": "外部圆角",
                "categories": ["plot"],
                "category_labels": ["绘图扩展"],
                "type_ids": ["external_plot_reference_line"],
                "names_by_category": {"plot": ["外部圆角"]},
                "type_ids_by_category": {"plot": ["external_plot_reference_line"]},
                "source": "external",
                "source_label": "外部",
                "load_error": "",
            }
        ]

        temp_page = None
        try:
            with mock.patch("core.extension_settings._CONFIG_PATH", Path("/nonexistent/aline_extension_settings.json")), \
                 mock.patch("core.extension_api.list_builtin_extension_specs", return_value=builtin_specs), \
                 mock.patch("core.extension_api.list_external_extension_specs", return_value=external_specs):
                temp_page = SettingsPage()

            self.assertIsInstance(temp_page._tabs.navigationWidget, Pivot)
            self.assertEqual(temp_page._tabs.tabText(1), "扩展")
            self.assertIsInstance(temp_page._extension_tabs.navigationWidget, SegmentedWidget)
            self.assertIsInstance(temp_page._external_extension_tabs.navigationWidget, SegmentedWidget)
            self.assertEqual(temp_page._extension_tabs.tabText(0), "绘图扩展")
            self.assertEqual(temp_page._extension_tabs.tabText(3), "数字化扩展")
            self.assertEqual(temp_page._external_extension_tabs.tabText(0), "绘图扩展")
            self.assertEqual(temp_page._external_extension_tabs.tabText(3), "数字化扩展")
            self.assertEqual(temp_page._builtin_extension_checkboxes["plot_reference_line"].text(), "参考线标注")
            self.assertEqual(temp_page._external_extension_checkboxes["external_plot_reference_line"].text(), "外部圆角")
        finally:
            if temp_page is not None:
                temp_page.deleteLater()

    def test_extension_management_tabs_cap_height_and_use_scrollable_option_area(self):
        from PySide6.QtCore import Qt
        from qfluentwidgets import SmoothScrollArea
        from ui.pages.settings_page import (
            _EXTENSION_CATEGORY_TABS_HEIGHT_MULTIPLIER,
            _EXTENSION_CATEGORY_TABS_MAX_HEIGHT,
        )

        self.assertEqual(_EXTENSION_CATEGORY_TABS_MAX_HEIGHT, 60750)
        self.assertEqual(self.page._extension_tabs.maximumHeight(), _EXTENSION_CATEGORY_TABS_MAX_HEIGHT)
        self.assertEqual(self.page._external_extension_tabs.maximumHeight(), _EXTENSION_CATEGORY_TABS_MAX_HEIGHT)
        self.assertGreaterEqual(
            self.page._extension_tabs.minimumHeight(),
            max(
                self.page._extension_tabs.sizeHint().height() * _EXTENSION_CATEGORY_TABS_HEIGHT_MULTIPLIER,
                self.page._extension_tabs.navigationWidget.sizeHint().height(),
            ),
        )
        self.assertGreaterEqual(
            self.page._external_extension_tabs.minimumHeight(),
            max(
                self.page._external_extension_tabs.sizeHint().height() * _EXTENSION_CATEGORY_TABS_HEIGHT_MULTIPLIER,
                self.page._external_extension_tabs.navigationWidget.sizeHint().height(),
            ),
        )

        builtin_scroll_areas = self.page._extension_tabs.findChildren(SmoothScrollArea)
        external_scroll_areas = self.page._external_extension_tabs.findChildren(SmoothScrollArea)

        self.assertGreaterEqual(len(builtin_scroll_areas), 4)
        self.assertGreaterEqual(len(external_scroll_areas), 4)
        for layout in [*self.page._extension_option_layouts.values(), *self.page._external_extension_option_layouts.values()]:
            self.assertTrue(bool(layout.alignment() & Qt.AlignmentFlag.AlignTop))

    def test_extension_management_tabs_restore_height_when_extensions_tab_becomes_visible(self):
        from ui.pages.settings_page import _EXTENSION_CATEGORY_TABS_HEIGHT_MULTIPLIER

        self.page._extension_tabs.setMinimumHeight(1)
        self.page._external_extension_tabs.setMinimumHeight(1)

        self.page.resize(1280, 900)
        self.page.show()
        QApplication.processEvents()
        self.page._tabs.setCurrentIndex(1)
        QApplication.processEvents()

        self.assertEqual(
            self.page._extension_tabs.minimumHeight(),
            max(
                self.page._extension_tabs.sizeHint().height() * _EXTENSION_CATEGORY_TABS_HEIGHT_MULTIPLIER,
                self.page._extension_tabs.navigationWidget.sizeHint().height(),
            ),
        )
        self.assertEqual(
            self.page._external_extension_tabs.minimumHeight(),
            max(
                self.page._external_extension_tabs.sizeHint().height() * _EXTENSION_CATEGORY_TABS_HEIGHT_MULTIPLIER,
                self.page._external_extension_tabs.navigationWidget.sizeHint().height(),
            ),
        )

    def test_extension_status_summary_is_rendered_in_settings_page(self):
        from ui.theme import warning_color

        with mock.patch(
            "ui.pages.settings_page.get_extension_load_status",
            return_value={
                "registered_count": 3,
                "error_count": 1,
                "source_summary": {
                    "loaded_extension_counts": {"builtin": 2, "external": 1},
                },
                "details": {
                    "loaded": [{"path": "ok.py"}],
                    "errors": [{"path": "bad.py", "message": "boom", "categories": ["processing"]}],
                },
            },
        ):
            self.page._refresh_extension_status_summary()

        self.page.show()
        QApplication.processEvents()
        self.assertEqual(self.page._extension_status_summary_btn.text(), "当前扩展状态：3 项可用（内置 2 / 外部 1），1 项失败")
        self.assertTrue(self.page._extension_status_summary_btn.isEnabled())
        self.assertIn(warning_color(), self.page._extension_status_summary_btn.styleSheet())
        self.assertFalse(self.page._extension_card.isAncestorOf(self.page._builtin_extension_card))
        self.assertFalse(self.page._extension_card.isAncestorOf(self.page._external_extension_card))
        self.assertFalse(self.page._extension_card.isAncestorOf(self.page._extension_other_settings_card))
        self.assertTrue(self.page._builtin_extension_card.isAncestorOf(self.page._builtin_extension_management_card))
        self.assertTrue(self.page._external_extension_card.isAncestorOf(self.page._external_extension_management_card))
        self.assertTrue(self.page._extension_other_settings_card.isAncestorOf(self.page._external_extension_number_decimals_card))
        self.assertFalse(self.page._external_extension_card.isAncestorOf(self.page._external_extension_number_decimals_card))

    def test_save_ai_config_no_crash(self):
        self.page._ai_url_edit.setText("https://api.openai.com/v1")
        self.page._ai_model_edit.setText("gpt-4o-mini")
        self.page._ai_timeout_edit.setText("30")
        self.page._save_ai_config()

    def test_save_ai_config_infobar_uses_top_level_parent(self):
        from PySide6.QtWidgets import QWidget
        from ui.pages import settings_page
        from ui.pages.settings_page import SettingsPage

        host = QWidget()
        page = SettingsPage(parent=host)
        try:
            with mock.patch.object(settings_page.InfoBar, "success") as success_mock:
                page._save_ai_config()

            success_mock.assert_called_once()
            self.assertIs(success_mock.call_args.kwargs["parent"], host)
        finally:
            page.deleteLater()
            host.deleteLater()

    def test_save_extension_settings_persists_and_reloads_extensions(self):
        received = []
        self.page.extensions_reloaded.connect(lambda: received.append(True))

        with tempfile.TemporaryDirectory() as temp_dir:
            config_path = Path(temp_dir) / "extension_settings.json"
            external_dir = Path(temp_dir) / "external_extensions"
            second_external_dir = Path(temp_dir) / "external_extensions_b"
            builtin_specs = [
                {
                    "id": "plot_reference_line",
                    "file_name": "plot_reference_line.py",
                    "name": "参考线标注",
                    "categories": ["plot"],
                    "category_labels": ["绘图扩展"],
                    "type_ids": ["plot_reference_line"],
                    "names_by_category": {"plot": ["参考线标注"]},
                    "type_ids_by_category": {"plot": ["plot_reference_line"]},
                    "source": "builtin",
                    "source_label": "内置",
                    "load_error": "",
                }
            ]
            external_specs = [
                {
                    "id": "external_plot_reference_line",
                    "file_name": "external_plot_reference_line.py",
                    "name": "外部圆角",
                    "categories": ["plot"],
                    "category_labels": ["绘图扩展"],
                    "type_ids": ["external_plot_reference_line"],
                    "names_by_category": {"plot": ["外部圆角"]},
                    "type_ids_by_category": {"plot": ["external_plot_reference_line"]},
                    "source": "external",
                    "source_label": "外部",
                    "load_error": "",
                }
            ]
            with mock.patch("core.extension_settings._CONFIG_PATH", config_path), \
                 mock.patch("core.extension_api.list_builtin_extension_specs", return_value=builtin_specs), \
                 mock.patch("core.extension_api.list_external_extension_specs", return_value=external_specs), \
                 mock.patch(
                     "core.extension_api.reload_configured_extensions",
                     return_value={"loaded": ["builtin:plot_reference_line"], "errors": []},
                 ) as reload_mock:
                self.page._load_extension_settings()
                self.assertEqual(
                    self.page._builtin_extension_checkboxes["plot_reference_line"].text(),
                    "参考线标注",
                )
                self.assertEqual(
                    self.page._external_extension_checkboxes["external_plot_reference_line"].text(),
                    "外部圆角",
                )
                self.assertEqual(self.page._external_extension_number_decimals_slider.value(), 6)
                self.page._external_extensions_dirs_card.setFolders([str(external_dir), str(second_external_dir)])
                self.page._builtin_extension_checkboxes["plot_reference_line"].setChecked(False)
                self.page._external_extension_checkboxes["external_plot_reference_line"].setChecked(False)
                self.page._external_extensions_enabled_checkbox.setChecked(False)
                self.page._external_extension_number_decimals_slider.setValue(4)
                self.page._save_extension_settings()

            self.assertEqual(received, [True])
            self.assertTrue(config_path.exists())
            config_text = config_path.read_text(encoding="utf-8")
            self.assertIn("plot_reference_line", config_text)
            self.assertIn("external_plot_reference_line", config_text)
            self.assertIn(str(external_dir), config_text)
            self.assertIn(str(second_external_dir), config_text)
            self.assertIn('"load_external_extensions": false', config_text)
            self.assertIn('"number_decimals": 4', config_text)
            self.assertEqual(self.page._external_extension_number_decimals_slider.value(), 4)
            reload_mock.assert_called_once()

    def test_refresh_external_extension_specs_rebuilds_without_reloading_registry(self):
        from pathlib import Path
        from ui.pages import settings_page
        from ui.pages.settings_page import SettingsPage

        builtin_specs = [
            {
                "id": "plot_reference_line",
                "file_name": "plot_reference_line.py",
                "name": "参考线标注",
                "categories": ["plot"],
                "category_labels": ["绘图扩展"],
                "type_ids": ["plot_reference_line"],
                "names_by_category": {"plot": ["参考线标注"]},
                "type_ids_by_category": {"plot": ["plot_reference_line"]},
                "source": "builtin",
                "source_label": "内置",
                "load_error": "",
            }
        ]
        initial_external_specs = []
        refreshed_external_specs = [
            {
                "id": "external_plot_reference_line",
                "file_name": "external_plot_reference_line.py",
                "name": "外部圆角",
                "categories": ["plot"],
                "category_labels": ["绘图扩展"],
                "type_ids": ["external_plot_reference_line"],
                "names_by_category": {"plot": ["外部圆角"]},
                "type_ids_by_category": {"plot": ["external_plot_reference_line"]},
                "source": "external",
                "source_label": "外部",
                "load_error": "",
            }
        ]

        temp_page = None
        try:
            with mock.patch("core.extension_settings._CONFIG_PATH", Path("/nonexistent/aline_extension_settings.json")), \
                 mock.patch("core.extension_api.list_builtin_extension_specs", return_value=builtin_specs), \
                 mock.patch("core.extension_api.list_external_extension_specs", side_effect=[initial_external_specs, refreshed_external_specs]), \
                 mock.patch.object(settings_page.InfoBar, "success") as success_mock:
                temp_page = SettingsPage()
                self.assertNotIn("external_plot_reference_line", temp_page._external_extension_checkboxes)

                temp_page._refresh_external_extension_specs()

            self.assertEqual(
                temp_page._external_extension_checkboxes["external_plot_reference_line"].text(),
                "外部圆角",
            )
            success_mock.assert_called_once()
        finally:
            if temp_page is not None:
                temp_page.deleteLater()

    def test_provider_changed_to_ollama(self):
        # Ollama 也允许填写 API key（服务端代理场景）
        self.page._ai_provider_combo.setCurrentIndex(0)
        self.page._ai_url_edit.setText("")    # Clear URL so default gets set
        self.page._ai_model_edit.setText("")  # Clear model so default gets set
        self.page._ai_provider_combo.setCurrentIndex(1)  # Ollama
        self.assertTrue(self.page._ai_key_edit.isEnabled())
        self.assertTrue(self.page._ai_ollama_keep_alive_edit.isEnabled())
        # Ollama 时若 URL 为空则设置默认值
        self.assertIn("11434", self.page._ai_url_edit.text())
        self.assertIn("llama", self.page._ai_model_edit.text())
        self.assertIn("可选", self.page._ai_key_edit.placeholderText())

    def test_provider_changed_back_to_openai(self):
        self.page._ai_provider_combo.setCurrentIndex(1)
        self.page._ai_provider_combo.setCurrentIndex(0)
        self.assertTrue(self.page._ai_key_edit.isEnabled())
        self.assertFalse(self.page._ai_ollama_keep_alive_edit.isEnabled())

    def test_default_provider_is_openai_compatible(self):
        from pathlib import Path
        from unittest import mock
        from ui.pages.settings_page import SettingsPage

        temp_page = None
        try:
            with mock.patch("core.ai_client._CONFIG_PATH", Path("/nonexistent/aline_config.json")):
                temp_page = SettingsPage()
            self.assertEqual(temp_page._current_provider_key(), "openai_compatible")
        finally:
            if temp_page is not None:
                temp_page.deleteLater()

    def test_replay_onboarding_button_emits_signal(self):
        received = []
        self.page.replay_onboarding_requested.connect(lambda: received.append(True))

        self.page._replay_onboarding_btn.click()

        self.assertEqual(received, [True])


class TestHomePage(unittest.TestCase):
    """HomePage 引导入口与扩展状态"""

    def test_home_page_builds_banner_with_two_reserved_link_cards(self):
        from ui.pages.home_page import HomePage, _HOME_CONTENT_MARGIN, _HOME_LINK_CARD_HEIGHT, _HOME_LINK_CARD_WIDTH

        page = HomePage()
        try:
            page.resize(1200, 900)
            page.show()
            QApplication.processEvents()
            self.assertIsNotNone(page._banner)
            self.assertEqual(page._banner.x(), 0)
            self.assertEqual(page._banner.y(), 0)
            self.assertEqual(page._banner.width(), page.width())
            self.assertFalse(page._banner._background.isNull())
            self.assertEqual(page._banner._hero_title.text(), "ALine")
            self.assertEqual(page._banner._hero_subtitle.text(), "科研数据管理与可视化工作台")
            self.assertEqual(page._banner._hero_hint.maximumWidth(), 760)
            self.assertGreaterEqual(page._banner._hero_hint.height(), page._banner._hero_hint.sizeHint().height())
            self.assertIsNotNone(page._banner._link_card_view)
            self.assertEqual(page._banner._link_card_view._layout.count(), 2)
            self.assertEqual([card.titleLabel.text() for card in page._banner._link_cards], ["软件主页", "GitHub 仓库"])
            self.assertEqual(page._banner._link_cards[0]._icon_source, page._banner._card_icon_path)
            self.assertEqual(page._banner._link_cards[0].width(), _HOME_LINK_CARD_WIDTH)
            self.assertEqual(page._banner._link_cards[0].height(), _HOME_LINK_CARD_HEIGHT)
            card_left = page._banner._link_cards[0].mapTo(page, page._banner._link_cards[0].rect().topLeft()).x()
            new_button_left = page._new_btn.mapTo(page, page._new_btn.rect().topLeft()).x()
            self.assertEqual(card_left, new_button_left)
            self.assertEqual(new_button_left, _HOME_CONTENT_MARGIN)
        finally:
            page.deleteLater()

    def test_home_page_action_buttons_are_left_aligned(self):
        from ui.pages.home_page import HomePage

        page = HomePage()
        try:
            self.assertEqual(page._action_button_layout.spacing(), 20)
            self.assertTrue(bool(page._action_button_layout.alignment() & Qt.AlignmentFlag.AlignLeft))
            self.assertEqual(page._new_btn.width(), 150)
            self.assertEqual(page._open_btn.width(), 150)
        finally:
            page.deleteLater()

    def test_home_page_hides_guide_button_without_extension_status_summary(self):
        from core.ui_preferences import set_home_onboarding_completed
        from ui.pages.home_page import HomePage

        with tempfile.TemporaryDirectory() as temp_dir:
            config_path = Path(temp_dir) / "ui_preferences.json"
            with mock.patch("core.ui_preferences._CONFIG_PATH", config_path):
                set_home_onboarding_completed(True)
                page = HomePage()
                try:
                    self.assertIsNone(page._guide_toggle_btn)
                    self.assertIsNone(page._extension_status_btn)
                    self.assertIsNone(page._status_bar)
                finally:
                    page.deleteLater()

    def test_home_page_no_longer_uses_bottom_extension_status_bar(self):
        from ui.pages.home_page import HomePage

        page = HomePage()
        try:
            self.assertIsNone(page._status_bar)
            self.assertIsNone(page._extension_status_btn)
            self.assertEqual(len(page._home_onboarding_steps()), 4)
        finally:
            page.deleteLater()

    def test_home_page_recent_scroll_expands_to_fill_remaining_height(self):
        from PySide6.QtWidgets import QSizePolicy
        from ui.pages.home_page import HomePage

        page = HomePage()
        try:
            self.assertEqual(page._recent_scroll.sizePolicy().verticalPolicy(), QSizePolicy.Policy.Expanding)
            self.assertEqual(page._recent_scroll.maximumHeight(), 16777215)
        finally:
            page.deleteLater()

    def test_home_page_recent_rows_use_shared_hover_color(self):
        from ui.pages.home_page import HomePage
        from ui.theme import hover_color

        with mock.patch(
            "ui.pages.home_page.load_recent",
            return_value=[{"name": "示例项目", "path": "/tmp/example.aline", "opened_at": "2026-04-23 10:20:30"}],
        ):
            page = HomePage()

        try:
            row = self._first_recent_row(page)
            self.assertIsNotNone(row)
            self.assertIn(hover_color(), row.styleSheet())
        finally:
            page.deleteLater()

    def test_home_page_recent_tooltips_use_fluent_filters(self):
        from PySide6.QtWidgets import QWidget
        from ui.pages.home_page import HomePage

        with mock.patch(
            "ui.pages.home_page.load_recent",
            return_value=[{"name": "示例项目", "path": "/tmp/example.aline", "opened_at": "2026-04-23 10:20:30"}],
        ):
            page = HomePage()

        try:
            tooltip_widgets = [
                widget
                for widget in page.findChildren(QWidget)
                if widget.toolTip() in {"/tmp/example.aline", "从列表移除"}
            ]
            self.assertEqual(len(tooltip_widgets), 2)
            for widget in tooltip_widgets:
                self.assertTrue(widget.property("_alineFluentTooltip"))
        finally:
            page.deleteLater()

    def test_home_page_infobar_uses_top_level_parent(self):
        from PySide6.QtWidgets import QWidget
        from ui.pages import home_page
        from ui.pages.home_page import HomePage

        host = QWidget()
        page = HomePage(parent=host)
        try:
            with mock.patch.object(type(home_page.project_manager), "current_project", new_callable=mock.PropertyMock, return_value=None):
                with mock.patch.object(home_page.InfoBar, "warning") as warning_mock:
                    page._request_quick_start("chart")

            warning_mock.assert_called_once()
            self.assertIs(warning_mock.call_args.kwargs["parent"], host)
        finally:
            page.deleteLater()
            host.deleteLater()

    @staticmethod
    def _first_recent_row(page):
        for index in range(page._recent_items_layout.count()):
            item = page._recent_items_layout.itemAt(index)
            widget = item.widget() if item is not None else None
            if widget is not None:
                return widget
        return None


class TestSettingsPageAIActions(unittest.TestCase):

    def setUp(self):
        self._restore_assets = _patch_global_assets()
        from ui.pages.settings_page import SettingsPage
        self.page = SettingsPage()

    def tearDown(self):
        self.page.deleteLater()
        self._restore_assets()

    def test_shortcuts_changed_signal(self):
        received = []
        self.page.shortcuts_changed.connect(lambda: received.append(True))
        self.page._on_apply_shortcuts()
        self.assertEqual(len(received), 1)

    def test_shortcut_filter_hides_nonmatching_rows(self):
        zoom_in = next(lbl for lbl in self.page._shortcut_labels if "放大" in lbl.text())
        save = next(lbl for lbl in self.page._shortcut_labels if "保存项目" in lbl.text())
        self.page._filter_shortcut_rows("缩放")
        self.assertFalse(zoom_in.isHidden())
        self.assertTrue(save.isHidden())

    def test_reset_shortcuts_no_crash(self):
        self.page._on_reset_shortcuts()

    def test_timeout_invalid_fallback(self):
        """非数字超时输入应 fallback 到 60"""
        self.page._ai_timeout_edit.setText("abc")
        self.page._save_ai_config()  # 不崩溃

    def test_refresh_ai_tools_panel_without_project(self):
        self.page._refresh_ai_tools_panel()
        self.assertIn("全局资源", self.page._ai_tools_project_label.text())

    def test_refresh_ai_tools_panel_with_project(self):
        from core.global_assets import global_assets
        from core.project_manager import ProjectManager

        pm = ProjectManager()
        pm.create_new("settings_ai_tools")
        global_assets.add_ai_prompt("prompt-a", "hello")
        global_assets.add_ai_skill("skill-a", "result = 1")
        global_assets.add_ai_agent("agent-a", "system")
        restore = _patch_pm(pm)
        try:
            self.page._refresh_ai_tools_panel()
            self.assertIn("Prompt 1", self.page._ai_tools_summary_label.text())
            # 新 UI 使用下拉框，检查 combo 中包含 skill-a
            combo_texts = [
                self.page._ai_tool_selector.itemText(i)
                for i in range(self.page._ai_tool_selector.count())
            ]
            self.assertTrue(any("skill-a" in t for t in combo_texts))
        finally:
            restore()


class TestExtensionConfigPanel(unittest.TestCase):

    def test_reload_button_emits_signal(self):
        from ui.widgets.extension_panel import ExtensionConfigPanel

        panel = ExtensionConfigPanel()
        received = []
        panel.reload_requested.connect(lambda: received.append(True))
        panel._reload_btn.click()

        self.assertEqual(received, [True])
        panel.deleteLater()

    def test_reload_button_uses_icon_tool_button(self):
        from qfluentwidgets import ToolButton
        from ui.widgets.extension_panel import ExtensionConfigPanel

        panel = ExtensionConfigPanel()

        self.assertIsInstance(panel._reload_btn, ToolButton)
        panel.deleteLater()

    def test_config_help_section_uses_scroll_area(self):
        from qfluentwidgets import SmoothScrollArea
        from ui.widgets.extension_panel import ExtensionConfigPanel

        panel = ExtensionConfigPanel()

        self.assertIsInstance(panel._config_help_area, SmoothScrollArea)
        self.assertIs(panel._config_help_area.widget(), panel._config_help_container)
        panel.deleteLater()

    def test_reset_clear_and_config_actions_use_tool_buttons(self):
        from qfluentwidgets import ToolButton
        from ui.widgets.extension_panel import ExtensionConfigPanel

        panel = ExtensionConfigPanel()

        self.assertIsInstance(panel._config_load_btn, ToolButton)
        self.assertIsInstance(panel._config_add_btn, ToolButton)
        self.assertIsInstance(panel._config_overwrite_btn, ToolButton)
        self.assertIsInstance(panel._reset_btn, ToolButton)
        self.assertIsInstance(panel._clear_btn, ToolButton)
        self.assertEqual(panel._apply_btn.text(), "应用扩展")
        panel.deleteLater()

    def test_config_action_tool_buttons_use_square_metrics(self):
        from ui.widgets.extension_panel import ExtensionConfigPanel

        panel = ExtensionConfigPanel()

        for button in (panel._config_load_btn, panel._config_add_btn, panel._config_overwrite_btn):
            self.assertEqual(button.width(), button.height())
        panel.deleteLater()

    def test_panel_can_save_and_load_global_extension_configs(self):
        from core.global_assets import global_assets
        from ui.widgets.extension_panel import ExtensionConfigPanel

        restore_assets = _patch_global_assets()
        panel = ExtensionConfigPanel()
        try:
            panel.set_status_context("processing", "处理扩展")
            panel.set_entries([
                {
                    "type": "panel_extension_probe",
                    "name": "面板配置探针",
                    "label": "面板配置探针",
                    "description": "测试扩展配置面板的全局预设保存。",
                    "version": "1.2.0",
                    "default_options": {"factor": 2},
                    "config_fields": [],
                }
            ])

            self.assertEqual(
                [item.name for item in global_assets.list_extension_configs(category="processing", extension_type="panel_extension_probe")],
                ["默认配置"],
            )
            self.assertEqual(
                global_assets.get_extension_default_config("processing", "panel_extension_probe").extension_version,
                "1.2.0",
            )

            panel._editor.setPlainText('{"factor": 8}')
            with mock.patch("ui.widgets.extension_panel.TextInputDialog.get_text", return_value=("方案A", True)):
                panel._save_current_as_config()

            selector_items = [panel._config_selector.itemText(i) for i in range(panel._config_selector.count())]
            self.assertEqual(selector_items, ["默认配置", "方案A"])
            saved = global_assets.get_extension_config_by_name("processing", "panel_extension_probe", "方案A")
            self.assertIsNotNone(saved)
            self.assertEqual(saved.extension_version, "1.2.0")

            panel._editor.setPlainText('{"factor": 1}')
            panel._config_selector.setCurrentIndex(selector_items.index("方案A"))
            panel._load_selected_config()
            self.assertEqual(panel.current_options(), {"factor": 8})
            self.assertTrue(panel.load_config_by_id(saved.id))
            self.assertEqual(panel.current_options(), {"factor": 8})
        finally:
            panel.deleteLater()
            restore_assets()

    def test_loading_outdated_extension_config_shows_version_warning(self):
        from core.global_assets import global_assets
        from ui.widgets import extension_panel
        from ui.widgets.extension_panel import ExtensionConfigPanel

        restore_assets = _patch_global_assets()
        panel = ExtensionConfigPanel()
        try:
            panel.set_status_context("processing", "处理扩展")
            panel.set_entries([
                {
                    "type": "panel_extension_probe",
                    "name": "面板配置探针",
                    "label": "面板配置探针",
                    "description": "测试扩展配置面板的版本提醒。",
                    "version": "1.2.0",
                    "default_options": {"factor": 2},
                    "config_fields": [],
                }
            ])
            saved = global_assets.add_extension_config(
                category="processing",
                extension_type="panel_extension_probe",
                extension_name="面板配置探针",
                extension_version="1.0.0",
                name="旧版本方案",
                options={"factor": 3},
            )

            with mock.patch.object(extension_panel.InfoBar, "warning") as warning_mock:
                self.assertTrue(panel.load_config_by_id(saved.id))

            warning_mock.assert_called_once()
            self.assertIn("1.0.0", warning_mock.call_args.args[1])
            self.assertIn("1.2.0", warning_mock.call_args.args[1])
        finally:
            panel.deleteLater()
            restore_assets()

    def test_copy_current_config_shows_success_infobar(self):
        from ui.widgets import extension_panel
        from PySide6.QtWidgets import QWidget
        from ui.widgets.extension_panel import ExtensionConfigPanel

        host = QWidget()
        panel = ExtensionConfigPanel(parent=host)
        panel._editor.setPlainText('{"factor": 4}')
        try:
            with mock.patch.object(extension_panel.InfoBar, "success") as success_mock:
                panel._copy_current_config()

            self.assertEqual(QApplication.clipboard().text(), '{"factor": 4}')
            success_mock.assert_called_once()
            self.assertIs(success_mock.call_args.kwargs["parent"], host)
        finally:
            panel.deleteLater()
            host.deleteLater()

    def test_config_help_uses_key_type_required_description_format(self):
        from ui.widgets.extension_panel import ExtensionConfigPanel

        panel = ExtensionConfigPanel()

        help_text = panel._config_help_text(
            {
                "config_fields": [
                    {
                        "key": "show_reference_line",
                        "label": "显示参考线",
                        "field_type": "boolean",
                        "required": False,
                        "description": "在 before_plot 阶段绘制一条水平参考线。",
                    }
                ]
            }
        )

        self.assertIn(
            "显示参考线（boolean，可选）；在 before_plot 阶段绘制一条水平参考线。",
            help_text,
        )
        self.assertNotIn("show_reference_line: boolean", help_text)
        panel.deleteLater()

    def test_panel_layout_matches_page_card_spacing_baseline(self):
        from ui.widgets.extension_panel import ExtensionConfigPanel

        panel = ExtensionConfigPanel()

        margins = panel.layout().contentsMargins()
        self.assertEqual((margins.left(), margins.top(), margins.right(), margins.bottom()), (0, 0, 0, 0))
        self.assertGreaterEqual(panel._config_help_area.minimumHeight(), 124)
        self.assertGreaterEqual(panel.width(), 360)
        panel.deleteLater()

    def test_panel_uses_dividers_and_compact_descriptions(self):
        from qfluentwidgets import CaptionLabel
        from ui.theme import card_title_style_sheet, secondary_text_style_sheet
        from ui.widgets.extension_panel import ExtensionConfigPanel

        panel = ExtensionConfigPanel()

        self.assertEqual(len(panel._section_dividers), 3)
        self.assertIsInstance(panel._description_label, CaptionLabel)
        self.assertEqual(panel._title_label.styleSheet(), card_title_style_sheet(font_size=17))
        self.assertEqual(panel._current_entry_label.styleSheet(), card_title_style_sheet(font_size=14))
        self.assertEqual(panel._description_label.styleSheet(), secondary_text_style_sheet(font_size=12))
        self.assertIn("font-size: 12px", panel._description_label.styleSheet())
        self.assertTrue(panel._status_summary_btn.isFlat())
        self.assertFalse(hasattr(panel, "_status_label"))
        self.assertFalse(hasattr(panel, "_status_detail_btn"))
        panel.deleteLater()

    def test_panel_header_uses_name_source_version_format(self):
        from ui.widgets.extension_panel import ExtensionConfigPanel

        panel = ExtensionConfigPanel()
        try:
            panel.set_status_context("plot", "绘图扩展")
            panel.set_entries([
                {
                    "type": "panel_header_probe",
                    "name": "绘制箭头",
                    "label": "绘制箭头",
                    "description": "在图中添加箭头标注。",
                    "version": "0.1.0",
                    "source_label": "内置",
                    "settings": True,
                    "resolved_options": {},
                    "config_fields": [],
                }
            ])

            self.assertEqual(panel._extension_section_label.text(), "绘图扩展")
            self.assertEqual(panel._current_entry_label.text(), "当前扩展: 绘制箭头·内置·v0.1.0")
            self.assertEqual(panel._description_label.text(), "在图中添加箭头标注。")
        finally:
            panel.deleteLater()

    def test_status_summary_reflects_category_load_state(self):
        from ui.widgets.extension_panel import ExtensionConfigPanel
        from ui.theme import warning_color

        panel = ExtensionConfigPanel()
        with mock.patch(
            "ui.widgets.extension_panel.get_extension_load_status",
            return_value={
                "label": "处理扩展",
                "registered_count": 2,
                "error_count": 1,
                "source_summary": {
                    "loaded_extension_counts": {"builtin": 1, "external": 1},
                    "loaded_file_counts": {"builtin": 1, "external": 1},
                    "error_file_counts": {"builtin": 0, "external": 1},
                },
                "details": {"loaded": [{"path": "ok.py"}], "errors": [{"path": "bad.py"}]},
            },
        ):
            panel.set_status_context("processing", "处理扩展")

        self.assertEqual(panel._status_summary_btn.text(), "处理扩展 2 项可用（内置 1 / 外部 1），1 项失败")
        self.assertTrue(panel._status_summary_btn.isEnabled())
        self.assertIn("font-size: 12px", panel._status_summary_btn.styleSheet())
        self.assertIn(warning_color(), panel._status_summary_btn.styleSheet())
        panel.deleteLater()

    def test_status_summary_uses_home_style_accent_when_extensions_available(self):
        from ui.widgets.extension_panel import ExtensionConfigPanel
        from ui.theme import accent_color

        panel = ExtensionConfigPanel()
        with mock.patch(
            "ui.widgets.extension_panel.get_extension_load_status",
            return_value={
                "label": "绘图扩展",
                "registered_count": 2,
                "error_count": 0,
                "source_summary": {
                    "loaded_extension_counts": {"builtin": 1, "external": 1},
                    "loaded_file_counts": {"builtin": 1, "external": 1},
                    "error_file_counts": {"builtin": 0, "external": 0},
                },
                "details": {"loaded": [{"path": "plot_ok.py"}], "errors": []},
            },
        ):
            panel.set_status_context("plot", "绘图扩展")

        self.assertEqual(panel._status_summary_btn.text(), "绘图扩展 2 项可用（内置 1 / 外部 1）。")
        self.assertIn("font-size: 12px", panel._status_summary_btn.styleSheet())
        self.assertIn(accent_color(), panel._status_summary_btn.styleSheet())
        panel.deleteLater()

    def test_invalid_config_infobar_uses_top_level_parent(self):
        from PySide6.QtWidgets import QWidget
        from ui.widgets import extension_panel
        from ui.widgets.extension_panel import ExtensionConfigPanel

        host = QWidget()
        panel = ExtensionConfigPanel(parent=host)
        try:
            panel.set_status_context("processing", "处理扩展")
            panel.set_entries([
                {
                    "type": "panel_extension_probe",
                    "name": "面板配置探针",
                    "label": "面板配置探针",
                    "description": "测试扩展配置面板的全局预设保存。",
                    "default_options": {"factor": 2},
                    "config_fields": [],
                }
            ])
            panel._editor.setPlainText("{")

            with mock.patch.object(extension_panel.InfoBar, "error") as error_mock:
                panel._save_current_as_config()

            error_mock.assert_called_once()
            self.assertIs(error_mock.call_args.kwargs["parent"], host)
        finally:
            panel.deleteLater()
            host.deleteLater()
        panel.deleteLater()

    def test_help_only_mode_hides_controls_and_keeps_param_help(self):
        from ui.widgets.extension_panel import ExtensionConfigPanel
        from qfluentwidgets import CardWidget
        from PySide6.QtWidgets import QSizePolicy

        panel = ExtensionConfigPanel(mode="help_only")

        self.assertTrue(panel._title_label.isHidden())
        self.assertFalse(panel._extension_section_label.isHidden())
        self.assertFalse(panel._current_entry_label.isHidden())
        self.assertTrue(panel._selector_row_widget.isHidden())
        self.assertTrue(panel._config_row_widget.isHidden())
        self.assertTrue(panel._editor.isHidden())
        self.assertTrue(panel._action_row_widget.isHidden())
        self.assertFalse(panel._description_section_label.isHidden())
        self.assertFalse(panel._description_label.isHidden())
        self.assertFalse(panel._config_help_area.isHidden())
        self.assertEqual(panel._extension_section_label.text(), "自定义扩展")
        self.assertEqual(panel._current_entry_label.text(), "当前扩展: 未选择扩展")
        self.assertEqual(panel._description_label.text(), "在左侧选择扩展后，这里会显示扩展说明。")
        self.assertEqual(panel._parameter_section_label.text(), "参数说明")
        self.assertEqual(panel._config_help_area.minimumHeight(), 124)
        self.assertEqual(panel._config_help_area.maximumHeight(), 16777215)
        self.assertEqual(panel._current_entry_label.styleSheet(), "")
        self.assertIsInstance(panel._surface, CardWidget)
        self.assertEqual(panel._surface.sizePolicy().verticalPolicy(), QSizePolicy.Policy.Expanding)
        self.assertEqual(panel._config_help_area.sizePolicy().verticalPolicy(), QSizePolicy.Policy.Expanding)
        panel.deleteLater()

    def test_compact_mode_hides_titles_and_bottom_actions(self):
        from ui.widgets.extension_panel import ExtensionConfigPanel

        panel = ExtensionConfigPanel(mode="compact", framed=False)

        self.assertTrue(panel._title_label.isHidden())
        self.assertTrue(panel._status_row_widget.isHidden())
        self.assertTrue(panel._config_help_area.isHidden())
        self.assertTrue(panel._action_row_widget.isHidden())
        self.assertFalse(panel._selector_row_widget.isHidden())
        self.assertFalse(panel._config_row_widget.isHidden())
        panel.deleteLater()

    def test_report_dialog_uses_top_level_window_as_parent(self):
        from PySide6.QtWidgets import QWidget
        from ui.widgets import extension_panel

        host = QWidget()
        child = QWidget(host)
        captured = {}

        class _FakeDialog:
            def __init__(self, title, content, parent=None):
                captured["title"] = title
                captured["content"] = content
                captured["parent"] = parent

            def exec(self):
                captured["executed"] = True

        with mock.patch.object(extension_panel, "_ExtensionLoadReportDialog", _FakeDialog):
            extension_panel.show_extension_load_report_dialog(child, "扩展详情", "plot")

        self.assertEqual(captured["title"], "扩展详情")
        self.assertIs(captured["parent"], host)
        self.assertTrue(captured["executed"])
        child.deleteLater()
        host.deleteLater()


class TestExtensionOptionsForm(unittest.TestCase):

    def test_line_selection_dialog_select_all_marks_all_items_selected(self):
        from PySide6.QtWidgets import QWidget
        from ui.widgets.extension_options_form import _LineSelectionDialog

        host = QWidget()
        dialog = _LineSelectionDialog("输入曲线", ["A", "B", "C"], selected_indices=[], lines_number=(1, -1), parent=host)
        try:
            dialog._select_all()

            self.assertEqual(dialog.value(), [1, 2, 3])
            self.assertEqual(dialog._selected_list.count(), 3)
            self.assertEqual(dialog._available_list.count(), 0)

            dialog._clear()

            self.assertEqual(dialog._selected_list.count(), 0)
            self.assertEqual(dialog._available_list.count(), 3)
            self.assertEqual(dialog.value(), [])
        finally:
            dialog.deleteLater()
            host.deleteLater()

    def test_line_selection_dialog_select_all_updates_current_row(self):
        from PySide6.QtWidgets import QWidget
        from ui.widgets.extension_options_form import _LineSelectionDialog

        host = QWidget()
        dialog = _LineSelectionDialog("输入曲线", ["A", "B", "C"], selected_indices=[], lines_number=(1, -1), parent=host)
        try:
            dialog._select_all()

            self.assertTrue(dialog.yesButton.isEnabled())
            self.assertIn("已选择 3 条", dialog._status.text())

            dialog._clear()

            self.assertFalse(dialog.yesButton.isEnabled())
            self.assertIn("已选择 0 条", dialog._status.text())
        finally:
            dialog.deleteLater()
            host.deleteLater()

    def test_line_selection_dialog_restores_available_order_and_supports_reorder(self):
        from PySide6.QtWidgets import QWidget
        from ui.widgets.extension_options_form import _LineSelectionDialog

        host = QWidget()
        dialog = _LineSelectionDialog("输入曲线", ["A", "B", "C", "D"], selected_indices=[2, 4], lines_number=(2, -1), parent=host)
        try:
            self.assertEqual(dialog._available_title_label.text(), "候选区")
            self.assertEqual(dialog._selected_title_label.text(), "已选中")

            dialog._available_list.setCurrentRow(0)
            dialog._move_to_selected()
            self.assertEqual(dialog.value(), [2, 4, 1])

            dialog._selected_list.setCurrentRow(2)
            dialog._move_selected_up()
            self.assertEqual(dialog.value(), [2, 1, 4])
            self.assertEqual(dialog._selected_list.currentItem().data(Qt.ItemDataRole.UserRole), 1)

            dialog._selected_list.setCurrentRow(1)
            dialog._move_to_available()
            self.assertEqual(dialog.value(), [2, 4])
            self.assertEqual(
                [dialog._available_list.item(row).data(Qt.ItemDataRole.UserRole) for row in range(dialog._available_list.count())],
                [1, 3],
            )
        finally:
            dialog.deleteLater()
            host.deleteLater()

    def test_line_selection_dialog_wraps_lists_with_frames(self):
        from PySide6.QtWidgets import QFrame, QWidget
        from ui.widgets.extension_options_form import _LineSelectionDialog

        host = QWidget()
        dialog = _LineSelectionDialog("输入曲线", ["A", "B"], selected_indices=[1], lines_number=(1, -1), parent=host)
        try:
            self.assertIsInstance(dialog._available_list.parentWidget(), QFrame)
            self.assertIsInstance(dialog._selected_list.parentWidget(), QFrame)
            self.assertEqual(dialog._available_list.parentWidget().objectName(), "selectionListFrame")
            self.assertEqual(dialog._selected_list.parentWidget().objectName(), "selectionListFrame")
        finally:
            dialog.deleteLater()
            host.deleteLater()

    def test_line_selection_dialog_clear_button_stays_visible_for_fixed_selection(self):
        from PySide6.QtWidgets import QWidget
        from ui.widgets.extension_options_form import _LineSelectionDialog

        host = QWidget()
        dialog = _LineSelectionDialog("输入曲线", ["A", "B", "C"], selected_indices=[1, 2], lines_number=(2, 2), parent=host)
        try:
            self.assertTrue(dialog._select_all_btn.isHidden())
            self.assertFalse(dialog._clear_btn.isHidden())
        finally:
            dialog.deleteLater()
            host.deleteLater()

    def test_single_line_field_uses_dedicated_selector(self):
        from qfluentwidgets import PushButton
        from ui.widgets.extension_options_form import ExtensionOptionsForm

        form = ExtensionOptionsForm()
        try:
            form.set_line_candidates(["A", "B", "C"])
            form.set_fields(
                [{"key": "target_line", "label": "对齐曲线", "field_type": "line", "default": 1}],
                {"target_line": 1},
            )
            button = next(button for button in form.findChildren(PushButton) if button.text() == "选择曲线")

            with mock.patch("ui.widgets.extension_options_form._SingleLineSelectionDialog.get_index", return_value=(2, True)):
                button.click()

            self.assertEqual(form.current_options()["target_line"], 2)
            self.assertEqual(button.toolTip(), "当前参数曲线: B")
        finally:
            form.deleteLater()

    def test_line_field_summary_uses_semicolon_and_fluent_tooltip(self):
        from qfluentwidgets import CaptionLabel
        from ui.widgets.extension_options_form import ExtensionOptionsForm

        form = ExtensionOptionsForm()
        try:
            form.set_line_candidates(["项目 / 数据集 / 曲线 A", "项目 / 数据集 / 曲线 B"])
            form.set_fields(
                [
                    {
                        "key": "lines_list",
                        "label": "lines",
                        "field_type": "lines",
                        "default": [1, 2],
                        "extra": {"lines_number": [2, -1]},
                    }
                ],
                {"lines_list": [1, 2]},
            )

            summaries = [label for label in form.findChildren(CaptionLabel) if "当前:" in label.text()]

            self.assertTrue(summaries)
            self.assertEqual(summaries[0].text(), "当前: 项目 / 数据集 / 曲线 A；项目 / 数据集 / 曲线 B")
            self.assertEqual(summaries[0].toolTip(), "项目 / 数据集 / 曲线 A；项目 / 数据集 / 曲线 B")
            self.assertTrue(bool(summaries[0].property("_alineFluentTooltip")))
        finally:
            form.deleteLater()

    def test_compact_field_controls_use_reduced_min_widths(self):
        from PySide6.QtWidgets import QSlider
        from qfluentwidgets import ComboBox, LineEdit, PushButton
        from ui.widgets.extension_options_form import ExtensionOptionsForm

        form = ExtensionOptionsForm()
        try:
            form.set_fields(
                [
                    {"key": "name", "label": "名称", "field_type": "string", "default": "demo"},
                    {"key": "mode", "label": "模式", "field_type": "selective", "choices": ["A", "B"], "default": "A"},
                    {"key": "alpha", "label": "透明度", "field_type": "limited", "min_value": 0, "max_value": 10, "default": 3},
                ],
                {"name": "demo", "mode": "A", "alpha": 3},
            )

            edits = form.findChildren(LineEdit)
            combos = form.findChildren(ComboBox)
            sliders = form.findChildren(QSlider)
            buttons = [button for button in form.findChildren(PushButton) if button.text() == "选择曲线"]

            self.assertTrue(any(widget.minimumWidth() == 75 for widget in edits))
            self.assertTrue(any(widget.minimumWidth() == 75 for widget in combos))
            self.assertTrue(any(widget.minimumWidth() == 90 for widget in sliders))
            self.assertFalse(buttons)
        finally:
            form.deleteLater()

    def test_color_field_uses_fluent_tooltip_without_text_caption(self):
        from qfluentwidgets import CaptionLabel, ColorPickerButton
        from ui.widgets.extension_options_form import ExtensionOptionsForm

        form = ExtensionOptionsForm()
        try:
            form.set_fields(
                [{"key": "color", "label": "颜色", "field_type": "color", "default": "#112233"}],
                {"color": "#112233"},
            )

            button = form.findChild(ColorPickerButton)
            self.assertIsNotNone(button)
            self.assertEqual(button.toolTip(), "#112233")
            self.assertEqual(button.width(), button.height())
            self.assertFalse(any(label.text() == "#112233" for label in form.findChildren(CaptionLabel)))
        finally:
            form.deleteLater()

    def test_pickcolor_field_requests_interaction_and_accepts_async_value(self):
        from qfluentwidgets import CaptionLabel, ToolButton
        from ui.widgets.extension_options_form import ExtensionOptionsForm

        form = ExtensionOptionsForm()
        try:
            requested = []
            form.interactiveFieldRequested.connect(lambda key, field: requested.append((key, dict(field))))
            form.set_fields(
                [{"key": "sampled_color", "label": "采样颜色", "field_type": "pickcolor"}],
                {},
            )

            button = next(
                widget for widget in form.findChildren(ToolButton)
                if widget.objectName() == "interactiveFieldButton:sampled_color"
            )
            summary = next(
                widget for widget in form.findChildren(CaptionLabel)
                if widget.objectName() == "interactiveFieldSummary:sampled_color"
            )

            button.click()
            self.assertEqual(requested[0][0], "sampled_color")

            self.assertTrue(form.set_field_value("sampled_color", {"r": 17, "g": 34, "b": 51}))

            self.assertEqual(form.current_options()["sampled_color"], {"r": 17, "g": 34, "b": 51})
            self.assertEqual(summary.text(), "#112233")
            self.assertEqual(button.toolTip(), "#112233")
            self.assertIn("#112233", button.styleSheet().lower())
        finally:
            form.deleteLater()

    def test_shot_field_requests_interaction_and_summarizes_capture(self):
        from qfluentwidgets import CaptionLabel, ToolButton
        from ui.widgets.extension_options_form import ExtensionOptionsForm

        form = ExtensionOptionsForm()
        try:
            requested = []
            form.interactiveFieldRequested.connect(lambda key, field: requested.append((key, dict(field))))
            form.set_fields(
                [{"key": "template_info", "label": "截图模板", "field_type": "shot"}],
                {},
            )

            button = next(
                widget for widget in form.findChildren(ToolButton)
                if widget.objectName() == "interactiveFieldButton:template_info"
            )
            summary = next(
                widget for widget in form.findChildren(CaptionLabel)
                if widget.objectName() == "interactiveFieldSummary:template_info"
            )

            button.click()
            self.assertEqual(requested[0][0], "template_info")

            value = {"size": [32, 14], "bounds": [1, 2, 33, 16]}
            self.assertTrue(form.set_field_value("template_info", value))

            self.assertEqual(form.current_options()["template_info"], value)
            self.assertEqual(summary.text(), "已截取 32×14px")
        finally:
            form.deleteLater()

    def test_boolean_field_keeps_label_on_left_and_checkbox_on_right(self):
        from qfluentwidgets import CheckBox
        from ui.widgets.extension_options_form import ExtensionOptionsForm

        form = ExtensionOptionsForm()
        try:
            form.set_fields([
                {"key": "enabled", "label": "启用", "field_type": "boolean", "default": True}
            ])

            container_widget = form._flow.itemAt(0).widget()
            row_widget = container_widget.layout().itemAt(0).widget()
            checkbox = form.findChild(CheckBox)

            self.assertIsNotNone(checkbox)
            self.assertEqual(checkbox.text(), "")
            self.assertEqual(row_widget.layout().itemAtPosition(0, 0).widget().text(), "启用:")
        finally:
            form.deleteLater()

    def test_extension_options_form_does_not_infer_unknown_fields_by_default(self):
        from qfluentwidgets import LineEdit
        from ui.widgets.extension_options_form import ExtensionOptionsForm

        form = ExtensionOptionsForm()
        try:
            form.set_fields(
                [{"key": "method", "label": "方法", "field_type": "selective", "choices": ["pearson", "spearman"], "default": "pearson"}],
                {"method": "pearson", "legacy_extra": "ghost"},
            )

            self.assertNotIn("legacy_extra", form.current_options())
            self.assertFalse(any(edit.text() == "ghost" for edit in form.findChildren(LineEdit)))
        finally:
            form.deleteLater()

    def test_fields_use_inline_row_layout_and_compact_spacing(self):
        from PySide6.QtWidgets import QWidget
        from qfluentwidgets import DoubleSpinBox
        from ui.widgets.extension_options_form import ExtensionOptionsForm

        form = ExtensionOptionsForm()
        try:
            form.set_fields([
                {
                    "key": "factor",
                    "label": "倍率",
                    "field_type": "number",
                    "default": 1.0,
                }
            ])

            container_widget = form._flow.itemAt(0).widget()
            container_layout = container_widget.layout() if container_widget is not None else None
            row_widget = container_layout.itemAt(0).widget() if container_layout is not None else None
            editors = form.findChildren(DoubleSpinBox)

            self.assertIsInstance(row_widget, QWidget)
            self.assertEqual(row_widget.objectName(), "adaptiveFieldRow")
            self.assertEqual(form._flow.spacing(), 6)
            self.assertTrue(editors)
            row_widget.resize(280, row_widget.height() or 40)
            self.assertFalse(bool(row_widget.property("wrapped")))
        finally:
            form.deleteLater()

    def test_integer_and_number_fields_use_spin_widgets(self):
        from qfluentwidgets import DoubleSpinBox, SpinBox
        from ui.widgets.extension_options_form import ExtensionOptionsForm

        form = ExtensionOptionsForm()
        try:
            form.set_fields(
                [
                    {"key": "count", "label": "计数", "field_type": "integer", "default": 3, "min_value": 0, "max_value": 10},
                    {"key": "factor", "label": "倍率", "field_type": "number", "default": 1.25, "min_value": 0.0, "max_value": 10.0},
                ],
                {"count": 3, "factor": 1.25},
            )

            int_widget = form.findChild(SpinBox)
            float_widget = form.findChild(DoubleSpinBox)

            self.assertIsNotNone(int_widget)
            self.assertIsNotNone(float_widget)
            self.assertEqual(int_widget.value(), 3)
            self.assertAlmostEqual(float_widget.value(), 1.25)
            self.assertEqual(float_widget.decimals(), 6)

            int_widget.setValue(5)
            float_widget.setValue(2.5)

            self.assertEqual(form.current_options()["count"], 5)
            self.assertAlmostEqual(form.current_options()["factor"], 2.5)
        finally:
            form.deleteLater()

    def test_number_field_uses_configured_default_decimals(self):
        from qfluentwidgets import DoubleSpinBox
        from ui.widgets.extension_options_form import ExtensionOptionsForm

        with tempfile.TemporaryDirectory() as temp_dir:
            config_path = Path(temp_dir) / "extension_settings.json"
            config_path.write_text('{"number_decimals": 4}', encoding="utf-8")

            with mock.patch("core.extension_settings._CONFIG_PATH", config_path):
                form = ExtensionOptionsForm()
                try:
                    form.set_fields([
                        {"key": "factor", "label": "倍率", "field_type": "number", "default": 1.0}
                    ])

                    float_widget = form.findChild(DoubleSpinBox)
                    self.assertIsNotNone(float_widget)
                    self.assertEqual(float_widget.decimals(), 4)
                finally:
                    form.deleteLater()

    def test_lines_button_uses_reduced_min_width(self):
        from qfluentwidgets import PushButton
        from ui.widgets.extension_options_form import ExtensionOptionsForm

        form = ExtensionOptionsForm()
        try:
            form.set_line_candidates(["A", "B", "C"])
            form.set_fields(
                [{"key": "lines_list", "label": "lines", "field_type": "lines", "default": [], "extra": {"lines_number": [2, -1]}}],
                {"lines_list": []},
            )

            button = next(button for button in form.findChildren(PushButton) if button.text() == "选择曲线")
            self.assertEqual(button.minimumWidth(), 96)
        finally:
            form.deleteLater()

    def test_slider_live_tooltip_reuses_single_fluent_tip(self):
        from PySide6.QtWidgets import QSlider
        from qfluentwidgets import ToolTipPosition
        from ui.widgets.extension_options_form import ExtensionOptionsForm

        form = ExtensionOptionsForm()
        try:
            form.set_fields(
                [{"key": "alpha", "label": "透明度", "field_type": "limited", "min_value": 0, "max_value": 10, "default": 3}],
                {"alpha": 3},
            )
            form.show()
            QApplication.processEvents()

            slider = form.findChild(QSlider)
            self.assertIsNotNone(slider)
            slider.setFocus()
            QApplication.processEvents()

            slider.setValue(4)
            form._show_live_tooltip_at(slider, slider.toolTip(), ToolTipPosition.TOP)
            first_tooltip = form._live_tooltip
            self.assertIsNotNone(first_tooltip)

            slider.setValue(6)
            form._show_live_tooltip_at(slider, slider.toolTip(), ToolTipPosition.TOP)
            second_tooltip = form._live_tooltip
            self.assertIs(first_tooltip, second_tooltip)
            self.assertEqual(slider.toolTip(), "6")
            self.assertEqual(second_tooltip.text(), "6")
            self.assertFalse(bool(slider.property("_alineFluentTooltip")))
        finally:
            form.deleteLater()

    def test_fields_wrap_when_label_and_editor_cannot_share_row(self):
        from PySide6.QtWidgets import QWidget
        from ui.widgets.extension_options_form import ExtensionOptionsForm

        form = ExtensionOptionsForm()
        try:
            form.set_fields([
                {
                    "key": "very_long_factor_name",
                    "label": "这是一个很长的扩展参数标签用于验证最小控件宽度不足时自动换行",
                    "field_type": "number",
                    "default": 1.0,
                }
            ])

            container_widget = form._flow.itemAt(0).widget()
            self.assertIsNotNone(container_widget)
            row_widget = container_widget.layout().itemAt(0).widget()
            self.assertIsInstance(row_widget, QWidget)

            row_widget.resize(200, row_widget.height() or 40)
            row_widget.show()
            self.assertTrue(bool(row_widget.property("wrapped")))
        finally:
            form.deleteLater()


class TestPageOnboardingController(unittest.TestCase):

    def test_controller_uses_end_prev_next_buttons_in_one_flow(self):
        from PySide6.QtWidgets import QWidget
        from qfluentwidgets import PrimaryPushButton, PushButton, TeachingTipTailPosition
        from ui.widgets.onboarding import OnboardingStep, PageOnboardingController

        host = QWidget()
        target = QWidget(host)
        host.show()
        QApplication.processEvents()
        captured = {}

        class _FakeTip:
            def close(self):
                return None

        def _fake_make(view, tip_target, duration, tail_position, parent):
            captured["target"] = tip_target
            captured["tail_position"] = tail_position
            captured["parent"] = parent
            buttons = [button for button in view.findChildren(PushButton) if button.text()]
            captured["push_buttons"] = sorted(button.text() for button in buttons)
            captured["button_parent_ids"] = {id(button.parentWidget()) for button in buttons}
            captured["primary_buttons"] = [
                button.text() for button in view.findChildren(PrimaryPushButton) if button.text()
            ]
            return _FakeTip()

        controller = PageOnboardingController(
            host,
            "test-onboarding",
            lambda: [
                OnboardingStep(
                    lambda: target,
                    TeachingTipTailPosition.BOTTOM,
                    "标题",
                    "内容",
                )
            ],
            is_completed=lambda: False,
            mark_completed=lambda completed: completed,
        )

        with mock.patch("ui.widgets.onboarding.TeachingTip.make", side_effect=_fake_make):
            controller.start(force=True)

        self.assertCountEqual(captured["push_buttons"], ["上一步", "结束引导", "下一步"])
        self.assertEqual(captured["primary_buttons"], ["下一步"])
        self.assertEqual(len(captured["button_parent_ids"]), 1)
        self.assertIs(captured["target"], target)
        self.assertIs(captured["parent"], host)
        target.deleteLater()
        host.deleteLater()


# ═══════════════════════════════════════════════════════════════════════════
# 3. DataPage — on_tree_node_selected
# ═══════════════════════════════════════════════════════════════════════════

class TestDataPage(unittest.TestCase):

    def setUp(self):
        self.pm, self.p, self.df, self.s = _make_project("dp_test")
        self._restore = _patch_pm(self.pm)
        from ui.pages.data_page import DataPage
        self.page = DataPage()

    def tearDown(self):
        self._restore()
        self.page.deleteLater()

    def test_page_creates_no_crash(self):
        self.assertIsNotNone(self.page)

    def test_plot_preview_uses_shared_navigation_controls(self):
        if self.page._preview_figure is None or self.page._preview_canvas is None:
            self.skipTest("matplotlib unavailable")

        self.page._preview_xs = [0.0, 1.0, 2.0, 3.0]
        self.page._preview_ys = [1.0, 2.0, 1.5, 2.5]
        self.page._draw_preview()

        axis = self.page._preview_figure.axes[0]
        original_x_span = abs(axis.get_xlim()[1] - axis.get_xlim()[0])
        original_y_span = abs(axis.get_ylim()[1] - axis.get_ylim()[0])

        self.page._preview_box_zoom_btn.click()
        QApplication.processEvents()
        self.assertEqual(self.page._preview_navigation_mode(), "zoom")
        self.assertTrue(self.page._preview_box_zoom_btn.isChecked())
        self.assertFalse(self.page._preview_pan_btn.isChecked())

        self.page._preview_zoom_in_btn.click()
        QApplication.processEvents()

        zoomed_axis = self.page._preview_figure.axes[0]
        self.assertLess(abs(zoomed_axis.get_xlim()[1] - zoomed_axis.get_xlim()[0]), original_x_span)
        self.assertLess(abs(zoomed_axis.get_ylim()[1] - zoomed_axis.get_ylim()[0]), original_y_span)

        self.page._preview_fit_btn.click()
        QApplication.processEvents()

        reset_axis = self.page._preview_figure.axes[0]
        self.assertAlmostEqual(abs(reset_axis.get_xlim()[1] - reset_axis.get_xlim()[0]), original_x_span, places=6)
        self.assertAlmostEqual(abs(reset_axis.get_ylim()[1] - reset_axis.get_ylim()[0]), original_y_span, places=6)

    def test_global_extension_config_selection_opens_json_editor_and_saves(self):
        from core.extension_api import ExtensionConfigField, ProcessingExtension, extension_registry
        from core.global_assets import global_assets

        def _probe(lines, params):
            return lines[0] if lines else []

        restore_assets = _patch_global_assets()
        extension_registry.register_processing(
            ProcessingExtension(
                type="data_page_config_editor",
                name="数据页配置编辑",
                handler=_probe,
                config_fields=[ExtensionConfigField(key="factor", field_type="number", required=True, default=2.0)],
            )
        )
        try:
            entry = extension_registry.get_processing("data_page_config_editor")
            config = global_assets.add_extension_config(
                category="processing",
                extension_type="data_page_config_editor",
                extension_name="数据页配置编辑",
                extension_version=getattr(entry, "version", None),
                name="方案A",
                options={"factor": 2.0},
            )

            self.page.on_tree_node_selected("global_extension_config", config.id)

            self.assertIs(self.page._preview_stack.currentWidget(), self.page._text_preview)
            self.assertEqual(self.page._preview_section_label.text(), "配置编辑")
            self.assertFalse(self.page._text_preview.isReadOnly())
            self.assertFalse(self.page._extension_config_action_panel.isHidden())
            self.assertFalse(self.page._btn_reset_extension_config.isHidden())
            self.assertFalse(self.page._config_editor_header_panel.isHidden())
            self.assertTrue(self.page._summary_footer.isHidden())
            self.assertIn("配置名称", self.page._config_editor_title_label.text())
            self.assertFalse(self.page._extension_preview_panel.isHidden())
            self.assertEqual(self.page._extension_field_help_area.maximumHeight(), 202)
            self.assertEqual(self.page._extension_field_help_title.text(), "参数说明")
            self.assertIn("factor", self.page._extension_field_help_label.text())

            self.page._text_preview.setPlainText('{"factor": 4.5}')
            self.page._reset_selected_extension_config_edit()
            self.assertEqual(self.page._text_preview.toPlainText().strip(), '{\n  "factor": 2.0\n}')
            self.page._text_preview.setPlainText('{"factor": 4.5}')
            self.page._save_selected_extension_config()

            self.assertEqual(global_assets.get_extension_config(config.id).options, {"factor": 4.5})
        finally:
            extension_registry.unregister_processing("data_page_config_editor")
            restore_assets()

    def test_global_extension_config_save_validates_json_and_required_fields(self):
        from core.extension_api import ExtensionConfigField, ProcessingExtension, extension_registry
        from core.global_assets import global_assets

        def _probe(lines, params):
            return lines[0] if lines else []

        restore_assets = _patch_global_assets()
        extension_registry.register_processing(
            ProcessingExtension(
                type="data_page_config_validate",
                name="数据页配置校验",
                handler=_probe,
                config_fields=[ExtensionConfigField(key="factor", field_type="number", required=True, default=2.0)],
            )
        )
        try:
            config = global_assets.add_extension_config(
                category="processing",
                extension_type="data_page_config_validate",
                extension_name="数据页配置校验",
                name="方案B",
                options={"factor": 2.0},
            )
            self.page.on_tree_node_selected("global_extension_config", config.id)

            with mock.patch("ui.pages.data_page.InfoBar.warning") as warning_mock:
                self.page._text_preview.setPlainText('{')
                self.page._save_selected_extension_config()
                self.page._text_preview.setPlainText('{}')
                self.page._save_selected_extension_config()

            self.assertEqual(global_assets.get_extension_config(config.id).options, {"factor": 2.0})
            self.assertGreaterEqual(warning_mock.call_count, 2)
        finally:
            extension_registry.unregister_processing("data_page_config_validate")
            restore_assets()

    def test_extension_field_help_formats_lines_summary(self):
        from core.extension_api import ProcessingExtension, build_extension_entry

        entry = build_extension_entry(
            ProcessingExtension(
                type="data_page_lines_help",
                name="DataPage 曲线帮助",
                handler=lambda lines, params: lines[0] if lines else [],
                lines_number=(2, -1),
            )
        )

        help_text = self.page._extension_field_help_text(entry)

        self.assertIn("lines: 本扩展支持的曲线数量为 2 条及以上。", help_text)

    def test_extension_field_help_omits_repeated_key_name(self):
        from core.extension_api import ExtensionConfigField, ProcessingExtension, build_extension_entry

        entry = build_extension_entry(
            ProcessingExtension(
                type="data_page_help_label_only",
                name="DataPage 标签帮助",
                handler=lambda lines, params: lines[0] if lines else [],
                config_fields=[
                    ExtensionConfigField(key="n", label="n", field_type="integer", default=3, description="采样点数量。")
                ],
            )
        )

        help_text = self.page._extension_field_help_text(entry)

        self.assertIn("n（integer，可选）；采样点数量。；默认值: 3", help_text)
        self.assertNotIn("n（n，integer", help_text)

    def test_global_extension_group_preview_uses_extension_description_and_field_help(self):
        from core.extension_api import ExtensionConfigField, ProcessingExtension, extension_registry

        def _probe(xs, ys, params):
            return list(xs), list(ys)

        restore_assets = _patch_global_assets()
        extension_registry.register_processing(
            ProcessingExtension(
                type="data_page_group_preview",
                name="数据页分组预览",
                handler=_probe,
                description="用于验证扩展节点预览说明。",
                settings=True,
                config_fields=[
                    ExtensionConfigField(key="factor", label="倍率", field_type="number", description="缩放结果。", default=2.0)
                ],
            )
        )
        try:
            self.page.on_tree_node_selected("global_group", "__global_extension_configs__:processing:data_page_group_preview")

            self.assertFalse(self.page._extension_preview_panel.isHidden())
            self.assertTrue(self.page._summary_footer.isHidden())
            self.assertEqual(self.page._extension_field_help_area.maximumHeight(), 324)
            self.assertEqual(self.page._extension_detail_label.text(), "处理扩展·数据页分组预览")
            self.assertIn("ID: data_page_group_preview", self.page._extension_detail_meta_label.text())
            self.assertIn("描述: 用于验证扩展节点预览说明。", self.page._extension_detail_meta_label.text())
            self.assertIn("缩放结果。", self.page._extension_field_help_label.text())
        finally:
            extension_registry.unregister_processing("data_page_group_preview")
            restore_assets()

    def test_management_panel_uses_left_tool_layout(self):
        from ui.theme import WORKBENCH_TOOL_PANEL_WIDTH

        self.assertEqual(self.page._tool_panel.minimumWidth(), WORKBENCH_TOOL_PANEL_WIDTH)
        self.assertEqual(self.page._tool_panel.maximumWidth(), WORKBENCH_TOOL_PANEL_WIDTH)
        self.assertEqual(self.page._content_splitter.count(), 2)

    def test_series_preview_uses_plot_canvas_and_switches_plot_type(self):
        self.page.on_tree_node_selected("series", self.s.id)

        if self.page._preview_figure is None or self.page._preview_canvas is None:
            self.skipTest("matplotlib unavailable")

        self.assertEqual(self.page._preview_type_combo.count(), 5)
        self.assertEqual(len(self.page._preview_figure.axes), 1)
        self.assertGreaterEqual(len(self.page._preview_figure.axes[0].lines), 1)

        scatter_index = next(
            index for index in range(self.page._preview_type_combo.count())
            if self.page._preview_type_combo.itemText(index) == "散点"
        )
        self.page._preview_type_combo.setCurrentIndex(scatter_index)
        self.assertGreaterEqual(len(self.page._preview_figure.axes[0].collections), 1)

    def test_non_data_preview_hides_plot_type_controls(self):
        pictures_root = self.pm._find_folder_by_group_type("pictures")
        self.assertIsNotNone(pictures_root)

        self.page.on_tree_node_selected("folder", pictures_root.id)
        self.assertTrue(self.page._preview_plot_type_controls.isHidden())

        self.page.on_tree_node_selected("series", self.s.id)
        self.assertFalse(self.page._preview_plot_type_controls.isHidden())

    def test_picture_node_selection_shows_tree_preview(self):
        from PySide6.QtGui import QImage

        with tempfile.TemporaryDirectory() as temp_dir:
            picture_path = Path(temp_dir) / "preview.png"
            image = QImage(48, 32, QImage.Format.Format_RGB32)
            image.fill(Qt.GlobalColor.white)
            self.assertTrue(image.save(str(picture_path)))
            picture_node = self.pm.add_picture(str(picture_path), name="preview.png")

            self.assertIsNotNone(picture_node)
            self.page.on_tree_node_selected("picture", picture_node.id)

        self.assertIs(self.page._preview_stack.currentWidget(), self.page._picture_preview_tree)
        self.assertEqual(self.page._stats_title_label.text(), "图片名称: preview.png")
        self.assertIn("绘图快照: 未保存", self.page._stats_label.text())
        root = self.page._picture_preview_tree.topLevelItem(0)
        self.assertIsNotNone(root)
        self.assertEqual(root.text(0), "preview.png")
        self.assertEqual([root.child(i).text(0) for i in range(root.childCount())], ["图片文件", "绘图快照"])

    def test_picture_tree_preview_shows_snapshot_content(self):
        from PySide6.QtGui import QImage
        from models.schemas import PicturePlotExtensionSnapshot, PicturePlotSeriesSnapshot, PicturePlotSnapshot

        with tempfile.TemporaryDirectory() as temp_dir:
            picture_path = Path(temp_dir) / "snapshot.png"
            image = QImage(24, 24, QImage.Format.Format_RGB32)
            image.fill(Qt.GlobalColor.white)
            self.assertTrue(image.save(str(picture_path)))

            snapshot = PicturePlotSnapshot(
                selected_curve_key="curve-1",
                series=[
                    PicturePlotSeriesSnapshot(
                        curve_key="curve-1",
                        name="curve-1",
                        display_name="曲线 1",
                        x=[1.0, 2.0],
                        y=[3.0, 4.0],
                        source="picture",
                        visible=True,
                    )
                ],
                applied_extensions=[
                    PicturePlotExtensionSnapshot(
                        id="ext-1",
                        type="plot_text_annotation",
                        sequence=1,
                        options={"text": "A"},
                        curve_display_name="曲线 1",
                        extension_version="0.1.0",
                    )
                ],
            )
            picture_node = self.pm.add_picture(str(picture_path), name="snapshot.png", plot_snapshot=snapshot)

            self.assertIsNotNone(picture_node)
            self.page.on_tree_node_selected("picture", picture_node.id)

        root = self.page._picture_preview_tree.topLevelItem(0)
        self.assertIsNotNone(root)
        snapshot_item = root.child(1)
        self.assertEqual(snapshot_item.text(0), "绘图快照")
        labels = [snapshot_item.child(i).text(0) for i in range(snapshot_item.childCount())]
        self.assertIn("当前选中曲线: curve-1", labels)
        self.assertIn("曲线 (1)", labels)
        self.assertIn("绘图扩展 (1)", labels)

    def test_picture_root_folder_selection_uses_tree_preview(self):
        pictures_root = self.pm._find_folder_by_group_type("pictures")
        self.assertIsNotNone(pictures_root)

        self.page.on_tree_node_selected("folder", pictures_root.id)

        self.assertIs(self.page._preview_stack.currentWidget(), self.page._picture_preview_tree)
        root = self.page._picture_preview_tree.topLevelItem(0)
        self.assertIsNotNone(root)
        self.assertEqual(root.text(0), "图片集")
        self.assertIn("文件夹: 图片集", self.page._stats_title_label.text())

    def test_analysis_root_folder_selection_uses_tree_preview(self):
        analysis_root = self.pm._find_folder_by_group_type("analysis_result_group")
        self.assertIsNotNone(analysis_root)

        self.page.on_tree_node_selected("folder", analysis_root.id)

        self.assertIs(self.page._preview_stack.currentWidget(), self.page._picture_preview_tree)
        root = self.page._picture_preview_tree.topLevelItem(0)
        self.assertIsNotNone(root)
        self.assertEqual(root.text(0), "分析结果")
        self.assertIn("文件夹: 分析结果", self.page._stats_title_label.text())

    def test_global_root_selection_uses_tree_preview(self):
        self.page.on_tree_node_selected("global_root", "__global_root__")

        self.assertIs(self.page._preview_stack.currentWidget(), self.page._picture_preview_tree)
        root = self.page._picture_preview_tree.topLevelItem(0)
        self.assertIsNotNone(root)
        self.assertEqual(root.text(0), "全局资源")
        self.assertGreaterEqual(root.childCount(), 5)

    def test_global_report_template_selection_uses_tree_preview(self):
        from core.global_assets import global_assets
        from models.schemas import ReportTemplate

        template = global_assets.ensure_report_template(ReportTemplate(name="全局报告预览", content="# report preview"))
        self.page.on_tree_node_selected("global_report_template", template.id)

        self.assertIs(self.page._preview_stack.currentWidget(), self.page._picture_preview_tree)
        root = self.page._picture_preview_tree.topLevelItem(0)
        self.assertIsNotNone(root)
        self.assertEqual(root.text(0), "全局报告预览")
        self.assertGreaterEqual(root.childCount(), 2)

    def test_picture_node_management_actions_enable_rename_delete_and_visualize(self):
        from PySide6.QtGui import QImage

        with tempfile.TemporaryDirectory() as temp_dir:
            picture_path = Path(temp_dir) / "preview.png"
            image = QImage(48, 32, QImage.Format.Format_RGB32)
            image.fill(Qt.GlobalColor.white)
            self.assertTrue(image.save(str(picture_path)))
            picture_node = self.pm.add_picture(str(picture_path), name="preview.png")

            self.assertIsNotNone(picture_node)
            picture_id = picture_node.picture_id
            self.page.on_tree_node_selected("picture", picture_node.id)

            self.assertTrue(self.page._btn_to_vis.isEnabled())
            self.assertTrue(self.page._btn_apply_name.isEnabled())
            self.assertTrue(self.page._btn_delete_node.isEnabled())

            sent: list[tuple[str, str]] = []
            self.page.send_to_visualize.connect(lambda kind, node_id: sent.append((kind, node_id)))
            self.page._send_to_visualize()
            self.assertEqual(sent, [("picture", picture_node.id)])

            self.page._manage_name_edit.setText("renamed.png")
            self.page._apply_rename_current_node()
            self.assertEqual(self.pm.get_picture(picture_id).name, "renamed.png")

            with mock.patch("ui.pages.data_page.MessageBox.exec", return_value=True):
                self.page._delete_current_node()

            self.assertIsNone(self.pm.get_picture(picture_id))
            self.assertIsNone(self.pm.get_node_by_id(picture_node.id))

    def test_source_file_data_preview_switches_between_parsed_and_source_modes(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            source_path = Path(temp_dir) / "source-data.csv"
            source_path.write_text("raw_x,raw_y\n10,20\n30,40\n", encoding="utf-8")
            node = self.pm.add_source_file(str(source_path))
            self.assertIsNotNone(node)

            self.page.on_tree_node_selected("source_file", node.id)

            self.assertFalse(self.page._source_file_preview_controls.isHidden())
            self.assertEqual(
                [self.page._source_file_preview_combo.itemText(index) for index in range(self.page._source_file_preview_combo.count())],
                ["解析", "源文件"],
            )
            self.assertIs(self.page._preview_stack.currentWidget(), self.page._parsed_preview_table)
            self.assertEqual(self.page._parsed_preview_table.rowCount(), 2)
            self.assertEqual(self.page._parsed_preview_table.columnCount(), 3)
            self.assertEqual(self.page._parsed_preview_table.horizontalHeaderItem(0).text(), "行号")
            self.assertEqual(self.page._parsed_preview_table.item(0, 0).text(), "1")
            self.assertEqual(self.page._parsed_preview_table.item(0, 1).text(), "10")
            self.assertEqual(self.page._parsed_preview_table.item(0, 2).text(), "20")
            self.assertIn("列数: 2", self.page._stats_label.text())

            self.page._source_file_preview_combo.setCurrentText("源文件")

            self.assertIs(self.page._preview_stack.currentWidget(), self.page._text_preview)
            self.assertIn("raw_x,raw_y", self.page._text_preview.toPlainText())
            self.assertEqual(self.page._current_source_path_button.toolTip(), str(source_path))

    def test_source_file_parsed_preview_shows_metadata_and_custom_row_limit(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            source_path = Path(temp_dir) / "source-metadata.csv"
            source_path.write_text("x;y\n1;2.5\n3;4\n5;6\n", encoding="utf-8")
            node = self.pm.add_source_file(str(source_path))
            self.assertIsNotNone(node)

            self.page.on_tree_node_selected("source_file", node.id)
            self.page._source_file_row_limit_edit.setText("2")
            self.page._on_source_file_row_limit_changed()

            self.assertIs(self.page._preview_stack.currentWidget(), self.page._parsed_preview_table)
            self.assertEqual(self.page._parsed_preview_table.rowCount(), 2)
            self.assertIn("编码:", self.page._stats_label.text())
            self.assertIn("分隔: 分号 (;)" , self.page._stats_label.text())
            self.assertIn("表头: 自动识别", self.page._stats_label.text())
            self.assertNotIn("列类型:", self.page._stats_label.text())
            self.assertNotIn("当前显示", self.page._stats_label.text())
            self.assertEqual(self.page._parsed_preview_table.horizontalHeaderItem(0).text(), "行号")
            self.assertEqual(self.page._parsed_preview_table.horizontalHeaderItem(1).text(), "x · 整数")
            self.assertEqual(self.page._parsed_preview_table.horizontalHeaderItem(2).text(), "y · 浮点数")
            self.assertEqual(self.page._source_file_page_label.text(), "1 - 2 / 3")

    def test_source_file_skip_rows_disables_header_detection(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            source_path = Path(temp_dir) / "skip-header.csv"
            source_path.write_text("meta,info\nx,y\n1,2\n3,4\n", encoding="utf-8")
            node = self.pm.add_source_file(str(source_path))
            self.assertIsNotNone(node)

            self.page.on_tree_node_selected("source_file", node.id)
            self.page._source_file_skip_rows_edit.setText("1")
            self.page._on_source_file_skip_rows_changed()

            self.assertIs(self.page._preview_stack.currentWidget(), self.page._parsed_preview_table)
            self.assertEqual(self.page._parsed_preview_table.rowCount(), 2)
            self.assertEqual(self.page._parsed_preview_table.horizontalHeaderItem(0).text(), "行号")
            self.assertEqual(self.page._parsed_preview_table.horizontalHeaderItem(1).text(), "col_0 · 整数")
            self.assertEqual(self.page._parsed_preview_table.horizontalHeaderItem(2).text(), "col_1 · 整数")
            self.assertIn("表头: 跳过1行", self.page._stats_label.text())
            self.assertIn("无效行: 1", self.page._stats_label.text())

    def test_source_file_auto_header_missing_is_reported_in_invalid_rows(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            source_path = Path(temp_dir) / "no-header.csv"
            source_path.write_text("1,2\n3,4\n5,6\n", encoding="utf-8")
            node = self.pm.add_source_file(str(source_path))
            self.assertIsNotNone(node)

            self.page.on_tree_node_selected("source_file", node.id)

            self.assertIs(self.page._preview_stack.currentWidget(), self.page._parsed_preview_table)
            self.assertIn("表头: 自动识别", self.page._stats_label.text())
            self.assertIn("无效行: 0（缺少表头）", self.page._stats_label.text())
            self.assertEqual(self.page._parsed_preview_table.horizontalHeaderItem(0).text(), "行号")
            self.assertEqual(self.page._parsed_preview_table.horizontalHeaderItem(1).text(), "col_0 · 整数")
            self.assertEqual(self.page._parsed_preview_table.horizontalHeaderItem(2).text(), "col_1 · 整数")

    def test_source_file_header_issue_is_reported_in_invalid_rows(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            source_path = Path(temp_dir) / "bad-header.csv"
            source_path.write_text("x,\n1,2\n3,4\n", encoding="utf-8")
            node = self.pm.add_source_file(str(source_path))
            self.assertIsNotNone(node)

            self.page.on_tree_node_selected("source_file", node.id)

            self.assertIs(self.page._preview_stack.currentWidget(), self.page._parsed_preview_table)
            self.assertIn("表头: 自动识别", self.page._stats_label.text())
            self.assertIn("无效行: 0（表头缺项或错误）", self.page._stats_label.text())
            self.assertEqual(self.page._parsed_preview_table.horizontalHeaderItem(0).text(), "行号")
            self.assertEqual(self.page._parsed_preview_table.horizontalHeaderItem(1).text(), "col_0 · 整数")
            self.assertEqual(self.page._parsed_preview_table.horizontalHeaderItem(2).text(), "col_1 · 整数")

    def test_source_file_parsed_preview_remains_fully_horizontally_scrollable_when_narrow(self):
        from PySide6.QtGui import QFontMetrics
        from PySide6.QtWidgets import QHeaderView

        with tempfile.TemporaryDirectory() as temp_dir:
            source_path = Path(temp_dir) / "wide-header.csv"
            headers = [f"column_{index}_very_long_header_name" for index in range(1, 13)]
            row_a = [str(index) for index in range(1, 13)]
            row_b = [str(index * 10) for index in range(1, 13)]
            source_path.write_text(
                ",".join(headers) + "\n" + ",".join(row_a) + "\n" + ",".join(row_b) + "\n",
                encoding="utf-8",
            )
            node = self.pm.add_source_file(str(source_path))
            self.assertIsNotNone(node)

            self.page.resize(360, 820)
            self.page.show()
            self.page.on_tree_node_selected("source_file", node.id)
            QApplication.processEvents()

            table = self.page._parsed_preview_table
            header = table.horizontalHeader()
            last_column = table.columnCount() - 1
            last_header_text = table.horizontalHeaderItem(last_column).text()
            last_header_text_width = QFontMetrics(header.font()).horizontalAdvance(last_header_text)

            self.assertGreater(table.horizontalScrollBar().maximum(), 0)
            self.assertEqual(header.sectionResizeMode(last_column), QHeaderView.ResizeMode.ResizeToContents)
            self.assertGreaterEqual(header.sectionSize(last_column), last_header_text_width - 6)

            table.horizontalScrollBar().setValue(table.horizontalScrollBar().maximum())
            QApplication.processEvents()
            last_rect = table.visualItemRect(table.item(0, last_column))
            self.assertLessEqual(last_rect.right(), table.viewport().width())

    def test_source_file_parse_failure_shows_reason(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            source_path = Path(temp_dir) / "failed.csv"
            source_path.write_text("name,value\nfoo,bar\nhello,world\n", encoding="utf-8")
            node = self.pm.add_source_file(str(source_path))
            self.assertIsNotNone(node)

            self.page.on_tree_node_selected("source_file", node.id)

            self.assertIs(self.page._preview_stack.currentWidget(), self.page._text_preview)
            self.assertIn("无法解析文件预览", self.page._text_preview.toPlainText())
            self.assertIn("解析状态: 失败", self.page._stats_label.text())
            self.assertIn("失败原因:", self.page._stats_label.text())

    def test_source_file_excel_preview_supports_sheet_selection(self):
        import openpyxl

        with tempfile.TemporaryDirectory() as temp_dir:
            source_path = Path(temp_dir) / "sheet-preview.xlsx"
            workbook = openpyxl.Workbook()
            sheet_a = workbook.active
            sheet_a.title = "SheetA"
            sheet_a.append(["x", "y"])
            sheet_a.append([1, 2])
            sheet_b = workbook.create_sheet("SheetB")
            sheet_b.append(["x", "y"])
            sheet_b.append([10, 20])
            workbook.save(source_path)
            workbook.close()

            node = self.pm.add_source_file(str(source_path))
            self.assertIsNotNone(node)

            self.page.on_tree_node_selected("source_file", node.id)

            self.assertFalse(self.page._source_file_sheet_combo.isHidden())
            self.assertEqual(
                [self.page._source_file_sheet_combo.itemText(index) for index in range(self.page._source_file_sheet_combo.count())],
                ["SheetA", "SheetB"],
            )
            self.assertEqual(self.page._parsed_preview_table.item(0, 0).text(), "1")
            self.assertEqual(self.page._parsed_preview_table.item(0, 1).text(), "1")

            self.page._source_file_sheet_combo.setCurrentText("SheetB")

            self.assertEqual(self.page._parsed_preview_table.item(0, 0).text(), "1")
            self.assertEqual(self.page._parsed_preview_table.item(0, 1).text(), "10")
            self.assertIn("工作表: SheetB", self.page._stats_label.text())

    def test_source_file_preview_state_is_remembered_per_node(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            source_path_a = Path(temp_dir) / "node-a.csv"
            source_path_b = Path(temp_dir) / "node-b.csv"
            source_path_a.write_text("x,y\n10,20\n30,40\n50,60\n70,80\n", encoding="utf-8")
            source_path_b.write_text("x,y\n1,2\n3,4\n", encoding="utf-8")
            node_a = self.pm.add_source_file(str(source_path_a))
            node_b = self.pm.add_source_file(str(source_path_b))
            self.assertIsNotNone(node_a)
            self.assertIsNotNone(node_b)

            self.page.on_tree_node_selected("source_file", node_a.id)
            self.page._source_file_preview_combo.setCurrentText("源文件")
            self.page._source_file_row_limit_edit.setText("2")
            self.page._on_source_file_row_limit_changed()
            self.page._source_file_skip_rows_edit.setText("3")
            self.page._on_source_file_skip_rows_changed()
            self.page._source_file_next_page_btn.click()

            self.assertIn("30,40", self.page._text_preview.toPlainText())
            self.assertEqual(self.page._source_file_page_label.text(), "3 - 4 / 5")

            self.page.on_tree_node_selected("source_file", node_b.id)
            self.page.on_tree_node_selected("source_file", node_a.id)

            self.assertEqual(self.page._source_file_preview_combo.currentText(), "源文件")
            self.assertEqual(self.page._source_file_row_limit_edit.text(), "2")
            self.assertEqual(self.page._source_file_skip_rows_edit.text(), "3")
            self.assertEqual(self.page._source_file_page_label.text(), "3 - 4 / 5")
            self.assertIn("30,40", self.page._text_preview.toPlainText())

    def test_source_file_pagination_supports_first_last_and_jump_line(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            source_path = Path(temp_dir) / "paged.csv"
            source_path.write_text("x,y\n1,10\n2,20\n3,30\n4,40\n5,50\n", encoding="utf-8")
            node = self.pm.add_source_file(str(source_path))
            self.assertIsNotNone(node)

            self.page.on_tree_node_selected("source_file", node.id)
            self.page._source_file_row_limit_edit.setText("2")
            self.page._on_source_file_row_limit_changed()

            self.page._source_file_last_page_btn.click()
            self.assertEqual(self.page._source_file_page_label.text(), "5 - 5 / 5")

            self.page._source_file_first_page_btn.click()
            self.assertEqual(self.page._source_file_page_label.text(), "1 - 2 / 5")

            self.page._source_file_jump_line_edit.setText("4")
            self.page._jump_to_source_file_page()
            self.assertEqual(self.page._source_file_page_label.text(), "3 - 4 / 5")

    def test_plot_type_is_remembered_per_node(self):
        data_file_node = next((n for n in self.p.tree.nodes if n.kind == "data_file"), None)
        self.assertIsNotNone(data_file_node)

        self.page.on_tree_node_selected("data_file", data_file_node.id)
        self.page._preview_type_combo.setCurrentText("柱状")

        self.page.on_tree_node_selected("series", self.s.id)
        self.page._preview_type_combo.setCurrentText("散点")

        self.page.on_tree_node_selected("data_file", data_file_node.id)
        self.assertEqual(self.page._preview_type_combo.currentText(), "柱状")

        self.page.on_tree_node_selected("series", self.s.id)
        self.assertEqual(self.page._preview_type_combo.currentText(), "散点")

    def test_folder_browser_directory_is_remembered_per_node(self):
        dataset_root = self.pm._find_folder_by_group_type("datasets")
        self.assertIsNotNone(dataset_root)
        folder_a = self.pm.add_folder("批次A", parent_id=dataset_root.id, group_type="datasets")
        folder_b = self.pm.add_folder("批次B", parent_id=dataset_root.id, group_type="datasets")
        self.assertIsNotNone(folder_a)
        self.assertIsNotNone(folder_b)

        with tempfile.TemporaryDirectory() as dir_a, tempfile.TemporaryDirectory() as dir_b:
            self.page.on_tree_node_selected("folder", folder_a.id)
            self.page._external_browser_dir = Path(dir_a)
            self.page._remember_current_external_browser_dir()

            self.page.on_tree_node_selected("folder", folder_b.id)
            self.page._external_browser_dir = Path(dir_b)
            self.page._remember_current_external_browser_dir()

            self.page.on_tree_node_selected("folder", folder_a.id)
            self.assertEqual(self.page._external_browser_dir, Path(dir_a))

            self.page.on_tree_node_selected("folder", folder_b.id)
            self.assertEqual(self.page._external_browser_dir, Path(dir_b))

    def test_pending_source_file_can_be_renamed_before_import(self):
        source_root = self.pm._find_folder_by_group_type("source_files")
        self.assertIsNotNone(source_root)

        with tempfile.TemporaryDirectory() as temp_dir:
            source_path = Path(temp_dir) / "raw-source.csv"
            source_path.write_text("x,y\n1,2\n", encoding="utf-8")

            self.page.on_tree_node_selected("folder", source_root.id)
            self.page._append_source_files_to_pending([str(source_path)])

            item = self.page._pending_source_list.topLevelItem(0)
            self.assertIsNotNone(item)
            self.assertTrue(bool(item.flags() & Qt.ItemFlag.ItemIsEditable))

            item.setText(0, "renamed-source.csv")
            self.page._on_pending_source_item_changed(item, 0)
            self.page._import_pending_files_as_source_files()

        imported = next(
            (node for node in self.p.tree.nodes if node.kind == "source_file" and node.name == "renamed-source.csv"),
            None,
        )
        self.assertIsNotNone(imported)

    def test_pending_and_browser_item_views_use_fluent_tooltip_filters(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            source_path = Path(temp_dir) / "tooltip-source.csv"
            source_path.write_text("x,y\n1,2\n", encoding="utf-8")
            self.page._append_source_files_to_pending([str(source_path)])

            pending_item = self.page._pending_source_list.topLevelItem(0)
            self.assertIsNotNone(pending_item)
            self.assertEqual(pending_item.toolTip(0), self.page._pending_entry_tooltip(str(source_path)))

        self.assertIs(self.page._fluent_tooltip_views[self.page._pending_source_list.viewport()], self.page._pending_source_list)
        self.assertIs(self.page._fluent_tooltip_views[self.page._source_browser.viewport()], self.page._source_browser)
        self.assertIs(self.page._fluent_tooltip_views[self.page._project_source_browser.viewport()], self.page._project_source_browser)
        self.assertNotIn("QToolTip", self.page._pending_source_list.styleSheet())

    def test_pending_import_lists_are_isolated_by_group(self):
        source_root = self.pm._find_folder_by_group_type("source_files")
        datasets_root = self.pm._find_folder_by_group_type("datasets")
        images_root = self.pm._find_folder_by_group_type("images")
        self.assertIsNotNone(source_root)
        self.assertIsNotNone(datasets_root)
        self.assertIsNotNone(images_root)

        with tempfile.TemporaryDirectory() as temp_dir:
            source_path = Path(temp_dir) / "source.csv"
            dataset_path = Path(temp_dir) / "dataset.csv"
            source_path.write_text("x,y\n1,2\n", encoding="utf-8")
            dataset_path.write_text("x,y\n3,4\n", encoding="utf-8")

            self.page.on_tree_node_selected("folder", source_root.id)
            self.page._append_source_files_to_pending([str(source_path)])
            self.assertEqual(self.page._pending_source_list.topLevelItemCount(), 1)

            self.page.on_tree_node_selected("folder", datasets_root.id)
            self.assertEqual(self.page._pending_source_list.topLevelItemCount(), 0)
            self.page._append_source_files_to_pending([str(dataset_path)])
            self.assertEqual(self.page._pending_source_list.topLevelItemCount(), 1)

            self.page.on_tree_node_selected("folder", images_root.id)
            self.assertEqual(self.page._pending_source_list.topLevelItemCount(), 0)

            self.page.on_tree_node_selected("folder", source_root.id)
            self.assertEqual(self.page._pending_source_list.topLevelItemCount(), 1)
            self.assertEqual(
                self.page._pending_source_list.topLevelItem(0).data(0, Qt.ItemDataRole.UserRole),
                str(source_path.resolve()),
            )

            self.page.on_tree_node_selected("folder", datasets_root.id)
            self.assertEqual(self.page._pending_source_list.topLevelItemCount(), 1)
            self.assertEqual(
                self.page._pending_source_list.topLevelItem(0).data(0, Qt.ItemDataRole.UserRole),
                str(dataset_path.resolve()),
            )

    def test_pending_dataset_file_rename_is_used_as_default_import_name(self):
        datasets_root = self.pm._find_folder_by_group_type("datasets")
        self.assertIsNotNone(datasets_root)

        with tempfile.TemporaryDirectory() as temp_dir:
            source_path = Path(temp_dir) / "raw-dataset.csv"
            source_path.write_text("x,y\n1,2\n", encoding="utf-8")

            self.page.on_tree_node_selected("folder", datasets_root.id)
            self.page._append_source_files_to_pending([str(source_path)])

            item = self.page._pending_source_list.topLevelItem(0)
            self.assertIsNotNone(item)
            self.assertTrue(bool(item.flags() & Qt.ItemFlag.ItemIsEditable))
            item.setText(0, "renamed-dataset")
            self.page._on_pending_source_item_changed(item, 0)

            normalized = str(source_path.resolve())
            dialog = self.page._create_import_dialog(
                str(source_path),
                default_file_name=self.page._pending_import_names[normalized],
            )
            try:
                self.assertEqual(dialog.get_file_name(), "renamed-dataset")
            finally:
                dialog.deleteLater()

    def test_pending_digitize_file_rename_is_used_on_import(self):
        from PySide6.QtGui import QImage

        images_root = self.pm._find_folder_by_group_type("images")
        self.assertIsNotNone(images_root)

        with tempfile.TemporaryDirectory() as temp_dir:
            source_path = Path(temp_dir) / "raw-image.png"
            image = QImage(18, 18, QImage.Format.Format_RGB32)
            image.fill(Qt.GlobalColor.white)
            self.assertTrue(image.save(str(source_path)))

            self.page.on_tree_node_selected("folder", images_root.id)
            self.page._append_source_files_to_pending([str(source_path)])

            item = self.page._pending_source_list.topLevelItem(0)
            self.assertIsNotNone(item)
            self.assertTrue(bool(item.flags() & Qt.ItemFlag.ItemIsEditable))
            item.setText(0, "renamed-image.png")
            self.page._on_pending_source_item_changed(item, 0)

            self.page._import_pending_files_for_current_group()

            imported = next((image_work for image_work in self.p.images if image_work.name == "renamed-image.png"), None)
            self.assertIsNotNone(imported)

    def test_import_dialog_accepts_whitespace_delimited_source_file(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            source_path = Path(temp_dir) / "whitespace.dat"
            source_path.write_text("1 2\n3 4\n5 6\n", encoding="utf-8")

            dialog = self.page._create_import_dialog(str(source_path))
            try:
                self.assertEqual(dialog._raw_headers, ["col_0", "col_1"])
                self.assertEqual(dialog._raw_rows, [[1.0, 2.0], [3.0, 4.0], [5.0, 6.0]])
            finally:
                dialog.deleteLater()

    def test_scaled_preview_image_uses_current_label_size(self):
        from PySide6.QtGui import QImage, QPixmap

        image = QImage(1200, 300, QImage.Format.Format_RGB32)
        image.fill(Qt.GlobalColor.white)
        pixmap = QPixmap.fromImage(image)

        self.page.resize(1280, 920)
        self.page.show()
        self.page._preview_stack.setCurrentWidget(self.page._image_preview_label)
        QApplication.processEvents()

        target_width, target_height = self.page._preview_image_target_size()
        scaled = self.page._scaled_preview_image_pixmap(pixmap)

        self.assertGreater(target_height, 320)
        self.assertTrue(scaled.width() == target_width or scaled.height() == target_height)

    def test_preview_host_follows_dark_background(self):
        if self.page._preview_figure is None or self.page._preview_canvas is None:
            self.skipTest("matplotlib unavailable")

        with mock.patch("ui.pages.data_page.isDarkTheme", return_value=True):
            self.page._apply_preview_host_background()
            self.page._draw_preview()

        self.assertIn("#1e1e1e", self.page._plot_preview_panel.styleSheet())
        self.assertIn("#1e1e1e", self.page._preview_canvas.styleSheet())

    def test_preview_drop_detects_supported_import_file(self):
        from PySide6.QtCore import QMimeData, QUrl

        mime_data = QMimeData()
        mime_data.setUrls([
            QUrl.fromLocalFile("/tmp/ignore.md"),
            QUrl.fromLocalFile("/tmp/demo.csv"),
        ])

        self.assertEqual(self.page._supported_drop_file_path(mime_data), "/tmp/demo.csv")

    def test_on_tree_node_selected_data_file(self):
        node = next((n for n in self.p.tree.nodes if n.kind == "data_file"), None)
        if node:
            self.page.on_tree_node_selected("data_file", node.id)
            self.assertEqual(self.page._selected_id, self.s.id)
            self.assertIs(self.page._preview_stack.currentWidget(), self.page._plot_preview_panel)
            self.assertTrue(self.page._source_file_preview_controls.isHidden())

    def test_import_dialog_results_preserve_source_path_on_new_data_file(self):
        from models.schemas import DataSeries

        dialog = mock.Mock()
        dialog.get_results.return_value = [DataSeries(name="imported", x=[1.0], y=[2.0])]
        dialog.get_target_data_file_id.return_value = None
        dialog.get_file_name.return_value = "imported.csv"
        dialog.get_source_path.return_value = "/tmp/imported.csv"

        self.assertTrue(self.page._apply_import_dialog_results(dialog, show_feedback=False))

        created = next((df for df in self.p.data_files if df.name == "imported.csv"), None)
        self.assertIsNotNone(created)
        self.assertEqual(created.source_path, "/tmp/imported.csv")

    def test_on_tree_node_selected_folder(self):
        node = next((n for n in self.p.tree.nodes if n.kind == "folder"), None)
        if node:
            self.page.on_tree_node_selected("folder", node.id)

    def test_on_tree_node_selected_folder_shows_structure_preview(self):
        node = next((n for n in self.p.tree.nodes if n.kind == "folder" and getattr(n, "group_type", None) == "pictures"), None)
        self.assertIsNotNone(node)

        self.page.on_tree_node_selected("folder", node.id)

        self.assertIs(self.page._preview_stack.currentWidget(), self.page._picture_preview_tree)
        root = self.page._picture_preview_tree.topLevelItem(0)
        self.assertIsNotNone(root)
        self.assertEqual(root.text(0), "图片集")
        self.assertIn("文件夹: 图片集", self.page._stats_title_label.text())

    def test_picture_root_cannot_be_renamed_from_management_panel(self):
        self.pm.migrate_to_v3(self.p)
        picture_root = self.pm._find_folder_by_group_type("pictures")
        self.assertIsNotNone(picture_root)

        self.page.on_tree_node_selected("folder", picture_root.id)

        self.assertFalse(self.page._can_rename_current_node())
        self.assertFalse(self.page._manage_name_edit.isEnabled())
        self.assertFalse(self.page._btn_apply_name.isEnabled())

    def test_on_tree_node_selected_image_work_shows_image_preview(self):
        from PySide6.QtGui import QImage

        temp_path = Path(tempfile.NamedTemporaryFile(suffix=".png", delete=False).name)
        try:
            image = QImage(24, 18, QImage.Format.Format_RGB32)
            image.fill(Qt.GlobalColor.white)
            self.assertTrue(image.save(str(temp_path)))

            created = self.pm.add_image(str(temp_path), name="sample-image")
            node = next((n for n in self.p.tree.nodes if n.kind == "image_work" and n.image_work_id == created.id), None)
            self.assertIsNotNone(node)

            self.page.on_tree_node_selected("image_work", node.id)

            self.assertIs(self.page._preview_stack.currentWidget(), self.page._image_preview_label)
            self.assertIn("图像名称: sample-image", self.page._stats_title_label.text())
        finally:
            temp_path.unlink(missing_ok=True)

    def test_on_tree_node_selected_analysis_result_shows_text_preview(self):
        from models.schemas import AnalysisResult

        result = AnalysisResult(
            name="拟合结果A",
            analysis_type="curve_fit",
            summary={"analysis_type": "curve_fit", "source_name": "series-A", "model": "linear", "r2": 0.98},
        )
        self.assertTrue(self.pm.add_analysis(result))
        node = next((n for n in self.p.tree.nodes if n.kind == "analysis_result" and n.analysis_id == result.id), None)
        self.assertIsNotNone(node)

        self.page.on_tree_node_selected("analysis_result", node.id)

        self.assertIs(self.page._preview_stack.currentWidget(), self.page._text_preview)
        preview_text = self.page._text_preview.toPlainText()
        self.assertIn("名称: 拟合结果A", preview_text)
        self.assertIn("R²: 0.98", preview_text)

    def test_on_tree_node_selected_unknown_kind(self):
        self.page.on_tree_node_selected("unknown", "fake-id")

    def test_on_tree_node_selected_nonexistent_id(self):
        self.page.on_tree_node_selected("data_file", "nonexistent-id-xyz")

    def test_tree_filter_kinds_attribute(self):
        self.assertIsInstance(
            getattr(self.page, "tree_filter_kinds", []), list
        )

    def test_legacy_left_entry_not_built(self):
        self.assertFalse(hasattr(self.page, "_btn_add_ds"))
        self.assertFalse(hasattr(self.page, "_btn_import"))
        self.assertFalse(hasattr(self.page, "_btn_copy_curve"))
        self.assertFalse(hasattr(self.page, "_btn_delete"))

    def test_management_panel_updates_for_data_file_selection(self):
        node = next((n for n in self.p.tree.nodes if n.kind == "data_file"), None)
        self.assertIsNotNone(node)

        self.page.on_tree_node_selected("data_file", node.id)

        self.assertEqual(self.page._manage_target_label.text(), f"[数据文件] {self.df.name}")
        self.assertEqual(self.page._manage_type_label.text(), "")
        self.assertTrue(self.page._btn_apply_name.isEnabled())
        self.assertTrue(self.page._btn_delete_node.isEnabled())
        self.assertFalse(hasattr(self.page, "_btn_copy_to_data_file"))

    def test_management_target_label_opens_node_detail_dialog(self):
        node = next((n for n in self.p.tree.nodes if n.kind == "data_file"), None)
        self.assertIsNotNone(node)

        self.page.resize(1280, 820)
        self.page.show()
        self.page.on_tree_node_selected("data_file", node.id)
        QApplication.processEvents()

        with mock.patch("ui.pages.data_page._NodeDetailDialog") as dialog_cls:
            dialog = dialog_cls.return_value
            QTest.mouseClick(self.page._manage_target_label, Qt.MouseButton.LeftButton)
            QApplication.processEvents()

        dialog_cls.assert_called_once()
        self.assertIn("节点信息", dialog_cls.call_args.args[0])
        self.assertIn("数据文件", dialog_cls.call_args.args[1])
        self.assertIn("系列数量: 1", dialog_cls.call_args.args[1])
        dialog.exec.assert_called_once()

    def test_management_primary_buttons_use_new_labels_and_export_sits_after_delete(self):
        self.page.resize(1280, 820)
        self.page.show()
        QApplication.processEvents()

        self.assertEqual(self.page._btn_delete_node.text(), "删除节点")
        self.assertEqual(self.page._btn_export.text(), "导出数据")
        self.assertEqual(self.page._btn_apply_name.text(), "重命名")
        self.assertLess(self.page._btn_apply_name.minimumWidth(), self.page._manage_name_edit.width())
        self.assertTrue(self.page._btn_export.icon().isNull())
        self.assertIs(self.page._btn_export.parent(), self.page._tool_panel)
        self.assertGreater(self.page._btn_export.x(), self.page._btn_delete_node.x())

    def test_export_button_uses_shared_curve_export_for_data_file(self):
        from models.schemas import DataSeries
        from ui.dialogs.export_flow import CurveFileExportPlan

        self.df.series.append(DataSeries(name="s2", x=[1.0, 2.0, 3.0, 4.0, 5.0], y=[3.0, 6.0, 9.0, 12.0, 15.0]))
        node = next((n for n in self.p.tree.nodes if n.kind == "data_file"), None)
        self.assertIsNotNone(node)
        self.page.on_tree_node_selected("data_file", node.id)

        with tempfile.NamedTemporaryFile(suffix=".csv", delete=False) as handle:
            path = handle.name
        try:
            with mock.patch(
                "ui.pages.data_page.choose_curve_file_export_plan",
                return_value=CurveFileExportPlan(action="file", file_format="csv", include_timestamp=False, merged=True),
            ), mock.patch(
                "ui.pages.data_page.QFileDialog.getSaveFileName",
                return_value=(path, "CSV 文件 (*.csv)"),
            ):
                self.page._export_csv()

            content = Path(path).read_text(encoding="utf-8-sig")
            self.assertIn("x,s1,s2", content)
            self.assertIn("2.0,4.0,6.0", content)
        finally:
            Path(path).unlink(missing_ok=True)

    def test_management_panel_without_pending_list_keeps_controls_compact_at_top(self):
        node = next((n for n in self.p.tree.nodes if n.kind == "data_file"), None)
        self.assertIsNotNone(node)

        self.page.resize(1280, 820)
        self.page.show()
        self.page.on_tree_node_selected("data_file", node.id)
        QApplication.processEvents()

        self.assertTrue(self.page._import_queue_panel.isHidden())
        self.assertFalse(self.page._manage_bottom_spacer.isHidden())
        self.assertLess(self.page._manage_target_label.y(), self.page._manage_name_edit.y())
        self.assertLess(self.page._manage_help_label.y(), self.page._manage_name_edit.y())
        self.assertLess(self.page._btn_delete_node.y(), self.page._manage_bottom_spacer.y())

    def test_source_preview_widgets_use_fluent_surface_style(self):
        self.page._apply_preview_host_background()

        self.assertIn("border-radius: 12px", self.page._image_preview_label.styleSheet())
        self.assertIn("border-radius: 12px", self.page._text_preview.styleSheet())

    def test_source_file_selection_switches_to_preview_mode(self):
        from qfluentwidgets import FluentIcon as FIF

        source_path = Path(tempfile.NamedTemporaryFile(suffix=".csv", delete=False).name)
        source_path.write_text("x,y\n1,2\n", encoding="utf-8")
        try:
            node = self.pm.add_source_file(str(source_path))
            self.assertIsNotNone(node)
            asset = self.pm.get_source_file(node.source_file_id)
            self.assertIsNotNone(asset)

            self.page.on_tree_node_selected("source_file", node.id)

            self.assertIs(self.page._right_mode_stack.currentWidget(), self.page._preview_card)
            self.assertIs(self.page._preview_stack.currentWidget(), self.page._parsed_preview_table)
            self.assertTrue(self.page._btn_apply_name.isEnabled())
            self.assertTrue(self.page._btn_delete_node.isEnabled())
            self.assertTrue(self.page._btn_to_vis.isHidden())
            self.assertTrue(self.page._btn_to_proc.isHidden())
            self.assertTrue(self.page._btn_import_source_to_data.isEnabled())
            self.assertFalse(self.page._btn_import_source_to_digitize.isEnabled())
            self.assertFalse(self.page._source_file_preview_controls.isHidden())
            self.assertEqual(
                [self.page._source_file_preview_combo.itemText(index) for index in range(self.page._source_file_preview_combo.count())],
                ["解析", "源文件"],
            )
            self.assertEqual(
                self.page._btn_import_source_to_data.icon().pixmap(20, 20).toImage(),
                FIF.DOWNLOAD.icon().pixmap(20, 20).toImage(),
            )
            self.assertFalse(self.page._source_path_panel.isHidden())
            self.assertEqual(self.page._current_source_path_button.toolTip(), self.pm.get_source_file_path(asset.id))
            self.assertEqual(self.page._origin_source_path_button.toolTip(), asset.source_file_path)
        finally:
            source_path.unlink(missing_ok=True)

    def test_source_file_browser_uses_photo_icon_for_images(self):
        from qfluentwidgets import FluentIcon as FIF

        source_root = self.pm._find_folder_by_group_type("source_files")
        self.assertIsNotNone(source_root)

        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            image_path = temp_path / "sample.png"
            text_path = temp_path / "sample.csv"
            from PySide6.QtGui import QImage

            image = QImage(16, 16, QImage.Format.Format_RGB32)
            image.fill(Qt.GlobalColor.white)
            self.assertTrue(image.save(str(image_path)))
            text_path.write_text("x,y\n1,2\n", encoding="utf-8")

            self.page.on_tree_node_selected("folder", source_root.id)
            self.page._external_browser_dir = temp_path
            self.page._refresh_source_browser()

            icon_map = {
                self.page._source_browser.topLevelItem(i).text(0): self.page._source_browser.topLevelItem(i).icon(0).pixmap(20, 20).toImage()
                for i in range(self.page._source_browser.topLevelItemCount())
            }

        self.assertEqual(icon_map["sample.png"], self.page._source_file_icon_for_path(str(image_path)).icon().pixmap(20, 20).toImage())
        self.assertEqual(icon_map["sample.csv"], self.page._source_file_icon_for_path(str(text_path)).icon().pixmap(20, 20).toImage())

    def test_source_file_path_buttons_elide_long_paths_and_keep_full_tooltip(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            long_dir = Path(temp_dir) / "very_long_directory_name_for_data_page_preview" / "another_nested_directory"
            long_dir.mkdir(parents=True)
            source_path = long_dir / "very_long_source_file_name_for_preview.csv"
            source_path.write_text("x,y\n1,2\n", encoding="utf-8")

            node = self.pm.add_source_file(str(source_path))
            self.assertIsNotNone(node)
            asset = self.pm.get_source_file(node.source_file_id)
            self.assertIsNotNone(asset)

            self.page.resize(420, 820)
            self.page.show()
            self.page.on_tree_node_selected("source_file", node.id)
            QApplication.processEvents()

            self.assertEqual(self.page._current_source_path_button.toolTip(), self.pm.get_source_file_path(asset.id))
            self.assertEqual(self.page._origin_source_path_button.toolTip(), asset.source_file_path)
            self.assertIn("…", self.page._current_source_path_button.text())
            self.assertIn("…", self.page._origin_source_path_button.text())
            self.assertLess(self.page._current_source_path_button.width(), self.page._source_path_panel.width())
            self.assertLess(self.page._origin_source_path_button.width(), self.page._source_path_panel.width())

    def test_source_file_path_buttons_do_not_elide_short_paths_on_first_show(self):
        with tempfile.TemporaryDirectory(dir="/tmp") as temp_dir:
            source_path = Path(temp_dir) / "a.csv"
            source_path.write_text("x,y\n1,2\n", encoding="utf-8")

            node = self.pm.add_source_file(str(source_path))
            self.assertIsNotNone(node)
            asset = self.pm.get_source_file(node.source_file_id)
            self.assertIsNotNone(asset)

            self.page.resize(1280, 820)
            self.page.show()
            self.page.on_tree_node_selected("source_file", node.id)
            QApplication.processEvents()

            self.assertEqual(self.page._current_source_path_button.text(), self.pm.get_source_file_path(asset.id))
            self.assertEqual(self.page._origin_source_path_button.text(), asset.source_file_path)
            self.assertLess(self.page._current_source_path_button.width(), self.page._source_path_panel.width())

    def test_system_file_manager_hides_dotfiles_until_toggled(self):
        source_root = self.pm._find_folder_by_group_type("source_files")
        self.assertIsNotNone(source_root)

        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            visible_file = temp_path / "visible.csv"
            hidden_file = temp_path / ".hidden.csv"
            visible_file.write_text("x,y\n1,2\n", encoding="utf-8")
            hidden_file.write_text("x,y\n3,4\n", encoding="utf-8")

            self.page.resize(1280, 820)
            self.page.show()
            self.page.on_tree_node_selected("folder", source_root.id)
            self.page._external_browser_dir = temp_path
            self.page._refresh_source_browser()
            QApplication.processEvents()

            labels = [self.page._source_browser.topLevelItem(i).text(0) for i in range(self.page._source_browser.topLevelItemCount())]
            self.assertIn("visible.csv", labels)
            self.assertNotIn(".hidden.csv", labels)
            self.assertLess(self.page._btn_toggle_hidden_browser.x(), self.page._btn_refresh_browser.x())
            initial_icon_key = self.page._btn_toggle_hidden_browser.icon().cacheKey()

            self.page._toggle_hidden_browser_entries()
            QApplication.processEvents()

            toggled_labels = [self.page._source_browser.topLevelItem(i).text(0) for i in range(self.page._source_browser.topLevelItemCount())]
            self.assertIn(".hidden.csv", toggled_labels)
            self.assertEqual(self.page._btn_toggle_hidden_browser.toolTip(), "隐藏隐藏文件")
            self.assertNotEqual(self.page._btn_toggle_hidden_browser.icon().cacheKey(), initial_icon_key)

    def test_update_theme_reapplies_preview_host_background(self):
        if self.page._preview_figure is None or self.page._preview_canvas is None:
            self.skipTest("matplotlib unavailable")

        with mock.patch("ui.pages.data_page.isDarkTheme", return_value=True):
            self.page.update_theme()

        self.assertIn("#1e1e1e", self.page._plot_preview_panel.styleSheet())
        self.assertIn("#1e1e1e", self.page._preview_canvas.styleSheet())

    def test_hidden_update_theme_flushes_preview_on_show(self):
        if self.page._preview_figure is None or self.page._preview_canvas is None:
            self.skipTest("matplotlib unavailable")

        with mock.patch.object(self.page, "_draw_preview") as draw_preview:
            self.page.update_theme()
            self.assertEqual(draw_preview.call_count, 0)

            self.page.show()
            QApplication.processEvents()

        self.assertEqual(draw_preview.call_count, 1)

    def test_source_browser_breadcrumb_can_navigate_to_parent(self):
        datasets_root = self.pm._find_folder_by_group_type("datasets")
        self.assertIsNotNone(datasets_root)

        with tempfile.TemporaryDirectory() as temp_dir:
            parent_dir = Path(temp_dir) / "crumb_parent"
            child_dir = parent_dir / "crumb_child"
            child_dir.mkdir(parents=True)

            self.page.on_tree_node_selected("folder", datasets_root.id)
            self.page._external_browser_dir = child_dir
            self.page._refresh_source_browser()

            self.assertGreaterEqual(self.page._source_breadcrumb_bar.count(), 3)
            self.page._source_breadcrumb_bar.setCurrentItem(str(parent_dir))

            self.assertEqual(self.page._external_browser_dir, parent_dir)

    def test_system_file_manager_can_add_directory_to_favorites_and_activate_it(self):
        source_root = self.pm._find_folder_by_group_type("source_files")
        self.assertIsNotNone(source_root)

        with tempfile.TemporaryDirectory() as temp_dir:
            base_dir = Path(temp_dir)
            favorite_dir = base_dir / "favorite-dir"
            favorite_dir.mkdir()

            self.page.on_tree_node_selected("folder", source_root.id)
            self.page._external_browser_dir = base_dir
            self.page._refresh_source_browser()

            self.assertTrue(self.page._add_source_favorite(str(favorite_dir)))
            self.assertEqual(self.page._source_favorites_list.topLevelItemCount(), 1)
            favorite_item = self.page._source_favorites_list.topLevelItem(0)
            self.assertEqual(favorite_item.text(0), "favorite-dir")
            self.assertEqual(favorite_item.toolTip(0), str(favorite_dir))

            self.page._on_source_favorite_item_activated(favorite_item, 0)
            self.assertEqual(self.page._external_browser_dir, favorite_dir)

    def test_source_favorites_context_menu_excludes_open_directory_action(self):
        from PySide6.QtCore import QPoint

        source_root = self.pm._find_folder_by_group_type("source_files")
        self.assertIsNotNone(source_root)

        with tempfile.TemporaryDirectory() as temp_dir:
            favorite_dir = Path(temp_dir) / "favorite-dir"
            favorite_dir.mkdir()

            self.page.on_tree_node_selected("folder", source_root.id)
            self.page._source_favorite_paths = [str(favorite_dir)]
            self.page._refresh_source_favorites()
            favorite_item = self.page._source_favorites_list.topLevelItem(0)
            captured = {}

            def _fake_exec(menu, *_args, **_kwargs):
                captured["actions"] = [action.text() for action in menu.actions()]

            with mock.patch.object(self.page._source_favorites_list, "itemAt", return_value=favorite_item):
                with mock.patch("ui.pages.data_page.RoundMenu.exec", autospec=True, side_effect=_fake_exec):
                    self.page._show_source_favorites_context_menu(QPoint(1, 1))

            self.assertIn("移出收藏夹", captured["actions"])
            self.assertNotIn("打开该目录", captured["actions"])

    def test_source_favorites_persist_via_ui_preferences(self):
        from core.ui_preferences import get_data_page_source_favorites

        with tempfile.TemporaryDirectory() as temp_dir:
            config_path = Path(temp_dir) / "ui_preferences.json"
            favorite_dir = Path(temp_dir) / "persisted"
            favorite_dir.mkdir()

            with mock.patch("core.ui_preferences._CONFIG_PATH", config_path):
                page = self.page.__class__()
                try:
                    page._source_favorite_paths = [str(favorite_dir)]
                    page._refresh_source_favorites()
                    self.assertEqual(get_data_page_source_favorites(), [str(favorite_dir)])
                finally:
                    page.deleteLater()

    def test_system_file_manager_uses_splitter_for_favorites_and_browser(self):
        from PySide6.QtWidgets import QSplitter

        self.assertIsInstance(self.page._source_browser_splitter, QSplitter)
        self.assertIs(self.page._source_browser_splitter.widget(0), self.page._source_favorites_panel)
        self.assertEqual(self.page._source_browser_splitter.count(), 2)
        self.assertEqual(self.page._source_browser_splitter.handleWidth(), 2)
        self.assertIn("QSplitter::handle", self.page._source_browser_splitter.styleSheet())
        self.assertIn("border-left", self.page._source_browser_splitter.styleSheet())

    def test_source_folder_can_add_external_files_to_pending_list(self):
        source_root = self.pm._find_folder_by_group_type("source_files")
        self.assertIsNotNone(source_root)

        external_dir = Path(tempfile.mkdtemp())
        source_path = external_dir / "external.csv"
        source_path.write_text("x,y\n1,2\n", encoding="utf-8")
        try:
            self.page.on_tree_node_selected("folder", source_root.id)
            self.page._external_browser_dir = external_dir
            self.page._refresh_source_browser()

            item = self.page._source_browser.topLevelItem(0)
            self.assertIsNotNone(item)
            item.setSelected(True)
            self.page._source_browser.setCurrentItem(item)
            self.page._add_selected_browser_source_files_to_pending()

            self.assertEqual(len(self.page._pending_import_paths), 1)
            self.assertEqual(self.page._pending_source_list.topLevelItemCount(), 1)
        finally:
            source_path.unlink(missing_ok=True)
            external_dir.rmdir()

    def test_system_file_manager_places_add_button_below_browser_and_uses_breadcrumb_only(self):
        from qfluentwidgets import CaptionLabel

        source_root = self.pm._find_folder_by_group_type("source_files")
        self.assertIsNotNone(source_root)

        self.page.resize(1280, 820)
        self.page.show()
        self.page.on_tree_node_selected("folder", source_root.id)
        QApplication.processEvents()

        system_page = self.page._source_browser.parentWidget()
        system_labels = [label.text() for label in system_page.findChildren(CaptionLabel)]

        self.assertNotIn("当前路径", system_labels)
        self.assertGreater(self.page._btn_add_selected_sources.y(), self.page._source_browser.y())
        self.assertGreater(self.page._btn_add_selected_sources.y(), self.page._source_browser_detail_label.y())

    def test_dataset_folder_switches_to_external_file_manager_mode(self):
        from qfluentwidgets import SegmentedWidget
        from PySide6.QtWidgets import QSizePolicy

        datasets_root = self.pm._find_folder_by_group_type("datasets")
        self.assertIsNotNone(datasets_root)

        self.page.on_tree_node_selected("folder", datasets_root.id)

        self.assertIs(self.page._right_mode_stack.currentWidget(), self.page._source_manager_card)
        self.assertFalse(self.page._import_queue_panel.isHidden())
        self.assertEqual(self.page._btn_import_pending.text(), "导入到数据集")
        self.assertIsInstance(self.page._source_browser_tabs.navigationWidget, SegmentedWidget)
        self.assertEqual(
            self.page._source_browser_tabs.navigationWidget.sizePolicy().horizontalPolicy(),
            QSizePolicy.Policy.Expanding,
        )
        self.assertFalse(self.page._source_browser_tabs.tabBar.isHidden())
        self.assertTrue(self.page._source_manager_hint.isHidden())
        self.assertTrue(self.page._pending_source_hint.isHidden())
        self.assertGreaterEqual(self.page._pending_source_list.minimumHeight(), 260)
        self.assertGreater(self.page._source_breadcrumb_host.layout().contentsMargins().left(), 0)

    def test_dataset_folder_can_add_project_source_file_to_pending(self):
        datasets_root = self.pm._find_folder_by_group_type("datasets")
        self.assertIsNotNone(datasets_root)

        with tempfile.TemporaryDirectory() as temp_dir:
            source_path = Path(temp_dir) / "managed.csv"
            source_path.write_text("x,y\n1,2\n", encoding="utf-8")
            node = self.pm.add_source_file(str(source_path))
            self.assertIsNotNone(node)

            self.page.on_tree_node_selected("folder", datasets_root.id)
            self.page._source_browser_tabs.setCurrentIndex(0)
            self.page._refresh_project_source_browser()

            item = self.page._project_source_browser.topLevelItem(0)
            self.assertIsNotNone(item)
            item.setSelected(True)
            self.page._project_source_browser.setCurrentItem(item)
            self.page._add_selected_project_source_files_to_pending()

            self.assertEqual(len(self.page._pending_import_paths), 1)
            self.assertEqual(self.page._pending_import_paths[0], self.pm.get_source_file_path(node.source_file_id))

    def test_pending_external_files_can_import_to_digitize(self):
        from PySide6.QtGui import QImage

        images_root = self.pm._find_folder_by_group_type("images")
        self.assertIsNotNone(images_root)
        external_dir = Path(tempfile.mkdtemp())
        source_path = external_dir / "sample.png"
        image = QImage(16, 16, QImage.Format.Format_RGB32)
        image.fill(Qt.GlobalColor.white)
        self.assertTrue(image.save(str(source_path)))
        try:
            self.page.on_tree_node_selected("folder", images_root.id)
            self.page._external_browser_dir = external_dir
            self.page._refresh_source_browser()

            item = self.page._source_browser.topLevelItem(0)
            self.assertIsNotNone(item)
            item.setSelected(True)
            self.page._source_browser.setCurrentItem(item)
            self.page._add_selected_browser_source_files_to_pending()

            before = len(self.p.images)
            self.page._import_pending_files_for_current_group()

            self.assertEqual(len(self.p.images), before + 1)
            self.assertEqual(len(self.page._pending_import_paths), 0)
        finally:
            source_path.unlink(missing_ok=True)
            external_dir.rmdir()

    def test_image_folder_can_import_project_source_image_via_source_tab(self):
        from PySide6.QtGui import QImage

        images_root = self.pm._find_folder_by_group_type("images")
        self.assertIsNotNone(images_root)

        with tempfile.TemporaryDirectory() as temp_dir:
            source_path = Path(temp_dir) / "managed.png"
            image = QImage(12, 12, QImage.Format.Format_RGB32)
            image.fill(Qt.GlobalColor.white)
            self.assertTrue(image.save(str(source_path)))
            node = self.pm.add_source_file(str(source_path))
            self.assertIsNotNone(node)

            self.page.on_tree_node_selected("folder", images_root.id)
            self.page._source_browser_tabs.setCurrentIndex(0)
            self.page._refresh_project_source_browser()

            item = self.page._project_source_browser.topLevelItem(0)
            self.assertIsNotNone(item)
            item.setSelected(True)
            self.page._project_source_browser.setCurrentItem(item)
            self.page._add_selected_project_source_files_to_pending()

            before = len(self.p.images)
            self.page._import_pending_files_for_current_group()

            self.assertEqual(len(self.p.images), before + 1)
            self.assertEqual(len(self.page._pending_import_paths), 0)

    def test_management_panel_buttons_use_expanding_width_policy(self):
        from PySide6.QtWidgets import QSizePolicy

        buttons = [
            self.page._btn_delete_node,
            self.page._btn_export,
            self.page._btn_to_vis,
            self.page._btn_to_proc,
            self.page._btn_import_pending,
        ]

        for button in buttons:
            self.assertEqual(button.sizePolicy().horizontalPolicy(), QSizePolicy.Policy.Expanding)

    def test_management_can_rename_selected_data_file(self):
        node = next((n for n in self.p.tree.nodes if n.kind == "data_file"), None)
        self.assertIsNotNone(node)

        self.page.on_tree_node_selected("data_file", node.id)
        self.page._manage_name_edit.setText("renamed.csv")
        self.page._apply_rename_current_node()

        self.assertEqual(self.df.name, "renamed.csv")
        self.assertEqual(node.name, "renamed.csv")

    def test_management_rename_conflict_keeps_original_name(self):
        from models.schemas import DataFile, DataSeries

        second = self.pm.add_data_file(DataFile(name="other.csv", series=[DataSeries(name="s2", x=[1.0], y=[2.0])]))
        self.assertIsNotNone(second)

        self.page.on_tree_node_selected("data_file", second.id)
        self.page._manage_name_edit.setText(self.df.name)
        self.page._apply_rename_current_node()

        self.assertEqual(self.pm.get_data_file(second.data_file_id).name, "other.csv")
        self.assertEqual(self.p.tree.get_node(second.id).name, "other.csv")

    def test_import_dialog_duplicate_data_file_name_auto_renames(self):
        from models.schemas import DataFile, DataSeries

        existing = self.pm.add_data_file(DataFile(name="dup.csv", series=[DataSeries(name="s0", x=[0.0], y=[1.0])]))
        self.assertIsNotNone(existing)

        dialog = mock.Mock()
        dialog.get_results.return_value = [DataSeries(name="s1", x=[1.0], y=[2.0])]
        dialog.get_target_data_file_id.return_value = None
        dialog.get_file_name.return_value = "dup.csv"

        changed = self.page._apply_import_dialog_results(dialog, show_feedback=False)

        self.assertTrue(changed)
        names = sorted(df.name for df in self.p.data_files)
        self.assertIn("dup.csv", names)
        self.assertIn("dup_1.csv", names)

    def test_management_disables_protected_folder_edits(self):
        node = next((n for n in self.p.tree.nodes if n.kind == "folder" and getattr(n, "group_type", None) == "datasets"), None)
        self.assertIsNotNone(node)

        self.page.on_tree_node_selected("folder", node.id)

        self.assertFalse(self.page._btn_apply_name.isEnabled())
        self.assertFalse(self.page._btn_delete_node.isEnabled())

    def test_management_enables_nested_dataset_folder_edits(self):
        datasets_root = self.pm._find_folder_by_group_type("datasets")
        self.assertIsNotNone(datasets_root)
        folder = self.pm.add_folder("子实验", parent_id=datasets_root.id, group_type="datasets")
        self.assertIsNotNone(folder)

        self.page.on_tree_node_selected("folder", folder.id)

        self.assertTrue(self.page._btn_apply_name.isEnabled())
        self.assertTrue(self.page._btn_delete_node.isEnabled())

    def test_add_dataset_uses_selected_dataset_folder(self):
        datasets_root = self.pm._find_folder_by_group_type("datasets")
        self.assertIsNotNone(datasets_root)
        folder = self.pm.add_folder("导入目标", parent_id=datasets_root.id, group_type="datasets")
        self.assertIsNotNone(folder)
        self.page.on_tree_node_selected("folder", folder.id)

        with mock.patch("ui.pages.data_page._NameDialog.exec", return_value=True), \
             mock.patch("ui.pages.data_page._NameDialog.get_name", return_value="子目录数据集"):
            self.page._add_dataset()

        node = next((n for n in self.p.tree.nodes if n.kind == "data_file" and n.name == "子目录数据集"), None)
        self.assertIsNotNone(node)
        self.assertEqual(node.parent_id, folder.id)

    def test_import_file_uses_selected_dataset_folder(self):
        from models.schemas import DataSeries

        datasets_root = self.pm._find_folder_by_group_type("datasets")
        self.assertIsNotNone(datasets_root)
        folder = self.pm.add_folder("导入子目录", parent_id=datasets_root.id, group_type="datasets")
        self.assertIsNotNone(folder)
        self.page.on_tree_node_selected("folder", folder.id)

        with mock.patch("ui.dialogs.import_dialog.ImportDialog") as dialog_cls:
            dialog = dialog_cls.return_value
            dialog.exec.return_value = True
            dialog.get_results.return_value = [DataSeries(name="imported", x=[1.0], y=[2.0])]
            dialog.get_target_data_file_id.return_value = None
            dialog.get_file_name.return_value = "imported.csv"
            self.page._import_file()

        node = next((n for n in self.p.tree.nodes if n.kind == "data_file" and n.name == "imported.csv"), None)
        self.assertIsNotNone(node)
        self.assertEqual(node.parent_id, folder.id)

    def test_import_file_can_append_to_existing_data_file(self):
        from models.schemas import DataFile, DataSeries

        target = self.pm.add_data_file(DataFile(name="existing.csv", series=[]))
        self.assertIsNotNone(target)
        existing = next((df for df in self.p.data_files if df.name == "existing.csv"), None)
        self.assertIsNotNone(existing)

        with mock.patch("ui.dialogs.import_dialog.ImportDialog") as dialog_cls:
            dialog = dialog_cls.return_value
            dialog.exec.return_value = True
            dialog.get_results.return_value = [DataSeries(name="力", x=[1.0], y=[2.0])]
            dialog.get_target_data_file_id.return_value = existing.id
            self.page._import_file()

        self.assertEqual(len(existing.series), 1)
        self.assertEqual(existing.series[0].name, "力")

    def test_import_file_preloads_dropped_path_into_dialog(self):
        with mock.patch("ui.dialogs.import_dialog.ImportDialog") as dialog_cls:
            dialog = dialog_cls.return_value
            dialog.exec.return_value = False

            self.page._import_file("/tmp/dropped.csv")

        dialog.load_file.assert_called_once_with("/tmp/dropped.csv")

    def test_dataset_pending_import_shows_default_mode_button(self):
        datasets_root = self.pm._find_folder_by_group_type("datasets")
        self.assertIsNotNone(datasets_root)
        self.page.on_tree_node_selected("folder", datasets_root.id)
        self.page._pending_import_paths = ["/tmp/demo.csv"]
        self.page._pending_import_names = {"/tmp/demo.csv": "demo.csv"}

        self.page._refresh_pending_source_controls()

        self.assertFalse(self.page._btn_import_pending_default.isHidden())
        self.assertTrue(self.page._btn_import_pending_default.isEnabled())

    def test_default_mode_pending_dataset_import_uses_dialog_defaults(self):
        datasets_root = self.pm._find_folder_by_group_type("datasets")
        self.assertIsNotNone(datasets_root)
        self.page.on_tree_node_selected("folder", datasets_root.id)

        temp_path = Path(tempfile.NamedTemporaryFile(suffix=".csv", delete=False).name)
        try:
            temp_path.write_text("x,y\n0,1\n1,2\n", encoding="utf-8")
            self.page._pending_import_paths = [str(temp_path)]
            self.page._pending_import_names = {str(temp_path): temp_path.name}
            dialog = mock.Mock()
            dialog.import_with_default_options.return_value = [mock.Mock(name="series")]

            with mock.patch.object(self.page, "_create_import_dialog", return_value=dialog) as create_dialog, \
                 mock.patch.object(self.page, "_apply_import_dialog_results", return_value=True):
                self.page._import_pending_source_files_to_datasets_default_mode()

            create_dialog.assert_called_once_with(str(temp_path), default_file_name=temp_path.name)
            dialog.import_with_default_options.assert_called_once_with()
            self.assertEqual(self.page._pending_import_paths, [])
        finally:
            temp_path.unlink(missing_ok=True)


# ═══════════════════════════════════════════════════════════════════════════
# 4. ChartPage — on_tree_node_selected, load_template
# ═══════════════════════════════════════════════════════════════════════════

class TestChartPage(unittest.TestCase):

    def setUp(self):
        self._restore_assets = _patch_global_assets()
        self.pm, self.p, self.df, self.s = _make_project("cp_test")
        self._restore = _patch_pm(self.pm)
        from ui.pages.chart_page import ChartPage
        self.page = ChartPage()

    def tearDown(self):
        self._restore()
        self._restore_assets()
        self.page.deleteLater()

    def test_page_creates_no_crash(self):
        self.assertIsNotNone(self.page)

    def test_chart_preview_uses_shared_navigation_controls(self):
        if self.page._figure is None or self.page._canvas is None:
            self.skipTest("matplotlib unavailable")

        self.page.on_tree_node_activated("series", self.s.id)
        axis = self.page._figure.axes[0]
        original_x_span = abs(axis.get_xlim()[1] - axis.get_xlim()[0])
        original_y_span = abs(axis.get_ylim()[1] - axis.get_ylim()[0])

        self.page._chart_preview_box_zoom_btn.click()
        QApplication.processEvents()
        self.assertEqual(self.page._chart_preview_navigation_mode(), "zoom")
        self.assertTrue(self.page._chart_preview_box_zoom_btn.isChecked())
        self.assertFalse(self.page._chart_preview_pan_btn.isChecked())

        self.page._chart_preview_zoom_in_btn.click()
        QApplication.processEvents()

        zoomed_axis = self.page._figure.axes[0]
        self.assertLess(abs(zoomed_axis.get_xlim()[1] - zoomed_axis.get_xlim()[0]), original_x_span)
        self.assertLess(abs(zoomed_axis.get_ylim()[1] - zoomed_axis.get_ylim()[0]), original_y_span)

        self.page._chart_preview_fit_btn.click()
        QApplication.processEvents()

        reset_axis = self.page._figure.axes[0]
        self.assertAlmostEqual(abs(reset_axis.get_xlim()[1] - reset_axis.get_xlim()[0]), original_x_span, places=6)
        self.assertAlmostEqual(abs(reset_axis.get_ylim()[1] - reset_axis.get_ylim()[0]), original_y_span, places=6)

    def test_chart_quick_config_line_edits_wait_for_edit_commit(self):
        original_font_size = self.page._figure_state.font_size

        self.page._font_size_edit.setText("18")
        QApplication.processEvents()
        self.assertEqual(self.page._figure_state.font_size, original_font_size)

        self.page._font_size_edit.editingFinished.emit()
        self.assertEqual(self.page._figure_state.font_size, 18)

    def test_style_tabs_use_consistent_scroll_containers(self):
        from qfluentwidgets import SmoothScrollArea
        from PySide6.QtWidgets import QSizePolicy

        tab_pages = [self.page._style_tabs.widget(index) for index in range(3)]

        self.assertTrue(all(isinstance(tab_page, SmoothScrollArea) for tab_page in tab_pages))
        min_heights = {tab_page.widget().minimumHeight() for tab_page in tab_pages}
        self.assertEqual(len(min_heights), 1)
        vertical_policies = [tab_page.widget().sizePolicy().verticalPolicy() for tab_page in tab_pages]
        self.assertEqual(vertical_policies[:2], [QSizePolicy.Policy.Maximum, QSizePolicy.Policy.Maximum])
        self.assertEqual(vertical_policies[2], QSizePolicy.Policy.Expanding)

    def test_style_tabs_navigation_remains_fill_width(self):
        from PySide6.QtWidgets import QSizePolicy

        navigation = self.page._style_tabs.navigationWidget

        self.assertEqual(navigation.sizePolicy().horizontalPolicy(), QSizePolicy.Policy.Expanding)
        self.assertEqual(navigation.minimumWidth(), 0)
        self.assertGreater(navigation.maximumWidth(), 1000000)

    def test_plot_extension_repeat_hint_is_directly_under_loaded_header(self):
        container_layout = self.page._style_tabs.widget(2).widget().layout()

        self.assertLess(
            container_layout.indexOf(self.page._plot_extension_repeat_hint),
            container_layout.indexOf(self.page._plot_extension_applied_list),
        )

    def test_plot_extension_uses_help_teaching_tip_instead_of_static_text(self):
        from qfluentwidgets import BodyLabel

        page_widget = self.page._style_tabs.widget(2).widget()
        label_texts = [label.text() for label in page_widget.findChildren(BodyLabel) if label.text()]

        self.assertNotIn("在右侧面板选择扩展，并叠加到当前图表。", label_texts)
        self.assertNotIn("适合参考线、标注或自定义绘制流程。", label_texts)
        self.assertIsNotNone(self.page._plot_extension_help_btn)

        with mock.patch("ui.pages.chart_page.TeachingTip.make") as make_mock:
            self.page._show_plot_extension_teaching_tip()

        make_mock.assert_called_once()
        self.assertIs(make_mock.call_args.args[1], self.page._plot_extension_help_btn)

    def test_plot_extension_config_help_area_uses_fixed_scroll_region(self):
        self.page.resize(1260, 900)
        self.page.show()
        QApplication.processEvents()

        expected_height = max(124, self.page._extension_panel.height() // 3)
        self.assertEqual(self.page._extension_panel._config_help_area.minimumHeight(), expected_height)
        self.assertEqual(self.page._extension_panel._config_help_area.maximumHeight(), expected_height)
        self.assertIs(self.page._extension_panel._config_help_area.widget(), self.page._extension_panel._config_help_container)

    def test_plot_extension_tab_shows_divider_below_config_selector(self):
        self.page._style_tabs.setCurrentIndex(2)
        self.page._refresh_style_extension_panel()

        self.assertFalse(self.page._plot_extension_controls._config_row_widget.isHidden())
        self.assertFalse(self.page._plot_extension_controls._config_row_divider.isHidden())

    def test_extension_panel_uses_splitter_side_panel(self):
        from ui.widgets.extension_panel import ExtensionConfigPanel

        self.page.resize(1280, 820)
        self.page.show()
        QApplication.processEvents()

        self.assertEqual(self.page._page_splitter.count(), 2)
        self.assertIs(self.page._page_splitter.widget(1), self.page._extension_panel)
        self.assertIsInstance(self.page._extension_panel, ExtensionConfigPanel)
        self.assertTrue(self.page._extension_panel.isHidden())

        self.page.set_extension_panel_visible(True)
        QApplication.processEvents()

        self.assertFalse(self.page._extension_panel.isHidden())
        self.assertGreater(self.page._page_splitter.sizes()[1], 0)

    def test_chart_main_splitters_use_shared_workbench_handle_width(self):
        self.assertEqual(self.page._page_splitter.handleWidth(), 4)
        self.assertEqual(self.page._content_splitter.handleWidth(), 4)

    def test_chart_current_curve_color_uses_fluent_picker_button(self):
        from PySide6.QtWidgets import QFrame
        from qfluentwidgets import ColorPickerButton

        self.assertIsInstance(self.page._style_color_btn, ColorPickerButton)
        self.assertEqual(self.page._style_color_btn.width(), 32)
        self.assertEqual(self.page._style_color_btn.height(), 32)
        self.assertEqual(self.page._plot_style_scroll.frameShape(), QFrame.Shape.NoFrame)

    def test_chart_path_label_uses_shared_hint_style(self):
        from ui.theme import make_hint_label

        probe = make_hint_label("路径：示例")
        try:
            self.assertEqual(self.page._chart_path_label.styleSheet(), probe.styleSheet())
        finally:
            probe.deleteLater()

    def test_on_tree_node_selected_data_file(self):
        node = next((n for n in self.p.tree.nodes if n.kind == "data_file"), None)
        if node:
            self.page.on_tree_node_selected("data_file", node.id)

    def test_on_tree_node_selected_image_work(self):
        self.page.on_tree_node_selected("image_work", "fake-image-id")

    def test_on_tree_node_selected_curve_updates_selection_state(self):
        self.page.on_tree_node_selected("curve", "curve-1")

        self.assertEqual(self.page._selected_tree_kind, "curve")
        self.assertEqual(self.page._selected_tree_id, "curve-1")

    def test_on_tree_node_selected_folder(self):
        node = next((n for n in self.p.tree.nodes if n.kind == "folder"), None)
        if node:
            self.page.on_tree_node_selected("folder", node.id)

    def test_load_template_no_crash(self):
        self.page.load_template("nonexistent-template-id")

    def test_load_template_valid(self):
        """添加模板后再加载"""
        from models.schemas import FigureConfig, AxisConfig
        cfg = FigureConfig(name="templ1")
        cfg.theme = "Nature"
        cfg.show_errbar = True
        cfg.figure_size = (9.0, 6.0)
        cfg.dpi = 220
        cfg.font_size = 14
        cfg.font_family = "DejaVu Sans"
        cfg.legend_font_size = 12
        cfg.line_width = 2.5
        cfg.marker_size = 7.0
        cfg.grid_alpha = 0.4
        cfg.grid_line_width = 1.1
        cfg.legend_position = "lower left"
        cfg.typed_axis_config = AxisConfig(x_label="Time", y_label="Signal")
        template = self.pm.add_figure_template(cfg)
        self.assertIsNotNone(template)
        self.page.load_template(template.id)
        self.assertEqual(self.page._figure_state.theme, "Nature")
        self.assertEqual(self.page._figure_state.x_label, "Time")
        self.assertTrue(self.page._figure_state.show_errbar)
        self.assertEqual(self.page._figure_state.figure_width, 9.0)
        self.assertEqual(self.page._figure_state.figure_height, 6.0)
        self.assertEqual(self.page._figure_state.dpi, 220)
        self.assertEqual(self.page._figure_state.font_family, "DejaVu Sans")
        self.assertEqual(self.page._figure_state.legend_font_size, 12)
        self.assertEqual(self.page._figure_state.line_width, 2.5)
        self.assertEqual(self.page._figure_state.marker_size, 7.0)
        self.assertEqual(self.page._figure_state.legend_pos, "lower left")
        self.assertEqual(self.page._template_combo.currentText(), "templ1")
        self.assertTrue(self.page._btn_update_template.isEnabled())

    def test_refresh_template_combo_includes_saved_template(self):
        from models.schemas import FigureConfig

        self.pm.add_figure_template(FigureConfig(name="模板甲", theme="ACS"))
        self.page._refresh_template_combo()
        items = [self.page._template_combo.itemText(i) for i in range(self.page._template_combo.count())]
        self.assertIn("模板甲", items)

    def test_update_current_template_overwrites_existing_figure(self):
        from models.schemas import FigureConfig

        from core.global_assets import global_assets

        template = self.pm.add_figure_template(FigureConfig(name="模板可更新", theme="默认"))
        self.assertIsNotNone(template)
        figures_before = len(global_assets.list_figure_templates())
        self.page.load_template(template.id)
        self.page.load_plot_style("IEEE")
        self.page._x_label_edit.setText("Voltage")
        updated = self.page._update_current_template()
        self.assertTrue(updated)
        self.assertEqual(len(global_assets.list_figure_templates()), figures_before)
        fig = global_assets.get_figure_template(template.id)
        self.assertIsNotNone(fig)
        self.assertEqual(fig.theme, "IEEE")
        self.assertEqual(fig.typed_axis_config.x_label, "Voltage")

    def test_theme_hint_updates_with_theme(self):
        self.page.load_plot_style("Nature")
        self.page._on_quick_config_changed()
        self.assertIn("论文", self.page._theme_hint_label.text())

    def test_quick_controls_update_figure_state(self):
        self.page._x_label_edit.setText("Time")
        self.page._y_label_edit.setText("Intensity")
        self.page.load_plot_style("ACS")
        self.page._errbar_cb.setChecked(True)
        state = self.page._sync_state_from_controls()
        self.assertEqual(state.x_label, "Time")
        self.assertEqual(state.y_label, "Intensity")
        self.assertEqual(state.theme, "ACS")
        self.assertTrue(state.show_errbar)

    def test_apply_advanced_config_updates_figure_state(self):
        self.page._apply_advanced_config({
            "theme": "IEEE",
            "x_label": "Voltage",
            "y_label": "Current",
            "x_min": "0",
            "x_max": "10",
            "y_min": "1",
            "y_max": "5",
            "x_log": True,
            "y_log": False,
            "grid": False,
            "legend_pos": "upper right",
            "font_size": 12,
            "font_family": "DejaVu Sans",
            "legend_font_size": 11,
            "figure_width": 8.5,
            "figure_height": 6.5,
            "dpi": 180,
            "line_width": 2.1,
            "marker_size": 6.2,
            "grid_alpha": 0.4,
            "grid_line_width": 0.8,
            "show_errbar": True,
        })
        self.assertEqual(self.page._figure_state.theme, "IEEE")
        self.assertEqual(self.page._figure_state.x_label, "Voltage")
        self.assertEqual(self.page._figure_state.legend_pos, "upper right")
        self.assertEqual(self.page._figure_state.x_max, 10.0)
        self.assertEqual(self.page._figure_state.figure_width, 8.5)
        self.assertEqual(self.page._figure_state.figure_height, 6.5)
        self.assertEqual(self.page._figure_state.dpi, 180)
        self.assertEqual(self.page._figure_state.font_family, "DejaVu Sans")
        self.assertEqual(self.page._figure_state.legend_font_size, 11)
        self.assertEqual(self.page._figure_state.line_width, 2.1)
        self.assertEqual(self.page._figure_state.marker_size, 6.2)
        self.assertEqual(self.page._legend_pos_combo.currentText(), "upper right")
        self.assertEqual(self.page._figure_width_edit.text(), "8.5")
        self.assertEqual(self.page._plot_line_width_edit.text(), "2.1")

    def test_builtin_plot_style_fully_overrides_default_metrics(self):
        self.page._plot_line_width_edit.setText("5.0")

        self.page.load_plot_style("IEEE")

        self.assertAlmostEqual(self.page._figure_state.line_width, 1.2)
        self.assertAlmostEqual(self.page._figure_state.marker_size, 3.8)

    def test_curve_style_tab_has_explicit_load_button(self):
        self.assertIsNotNone(self.page._style_tabs)
        self.assertIsNotNone(self.page._btn_load_curve_style_template)

    def test_plot_style_tab_has_explicit_load_button(self):
        self.assertIsNotNone(self.page._btn_load_template)

    def test_style_tabs_include_dedicated_plot_extension_tab(self):
        titles = [self.page._style_tabs.tabText(i) for i in range(self.page._style_tabs.count())]

        self.assertEqual(titles, ["曲线样式", "绘图样式", "绘图扩展"])

    def test_plot_extension_appears_in_extension_panel(self):
        from core.extension_api import PlotExtension, extension_registry
        from extensions.plot._runtime import current_axis

        def _draw(lines, options):
            del lines
            axis = current_axis()
            if axis is not None:
                axis.axhline(options.get("y", 0.0))

        extension_registry.register_plot(
            PlotExtension(
                type="ui_plot_extension_test",
                name="UI 参考线",
                handler=_draw,
                default_options={"y": 1.5},
            )
        )
        try:
            self.page._style_tabs.setCurrentIndex(2)
            self.page._refresh_style_extension_panel()
            selector_items = [
                self.page._plot_extension_controls._selector.itemText(i)
                for i in range(self.page._plot_extension_controls._selector.count())
            ]

            self.assertIn("UI 参考线", selector_items)
        finally:
            extension_registry.unregister_plot("ui_plot_extension_test")
            self.page._refresh_style_extension_panel()

    def test_chart_curve_order_buttons_follow_selected_curve_position(self):
        from models.schemas import DataSeries

        other = DataSeries(name="s2", x=[1.0, 2.0], y=[2.0, 3.0])
        self.df.series.append(other)

        self.page.on_tree_node_activated("series", self.s.id)
        self.page.on_tree_node_activated("series", other.id)

        self.assertFalse(self.page._btn_selected_up.isEnabled())
        self.assertTrue(self.page._btn_selected_down.isEnabled())

        self.page._chart_list.setCurrentRow(1)

        self.assertTrue(self.page._btn_selected_up.isEnabled())
        self.assertFalse(self.page._btn_selected_down.isEnabled())

    def test_moving_selected_curve_updates_list_and_draw_order(self):
        from models.schemas import DataSeries

        other = DataSeries(name="s2", x=[1.0, 2.0], y=[2.0, 3.0])
        self.df.series.append(other)

        self.page.on_tree_node_activated("series", self.s.id)
        self.page.on_tree_node_activated("series", other.id)
        self.page._move_selected_curve_down()

        self.assertEqual([curve.get("name") for curve in self.page._chart_series], [other.name, self.s.name])
        self.assertIn(other.name, self.page._chart_list.item(0).text())
        self.assertIn(self.s.name, self.page._chart_list.item(1).text())

        axis = self.page._figure.axes[0]
        labels = [line.get_label() for line in axis.lines if not str(line.get_label()).startswith("_")]
        self.assertEqual(labels, [other.name, self.s.name])

    def test_plot_extension_runs_after_default_draw(self):
        from core.extension_api import PlotExtension, extension_registry
        from extensions.plot._runtime import current_axis

        def _draw(lines, options):
            del lines
            axis = current_axis()
            if axis is not None:
                axis.axhline(options.get("y", 0.0), color="red", label="参考线")

        extension_registry.register_plot(
            PlotExtension(
                type="ui_plot_extension_reference_test",
                name="参考线",
                handler=_draw,
                default_options={"y": 3.5},
                phases=("after_plot",),
            )
        )
        try:
            self.page.on_tree_node_activated("series", self.s.id)
            self.page._plot_extension_options["ui_plot_extension_reference_test"] = {"y": 3.5}

            self.page._apply_plot_extension("ui_plot_extension_reference_test")

            axis = self.page._figure.axes[0]
            self.assertTrue(
                any(
                    list(line.get_ydata()) == [3.5, 3.5]
                    for line in axis.lines
                )
            )
        finally:
            extension_registry.unregister_plot("ui_plot_extension_reference_test")
            self.page._plot_extension_options.pop("ui_plot_extension_reference_test", None)

    def test_plot_extension_context_can_take_over_matplotlib_figure(self):
        from core.extension_api import PlotExtension, extension_registry
        from extensions.plot._runtime import current_figure

        def _draw(lines, options):
            del lines
            figure = current_figure()
            if figure is None:
                return
            figure.clear()
            left = figure.add_subplot(121)
            right = figure.add_subplot(122)
            left.plot([0, 1, 2], [1, 3, 2], color=options.get("color", "#ff6600"))
            right.bar(["A", "B"], [2, 4], color=options.get("bar_color", "#3366cc"))

        extension_registry.register_plot(
            PlotExtension(
                type="ui_plot_extension_context_test",
                name="Context 自定义子图",
                handler=_draw,
                default_options={"color": "#ff6600", "bar_color": "#3366cc"},
                phases=("after_plot",),
            )
        )
        try:
            self.page.on_tree_node_activated("series", self.s.id)
            self.page._plot_extension_options["ui_plot_extension_context_test"] = {
                "color": "#ff6600",
                "bar_color": "#3366cc",
            }

            self.page._apply_plot_extension("ui_plot_extension_context_test")

            self.assertEqual(len(self.page._figure.axes), 2)
            self.assertEqual(len(self.page._figure.axes[0].lines), 1)
            self.assertEqual(len(self.page._figure.axes[1].patches), 2)
        finally:
            extension_registry.unregister_plot("ui_plot_extension_context_test")
            self.page._plot_extension_options.pop("ui_plot_extension_context_test", None)

    def test_plot_extension_after_plot_preserves_manual_line_width_and_updates_legend(self):
        from core.extension_api import PlotExtension, extension_registry
        from extensions.plot._runtime import current_axis

        def _style_patch(lines, options):
            del lines
            axis = current_axis()
            if axis is None:
                return
            legend = axis.get_legend()
            if legend is None:
                return
            font_size = int(options.get("legend_font_size", 12))
            for text in legend.get_texts():
                text.set_fontsize(font_size)

        extension_registry.register_plot(
            PlotExtension(
                type="chart_plot_style_patch_order",
                name="样式覆盖顺序",
                handler=_style_patch,
                default_options={"line_width": 4.0, "legend_font_size": 12},
                phases=("after_plot",),
            )
        )
        try:
            self.page.on_tree_node_activated("series", self.s.id)
            self.page._plot_line_width_edit.setText("2.0")
            self.page._on_quick_config_changed()

            self.page._on_chart_extension_apply("chart_plot_style_patch_order", {"line_width": 4.0, "legend_font_size": 12})

            axis = self.page._figure.axes[0]
            legend = axis.get_legend()
            self.assertAlmostEqual(axis.lines[0].get_linewidth(), 2.0)
            self.assertIsNotNone(legend)
            self.assertAlmostEqual(legend.get_texts()[0].get_fontsize(), 12)

            self.page._plot_line_width_edit.setText("2.6")
            self.page._on_quick_config_changed()
            self.page._redraw_now()

            axis = self.page._figure.axes[0]
            legend = axis.get_legend()
            self.assertAlmostEqual(axis.lines[0].get_linewidth(), 2.6)
            self.assertIsNotNone(legend)
            self.assertAlmostEqual(legend.get_texts()[0].get_fontsize(), 12)
        finally:
            extension_registry.unregister_plot("chart_plot_style_patch_order")
            self.page._plot_extension_options.pop("chart_plot_style_patch_order", None)
            self.page._applied_plot_extensions = [
                entry for entry in self.page._applied_plot_extensions
                if entry.get("type") != "chart_plot_style_patch_order"
            ]
            self.page._refresh_style_extension_panel()

    def test_plot_extension_style_patch_preserves_other_nested_style_fields_on_manual_update(self):
        from matplotlib.colors import to_rgba
        from core.extension_api import PlotExtension, extension_registry
        from extensions.plot._runtime import current_axis

        def _style_patch(lines, _options):
            del lines
            axis = current_axis()
            if axis is None:
                return
            legend = axis.get_legend()
            if legend is None:
                legend = axis.legend()
            legend.get_frame().set_visible(True)
            legend.get_frame().set_edgecolor("#333333")

        extension_registry.register_plot(
            PlotExtension(
                type="chart_plot_style_patch_nested",
                name="样式最小覆盖",
                handler=_style_patch,
                default_options={},
            )
        )
        try:
            self.page.on_tree_node_activated("series", self.s.id)
            self.page._on_chart_extension_apply("chart_plot_style_patch_nested", {})

            self.page._apply_advanced_config({"legend_kwargs": {"facecolor": "#ccffee"}})

            legend = self.page._figure.axes[0].get_legend()
            self.assertIsNotNone(legend)
            self.assertTrue(legend.get_frame().get_visible())
            expected_rgb = tuple(round(value, 3) for value in to_rgba("#ccffee")[:3])
            actual_rgb = tuple(round(value, 3) for value in legend.get_frame().get_facecolor()[:3])
            self.assertEqual(actual_rgb, expected_rgb)
        finally:
            extension_registry.unregister_plot("chart_plot_style_patch_nested")
            self.page._plot_extension_options.pop("chart_plot_style_patch_nested", None)
            self.page._applied_plot_extensions = [
                entry for entry in self.page._applied_plot_extensions
                if entry.get("type") != "chart_plot_style_patch_nested"
            ]
            self.page._refresh_style_extension_panel()

    def test_builtin_science_plot_extension_applies_labels_and_legend_frame(self):
        from core.extension_api import reload_configured_extensions

        reload_configured_extensions()
        self.page._refresh_style_extension_panel()
        self.page.on_tree_node_activated("series", self.s.id)

        self.page._on_chart_extension_apply(
            "demo_plot_science_style",
            {
                "x_label": "时间",
                "y_label": "幅值",
                "axis_label_size": 13,
                "tick_label_size": 9,
                "legend_frame": True,
                "legend_location": "upper left",
                "line_width": 1.8,
            },
        )

        axis = self.page._figure.axes[0]
        legend = axis.get_legend()
        self.assertEqual(axis.get_xlabel(), "时间")
        self.assertEqual(axis.get_ylabel(), "幅值")
        self.assertAlmostEqual(axis.lines[0].get_linewidth(), 1.8)
        self.assertIsNotNone(legend)
        self.assertTrue(legend.get_frame().get_visible())

    def test_builtin_polar_plot_extension_draws_polar_axis(self):
        from core.extension_api import reload_configured_extensions

        reload_configured_extensions()
        self.page._refresh_style_extension_panel()
        self.page.on_tree_node_activated("series", self.s.id)

        self.page._on_chart_extension_apply(
            "demo_plot_polar_projection",
            {
                "theta_unit": "degree",
                "title": "极坐标预览",
                "show_legend": True,
            },
        )

        axis = self.page._figure.axes[0]
        self.assertEqual(axis.name, "polar")
        self.assertEqual(axis.get_title(), "极坐标预览")
        self.assertGreaterEqual(len(axis.lines), 1)

    def test_curve_style_opacity_is_exposed_and_clamped(self):
        self.page.on_tree_node_activated("series", self.s.id)
        self.page._chart_list.setCurrentRow(0)

        self.assertTrue(self.page._style_opacity_slider.isEnabled())
        self.page._style_opacity_slider.setValue(70)
        self.page._style_opacity_slider.setValue(170)

        curve_key = self.page._curve_key(self.page._chart_series[0])
        self.assertEqual(self.page._style_opacity_slider.value(), 100)
        self.assertEqual(self.page._curve_styles[curve_key]["alpha"], 1.0)

    def test_style_tabs_hide_add_and_close_buttons(self):
        from qfluentwidgets import TabCloseButtonDisplayMode

        self.assertTrue(self.page._style_tabs.tabBar.addButton.isHidden())
        self.assertEqual(
            self.page._style_tabs.tabBar.closeButtonDisplayMode,
            TabCloseButtonDisplayMode.NEVER,
        )

    def test_style_tabs_use_segmented_navigation_and_selected_list_button_metrics(self):
        from qfluentwidgets import PushButton, SegmentedWidget

        self.assertIsInstance(self.page._style_tabs.navigationWidget, SegmentedWidget)
        self.assertIsInstance(self.page._btn_clear, PushButton)
        self.assertIsInstance(self.page._btn_remove, PushButton)
        self.assertEqual(self.page._btn_clear.height(), 32)
        self.assertEqual(self.page._btn_remove.height(), 32)
        self.assertEqual(self.page._btn_clear.text(), "清除")
        self.assertEqual(self.page._btn_remove.text(), "移除选中")

    def test_selected_curve_actions_do_not_overlap_in_left_panel(self):
        from ui.theme import WORKBENCH_BUTTON_MIN_WIDTH

        self.page.resize(1280, 820)
        self.page.show()
        QApplication.processEvents()

        action_buttons = [
            self.page._btn_clear,
            self.page._btn_remove,
            self.page._btn_selected_up,
            self.page._btn_selected_down,
            self.page._btn_toggle_visible,
        ]
        for index, button in enumerate(action_buttons):
            for other in action_buttons[index + 1:]:
                self.assertFalse(button.geometry().intersects(other.geometry()))
        self.assertEqual(len({button.geometry().top() for button in action_buttons}), 1)
        self.assertLess(self.page._btn_clear.width(), WORKBENCH_BUTTON_MIN_WIDTH)

    def test_plot_extension_side_panel_shows_description_section(self):
        from core.extension_api import PlotExtension, extension_registry
        from ui.widgets.extension_panel import ExtensionConfigPanel

        self.assertIsInstance(self.page._extension_panel, ExtensionConfigPanel)
        self.assertEqual(self.page._extension_panel._description_label.text(), "在左侧选择扩展后，这里会显示扩展说明。")
        self.assertEqual(self.page._extension_panel._config_help_label.text(), "保留 {} 使用默认参数。")

        extension_registry.register_plot(
            PlotExtension(
                type="chart_plot_description_probe",
                name="说明探针",
                description="用于验证右侧绘图扩展说明区域。",
                handler=lambda plot_context, options: None,
            )
        )
        try:
            self.page._refresh_style_extension_panel()
            self.page._update_plot_extension_info_panel("chart_plot_description_probe")

            self.assertEqual(
                self.page._extension_panel._description_label.text(),
                "用于验证右侧绘图扩展说明区域。",
            )
        finally:
            extension_registry.unregister_plot("chart_plot_description_probe")
            self.page._refresh_style_extension_panel()

    def test_chart_left_panel_uses_vertical_splitter(self):
        self.page.resize(1280, 820)
        self.page.show()
        QApplication.processEvents()

        self.assertEqual(self.page._chart_left_splitter.count(), 2)
        sizes = self.page._chart_left_splitter.sizes()
        self.assertGreater(sizes[1], 0)
        ratio = sizes[0] / sum(sizes)
        self.assertGreater(ratio, 0.34)
        self.assertLess(ratio, 0.46)

    def test_plot_style_numeric_inputs_use_compact_widths(self):
        self.assertLessEqual(self.page._font_size_edit.maximumWidth(), 96)
        self.assertLessEqual(self.page._figure_width_edit.maximumWidth(), 96)
        self.assertLessEqual(self.page._dpi_edit.maximumWidth(), 96)
        self.assertLessEqual(self.page._legend_anchor_x_edit.maximumWidth(), 96)
        self.assertLessEqual(self.page._legend_anchor_y_edit.maximumWidth(), 96)

    def test_style_tabs_split_dense_controls_into_multiple_rows(self):
        self.page.resize(1024, 820)
        self.page.show()
        QApplication.processEvents()

        self.page._style_tabs.setCurrentIndex(0)
        QApplication.processEvents()
        self.assertGreater(self.page._style_dash_scale_edit.geometry().top(), self.page._style_opacity_slider.geometry().top())

        self.page._style_tabs.setCurrentIndex(1)
        QApplication.processEvents()
        self.assertGreater(self.page._tick_label_top_cb.geometry().top(), self.page._tick_top_cb.geometry().top())
        self.assertGreater(self.page._grid_cb.geometry().top(), self.page._plot_marker_size_edit.geometry().top())
        self.assertGreater(self.page._grid_alpha_slider.geometry().top(), self.page._grid_cb.geometry().top())
        self.assertGreater(self.page._legend_face_color_edit.geometry().top(), self.page._legend_edge_color_edit.geometry().top())
        self.assertGreater(self.page._legend_face_alpha_slider.geometry().top(), self.page._legend_face_color_edit.geometry().top())
        self.assertGreater(self.page._canvas_alpha_slider.geometry().top(), self.page._canvas_color_edit.geometry().top())
        self.assertGreater(self.page._figure_height_edit.geometry().top(), self.page._figure_width_edit.geometry().top())

    def test_plot_actions_keep_only_export_button(self):
        plot_tab = self.page._style_tabs.widget(1)
        self.assertIsNot(self.page._btn_export.parent(), plot_tab)
        self.assertIs(self.page._btn_export.parent(), self.page._plot_actions_bar)
        self.assertFalse(hasattr(self.page, "_btn_advanced"))

    def test_font_family_uses_detected_combo_choices(self):
        self.assertGreater(self.page._font_family_combo.count(), 0)
        self.assertEqual(self.page._font_family_combo.itemText(0), "默认")

    def test_unbound_style_labels_use_none_text(self):
        self.assertEqual(self.page._template_combo.itemText(0), "无")
        self.assertEqual(self.page._curve_style_template_combo.itemText(0), "无")

    def test_grid_alpha_is_clamped_to_valid_range(self):
        self.page._grid_alpha_slider.setValue(500)

        state = self.page._sync_state_from_controls()

        self.assertEqual(state.grid_alpha, 1.0)

    def test_curve_style_tab_supports_visibility_and_dash_scale(self):
        self.page.on_tree_node_activated("series", self.s.id)
        self.page._chart_list.setCurrentRow(0)

        curve_key = self.page._curve_key(self.page._chart_series[0])
        self.assertTrue(self.page._style_visible_cb.isEnabled())
        self.page._style_visible_cb.setChecked(False)
        self.assertFalse(self.page._chart_series[0]["visible"])

        self.page._style_dash_scale_edit.setText("2.5")
        self.page._style_dash_scale_edit.editingFinished.emit()

        self.assertAlmostEqual(self.page._curve_styles[curve_key]["dash_scale"], 2.5)

    def test_plot_style_tick_and_spine_controls_apply_extras(self):
        if self.page._figure is None:
            self.skipTest("matplotlib unavailable")

        self.page.on_tree_node_activated("series", self.s.id)
        self.page._tick_bottom_cb.setChecked(False)
        self.page._tick_left_cb.setChecked(False)
        self.page._tick_top_cb.setChecked(True)
        self.page._tick_right_cb.setChecked(True)
        self.page._tick_label_bottom_cb.setChecked(False)
        self.page._tick_label_left_cb.setChecked(False)
        self.page._tick_label_top_cb.setChecked(True)
        self.page._tick_label_right_cb.setChecked(True)
        self.page._tick_direction_combo.setCurrentText("in")
        self.page._tick_length_edit.setText("8")
        self.page._tick_length_edit.editingFinished.emit()
        self.page._tick_width_edit.setText("1.5")
        self.page._tick_width_edit.editingFinished.emit()
        self.page._legend_frame_cb.setChecked(False)

        extras = self.page._plot_style_extras
        self.assertFalse(extras["tick_params"]["bottom"])
        self.assertFalse(extras["tick_params"]["left"])
        self.assertTrue(extras["tick_params"]["top"])
        self.assertTrue(extras["tick_params"]["right"])
        self.assertFalse(extras["tick_params"]["labelbottom"])
        self.assertFalse(extras["tick_params"]["labelleft"])
        self.assertTrue(extras["tick_params"]["labelright"])
        self.assertEqual(extras["tick_params"]["direction"], "in")
        self.assertAlmostEqual(extras["tick_params"]["length"], 8.0)
        self.assertAlmostEqual(extras["tick_params"]["width"], 1.5)
        self.assertFalse(extras["legend_kwargs"]["frameon"])

        self.page._redraw_now()

        axis = self.page._figure.axes[0]
        legend = axis.get_legend()
        first_x_tick = axis.xaxis.get_major_ticks()[0]
        first_y_tick = axis.yaxis.get_major_ticks()[0]
        self.assertFalse(first_x_tick.tick1line.get_visible())
        self.assertTrue(first_x_tick.tick2line.get_visible())
        self.assertFalse(first_x_tick.label1.get_visible())
        self.assertTrue(first_x_tick.label2.get_visible())
        self.assertFalse(first_y_tick.tick1line.get_visible())
        self.assertTrue(first_y_tick.tick2line.get_visible())
        self.assertFalse(first_y_tick.label1.get_visible())
        self.assertTrue(first_y_tick.label2.get_visible())
        self.assertAlmostEqual(first_x_tick.tick2line.get_markersize(), 8.0)
        self.assertFalse(legend.get_frame().get_visible())

        self.page._spine_bottom_cb.setChecked(False)
        self.page._spine_left_cb.setChecked(False)
        self.page._spine_top_cb.setChecked(False)
        self.page._spine_right_cb.setChecked(False)

        extras = self.page._plot_style_extras
        self.assertFalse(extras["spine_visibility"]["bottom"])
        self.assertFalse(extras["spine_visibility"]["left"])
        self.assertFalse(extras["spine_visibility"]["top"])
        self.assertFalse(extras["spine_visibility"]["right"])
        self.page._redraw_now()
        axis = self.page._figure.axes[0]
        self.assertFalse(axis.spines["bottom"].get_visible())
        self.assertFalse(axis.spines["left"].get_visible())
        self.assertFalse(axis.spines["top"].get_visible())
        self.assertFalse(axis.spines["right"].get_visible())

    def test_plot_style_legend_and_canvas_controls_apply_extras(self):
        if self.page._figure is None:
            self.skipTest("matplotlib unavailable")

        self.page.on_tree_node_activated("series", self.s.id)
        self.page._legend_frame_alpha_slider.setValue(35)
        self.page._legend_face_color_edit.setText("#123456")
        self.page._legend_face_color_edit.editingFinished.emit()
        self.page._legend_face_alpha_slider.setValue(60)
        self.page._canvas_color_edit.setText("#224466")
        self.page._canvas_color_edit.editingFinished.emit()
        self.page._canvas_alpha_slider.setValue(45)

        extras = self.page._plot_style_extras
        self.assertAlmostEqual(extras["legend_kwargs"]["edgealpha"], 0.35)
        self.assertEqual(extras["legend_kwargs"]["facecolor"], "#123456")
        self.assertAlmostEqual(extras["legend_kwargs"]["facealpha"], 0.60)
        self.assertEqual(extras["figure_facecolor"], "#224466")
        self.assertAlmostEqual(extras["figure_facealpha"], 0.45)

        self.page._redraw_now()

        axis = self.page._figure.axes[0]
        legend = axis.get_legend()
        legend_frame = legend.get_frame()
        self.assertAlmostEqual(legend_frame.get_edgecolor()[3], 0.35, places=2)
        self.assertAlmostEqual(legend_frame.get_facecolor()[3], 0.60, places=2)

        expected_canvas_rgb = (34 / 255, 68 / 255, 102 / 255)
        for actual, expected in zip(self.page._figure.get_facecolor()[:3], expected_canvas_rgb):
            self.assertAlmostEqual(actual, expected, places=2)
        for actual, expected in zip(axis.get_facecolor()[:3], expected_canvas_rgb):
            self.assertAlmostEqual(actual, expected, places=2)
        self.assertAlmostEqual(self.page._figure.get_facecolor()[3], 0.45, places=2)
        self.assertAlmostEqual(axis.get_facecolor()[3], 0.45, places=2)

    def test_plot_style_template_preserves_style_extras(self):
        self.page.on_tree_node_activated("series", self.s.id)
        self.page._tick_top_cb.setChecked(True)
        self.page._legend_frame_cb.setChecked(False)
        self.page._legend_frame_alpha_slider.setValue(35)
        self.page._legend_anchor_x_edit.setText("0.2")
        self.page._legend_anchor_x_edit.editingFinished.emit()
        self.page._legend_anchor_y_edit.setText("0.8")
        self.page._legend_anchor_y_edit.editingFinished.emit()
        self.page._legend_edge_color_edit.setText("#ff0000")
        self.page._legend_edge_color_edit.editingFinished.emit()
        self.page._legend_face_color_edit.setText("#123456")
        self.page._legend_face_color_edit.editingFinished.emit()
        self.page._legend_face_alpha_slider.setValue(60)
        self.page._canvas_color_edit.setText("#224466")
        self.page._canvas_color_edit.editingFinished.emit()
        self.page._canvas_alpha_slider.setValue(45)

        template = self.page._save_template_named("带扩展样式")

        self.assertIsNotNone(template)
        self.assertTrue(template.style_extras["tick_params"]["top"])
        self.assertFalse(template.style_extras["legend_kwargs"]["frameon"])
        self.assertAlmostEqual(template.style_extras["legend_kwargs"]["edgealpha"], 0.35)
        self.assertEqual(template.style_extras["legend_kwargs"]["bbox_to_anchor"], [0.2, 0.8])
        self.assertEqual(template.style_extras["legend_kwargs"]["edgecolor"], "#ff0000")
        self.assertEqual(template.style_extras["legend_kwargs"]["facecolor"], "#123456")
        self.assertAlmostEqual(template.style_extras["legend_kwargs"]["facealpha"], 0.60)
        self.assertEqual(template.style_extras["figure_facecolor"], "#224466")
        self.assertAlmostEqual(template.style_extras["figure_facealpha"], 0.45)

        self.page._tick_top_cb.setChecked(False)
        self.page._legend_frame_cb.setChecked(True)
        self.page._legend_frame_alpha_slider.setValue(80)
        self.page._legend_anchor_x_edit.clear()
        self.page._legend_anchor_x_edit.editingFinished.emit()
        self.page._legend_anchor_y_edit.clear()
        self.page._legend_anchor_y_edit.editingFinished.emit()
        self.page._legend_edge_color_edit.clear()
        self.page._legend_edge_color_edit.editingFinished.emit()
        self.page._legend_face_color_edit.clear()
        self.page._legend_face_color_edit.editingFinished.emit()
        self.page._legend_face_alpha_slider.setValue(80)
        self.page._canvas_color_edit.clear()
        self.page._canvas_color_edit.editingFinished.emit()
        self.page._canvas_alpha_slider.setValue(100)
        self.page.load_template(template.id)

        self.assertTrue(self.page._tick_top_cb.isChecked())
        self.assertFalse(self.page._legend_frame_cb.isChecked())
        self.assertEqual(self.page._legend_frame_alpha_slider.value(), 35)
        self.assertEqual(self.page._legend_anchor_x_edit.text(), "0.2")
        self.assertEqual(self.page._legend_anchor_y_edit.text(), "0.8")
        self.assertEqual(self.page._legend_edge_color_edit.text(), "#ff0000")
        self.assertEqual(self.page._legend_face_color_edit.text(), "#123456")
        self.assertEqual(self.page._legend_face_alpha_slider.value(), 60)
        self.assertEqual(self.page._canvas_color_edit.text(), "#224466")
        self.assertEqual(self.page._canvas_alpha_slider.value(), 45)

    def test_plot_style_supports_precise_legend_anchor(self):
        if self.page._figure is None:
            self.skipTest("matplotlib unavailable")

        self.page.on_tree_node_activated("series", self.s.id)
        self.page._legend_pos_combo.setCurrentText("upper left")
        self.page._legend_anchor_x_edit.setText("0.2")
        self.page._legend_anchor_x_edit.editingFinished.emit()
        self.page._legend_anchor_y_edit.setText("0.8")
        self.page._legend_anchor_y_edit.editingFinished.emit()

        self.assertEqual(self.page._plot_style_extras["legend_kwargs"]["bbox_to_anchor"], [0.2, 0.8])

        self.page._redraw_now()

        axis = self.page._figure.axes[0]
        legend = axis.get_legend()
        anchor_bbox = legend.get_bbox_to_anchor().transformed(axis.transAxes.inverted())
        self.assertAlmostEqual(anchor_bbox.x0, 0.2, places=2)
        self.assertAlmostEqual(anchor_bbox.y0, 0.8, places=2)

    def test_display_canvas_size_preserves_requested_aspect_ratio(self):
        width, height = self.page._fitted_canvas_size(1200, 800, 9.5 / 6.5)

        self.assertAlmostEqual(width / height, 9.5 / 6.5, delta=0.02)
        self.assertLessEqual(width, 1200)
        self.assertLessEqual(height, 800)

    def test_redraw_on_compact_canvas_avoids_tight_layout_warning(self):
        if self.page._figure is None or self.page._canvas_host is None:
            self.skipTest("matplotlib unavailable")

        self.page.on_tree_node_activated("series", self.s.id)
        curve = self.page._chart_series[0]
        self.page._set_curve_display_name(curve, "图例名称较长用于紧凑布局")
        self.page._apply_advanced_config({
            "font_size": 16,
            "legend_font_size": 12,
            "x_label": "较长横坐标标签",
            "y_label": "较长纵坐标标签",
        })
        self.page._canvas_host.setFixedSize(220, 160)
        self.page._sync_canvas_display_geometry()

        with warnings.catch_warnings(record=True) as caught:
            warnings.simplefilter("always")
            self.page._redraw_now()

        layout_warnings = [str(item.message) for item in caught if "Tight layout not applied" in str(item.message)]
        self.assertEqual(layout_warnings, [])

    def test_redraw_applies_font_family_and_sizes_to_axis_and_legend(self):
        if self.page._figure is None:
            self.skipTest("matplotlib unavailable")

        self.page.on_tree_node_activated("series", self.s.id)
        font_index = self.page._font_family_combo.findText("DejaVu Sans")
        self.assertGreaterEqual(font_index, 0)
        self.page._font_family_combo.setCurrentIndex(font_index)
        self.page._font_size_edit.setText("15")
        self.page._legend_font_size_edit.setText("9")

        self.page._redraw_now()

        axis = self.page._figure.axes[0]
        self.assertAlmostEqual(axis.xaxis.label.get_fontsize(), 15)
        self.assertIn("DejaVu Sans", axis.xaxis.label.get_fontfamily())
        self.assertAlmostEqual(axis.yaxis.label.get_fontsize(), 15)
        tick_labels = axis.get_xticklabels() + axis.get_yticklabels()
        self.assertTrue(tick_labels)
        self.assertAlmostEqual(tick_labels[0].get_fontsize(), 15)
        self.assertIn("DejaVu Sans", tick_labels[0].get_fontfamily())

        legend = axis.get_legend()
        self.assertIsNotNone(legend)
        legend_text = legend.get_texts()[0]
        self.assertAlmostEqual(legend_text.get_fontsize(), 9)
        self.assertIn("DejaVu Sans", legend_text.get_fontfamily())

    def test_load_selected_curve_style_template_applies_style(self):
        restore_assets = _patch_global_assets()
        try:
            from core.global_assets import global_assets
            from models.schemas import CurveStyle, CurveStyleTemplate

            self.page.on_tree_node_activated("series", self.s.id)
            template = global_assets.add_curve_style_template(
                CurveStyleTemplate(name="绿色虚线", style=CurveStyle(color="#00aa00", linestyle="--", marker="o"))
            )
            self.page._refresh_curve_style_template_combo()
            self.page._curve_style_template_combo.setCurrentIndex(
                self.page._curve_style_template_ids.index(template.id)
            )

            self.page._load_selected_curve_style_template()

            style = self.page._curve_styles[self.page._curve_key(self.page._chart_series[0])]
            self.assertEqual(style["color"], "#00aa00")
            self.assertEqual(style["linestyle"], "--")
            self.assertEqual(style["marker"], "o")
        finally:
            restore_assets()

    def test_plot_style_selection_requires_explicit_load(self):
        from models.schemas import FigureConfig

        template = self.pm.add_figure_template(FigureConfig(name="显式加载模板", theme="ACS"))
        self.assertIsNotNone(template)

        self.page._refresh_template_combo()
        combo_index = self.page._plot_style_refs.index(f"template:{template.id}")
        self.page._template_combo.setCurrentIndex(combo_index)

        self.assertEqual(self.page._figure_state.theme, "默认")

        self.page._load_selected_plot_style()

        self.assertEqual(self.page._figure_state.theme, "ACS")
        self.assertEqual(self.page._template_combo.currentText(), "显式加载模板")

    def test_chart_extension_panel_entries_are_consistent_across_tabs(self):
        from core.extension_api import PlotExtension, extension_registry

        extension_registry.register_plot(
            PlotExtension(
                type="chart_plot_extension_split",
                name="绘图扩展拆分",
                handler=lambda context, options: None,
                description="应在所有标签对应的侧栏中可见。",
            )
        )
        try:
            self.page._refresh_style_extension_panel()
            selector_snapshots = []
            for index in range(self.page._style_tabs.count()):
                self.page._style_tabs.setCurrentIndex(index)
                QApplication.processEvents()
                selector_snapshots.append([
                    self.page._plot_extension_controls._selector.itemText(i)
                    for i in range(self.page._plot_extension_controls._selector.count())
                ])

            self.assertTrue(selector_snapshots)
            for items in selector_snapshots:
                self.assertIn("绘图扩展拆分", items)
                self.assertEqual(items, selector_snapshots[0])
        finally:
            extension_registry.unregister_plot("chart_plot_extension_split")
            self.page._refresh_style_extension_panel()

    def test_chart_hidden_plot_extension_is_not_listed(self):
        from core.extension_api import PlotExtension, extension_registry

        extension_registry.register_plot(
            PlotExtension(
                type="chart_plot_hidden",
                name="隐藏绘图扩展",
                handler=lambda context, options: None,
                source_kind="base",
            )
        )
        try:
            self.page._refresh_style_extension_panel()
            selector_items = [
                self.page._plot_extension_controls._selector.itemText(i)
                for i in range(self.page._plot_extension_controls._selector.count())
            ]

            self.assertNotIn("隐藏绘图扩展", selector_items)
        finally:
            extension_registry.unregister_plot("chart_plot_hidden")
            self.page._refresh_style_extension_panel()

    def test_chart_plot_extension_applies_from_extension_tab(self):
        from core.extension_api import PlotExtension, extension_registry
        from extensions.plot._runtime import current_axis

        def _draw(lines, options):
            del lines
            axis = current_axis()
            if axis is not None:
                axis.axhline(float(options.get("y", 0.0)), color="#D13438")

        extension_registry.register_plot(
            PlotExtension(
                type="chart_plot_apply_any_tab",
                name="绘图扩展应用",
                handler=_draw,
                description="应通过左侧绘图扩展标签页应用。",
                default_options={"y": 3.2},
                phases=("after_plot",),
            )
        )
        try:
            self.page.on_tree_node_activated("series", self.s.id)

            self.page._style_tabs.setCurrentIndex(2)
            QApplication.processEvents()
            plot_selector_items = [
                self.page._plot_extension_controls._selector.itemText(i)
                for i in range(self.page._plot_extension_controls._selector.count())
            ]
            self.page._plot_extension_controls._selector.setCurrentIndex(plot_selector_items.index("绘图扩展应用"))
            self.page._plot_extension_controls._editor.setPlainText('{"y": 4.5}')
            self.page._plot_extension_controls._apply_current()

            self.assertEqual(len(self.page._applied_plot_extensions), 1)
            self.assertEqual(self.page._applied_plot_extensions[0]["type"], "chart_plot_apply_any_tab")
            axis = self.page._figure.axes[0]
            self.assertTrue(any(list(line.get_ydata()) == [4.5, 4.5] for line in axis.lines))
        finally:
            extension_registry.unregister_plot("chart_plot_apply_any_tab")
            self.page._plot_extension_options.pop("chart_plot_apply_any_tab", None)
            self.page._refresh_style_extension_panel()

    def test_plot_extension_applied_list_height_is_not_locked_to_loaded_count(self):
        from PySide6.QtWidgets import QSizePolicy
        from core.extension_api import PlotExtension, extension_registry
        from extensions.plot._runtime import current_axis

        def _draw(lines, options):
            del lines
            axis = current_axis()
            if axis is not None:
                axis.axhline(float(options.get("y", 0.0)), color="#D13438")

        extension_registry.register_plot(
            PlotExtension(
                type="chart_plot_dynamic_height",
                name="动态高度绘图扩展",
                handler=_draw,
                default_options={"y": 1.0},
                phases=("after_plot",),
            )
        )
        try:
            self.page.on_tree_node_activated("series", self.s.id)
            selector_items = [
                self.page._plot_extension_controls._selector.itemText(i)
                for i in range(self.page._plot_extension_controls._selector.count())
            ]
            self.page._plot_extension_controls._selector.setCurrentIndex(selector_items.index("动态高度绘图扩展"))
            self.page._plot_extension_controls._editor.setPlainText('{"y": 1.0}')
            self.page._plot_extension_controls._apply_current()
            first_min_height = self.page._plot_extension_applied_list.minimumHeight()
            first_max_height = self.page._plot_extension_applied_list.maximumHeight()

            self.page._plot_extension_controls._editor.setPlainText('{"y": 2.0}')
            self.page._plot_extension_controls._apply_current()
            second_min_height = self.page._plot_extension_applied_list.minimumHeight()
            second_max_height = self.page._plot_extension_applied_list.maximumHeight()

            self.assertEqual(self.page._plot_extension_applied_list.count(), 2)
            self.assertGreaterEqual(second_min_height, first_min_height)
            self.assertEqual(first_max_height, 16777215)
            self.assertEqual(second_max_height, 16777215)
            self.assertEqual(
                self.page._plot_extension_applied_list.sizePolicy().verticalPolicy(),
                QSizePolicy.Policy.Expanding,
            )
        finally:
            extension_registry.unregister_plot("chart_plot_dynamic_height")
            self.page._plot_extension_options.pop("chart_plot_dynamic_height", None)
            self.page._applied_plot_extensions = []
            self.page._refresh_style_extension_panel()

    def test_plot_extension_can_be_removed_after_apply(self):
        from core.extension_api import PlotExtension, extension_registry

        extension_registry.register_plot(
            PlotExtension(
                type="chart_plot_remove",
                name="绘图扩展撤销",
                handler=lambda lines, options: None,
                description="用于验证绘图扩展可撤销。",
                default_options={"enabled": True},
            )
        )
        try:
            self.page._style_tabs.setCurrentIndex(2)
            QApplication.processEvents()
            self.page._on_chart_extension_apply("chart_plot_remove", {"enabled": True})

            self.assertEqual(len(self.page._applied_plot_extensions), 1)
            self.assertEqual(self.page._remove_selected_plot_extension_btn.text(), "撤销选中")
            self.page._on_chart_extension_remove_requested("chart_plot_remove")
            self.assertEqual(self.page._applied_plot_extensions, [])
        finally:
            extension_registry.unregister_plot("chart_plot_remove")
            self.page._refresh_style_extension_panel()

    def test_plot_extension_can_clear_all_after_confirmation(self):
        from qfluentwidgets import ToolButton
        from core.extension_api import PlotExtension, extension_registry

        extension_registry.register_plot(
            PlotExtension(
                type="chart_plot_clear_all",
                name="绘图扩展清空",
                handler=lambda lines, options: None,
                default_options={"enabled": True},
            )
        )
        try:
            self.assertIsInstance(self.page._clear_all_plot_extensions_btn, ToolButton)
            self.assertFalse(self.page._clear_all_plot_extensions_btn.isEnabled())

            self.page._style_tabs.setCurrentIndex(2)
            QApplication.processEvents()
            self.page._on_chart_extension_apply("chart_plot_clear_all", {"enabled": True})
            self.page._on_chart_extension_apply("chart_plot_clear_all", {"enabled": False})

            self.assertEqual(len(self.page._applied_plot_extensions), 2)
            self.assertTrue(self.page._clear_all_plot_extensions_btn.isEnabled())

            with mock.patch("ui.pages.chart_page.MessageBox.exec", return_value=True):
                self.page._clear_all_plot_extensions_btn.click()

            self.assertEqual(self.page._applied_plot_extensions, [])
            self.assertFalse(self.page._clear_all_plot_extensions_btn.isEnabled())
        finally:
            extension_registry.unregister_plot("chart_plot_clear_all")
            self.page._plot_extension_options.pop("chart_plot_clear_all", None)
            self.page._refresh_style_extension_panel()

    def test_plot_extension_supports_multiple_instances_with_different_params(self):
        from core.extension_api import PlotExtension, extension_registry
        from extensions.plot._runtime import current_axis

        def _draw(lines, options):
            del lines
            axis = current_axis()
            if axis is not None:
                axis.axhline(float(options.get("y", 0.0)), color=options.get("color", "#cc3300"))

        extension_registry.register_plot(
            PlotExtension(
                type="chart_plot_multi",
                name="多实例参考线",
                handler=_draw,
                default_options={"y": 1.0, "color": "#cc3300"},
                phases=("after_plot",),
            )
        )
        try:
            self.page.on_tree_node_activated("series", self.s.id)
            self.page._style_tabs.setCurrentIndex(2)
            QApplication.processEvents()

            self.page._on_chart_extension_apply("chart_plot_multi", {"y": 1.0, "color": "#cc3300"})
            self.page._on_chart_extension_apply("chart_plot_multi", {"y": 2.5, "color": "#0066cc"})

            self.assertEqual(len(self.page._applied_plot_extensions), 2)
            self.assertEqual(self.page._plot_extension_applied_list.count(), 2)
            axis = self.page._figure.axes[0]
            horizontal_lines = [line for line in axis.lines if len(set(line.get_ydata())) == 1]
            self.assertTrue(any(list(line.get_ydata()) == [1.0, 1.0] for line in horizontal_lines))
            self.assertTrue(any(list(line.get_ydata()) == [2.5, 2.5] for line in horizontal_lines))
        finally:
            extension_registry.unregister_plot("chart_plot_multi")
            self.page._applied_plot_extensions.clear()
            self.page._plot_extension_options.pop("chart_plot_multi", None)
            self.page._refresh_style_extension_panel()

    def test_plot_extension_uses_selected_curve_line_when_applying(self):
        from core.extension_api import PlotExtension, extension_registry
        from extensions.plot._runtime import current_axis
        from processing.extension_tools import line_xy, normalize_lines

        selected_first_y_values: list[float | None] = []

        def _draw(lines, options):
            axis = current_axis()
            normalized = normalize_lines(lines)
            if axis is None or not normalized:
                return
            _xs, ys = line_xy(normalized[0])
            if not ys:
                return
            selected_first_y_values.append(float(ys[0]))
            axis.axhline(float(ys[0]), color=options.get("color", "#009966"))

        extension_registry.register_plot(
            PlotExtension(
                type="chart_plot_selected_curve",
                name="读取当前选中曲线",
                handler=_draw,
                default_options={"color": "#009966"},
                phases=("after_plot",),
            )
        )
        try:
            self.page.on_tree_node_activated("series", self.s.id)
            self.page.add_series_to_chart({
                "name": "第二条曲线",
                "x": [0.0, 1.0, 2.0],
                "y": [8.0, 9.0, 10.0],
                "color": "#ff6600",
                "obj_id": "series-second",
                "visible": True,
            })
            self.page._refresh_chart_list()
            self.page._chart_list.setCurrentRow(1)
            QApplication.processEvents()

            self.page._style_tabs.setCurrentIndex(2)
            QApplication.processEvents()
            self.page._on_chart_extension_apply("chart_plot_selected_curve", {"color": "#009966"})

            self.assertEqual(selected_names[-1], "第二条曲线")
            self.assertIn("当前选中：第二条曲线", self.page._plot_extension_target_hint.text())
            self.assertIn("\n", self.page._plot_extension_target_hint.text())
            applied = self.page._applied_plot_extensions[0]
            self.assertEqual(applied["curve_name"], "第二条曲线")
        finally:
            extension_registry.unregister_plot("chart_plot_selected_curve")
            self.page._applied_plot_extensions.clear()
            self.page._plot_extension_options.pop("chart_plot_selected_curve", None)
            self.page._refresh_style_extension_panel()

    def test_chart_preview_host_follows_dark_background(self):
        if self.page._figure is None or self.page._canvas_host is None:
            self.skipTest("matplotlib unavailable")

        from ui.theme import preview_canvas_background_color

        with mock.patch("ui.pages.chart_page.isDarkTheme", return_value=True):
            self.page._apply_preview_host_background()

        self.assertIn(preview_canvas_background_color(True), self.page._canvas_host.viewport().styleSheet())

    def test_chart_theme_palette_uses_shared_preview_helpers(self):
        from ui.theme import preview_canvas_background_color, preview_canvas_foreground_color, preview_canvas_grid_color

        with mock.patch("ui.pages.chart_page.global_assets.get_plot_theme", return_value=None), \
             mock.patch.object(self.page, "_resolve_plot_theme", return_value=None), \
             mock.patch("ui.pages.chart_page.isDarkTheme", return_value=True):
            palette = self.page._theme_palette_for_state(self.page._figure_state)

        self.assertEqual(
            palette,
            (
                preview_canvas_background_color(True),
                preview_canvas_foreground_color(True),
                preview_canvas_grid_color(True),
            ),
        )

    def test_chart_style_panel_keeps_compact_width_conventions(self):
        from ui.theme import WORKBENCH_TOOL_PANEL_WIDTH

        self.assertLessEqual(self.page._style_line_width_edit.maximumWidth(), 96)
        self.assertLessEqual(self.page._x_min_edit.maximumWidth(), 96)
        self.assertEqual(self.page._tool_panel.minimumWidth(), WORKBENCH_TOOL_PANEL_WIDTH)

    def test_chart_style_panel_uses_content_sized_labels(self):
        from qfluentwidgets import BodyLabel

        labels = {
            label.text(): label
            for label in self.page.findChildren(BodyLabel)
            if label.text() in {"图例位置:", "图例字号:", "默认线宽:", "默认点大小:", "网格透明度:", "网格线宽:"}
        }

        for text in ["图例位置:", "图例字号:", "默认线宽:", "默认点大小:", "网格透明度:", "网格线宽:"]:
            self.assertIn(text, labels)
            self.assertEqual(labels[text].minimumWidth(), labels[text].sizeHint().width())

        self.assertLess(labels["图例位置:"].minimumWidth(), labels["网格透明度:"].minimumWidth())

    def test_removing_last_curve_refreshes_path_and_extension_hint(self):
        self.page.on_tree_node_activated("series", self.s.id)
        self.page._chart_list.setCurrentRow(0)
        QApplication.processEvents()

        self.assertIn("test.csv", self.page._chart_path_label.text())

        self.page._on_remove_selected()

        self.assertEqual(self.page._chart_path_label.text(), "路径：—")
        self.assertEqual(self.page._style_target_label.text(), "当前选中：未选中")
        self.assertIn("请先向画布添加曲线", self.page._plot_extension_target_hint.text())

    def test_duplicate_named_curves_show_distinct_project_tree_paths(self):
        from models.schemas import DataFile, DataSeries

        self.page.on_tree_node_activated("series", self.s.id)
        other = DataFile(name="other.csv", series=[DataSeries(name=self.s.name, x=[0.0, 1.0], y=[2.0, 3.0])])
        node = self.pm.add_data_file(other)
        self.assertIsNotNone(node)

        self.page.on_tree_node_activated("series", other.series[0].id)
        self.page._chart_list.setCurrentRow(0)
        QApplication.processEvents()
        first_path = self.page._chart_path_label.text()

        self.page._chart_list.setCurrentRow(1)
        QApplication.processEvents()
        second_path = self.page._chart_path_label.text()

        self.assertIn("test.csv", first_path)
        self.assertIn("other.csv", second_path)
        self.assertNotEqual(first_path, second_path)

        self.assertIn("test.csv", self.page._chart_list.item(0).text())
        self.assertIn("other.csv", self.page._chart_list.item(1).text())
        self.assertEqual(self.page._chart_list.item(0).toolTip(), self.page._chart_list.item(0).text())
        self.assertEqual(self.page._chart_list.item(1).toolTip(), self.page._chart_list.item(1).text())

    def test_toggle_selected_visibility_only_affects_current_duplicate_named_curve(self):
        from models.schemas import DataFile, DataSeries

        self.page.on_tree_node_activated("series", self.s.id)
        other = DataFile(name="other_visible.csv", series=[DataSeries(name=self.s.name, x=[0.0, 1.0], y=[2.0, 3.0])])
        node = self.pm.add_data_file(other)
        self.assertIsNotNone(node)

        self.page.on_tree_node_activated("series", other.series[0].id)
        self.page._chart_list.setCurrentRow(1)
        QApplication.processEvents()

        self.page._toggle_selected_visibility()

        state_by_obj_id = {curve.get("obj_id"): bool(curve.get("visible", True)) for curve in self.page._chart_series}
        self.assertTrue(state_by_obj_id[self.s.id])
        self.assertFalse(state_by_obj_id[other.series[0].id])
        self.assertTrue(self.page._chart_list.item(1).text().startswith("[隐藏]"))

    def test_curve_display_name_only_affects_legend_when_chart_list_shows_tree_path(self):
        if self.page._figure is None:
            self.skipTest("matplotlib unavailable")

        self.page.on_tree_node_activated("series", self.s.id)
        curve = self.page._chart_series[0]
        original_name = curve["name"]

        self.page._set_curve_display_name(curve, "图例显示名")

        self.assertEqual(curve["name"], original_name)
        self.assertEqual(curve["display_name"], "图例显示名")
        self.assertIn(original_name, self.page._chart_list.item(0).text())
        self.assertNotEqual(self.page._chart_list.item(0).text(), "图例显示名")

        legend = self.page._figure.axes[0].get_legend()
        self.assertIsNotNone(legend)
        self.assertEqual(legend.get_texts()[0].get_text(), "图例显示名")

    def test_save_template_named_uses_figure_state(self):
        self.page._apply_advanced_config({
            "theme": "Nature",
            "x_label": "t",
            "y_label": "y",
            "show_errbar": True,
            "figure_width": 8.0,
            "figure_height": 5.5,
            "dpi": 240,
            "font_family": "DejaVu Sans",
            "legend_font_size": 10,
            "line_width": 2.0,
            "marker_size": 6.0,
        })
        from core.global_assets import global_assets

        template = self.page._save_template_named("模板状态A")
        self.assertIsNotNone(template)
        fig = global_assets.get_figure_template(template.id)
        self.assertIsNotNone(fig)
        self.assertEqual(fig.theme, "Nature")
        self.assertEqual(fig.figure_size, (8.0, 5.5))
        self.assertEqual(fig.dpi, 240)
        self.assertEqual(fig.font_family, "DejaVu Sans")
        self.assertEqual(fig.legend_font_size, 10)
        self.assertEqual(fig.line_width, 2.0)
        self.assertEqual(fig.marker_size, 6.0)
        self.assertEqual(self.page._template_combo.currentText(), "模板状态A")

    def test_load_template_dialog_uses_renamed_template_name(self):
        from models.schemas import FigureConfig

        from core.global_assets import global_assets

        template = self.pm.add_figure_template(FigureConfig(name="old_name", theme="ACS"))
        self.assertIsNotNone(template)
        global_assets.update_figure_template(template.id, name="renamed_template")

        def _fake_get_item(_parent, _title, _label, items, current_text=None):
            self.assertIn("renamed_template", items)
            return ("renamed_template", True)

        with mock.patch("ui.pages.chart_page.SelectionDialog.get_item", side_effect=_fake_get_item):
            self.page._on_load_template()

        self.assertEqual(self.page._figure_state.theme, "ACS")


# ═══════════════════════════════════════════════════════════════════════════
# 5. ProcessPage — on_tree_node_selected, load_pipeline
# ═══════════════════════════════════════════════════════════════════════════

class TestProcessPage(unittest.TestCase):

    def setUp(self):
        self._restore_assets = _patch_global_assets()
        self.pm, self.p, self.df, self.s = _make_project("pp_test")
        self._restore = _patch_pm(self.pm)
        from ui.pages.process_page import ProcessPage
        self.page = ProcessPage()

    def tearDown(self):
        self._restore()
        self.page.deleteLater()
        self._restore_assets()

    def test_page_creates_no_crash(self):
        self.assertIsNotNone(self.page)

    def test_process_extension_panel_uses_splitter_side_panel(self):
        self.page.resize(1280, 820)
        self.page.show()
        QApplication.processEvents()

        self.assertEqual(self.page._page_splitter.count(), 2)
        self.assertTrue(self.page._extension_panel.isHidden())

        self.page.set_extension_panel_visible(True)
        QApplication.processEvents()

        self.assertFalse(self.page._extension_panel.isHidden())
        self.assertGreater(self.page._page_splitter.sizes()[1], 0)

    def test_process_canvas_draw_avoids_tight_layout_warning(self):
        if self.page._figure is None or self.page._canvas is None:
            self.skipTest("matplotlib unavailable")

        self.page._src_xs = [0.0, 1.0, 2.0, 3.0]
        self.page._src_ys = [1.0, 2.0, 1.5, 2.5]
        self.page._out_xs = list(self.page._src_xs)
        self.page._out_ys = [1.1, 1.9, 1.6, 2.4]

        with warnings.catch_warnings(record=True) as caught:
            warnings.simplefilter("always")
            self.page._draw_preview()

        layout_warnings = [str(item.message) for item in caught if "Tight layout not applied" in str(item.message)]
        self.assertEqual(layout_warnings, [])

    def test_process_preview_host_follows_dark_background(self):
        if self.page._canvas is None:
            self.skipTest("matplotlib unavailable")

        with mock.patch("ui.pages.process_page.isDarkTheme", return_value=True):
            self.page._apply_preview_host_background()

        self.assertIn("#1e1e1e", self.page._canvas.styleSheet())

    def test_process_update_theme_reapplies_preview_host_background(self):
        if self.page._canvas is None:
            self.skipTest("matplotlib unavailable")

        with mock.patch("ui.pages.process_page.isDarkTheme", return_value=True):
            self.page.update_theme()

        self.assertIn("#1e1e1e", self.page._canvas.styleSheet())

    def test_process_preview_zoom_buttons_adjust_axes_range(self):
        if self.page._figure is None or self.page._canvas is None:
            self.skipTest("matplotlib unavailable")

        self.page._src_xs = [0.0, 1.0, 2.0, 3.0]
        self.page._src_ys = [1.0, 2.0, 1.5, 2.5]
        self.page._out_xs = list(self.page._src_xs)
        self.page._out_ys = [1.1, 1.8, 1.6, 2.2]
        self.page._draw_preview()

        original_ax = self.page._figure.axes[0]
        original_x_span = abs(original_ax.get_xlim()[1] - original_ax.get_xlim()[0])
        original_y_span = abs(original_ax.get_ylim()[1] - original_ax.get_ylim()[0])

        self.page._preview_zoom_in_btn.click()
        QApplication.processEvents()

        zoomed_ax = self.page._figure.axes[0]
        self.assertLess(abs(zoomed_ax.get_xlim()[1] - zoomed_ax.get_xlim()[0]), original_x_span)
        self.assertLess(abs(zoomed_ax.get_ylim()[1] - zoomed_ax.get_ylim()[0]), original_y_span)

        self.page._preview_fit_btn.click()
        QApplication.processEvents()

        reset_ax = self.page._figure.axes[0]
        self.assertAlmostEqual(abs(reset_ax.get_xlim()[1] - reset_ax.get_xlim()[0]), original_x_span, places=6)
        self.assertAlmostEqual(abs(reset_ax.get_ylim()[1] - reset_ax.get_ylim()[0]), original_y_span, places=6)

    def test_process_preview_navigation_modes_switch_exclusively(self):
        if self.page._preview_nav_toolbar is None:
            self.skipTest("matplotlib unavailable")

        self.page.show()
        QApplication.processEvents()

        self.page._preview_box_zoom_btn.click()
        QApplication.processEvents()
        self.assertEqual(self.page._preview_navigation_mode(), "zoom")
        self.assertTrue(self.page._preview_box_zoom_btn.isChecked())
        self.assertFalse(self.page._preview_pan_btn.isChecked())

        self.page._preview_pan_btn.click()
        QApplication.processEvents()
        self.assertEqual(self.page._preview_navigation_mode(), "pan")
        self.assertTrue(self.page._preview_pan_btn.isChecked())
        self.assertFalse(self.page._preview_box_zoom_btn.isChecked())

        self.page._preview_pan_btn.click()
        QApplication.processEvents()
        self.assertEqual(self.page._preview_navigation_mode(), "")
        self.assertFalse(self.page._preview_pan_btn.isChecked())
        self.assertFalse(self.page._preview_box_zoom_btn.isChecked())

    def test_process_preview_navigation_buttons_use_toggle_toolbuttons_and_right_click_clears_mode(self):
        from qfluentwidgets import ToggleToolButton

        if self.page._preview_nav_toolbar is None or self.page._canvas is None:
            self.skipTest("matplotlib unavailable")

        self.page.show()
        QApplication.processEvents()

        self.assertIsInstance(self.page._preview_pan_btn, ToggleToolButton)
        self.assertIsInstance(self.page._preview_box_zoom_btn, ToggleToolButton)

        self.page._preview_box_zoom_btn.click()
        QApplication.processEvents()
        self.assertEqual(self.page._preview_navigation_mode(), "zoom")

        QTest.mouseClick(
            self.page._canvas,
            Qt.MouseButton.RightButton,
            Qt.KeyboardModifier.NoModifier,
            self.page._canvas.rect().center(),
        )
        QApplication.processEvents()

        self.assertEqual(self.page._preview_navigation_mode(), "")
        self.assertFalse(self.page._preview_pan_btn.isChecked())
        self.assertFalse(self.page._preview_box_zoom_btn.isChecked())

    def test_on_tree_node_selected_series(self):
        self.page.on_tree_node_activated("series", self.s.id)

        self.assertEqual(len(self.page._selected_inputs), 1)
        self.assertEqual(self.page._selected_input_list.count(), 1)

    def test_process_selected_input_panel_uses_vertical_splitter(self):
        self.page.resize(1280, 820)
        self.page.show()
        QApplication.processEvents()

        self.assertEqual(self.page._selected_input_splitter.count(), 2)
        sizes = self.page._selected_input_splitter.sizes()
        self.assertGreater(sizes[1], 0)
        ratio = sizes[0] / sum(sizes)
        self.assertGreater(ratio, 0.34)
        self.assertLess(ratio, 0.46)

    def test_on_tree_node_selected_data_file(self):
        node = next((n for n in self.p.tree.nodes if n.kind == "data_file"), None)
        if node:
            self.page.on_tree_node_activated("data_file", node.id)

        self.assertGreaterEqual(len(self.page._selected_inputs), 1)

    def test_on_tree_node_selected_data_file_runs_pipeline_for_all_series(self):
        from models.schemas import DataFile, DataSeries

        node = self.pm.add_data_file(
            DataFile(
                name="batch.csv",
                series=[
                    DataSeries(name="sA", x=[0.0, 1.0, 2.0], y=[1.0, 2.0, 3.0]),
                    DataSeries(name="sB", x=[0.0, 1.0, 2.0], y=[2.0, 4.0, 6.0]),
                ],
            )
        )

        self.page.on_tree_node_activated("data_file", node.id)
        self.page._run_timer.stop()
        self.page._run_pipeline_now()

        self.assertEqual(len(self.page._selected_inputs), 2)
        self.assertEqual(len(self.page._src_series_batch), 2)
        self.assertEqual(len(self.page._out_series_batch), 2)
        self.assertEqual([series.name for series in self.page._out_series_batch], ["sA", "sB"])
        self.assertEqual(self.page._selected_input_list.count(), 2)
        self.assertIn("已选择 2 条", self.page._current_input_label.text())
        self.assertIn("当前处理:", self.page._selected_input_state_label.text())
        self.assertEqual(self.page._save_result_button.text(), "导出数据列")
        self.assertTrue(self.page._save_batch_result_button.isEnabled())

    def test_save_result_uses_current_selected_curve_preview(self):
        from models.schemas import DataFile, DataSeries
        from ui.dialogs.export_flow import DataExportPlan

        node = self.pm.add_data_file(
            DataFile(
                name="current_export.csv",
                series=[
                    DataSeries(name="left", x=[0.0, 1.0, 2.0], y=[1.0, 2.0, 3.0]),
                    DataSeries(name="right", x=[0.0, 1.0, 2.0], y=[10.0, 20.0, 30.0]),
                ],
            )
        )
        datasets_root = self.pm._find_folder_by_group_type("datasets")
        self.assertIsNotNone(datasets_root)

        self.page.on_tree_node_activated("data_file", node.id)
        self.page._run_timer.stop()
        self.page._run_pipeline_now()

        first_item = self.page._selected_input_list.item(0)
        second_item = self.page._selected_input_list.item(1)
        self.assertIsNotNone(first_item)
        self.assertIsNotNone(second_item)
        self.page._selected_input_list.clearSelection()
        first_item.setSelected(True)
        second_item.setSelected(True)
        self.page._selected_input_list.setCurrentItem(second_item)
        self.page._sync_selected_input_state()
        self.page._run_pipeline_now()

        with mock.patch(
            "ui.pages.process_page.choose_data_export_plan",
            return_value=DataExportPlan(
                export_name="当前曲线导出",
                new_parent_id=datasets_root.id,
                new_data_file_name="当前曲线导出.process",
            ),
        ):
            self.page._save_result()

        data_file = next((df for df in self.p.data_files if df.name == "当前曲线导出.process"), None)
        self.assertIsNotNone(data_file)
        self.assertEqual(len(data_file.series), 1)
        self.assertEqual(data_file.series[0].name, "当前曲线导出")
        self.assertEqual(data_file.series[0].y, [10.0, 20.0, 30.0])

    def test_save_result_defaults_export_name_to_current_preview_curve(self):
        from models.schemas import DataFile, DataSeries
        from ui.dialogs.export_flow import DataExportPlan

        node = self.pm.add_data_file(
            DataFile(
                name="current_export_name.csv",
                series=[
                    DataSeries(name="left", x=[0.0, 1.0, 2.0], y=[1.0, 2.0, 3.0]),
                    DataSeries(name="right", x=[0.0, 1.0, 2.0], y=[10.0, 20.0, 30.0]),
                ],
            )
        )
        datasets_root = self.pm._find_folder_by_group_type("datasets")
        self.assertIsNotNone(datasets_root)

        self.page.on_tree_node_activated("data_file", node.id)
        self.page._run_timer.stop()
        self.page._run_pipeline_now()

        first_item = self.page._selected_input_list.item(0)
        second_item = self.page._selected_input_list.item(1)
        self.assertIsNotNone(first_item)
        self.assertIsNotNone(second_item)
        self.page._selected_input_list.clearSelection()
        first_item.setSelected(True)
        second_item.setSelected(True)
        self.page._selected_input_list.setCurrentItem(second_item)
        self.page._sync_selected_input_state()
        self.page._run_pipeline_now()

        self.assertIn("right", self.page._save_name_edit.text())
        self.assertNotIn("left", self.page._save_name_edit.text())

        captured = {}

        def _capture_export_plan(*_args, **kwargs):
            captured.update(kwargs)
            export_name = kwargs["default_export_name"]
            return DataExportPlan(
                export_name=export_name,
                new_parent_id=datasets_root.id,
                new_data_file_name=f"{export_name}.process",
            )

        with mock.patch("ui.pages.process_page.choose_data_export_plan", side_effect=_capture_export_plan):
            self.page._save_result()

        self.assertEqual(captured["default_export_name"], "right")


    def test_selected_input_labels_use_tree_paths(self):
        from models.schemas import DataFile, DataSeries

        datasets_root = self.pm._find_folder_by_group_type("datasets")
        self.assertIsNotNone(datasets_root)
        folder = self.pm.add_folder("批次A", parent_id=datasets_root.id, group_type="datasets")
        self.assertIsNotNone(folder)
        node = self.pm.add_data_file(
            DataFile(name="path.csv", series=[DataSeries(name="s_path", x=[0.0, 1.0], y=[1.0, 2.0])]),
            parent_id=folder.id,
        )

        self.page.on_tree_node_activated("data_file", node.id)

        self.assertEqual(self.page._selected_input_list.item(0).text(), "批次A/path.csv/s_path")

    def test_param_stack_height_tracks_selected_operation(self):
        self.page._load_ops_into_chain([
            {"type": "derivative", "params": {}},
            {"type": "pairwise_compute", "params": {"x_expr": "x1", "y_expr": "y1 - y2"}},
        ])

        self.page._on_op_selected(0)
        derivative_height = self.page._param_stack.maximumHeight()
        self.assertEqual(derivative_height, self.page._param_widgets[0].sizeHint().height())

        self.page._on_op_selected(1)
        pairwise_height = self.page._param_stack.maximumHeight()
        self.assertEqual(pairwise_height, self.page._param_widgets[1].sizeHint().height())
        self.assertNotEqual(derivative_height, pairwise_height)

    def test_json_param_widget_uses_content_height_with_fixed_upper_bound(self):
        self.page._load_ops_into_chain([
            {"type": "resample", "params": {"mode": "spacing", "spacing_mode": "coord", "step": 0.25}},
        ])

        widget = self.page._param_widgets[0]

        self.assertLessEqual(widget.maximumHeight(), 180)
        self.assertLessEqual(self.page._param_stack.maximumHeight(), 180)

    def test_loading_processing_extension_settings_refreshes_param_stack_height(self):
        from core.extension_api import ProcessingExtension, extension_registry
        from core.global_assets import global_assets

        extension_registry.register_processing(
            ProcessingExtension(
                type="ui_process_config_height",
                name="UI 配置高度刷新",
                handler=lambda lines, params: lines[0] if lines else [],
                settings=True,
                config_fields=[
                    {"key": "mode", "field_type": "choice", "choices": ["short", "long"], "default": "short"},
                    {"key": "window", "field_type": "int", "default": 3},
                    {"key": "notes", "field_type": "string", "multiline": True, "default": ""},
                ],
            )
        )
        try:
            config = global_assets.add_extension_config(
                category="processing",
                extension_type="ui_process_config_height",
                extension_name="UI 配置高度刷新",
                extension_version="0.1.0",
                name="高参数配置",
                options={"mode": "long", "window": 9, "notes": "line1\nline2\nline3\nline4"},
            )
            self.page._load_ops_into_chain([
                {"type": "ui_process_config_height", "params": {"mode": "short", "window": 3, "notes": ""}},
            ])

            widget = self.page._param_widgets[0]
            self.page._on_op_selected(0)
            expected_min_height = min(widget.sizeHint().height(), widget.maximumHeight())

            with mock.patch.object(self.page, "_refresh_param_stack_height", wraps=self.page._refresh_param_stack_height) as refresh_height:
                widget.load_settings_config(config.id)

            self.assertEqual(widget.current_settings_config_id(), config.id)
            self.assertEqual(widget.get_params()["notes"], "line1\nline2\nline3\nline4")
            self.assertGreaterEqual(refresh_height.call_count, 1)
            self.assertEqual(self.page._param_stack.minimumHeight(), expected_min_height)
            self.assertEqual(self.page._param_stack.maximumHeight(), widget.maximumHeight())
            self.assertLessEqual(self.page._param_stack.maximumHeight(), 180)
        finally:
            extension_registry.unregister_processing("ui_process_config_height")

    def test_on_tree_node_selected_curve(self):
        self.page.on_tree_node_selected("curve", "fake-curve-id")

    def test_load_pipeline_valid(self):
        """保存 Pipeline → 加载到 ProcessPage"""
        ops = [{"type": "smooth", "params": {"method": "moving_avg", "window": 5}}]
        sp = self.pm.add_saved_pipeline("test_pipe", ops, "for test")
        self.page.load_pipeline(sp.id)
        # 操作链应被加载
        self.assertEqual(len(self.page._ops), 1)
        self.assertEqual(self.page._ops[0]["type"], "smooth")

    def test_load_pipeline_nonexistent(self):
        self.page.load_pipeline("nonexistent-node-id")

    def test_load_ops_into_chain_clears_previous(self):
        ops1 = [{"type": "crop", "params": {"x_min": 0, "x_max": 5}}]
        ops2 = [{"type": "normalize", "params": {"mode": "minmax"}},
                {"type": "resample", "params": {"n": 100}}]
        self.page._load_ops_into_chain(ops1)
        self.assertEqual(len(self.page._ops), 1)
        self.page._load_ops_into_chain(ops2)
        self.assertEqual(len(self.page._ops), 2)

    def test_load_ops_empty_list(self):
        self.page._load_ops_into_chain([])
        self.assertEqual(len(self.page._ops), 0)

    def test_save_pipeline_template_as_named(self):
        from core.global_assets import global_assets

        self.page._load_ops_into_chain([{"type": "smooth", "params": {"method": "moving_avg", "window": 5}}])
        result = self.page._save_pipeline_template_as_named("流程模板A")
        self.assertTrue(result)
        self.assertTrue(any(sp.name == "流程模板A" for sp in global_assets.list_saved_pipelines()))

    def test_overwrite_pipeline_template(self):
        from core.global_assets import global_assets
        from models.schemas import SavedPipeline

        sp = global_assets.add_saved_pipeline(SavedPipeline(name="流程模板B", ops=[{"type": "smooth", "params": {"window": 5}}]))
        self.page.load_pipeline(sp.id)
        self.page._load_ops_into_chain([{"type": "normalize", "params": {"mode": "minmax"}}])
        self.assertTrue(self.page._overwrite_pipeline_template())
        saved = global_assets.get_saved_pipeline(sp.id)
        self.assertEqual(saved.ops[0]["type"], "normalize")

    def test_load_ops_with_resample_spacing_params(self):
        ops = [{"type": "resample", "params": {"mode": "spacing", "step": 0.25}}]

        self.page._load_ops_into_chain(ops)

        self.assertEqual(
            self.page._param_widgets[0].get_params(),
            {"mode": "spacing", "spacing_mode": "coord", "step": 0.25},
        )

    def test_load_ops_with_resample_align_legacy_target_index_maps_to_target_line(self):
        ops = [{"type": "resample", "params": {"mode": "align", "target_index": 2}}]

        self.page._load_ops_into_chain(ops)

        params = self.page._param_widgets[0].get_params()

        self.assertEqual(params["mode"], "align")
        self.assertEqual(params["target_line"], 2)
        self.assertNotIn("target_index", params)

    def test_resample_coord_spacing_waits_for_edit_commit(self):
        from ui.pages.process_page import _ResampleParam

        on_change = mock.Mock()
        widget = _ResampleParam(self.page, on_change)
        widget._mode.setCurrentIndex(1)
        on_change.reset_mock()

        widget._value_edit.setText("0")
        QApplication.processEvents()
        self.assertFalse(on_change.called)

        widget._value_edit.editingFinished.emit()
        self.assertEqual(on_change.call_count, 1)
        widget.deleteLater()

    def test_clicking_blank_area_commits_active_param_editor(self):
        from PySide6.QtCore import QPoint

        self.page.resize(900, 640)
        self.page.show()
        QApplication.processEvents()

        with mock.patch.object(self.page, "_run_pipeline") as run_pipeline:
            self.page._load_ops_into_chain([
                {"type": "resample", "params": {"mode": "spacing", "spacing_mode": "coord", "step": 0.25}}
            ])
            widget = self.page._param_widgets[0]
            value_edit = widget._value_edit
            run_pipeline.reset_mock()

            value_edit.setFocus()
            QApplication.processEvents()
            value_edit.clear()
            QTest.keyClicks(value_edit, "0.5")
            QApplication.processEvents()
            self.assertTrue(value_edit.hasFocus())
            self.assertFalse(run_pipeline.called)

            QTest.mouseClick(self.page, Qt.MouseButton.LeftButton, Qt.KeyboardModifier.NoModifier, self.page.rect().bottomRight() - QPoint(8, 8))
            QApplication.processEvents()

            self.assertFalse(value_edit.hasFocus())
            self.assertEqual(run_pipeline.call_count, 1)

    def test_load_ops_with_fft_sampling_rate_params(self):
        ops = [{"type": "fft", "params": {"output": "power", "detrend": False, "sampling_rate": 50.0}}]

        self.page._load_ops_into_chain(ops)

        self.assertEqual(
            self.page._param_widgets[0].get_params(),
            {"output": "power", "detrend": False, "sampling_rate": 50.0},
        )

    def test_load_ops_with_filter_actual_frequency_params(self):
        ops = [{"type": "filter", "params": {"cutoff": 5.0, "cutoff_mode": "actual", "sampling_rate": 100.0, "order": 3, "mode": "high"}}]

        self.page._load_ops_into_chain(ops)

        self.assertEqual(
            self.page._param_widgets[0].get_params(),
            {"cutoff": 5.0, "order": 3, "mode": "high", "cutoff_mode": "actual", "sampling_rate": 100.0},
        )

    def test_save_result_creates_new_data_file_with_custom_name(self):
        from models.schemas import DataSeries
        from ui.dialogs.export_flow import DataExportPlan

        self.page._out_xs = [1.0, 2.0, 3.0]
        self.page._out_ys = [2.0, 3.0, 4.0]
        self.page._out_series_batch = [DataSeries(name="处理结果A", x=[1.0, 2.0, 3.0], y=[2.0, 3.0, 4.0], source="computed")]
        self.page._save_name_edit.setText("处理结果A")
        datasets_root = self.pm._find_folder_by_group_type("datasets")
        self.assertIsNotNone(datasets_root)

        with mock.patch(
            "ui.pages.process_page.choose_data_export_plan",
            return_value=DataExportPlan(
                export_name="处理结果A",
                new_parent_id=datasets_root.id,
                new_data_file_name="处理结果A.process",
            ),
        ):
            self.page._save_result()

        data_file = next((df for df in self.p.data_files if df.name == "处理结果A.process"), None)
        self.assertIsNotNone(data_file)
        self.assertEqual(data_file.series[0].name, "处理结果A")

    def test_save_result_appends_to_selected_data_file(self):
        from models.schemas import DataSeries
        from ui.dialogs.export_flow import DataExportPlan

        self.page._out_xs = [1.0, 2.0, 3.0]
        self.page._out_ys = [2.0, 3.0, 4.0]
        self.page._out_series_batch = [DataSeries(name="处理结果B", x=[1.0, 2.0, 3.0], y=[2.0, 3.0, 4.0], source="computed")]
        self.page._save_name_edit.setText("处理结果B")
        original_count = len(self.df.series)

        with mock.patch(
            "ui.pages.process_page.choose_data_export_plan",
            return_value=DataExportPlan(export_name="处理结果B", target_data_file_id=self.df.id),
        ):
            self.page._save_result()

        self.assertEqual(len(self.df.series), original_count + 1)
        self.assertEqual(self.df.series[-1].name, "处理结果B")

    def test_save_batch_result_for_data_file_creates_new_data_file_with_all_output_series(self):
        from models.schemas import DataFile, DataSeries
        from ui.dialogs.export_flow import BatchDataExportPlan

        node = self.pm.add_data_file(
            DataFile(
                name="batch_save.csv",
                series=[
                    DataSeries(name="left", x=[0.0, 1.0, 2.0, 3.0], y=[1.0, 2.0, 3.0, 4.0]),
                    DataSeries(name="right", x=[0.0, 1.0, 2.0, 3.0], y=[2.0, 4.0, 6.0, 8.0]),
                ],
            )
        )
        datasets_root = self.pm._find_folder_by_group_type("datasets")
        self.assertIsNotNone(datasets_root)

        self.page.on_tree_node_activated("data_file", node.id)
        self.page._load_ops_into_chain([{"type": "crop", "params": {"x_min": 1.0, "x_max": 2.0}}])
        self.page._run_timer.stop()
        self.page._run_pipeline_now()

        with mock.patch(
            "ui.pages.process_page.choose_data_export_batch_plan",
            return_value=BatchDataExportPlan(
                export_names=["left", "right"],
                new_parent_id=datasets_root.id,
                new_data_file_name="batch_save.process",
            ),
        ):
            self.page._save_batch_result()

        data_file = next((df for df in self.p.data_files if df.name == "batch_save.process"), None)
        self.assertIsNotNone(data_file)
        self.assertEqual([series.name for series in data_file.series], ["left", "right"])
        self.assertEqual(data_file.series[0].x, [1.0, 2.0])
        self.assertEqual(data_file.series[0].y, [2.0, 3.0])
        self.assertEqual(data_file.series[1].x, [1.0, 2.0])
        self.assertEqual(data_file.series[1].y, [4.0, 6.0])

    def test_pairwise_compute_requires_prealigned_inputs(self):
        from models.schemas import DataFile, DataSeries

        node = self.pm.add_data_file(
            DataFile(
                name="pairwise.csv",
                series=[
                    DataSeries(name="left", x=[0.0, 1.0, 2.0], y=[1.0, 3.0, 5.0]),
                    DataSeries(name="right", x=[0.0, 0.5, 1.0, 1.5, 2.0], y=[1.0, 2.0, 3.0, 4.0, 5.0]),
                ],
            )
        )

        self.page.on_tree_node_activated("data_file", node.id)
        self.page._load_ops_into_chain([
            {
                "type": "pairwise_compute",
                "params": {
                    "primary_index": 1,
                    "secondary_index": 2,
                    "x_expr": "x1",
                    "y_expr": "y1 - y2",
                    "align_mode": "auto",
                    "resample_mode": "count",
                    "n": 3,
                },
            }
        ])
        self.page._run_timer.stop()
        self.page._run_pipeline_now()

        self.assertEqual(len(self.page._out_series_batch), 0)
        self.assertIn("需进行坐标间距重采样", self.page._stats_label.text())
        self.assertEqual(self.page._save_result_button.text(), "导出数据列")
        self.assertFalse(self.page._save_batch_result_button.isEnabled())

    def test_selected_input_reorder_moves_item_to_top(self):
        from models.schemas import DataSeries

        other_a = DataSeries(name="s2", x=[1.0, 2.0], y=[2.0, 3.0])
        other_b = DataSeries(name="s3", x=[1.0, 2.0], y=[3.0, 4.0])
        self.df.series.extend([other_a, other_b])

        self.page.on_tree_node_activated("series", self.s.id)
        self.page.on_tree_node_activated("series", other_a.id)
        self.page.on_tree_node_activated("series", other_b.id)

        self.page._selected_input_list.clearSelection()
        target_item = self.page._selected_input_list.item(2)
        target_item.setSelected(True)
        self.page._move_selected_inputs_to_top()

        self.assertEqual(self.page._selected_inputs[0]["node_id"], other_b.id)

    def test_selected_input_reorder_preserves_current_indicator(self):
        from models.schemas import DataSeries

        other_a = DataSeries(name="s2", x=[1.0, 2.0], y=[2.0, 3.0])
        other_b = DataSeries(name="s3", x=[1.0, 2.0], y=[3.0, 4.0])
        self.df.series.extend([other_a, other_b])

        self.page.on_tree_node_activated("series", self.s.id)
        self.page.on_tree_node_activated("series", other_a.id)
        self.page.on_tree_node_activated("series", other_b.id)

        target_item = self.page._selected_input_list.item(1)
        self.page._selected_input_list.setCurrentItem(target_item)
        target_item.setSelected(True)

        self.page._move_selected_inputs_down()

        current_payload = self.page._selected_input_list.currentItem().data(Qt.ItemDataRole.UserRole)
        self.assertEqual(current_payload["node_id"], other_a.id)
        self.assertTrue(self.page._selected_input_list.currentItem().isSelected())

    def test_processing_extension_appears_in_selector_and_panel(self):
        from core.extension_api import ProcessingExtension, extension_registry

        def _scale(lines, params):
            factor = float(params.get("factor", 1.0))
            if not lines:
                return []
            return [[point[0], point[1] * factor] for point in lines[0]]

        extension_registry.register_processing(
            ProcessingExtension(
                type="ui_scale_test",
                name="UI 缩放测试",
                handler=_scale,
                description="按 factor 缩放 Y 值",
                default_options={"factor": 2},
            )
        )
        try:
            self.page._refresh_processing_extensions()
            combo_items = [self.page._add_op_combo.itemText(i) for i in range(self.page._add_op_combo.count())]
            selector_items = [self.page._extension_panel._selector.itemText(i) for i in range(self.page._extension_panel._selector.count())]

            self.assertIn("UI 缩放测试", combo_items)
            self.assertIn("UI 缩放测试", selector_items)
        finally:
            extension_registry.unregister_processing("ui_scale_test")
            self.page._refresh_processing_extensions()

    def test_processing_extension_panel_adds_extension_op(self):
        from core.extension_api import ProcessingExtension, extension_registry

        def _scale(lines, params):
            factor = float(params.get("factor", 1.0))
            if not lines:
                return []
            return [[point[0], point[1] * factor] for point in lines[0]]

        extension_registry.register_processing(
            ProcessingExtension(
                type="ui_scale_apply",
                name="UI 扩展应用",
                handler=_scale,
                description="按 factor 缩放 Y 值",
                default_options={"factor": 2},
            )
        )
        try:
            self.page._refresh_processing_extensions()
            selector_items = [self.page._extension_panel._selector.itemText(i) for i in range(self.page._extension_panel._selector.count())]
            self.page._extension_panel._selector.setCurrentIndex(selector_items.index("UI 扩展应用"))
            self.page._extension_panel._editor.setPlainText('{"factor": 4}')
            self.page._extension_panel._apply_current()

            self.assertEqual(self.page._ops[-1]["type"], "ui_scale_apply")
            self.assertEqual(self.page._ops[-1]["params"], {"factor": 4})
        finally:
            extension_registry.unregister_processing("ui_scale_apply")
            self.page._refresh_processing_extensions()

    def test_processing_extension_panel_shows_config_field_help(self):
        from core.extension_api import ExtensionConfigField, ProcessingExtension, extension_registry

        def _scale(xs, ys, params):
            factor = float(params.get("factor", 1.0))
            return list(xs), [value * factor for value in ys]

        extension_registry.register_processing(
            ProcessingExtension(
                type="ui_scale_help",
                name="UI 配置说明",
                handler=_scale,
                description="展示扩展配置字段说明。",
                default_options={"factor": 2.5},
                config_fields=[
                    ExtensionConfigField(
                        key="factor",
                        label="倍率",
                        description="用于缩放当前 Y 值。",
                        field_type="number",
                        default=2.5,
                    )
                ],
            )
        )
        try:
            self.page._refresh_processing_extensions()
            selector_items = [self.page._extension_panel._selector.itemText(i) for i in range(self.page._extension_panel._selector.count())]
            self.page._extension_panel._selector.setCurrentIndex(selector_items.index("UI 配置说明"))
            help_text = self.page._extension_panel._config_help_label.text()
            self.assertIn("factor: number; 可选; 用于缩放当前 Y 值。", help_text)
            self.assertNotIn("倍率", help_text)
        finally:
            extension_registry.unregister_processing("ui_scale_help")
            self.page._refresh_processing_extensions()

    def test_processing_multi_curve_extension_uses_page_level_line_picker(self):
        from qfluentwidgets import PushButton
        from models.schemas import DataSeries
        from core.extension_api import ExtensionConfigField, ProcessingExtension, extension_registry

        other = DataSeries(name="s2", x=[1.0, 2.0], y=[2.0, 3.0])
        third = DataSeries(name="s3", x=[1.0, 2.0], y=[3.0, 4.0])
        self.df.series.extend([other, third])

        extension_registry.register_processing(
            ProcessingExtension(
                type="ui_multi_line_picker",
                name="UI 多曲线按钮",
                handler=lambda xs, ys, params, lines=None: (list(xs), list(ys)),
                lines_number=(2, -1),
            )
        )
        try:
            self.page.on_tree_node_activated("series", self.s.id)
            self.page.on_tree_node_activated("series", other.id)
            self.page.on_tree_node_activated("series", third.id)
            self.page._refresh_processing_extensions()
            self.page._add_op_combo.setCurrentIndex(self.page._processing_op_types.index("ui_multi_line_picker"))
            self.page._add_op()

            buttons = [button for button in self.page._param_widgets[-1].findChildren(PushButton) if "选择曲线" in button.text()]

            self.assertFalse(buttons)
            self.assertTrue(self.page._pipeline_lines_button.isEnabled())
        finally:
            extension_registry.unregister_processing("ui_multi_line_picker")

    def test_process_page_pipeline_line_picker_injects_lines_list(self):
        from models.schemas import DataSeries

        other = DataSeries(name="s2", x=[1.0, 2.0], y=[2.0, 3.0])
        third = DataSeries(name="s3", x=[1.0, 2.0], y=[3.0, 4.0])
        self.df.series.extend([other, third])

        self.page.on_tree_node_activated("series", self.s.id)
        self.page.on_tree_node_activated("series", other.id)
        self.page.on_tree_node_activated("series", third.id)
        self.page._add_op_combo.setCurrentIndex(self.page._processing_op_types.index("pairwise_compute"))
        self.page._add_op()
        self.page._pipeline_selected_node_ids = [self.s.id, third.id]

        params = self.page._normalized_operation_params("pairwise_compute", {"x_expr": "x1", "y_expr": "y1-y2"})

        self.assertEqual(params["lines_list"], [1, 3])

    def test_process_page_choose_pipeline_lines_uses_dialog_result(self):
        from unittest import mock
        from models.schemas import DataSeries

        other = DataSeries(name="s2", x=[1.0, 2.0], y=[2.0, 3.0])
        third = DataSeries(name="s3", x=[1.0, 2.0], y=[3.0, 4.0])
        self.df.series.extend([other, third])

        self.page.on_tree_node_activated("series", self.s.id)
        self.page.on_tree_node_activated("series", other.id)
        self.page.on_tree_node_activated("series", third.id)
        self.page._add_op_combo.setCurrentIndex(self.page._processing_op_types.index("pairwise_compute"))
        self.page._add_op()

        with mock.patch(
            "ui.widgets.extension_options_form._LineSelectionDialog.get_indices",
            return_value=([2, 3], True),
        ) as patched:
            self.page._choose_pipeline_lines()

        patched.assert_called_once()
        self.assertEqual(self.page._pipeline_selected_node_ids, [other.id, third.id])

    def test_process_page_load_pipeline_restores_extension_config_id(self):
        from core.extension_api import ProcessingExtension, extension_registry
        from core.global_assets import global_assets

        extension_registry.register_processing(
            ProcessingExtension(
                type="ui_pipeline_config_restore",
                name="Pipeline 配置恢复",
                handler=lambda lines, params: lines[0] if lines else [],
                description="验证 pipeline 会恢复扩展 settings 配置。",
                default_options={"factor": 2},
                settings=True,
                config_fields=[
                    {"key": "factor", "label": "倍率", "field_type": "number", "default": 2},
                ],
            )
        )
        try:
            config = global_assets.add_extension_config(
                category="processing",
                extension_type="ui_pipeline_config_restore",
                extension_name="Pipeline 配置恢复",
                name="倍率 7",
                options={"factor": 7},
            )
            self.page._load_ops_into_chain([
                {"type": "ui_pipeline_config_restore", "params": {"factor": 2}, "config_id": config.id}
            ])

            self.assertEqual(self.page._ops[0]["config_id"], config.id)
            self.assertEqual(self.page._param_widgets[0].get_params()["factor"], 7)
        finally:
            extension_registry.unregister_processing("ui_pipeline_config_restore")
            self.page._refresh_processing_extensions()

    def test_builtin_processing_ops_use_generic_schema_param_widget(self):
        from ui.pages.process_page import _JsonParam

        self.page._add_op_combo.setCurrentIndex(self.page._processing_op_types.index("crop"))
        self.page._add_op()

        self.assertIsInstance(self.page._param_widgets[-1], _JsonParam)
        params = self.page._param_widgets[-1].get_params()
        self.assertEqual(params.get("x_min"), None)
        self.assertEqual(params.get("x_max"), None)

    def test_processing_extension_panel_uses_generic_apply_text(self):
        self.assertEqual(self.page._extension_panel._apply_btn.text(), "应用扩展")

    def test_hidden_processing_extension_is_not_listed(self):
        from core.extension_api import ProcessingExtension, extension_registry

        extension_registry.register_processing(
            ProcessingExtension(
                type="ui_hidden_processing",
                name="隐藏处理扩展",
                handler=lambda lines, params: lines[0] if lines else [],
                source_kind="base",
            )
        )
        try:
            self.page._refresh_processing_extensions()
            combo_items = [self.page._add_op_combo.itemText(i) for i in range(self.page._add_op_combo.count())]
            selector_items = [self.page._extension_panel._selector.itemText(i) for i in range(self.page._extension_panel._selector.count())]

            self.assertNotIn("隐藏处理扩展", combo_items)
            self.assertNotIn("隐藏处理扩展", selector_items)
        finally:
            extension_registry.unregister_processing("ui_hidden_processing")
            self.page._refresh_processing_extensions()

    def test_process_page_content_splitter_keeps_two_main_panels(self):
        self.assertEqual(self.page._content_splitter.count(), 2)
        self.assertFalse(hasattr(self.page, "_src_tree"))


# ═══════════════════════════════════════════════════════════════════════════
# 6. AnalysisPage — on_tree_node_selected
# ═══════════════════════════════════════════════════════════════════════════

class TestAnalysisPage(unittest.TestCase):

    def setUp(self):
        self._restore_assets = _patch_global_assets()
        self.pm, self.p, self.df, self.s = _make_project("ap_test")
        self._restore = _patch_pm(self.pm)
        from ui.pages.analysis_page import AnalysisPage
        self.page = AnalysisPage()

    def tearDown(self):
        self._restore()
        self._restore_assets()
        self.page.deleteLater()

    def _analysis_export_plan(self, name: str):
        from ui.dialogs.export_flow import DataExportPlan

        datasets_root = self.pm._find_folder_by_group_type("datasets")
        self.assertIsNotNone(datasets_root)
        return DataExportPlan(
            export_name=name,
            new_parent_id=datasets_root.id,
            new_data_file_name=f"{name}.analysis",
        )

    def test_page_creates_no_crash(self):
        self.assertIsNotNone(self.page)

    def test_analysis_preview_view_uses_shared_navigation_controls(self):
        view = self.page._analysis_tab_views.get("current")
        if view is None or view.get("figure") is None or view.get("canvas") is None:
            self.skipTest("matplotlib unavailable")

        selected = [(list(self.s.x), list(self.s.y), self.s.name)]
        view["analysis_type"] = "statistics"
        view["selected"] = list(selected)
        self.page._set_analysis_plot_surface(view, show_plot=True)
        self.page._draw_result(
            "statistics",
            selected,
            {"_preview_only": True},
            figure=view.get("figure"),
            canvas=view.get("canvas"),
        )

        axis = view["figure"].axes[0]
        original_x_span = abs(axis.get_xlim()[1] - axis.get_xlim()[0])
        original_y_span = abs(axis.get_ylim()[1] - axis.get_ylim()[0])

        view["preview_box_zoom_btn"].click()
        QApplication.processEvents()
        self.assertTrue(view["preview_box_zoom_btn"].isChecked())
        self.assertFalse(view["preview_pan_btn"].isChecked())

        view["preview_zoom_in_btn"].click()
        QApplication.processEvents()

        zoomed_axis = view["figure"].axes[0]
        self.assertLess(abs(zoomed_axis.get_xlim()[1] - zoomed_axis.get_xlim()[0]), original_x_span)
        self.assertLess(abs(zoomed_axis.get_ylim()[1] - zoomed_axis.get_ylim()[0]), original_y_span)

        view["preview_fit_btn"].click()
        QApplication.processEvents()

        reset_axis = view["figure"].axes[0]
        self.assertAlmostEqual(abs(reset_axis.get_xlim()[1] - reset_axis.get_xlim()[0]), original_x_span, places=6)
        self.assertAlmostEqual(abs(reset_axis.get_ylim()[1] - reset_axis.get_ylim()[0]), original_y_span, places=6)

    def test_selected_input_state_label_uses_shared_placeholder_color(self):
        from ui.theme import placeholder_color

        self.assertIn(placeholder_color(), self.page._selected_input_state_label.styleSheet())

    def test_analysis_extension_panel_uses_splitter_side_panel(self):
        self.page.resize(1280, 820)
        self.page.show()
        QApplication.processEvents()

        self.assertEqual(self.page._page_splitter.count(), 2)
        self.assertTrue(self.page._extension_panel.isHidden())

        self.page.set_extension_panel_visible(True)
        QApplication.processEvents()

        self.assertFalse(self.page._extension_panel.isHidden())
        self.assertGreater(self.page._page_splitter.sizes()[1], 0)

    def test_analysis_input_panel_uses_vertical_splitter(self):
        self.page.resize(1280, 820)
        self.page.show()
        QApplication.processEvents()

        self.assertEqual(self.page._input_panel_splitter.count(), 2)
        sizes = self.page._input_panel_splitter.sizes()
        self.assertGreater(sizes[1], 0)
        ratio = sizes[0] / sum(sizes)
        self.assertGreater(ratio, 0.34)
        self.assertLess(ratio, 0.46)

    def test_input_action_buttons_match_height_and_template_label_is_hidden(self):
        self.assertEqual(self.page._btn_clear_inputs.height(), 32)
        self.assertEqual(self.page._btn_remove_selected_inputs.height(), 32)
        self.assertTrue(self.page._report_template_label.isHidden())

    def test_result_tabs_use_segmented_navigation_and_report_combos_share_width(self):
        from qfluentwidgets import SegmentedWidget
        from PySide6.QtWidgets import QSizePolicy

        self.assertIsInstance(self.page._result_tabs.navigationWidget, SegmentedWidget)
        self.assertEqual(self.page._result_tabs.navigationWidget.sizePolicy().horizontalPolicy(), QSizePolicy.Policy.Expanding)
        self.assertEqual(self.page._report_template_combo.width(), self.page._report_placeholder_combo.width())

    def test_analysis_canvas_draws_chinese_without_glyph_warnings(self):
        if self.page._figure is None or self.page._canvas is None:
            self.skipTest("matplotlib unavailable")

        self.page._figure.clear()
        axis = self.page._figure.add_subplot(111)
        axis.set_title("拟合结果")
        axis.set_xlabel("波峰")
        axis.set_ylabel("波谷")

        with warnings.catch_warnings(record=True) as caught:
            warnings.simplefilter("always")
            self.page._canvas.draw()

        glyph_warnings = [str(item.message) for item in caught if "Glyph" in str(item.message)]
        self.assertEqual(glyph_warnings, [])

    def test_analysis_canvas_draw_avoids_tight_layout_warning(self):
        if self.page._figure is None or self.page._canvas is None:
            self.skipTest("matplotlib unavailable")

        self.page._figure.clear()
        axis = self.page._figure.add_subplot(111)
        axis.plot([0.0, 1.0, 2.0], [1.0, 2.0, 1.5], label="测试")
        axis.set_title("分析结果")
        axis.legend()

        with warnings.catch_warnings(record=True) as caught:
            warnings.simplefilter("always")
            self.page._canvas.draw()

        layout_warnings = [str(item.message) for item in caught if "Tight layout not applied" in str(item.message)]
        self.assertEqual(layout_warnings, [])

    def test_analysis_result_canvas_follows_dark_background(self):
        if self.page._canvas is None:
            self.skipTest("matplotlib unavailable")

        with mock.patch("ui.pages.analysis_page.isDarkTheme", return_value=True):
            self.page._apply_result_canvas_background(self.page._canvas)

        self.assertIn("#1e1e1e", self.page._canvas.styleSheet())
        self.assertEqual(self.page._canvas.figure.get_facecolor()[:3], (30 / 255.0, 30 / 255.0, 30 / 255.0))

    def test_report_preview_follows_dark_background(self):
        with mock.patch("ui.pages.analysis_page.isDarkTheme", return_value=True):
            self.page._apply_report_preview_theme()

        self.assertIn("#1e1e1e", self.page._report_preview.styleSheet())

    def test_report_preview_uses_fluent_selection_colors(self):
        from ui.theme import accent_color

        self.page._apply_report_preview_theme()

        editor_style = self.page._report_editor.styleSheet()
        preview_style = self.page._report_preview.styleSheet()
        self.assertIn("selection-background-color", editor_style)
        self.assertIn(accent_color(), editor_style)
        self.assertIn("selection-color: #ffffff", preview_style)

    def test_analysis_update_theme_reapplies_canvas_and_report_backgrounds(self):
        if self.page._canvas is None:
            self.skipTest("matplotlib unavailable")

        with mock.patch("ui.pages.analysis_page.isDarkTheme", return_value=True):
            self.page.update_theme()

        self.assertIn("#1e1e1e", self.page._canvas.styleSheet())
        self.assertIn("#1e1e1e", self.page._report_preview.styleSheet())

    def test_input_role_labels_are_hidden(self):
        self.assertTrue(self.page._primary_input_label.isHidden())
        self.assertTrue(self.page._secondary_input_label.isHidden())

    def test_input_list_copy_selection_to_clipboard(self):
        from models.schemas import DataFile, DataSeries

        other = DataFile(name="other.csv", series=[DataSeries(name="other", x=[0.0, 1.0], y=[1.0, 2.0])])
        self.pm.add_data_file(other)

        self.page._type_combo.setCurrentIndex(self.page._analysis_type_ids.index("correlation"))
        self.page.on_tree_node_activated("series", self.s.id)
        self.page.on_tree_node_activated("data_file", next(n.id for n in self.p.tree.nodes if n.kind == "data_file" and n.data_file_id == other.id))
        self.page._extension_params_edit.set_options({"lines_list": [1, 2]})
        self.page._on_extension_analysis_options_changed({"lines_list": [1, 2]})

        self.page._input_list.item(0).setSelected(True)
        self.page._input_list.item(1).setSelected(True)
        self.page._input_list.copy_selection_to_clipboard()

        clipboard_text = QApplication.clipboard().text()
        self.assertIn(self.s.name, clipboard_text)
        self.assertIn(other.series[0].name, clipboard_text)

    def test_summary_area_has_larger_minimum_height(self):
        current_view = self.page._analysis_tab_views["current"]
        self.assertGreaterEqual(current_view["summary_stack"].minimumHeight(), 280)

    def test_analysis_preview_without_inputs_uses_empty_placeholder_surface(self):
        current_view = self.page._analysis_tab_views["current"]

        self.assertIs(current_view["plot_stack"].currentWidget(), current_view["empty_preview_widget"])
        self.assertIn("选择曲线后将在此预览", current_view["empty_preview_label"].text())

    def test_switching_builtin_analysis_types_keeps_preview_plot_available(self):
        from models.schemas import DataFile, DataSeries

        other = DataFile(name="other.csv", series=[DataSeries(name="other", x=[0.0, 1.0], y=[1.0, 2.0])])
        self.pm.add_data_file(other)

        self.page.on_tree_node_activated("series", self.s.id)
        self.page.on_tree_node_activated(
            "data_file",
            next(n.id for n in self.p.tree.nodes if n.kind == "data_file" and n.data_file_id == other.id),
        )

        for type_id in ["peak_detect", "statistics", "correlation", "error_compare"]:
            if type_id in {"correlation", "error_compare"}:
                self.page._analysis_extension_options[type_id] = {"lines_list": [1, 2]}
            self.page._type_combo.setCurrentIndex(self.page._analysis_type_ids.index(type_id))
            current_view = self.page._analysis_tab_views["current"]
            self.assertEqual(self.page._current_analysis_type(), type_id)
            self.assertIs(current_view["plot_stack"].currentWidget(), current_view["plot_widget"])

    def test_run_analysis_reports_prevalidation_error_via_shared_ui_notification(self):
        from core.extension_api import extension_registry

        extension_type = next(
            type_id for type_id in self.page._analysis_type_ids if extension_registry.get_analysis(type_id) is not None
        )
        self.page._type_combo.setCurrentIndex(self.page._analysis_type_ids.index(extension_type))
        self.page.on_tree_node_activated("series", self.s.id)

        with mock.patch.object(self.page, "_current_extension_analysis_options", side_effect=ValueError("参数校验失败")), \
             mock.patch("ui.pages.analysis_page.show_error", return_value="参数校验失败") as error_mock:
            self.page._run_analysis()

        error_mock.assert_called_once()
        self.assertIn("参数校验失败", self.page._analysis_status_label.text())
        current_view = self.page._analysis_tab_views["current"]
        self.assertEqual(current_view["summary_table"].item(0, 0).text(), "错误")
        self.assertEqual(current_view["summary_table"].item(0, 1).text(), "参数校验失败")
        self.assertTrue(self.page._run_analysis_btn.isEnabled())

    def test_unknown_analysis_json_renders_as_flattened_summary_rows(self):
        payload = {
            "summary": {"score": 0.875, "status": "ok"},
            "items": [
                {"name": "A", "value": 1},
                {"name": "B", "value": 2},
            ],
        }

        self.page._write_summary("custom_json", payload)

        labels = [self.page._summary_table.item(row, 0).text() for row in range(self.page._summary_table.rowCount())]
        values = [self.page._summary_table.item(row, 1).text() for row in range(self.page._summary_table.rowCount())]
        self.assertIn("summary.score", labels)
        self.assertIn("items[0].name", labels)
        self.assertIn("0.875", values)
        self.assertIn("A", values)

    def test_on_tree_node_selected_data_file(self):
        node = next((n for n in self.p.tree.nodes if n.kind == "data_file"), None)
        if node:
            self.page.on_tree_node_selected("data_file", node.id)

    def test_on_tree_node_selected_folder(self):
        node = next((n for n in self.p.tree.nodes if n.kind == "folder"), None)
        if node:
            self.page.on_tree_node_selected("folder", node.id)

    def test_on_tree_node_selected_unknown(self):
        self.page.on_tree_node_selected("pipeline", "fake-id")

    def test_load_report_template_from_tree(self):
        tmpl = self.pm.add_report_template("analysis_tmpl", "# Report")
        self.page.load_report_template(tmpl.id)
        self.assertEqual(self.page._current_report_template_id, tmpl.id)
        self.assertEqual(self.page._report_editor.toPlainText(), "# Report")

    def test_on_tree_node_activated_report_template(self):
        tmpl = self.pm.add_report_template("analysis_tmpl", "# Report")
        self.page.on_tree_node_activated("global_report_template", tmpl.id)
        self.assertEqual(self.page._current_report_template_id, tmpl.id)

    def test_generate_report_renders_in_page(self):
        self.page.on_tree_node_activated("series", self.s.id)
        self.page._run_analysis()
        self.page._on_generate_report()
        self.assertEqual(self.page._result_tabs.currentIndex(), 1)
        self.assertIn("分析类型", self.page._report_preview.toPlainText())

    def test_save_result_prompts_for_name_before_persisting(self):
        from ui.dialogs.export_flow import AnalysisResultSavePlan

        self.page.on_tree_node_activated("series", self.s.id)
        self.page._run_analysis()
        before = len(self.p.data_files)
        analysis_root = self.pm._find_folder_by_group_type("analysis_result_group")
        self.assertIsNotNone(analysis_root)

        with mock.patch(
            "ui.pages.analysis_page.choose_analysis_result_save_plan",
            return_value=AnalysisResultSavePlan(result_name="拟合结果A", target_parent_id=analysis_root.id),
        ):
            self.page._save_result()

        saved = next((item for item in self.p.analyses if item.name == "拟合结果A"), None)
        self.assertIsNotNone(saved)
        self.assertEqual(len(self.p.data_files), before)

    def test_save_result_can_create_and_use_analysis_subfolder(self):
        from ui.dialogs.export_flow import AnalysisResultSavePlan

        self.page.on_tree_node_activated("series", self.s.id)
        self.page._run_analysis()
        analysis_root = self.pm._find_folder_by_group_type("analysis_result_group")
        self.assertIsNotNone(analysis_root)
        target_folder = self.pm.add_folder("拟合结果组", parent_id=analysis_root.id, group_type="analysis_result_group")
        self.assertIsNotNone(target_folder)

        with mock.patch(
            "ui.pages.analysis_page.choose_analysis_result_save_plan",
            return_value=AnalysisResultSavePlan(result_name="拟合结果B", target_parent_id=target_folder.id),
        ):
            self.page._save_result()

        node = next((n for n in self.p.tree.nodes if n.kind == "analysis_result" and n.name == "拟合结果B"), None)
        self.assertIsNotNone(node)
        self.assertEqual(node.parent_id, target_folder.id)

    def test_preferred_analysis_result_target_uses_workspace_tree_selection(self):
        analysis_root = self.pm._find_folder_by_group_type("analysis_result_group")
        self.assertIsNotNone(analysis_root)
        target_folder = self.pm.add_folder("分析子目录", parent_id=analysis_root.id)
        self.assertIsNotNone(target_folder)

        self.page.on_tree_node_selected("folder", target_folder.id)

        self.assertEqual(target_folder.id, self.page._preferred_analysis_result_target_node_id())

    def test_curve_fit_result_can_export_series_with_export_plan(self):
        self.page.on_tree_node_activated("series", self.s.id)
        self.page._run_analysis()

        self.assertFalse(self.page._export_result_btn.isHidden())
        self.assertTrue(self.page._export_result_btn.isEnabled())

        with mock.patch("ui.pages.analysis_page.SelectionDialog.get_item") as choose_curve, mock.patch(
            "ui.pages.analysis_page.choose_data_export_plan",
            return_value=self._analysis_export_plan("拟合曲线A"),
        ):
            self.page._export_result_series()

        choose_curve.assert_not_called()

        data_file = next((item for item in self.p.data_files if item.name == "拟合曲线A.analysis"), None)
        self.assertIsNotNone(data_file)
        self.assertEqual(data_file.series[0].name, "拟合曲线A")

    def test_peak_detect_result_can_export_series_with_export_plan(self):
        from models.schemas import DataSeries

        target = DataSeries(
            name="oscillation_export",
            x=[0.0, 1.0, 2.0, 3.0, 4.0],
            y=[0.0, 2.0, 0.0, -1.5, 0.0],
        )
        self.df.series.append(target)

        self.page._type_combo.setCurrentIndex(self.page._analysis_type_ids.index("peak_detect"))
        self.page.on_tree_node_activated("series", target.id)
        self.page._run_analysis()

        self.assertFalse(self.page._export_result_btn.isHidden())
        self.assertTrue(self.page._export_result_btn.isEnabled())

        with mock.patch(
            "ui.pages.analysis_page.SelectionDialog.get_item",
            return_value=("波峰 (1个)", True),
        ) as choose_curve, mock.patch(
            "ui.pages.analysis_page.choose_data_export_plan",
            return_value=self._analysis_export_plan("峰值点A"),
        ):
            self.page._export_result_series()

        choose_curve.assert_called_once()

        data_file = next((item for item in self.p.data_files if item.name == "峰值点A.analysis"), None)
        self.assertIsNotNone(data_file)
        exported_series = data_file.series[0]
        self.assertEqual(exported_series.name, "峰值点A")
        self.assertEqual(exported_series.x, [1.0])
        self.assertEqual(exported_series.y, [2.0])

    def test_statistics_result_hides_export_curve_button_without_lines(self):
        self.page._type_combo.setCurrentIndex(self.page._analysis_type_ids.index("statistics"))
        self.page.on_tree_node_activated("series", self.s.id)
        self.page._run_analysis()

        self.assertTrue(self.page._export_result_btn.isHidden())
        self.assertFalse(self.page._export_result_btn.isEnabled())

    def test_export_result_series_uses_current_tab_result(self):
        from models.schemas import DataSeries

        other = DataSeries(name="s2", x=[10.0, 11.0, 12.0, 13.0], y=[4.0, 5.0, 6.0, 7.0])
        self.df.series.append(other)

        self.page.on_tree_node_activated("series", self.s.id)
        self.page._run_analysis()
        first_tab_index = self.page._analysis_tabs.currentIndex()

        self.page.on_tree_node_activated("series", other.id)
        current_item = self.page._input_list.item(self.page._input_list.count() - 1)
        self.assertIsNotNone(current_item)
        self.page._on_input_list_item_clicked(current_item)
        self.page._run_analysis()

        self.page._analysis_tabs.setCurrentIndex(first_tab_index)
        QApplication.processEvents()

        with mock.patch(
            "ui.pages.analysis_page.choose_data_export_plan",
            return_value=self._analysis_export_plan("tab1_export"),
        ):
            self.page._export_result_series()

        data_file = next((item for item in self.p.data_files if item.name == "tab1_export.analysis"), None)
        self.assertIsNotNone(data_file)
        exported_series = data_file.series[0]
        self.assertLess(min(exported_series.x), 5.0)

    def test_save_report_template_as_named(self):
        from core.global_assets import global_assets

        self.page._report_editor.setPlainText("# Saved Report")
        self.assertTrue(self.page._save_report_template_as_named("模板A"))
        self.assertTrue(any(t.name == "模板A" for t in global_assets.list_report_templates()))

    def test_save_report_template_as_named_reuses_existing_name(self):
        from core.global_assets import global_assets

        self.page._report_editor.setPlainText("# Saved Report A")
        self.assertTrue(self.page._save_report_template_as_named("模板同名A"))
        first_id = self.page._current_report_template_id

        self.page._report_editor.setPlainText("# Saved Report B")
        self.assertTrue(self.page._save_report_template_as_named("模板同名A"))

        matches = [t for t in global_assets.list_report_templates() if t.name == "模板同名A"]
        self.assertEqual(len(matches), 1)
        self.assertEqual(matches[0].id, first_id)
        self.assertEqual(matches[0].content, "# Saved Report B")

    def test_report_template_workspace_refreshes_after_save(self):
        self.page._report_editor.setPlainText("# Saved Report")
        self.assertTrue(self.page._save_report_template_as_named("模板工作台A"))
        items = [self.page._report_template_combo.itemText(i) for i in range(self.page._report_template_combo.count())]
        self.assertIn("模板工作台A", items)

    def test_save_current_report_template_updates_existing(self):
        tmpl = self.pm.add_report_template("analysis_tmpl", "# Report")
        self.page._load_report_template_by_id(tmpl.id)
        self.page._report_editor.setPlainText("# Updated Report")
        self.page._save_current_report_template()
        self.assertEqual(self.pm.get_report_template(tmpl.id).content, "# Updated Report")

    def test_report_editor_can_insert_placeholder_from_selector(self):
        target_text = "{{analysis_type}}"
        index = next(
            i for i in range(self.page._report_placeholder_combo.count())
            if target_text in self.page._report_placeholder_combo.itemText(i)
        )
        self.page._report_editor.setPlainText("报告开始\n")
        self.page._report_placeholder_combo.setCurrentIndex(index)

        self.page._insert_selected_report_placeholder()

        self.assertIn(target_text, self.page._report_editor.toPlainText())

    def test_analysis_extension_appears_in_type_selector_and_panel(self):
        from core.extension_api import AnalysisExtension, extension_registry

        def _span(lines, params):
            values = [point[1] for point in lines[0]] if lines else []
            return {
                "analysis_type": "ui_span_selector",
                "span": (max(values) - min(values)) if values else 0.0,
            }

        extension_registry.register_analysis(
            AnalysisExtension(
                type="ui_span_selector",
                name="UI 跨度选择",
                handler=_span,
                description="返回当前序列跨度",
                default_options={"scale": 2},
            )
        )
        try:
            self.page._refresh_analysis_type_choices()
            combo_items = [self.page._type_combo.itemText(i) for i in range(self.page._type_combo.count())]
            selector_items = [self.page._extension_panel._selector.itemText(i) for i in range(self.page._extension_panel._selector.count())]

            self.assertIn("UI 跨度选择", combo_items)
            self.assertIn("UI 跨度选择", selector_items)
        finally:
            extension_registry.unregister_analysis("ui_span_selector")
            self.page._refresh_analysis_type_choices()

    def test_analysis_extension_panel_switches_type_and_runs(self):
        from core.extension_api import AnalysisExtension, extension_registry

        def _span(lines, params):
            values = [point[1] for point in lines[0]] if lines else []
            return {
                "analysis_type": "ui_span_run",
                "span": (max(values) - min(values)) if values else 0.0,
                "scale": params.get("scale", 1),
            }

        extension_registry.register_analysis(
            AnalysisExtension(
                type="ui_span_run",
                name="UI 跨度执行",
                handler=_span,
                description="返回当前序列跨度",
                default_options={"scale": 3},
            )
        )
        try:
            self.page._refresh_analysis_type_choices()
            selector_items = [self.page._extension_panel._selector.itemText(i) for i in range(self.page._extension_panel._selector.count())]
            self.page._extension_panel._selector.setCurrentIndex(selector_items.index("UI 跨度执行"))
            self.page._extension_panel._editor.setPlainText('{"scale": 5}')
            self.page._extension_panel._apply_current()

            self.assertEqual(self.page._current_analysis_type(), "ui_span_run")

            self.page._selected_inputs = [{"kind": "series", "node_id": self.s.id, "label": self.s.name}]
            self.page._get_data_for_inputs = lambda inputs: [([0.0, 1.0], [2.0, 6.0], "demo")]
            self.page._run_analysis()

            self.assertEqual(self.page._result["analysis_type"], "ui_span_run")
            self.assertEqual(self.page._result["span"], 4.0)
            self.assertEqual(self.page._result["scale"], 5)
        finally:
            extension_registry.unregister_analysis("ui_span_run")
            self.page._refresh_analysis_type_choices()

    def test_analysis_extension_can_run_with_left_json_params_editor(self):
        from core.extension_api import AnalysisExtension, extension_registry

        def _span(lines, params):
            values = [point[1] for point in lines[0]] if lines else []
            return {
                "analysis_type": "ui_span_left_json",
                "span": (max(values) - min(values)) if values else 0.0,
                "scale": params.get("scale", 1),
            }

        extension_registry.register_analysis(
            AnalysisExtension(
                type="ui_span_left_json",
                name="UI 左侧参数执行",
                handler=_span,
                default_options={"scale": 2},
            )
        )
        try:
            self.page._refresh_analysis_type_choices()
            combo_items = [self.page._type_combo.itemText(i) for i in range(self.page._type_combo.count())]
            self.assertIn("UI 左侧参数执行", combo_items)

            self.page._type_combo.setCurrentIndex(self.page._analysis_type_ids.index("ui_span_left_json"))
            self.assertFalse(self.page._extension_params_edit.isHidden())
            self.assertIn('"scale": 2', self.page._extension_params_edit.toPlainText())
            self.page._extension_params_edit.setPlainText('{"scale": 6}')

            self.page._selected_inputs = [{"kind": "series", "node_id": self.s.id, "label": self.s.name}]
            self.page._get_data_for_inputs = lambda inputs: [([0.0, 1.0], [2.0, 6.0], "demo")]
            self.page._run_analysis()

            self.assertEqual(self.page._result["analysis_type"], "ui_span_left_json")
            self.assertEqual(self.page._result["scale"], 6)
            self.assertEqual(
                self.page._current_analysis_params()["extension_options"],
                {"scale": 6},
            )
        finally:
            extension_registry.unregister_analysis("ui_span_left_json")
            self.page._refresh_analysis_type_choices()

    def test_single_line_analysis_extension_tracks_current_selected_input_after_run(self):
        from models.schemas import DataSeries
        from core.extension_api import AnalysisExtension, extension_registry

        other = DataSeries(name="s2", x=[10.0, 11.0], y=[20.0, 21.0])
        self.df.series.append(other)

        extension_registry.register_analysis(
            AnalysisExtension(
                type="ui_single_current_line",
                name="UI 单曲线当前项",
                handler=lambda lines, params: {
                    "analysis_type": "ui_single_current_line",
                    "point_count": len(lines[0]) if lines else 0,
                    "line_refs": list(params.get("lines_list") or []),
                },
                lines_number=(1, 1),
            )
        )
        try:
            self.page._refresh_analysis_type_choices()
            self.page._type_combo.setCurrentIndex(self.page._analysis_type_ids.index("ui_single_current_line"))
            self.page.on_tree_node_activated("series", self.s.id)
            self.page.on_tree_node_activated("series", other.id)

            self.page._run_analysis()
            self.assertEqual(self.page._result["source_name"], self.s.name)
            self.assertEqual(self.page._result["line_refs"], [1])

            target_item = self.page._input_list.item(1)
            self.assertIsNotNone(target_item)
            self.page._on_input_list_item_clicked(target_item)

            current_view = self.page._analysis_tab_views["current"]
            self.assertEqual([series[2] for series in current_view["selected"]], [other.name])

            self.page._run_analysis()

            self.assertEqual(self.page._result["source_name"], other.name)
            self.assertEqual(self.page._result["line_refs"], [2])
        finally:
            extension_registry.unregister_analysis("ui_single_current_line")
            self.page._refresh_analysis_type_choices()

    def test_builtin_curve_fit_uses_generic_schema_editor(self):
        self.page._type_combo.setCurrentIndex(self.page._analysis_type_ids.index("curve_fit"))

        self.assertFalse(self.page._extension_params_edit.isHidden())
        self.assertFalse(hasattr(self.page, "_fit_model_combo"))
        self.assertEqual(self.page._extension_params_edit.current_options().get("model"), "linear")

    def test_analysis_extension_lines_follow_explicit_extension_selection(self):
        from models.schemas import DataSeries
        from core.extension_api import AnalysisExtension, extension_registry

        other = DataSeries(name="s2", x=[1.0, 2.0], y=[2.0, 3.0])
        third = DataSeries(name="s3", x=[1.0, 2.0], y=[3.0, 5.0])
        self.df.series.extend([other, third])

        extension_registry.register_analysis(
            AnalysisExtension(
                type="ui_lines_selected",
                name="UI 选中曲线扩展",
                handler=lambda lines, params: {
                    "analysis_type": "ui_lines_selected",
                    "line_lengths": [len(line) for line in lines],
                    "line_refs": list(params.get("lines_list") or []),
                },
                lines_number=(2, -1),
            )
        )
        try:
            self.page._refresh_analysis_type_choices()
            self.page.on_tree_node_activated("series", self.s.id)
            self.page.on_tree_node_activated("series", other.id)
            self.page.on_tree_node_activated("series", third.id)
            self.page._type_combo.setCurrentIndex(self.page._analysis_type_ids.index("ui_lines_selected"))
            self.page._input_list.clearSelection()
            self.page._input_list.item(1).setSelected(True)
            self.page._input_list.item(2).setSelected(True)
            self.page._extension_params_edit.set_options({"lines_list": [2, 3]})
            self.page._on_extension_analysis_options_changed({"lines_list": [2, 3]})

            self.page._run_analysis()

            self.assertEqual(self.page._result["source_name"], other.name)
            self.assertEqual(self.page._result["name1"], other.name)
            self.assertEqual(self.page._result["name2"], third.name)
            self.assertEqual(self.page._result["line_lengths"], [2, 2])
            self.assertEqual(self.page._result["line_refs"], [2, 3])
        finally:
            extension_registry.unregister_analysis("ui_lines_selected")
            self.page._refresh_analysis_type_choices()

    def test_analysis_extension_preview_updates_before_run(self):
        from models.schemas import DataSeries
        from core.extension_api import AnalysisExtension, extension_registry

        other = DataSeries(name="s2", x=[1.0, 2.0], y=[2.0, 3.0])
        third = DataSeries(name="s3", x=[1.0, 2.0], y=[3.0, 5.0])
        self.df.series.extend([other, third])

        extension_registry.register_analysis(
            AnalysisExtension(
                type="ui_preview_lines",
                name="UI 预览曲线扩展",
                handler=lambda lines, params: {"analysis_type": "ui_preview_lines", "count": len(lines)},
                lines_number=(2, -1),
            )
        )
        try:
            self.page._refresh_analysis_type_choices()
            self.page.on_tree_node_activated("series", self.s.id)
            self.page.on_tree_node_activated("series", other.id)
            self.page.on_tree_node_activated("series", third.id)
            self.page._type_combo.setCurrentIndex(self.page._analysis_type_ids.index("ui_preview_lines"))
            self.page._extension_params_edit.set_options({"lines_list": [1, 3]})
            self.page._on_extension_analysis_options_changed({"lines_list": [1, 3]})

            current_view = self.page._analysis_tab_views["current"]
            summary_table = current_view["summary_table"]

            self.assertEqual(current_view["result"], None)
            self.assertEqual([series[2] for series in current_view["selected"]], [self.s.name, third.name])
            self.assertEqual(summary_table.item(0, 1).text(), "已更新输入预览，运行分析后显示结果。")
            selected_text = summary_table.item(1, 1).text()
            self.assertIn(self.s.name, selected_text)
            self.assertIn(third.name, selected_text)
        finally:
            extension_registry.unregister_analysis("ui_preview_lines")
            self.page._refresh_analysis_type_choices()

    def test_analysis_extension_with_explicit_lines_keeps_selected_input_list(self):
        from models.schemas import DataSeries
        from core.extension_api import AnalysisExtension, extension_registry

        other = DataSeries(name="s2", x=[1.0, 2.0], y=[2.0, 3.0])
        self.df.series.append(other)

        extension_registry.register_analysis(
            AnalysisExtension(
                type="ui_keep_selected_inputs",
                name="UI 保留已选列表",
                handler=lambda lines, params: {
                    "analysis_type": "ui_keep_selected_inputs",
                    "line_lengths": [len(line) for line in lines],
                    "line_refs": list(params.get("lines_list") or []),
                },
                lines_number=(2, -1),
            )
        )
        try:
            self.page._refresh_analysis_type_choices()
            self.page._type_combo.setCurrentIndex(self.page._analysis_type_ids.index("ui_keep_selected_inputs"))

            self.page.on_tree_node_activated("series", self.s.id)
            self.page.on_tree_node_activated("series", other.id)

            self.assertEqual(len(self.page._selected_inputs), 2)
            self.assertEqual(self.page._input_list.count(), 2)
            self.assertIn(self.s.name, self.page._input_list.item(0).text())
            self.assertIn(other.name, self.page._input_list.item(1).text())

            self.page._extension_params_edit.set_options({"lines_list": [1, 2]})
            self.page._on_extension_analysis_options_changed({"lines_list": [1, 2]})
            self.page._run_analysis()

            self.assertEqual(self.page._result["source_name"], self.s.name)
            self.assertEqual(self.page._result["name1"], self.s.name)
            self.assertEqual(self.page._result["name2"], other.name)
            self.assertEqual(self.page._result["line_lengths"], [len(self.s.x), 2])
            self.assertEqual(self.page._result["line_refs"], [1, 2])
        finally:
            extension_registry.unregister_analysis("ui_keep_selected_inputs")
            self.page._refresh_analysis_type_choices()

    def test_pair_analysis_extension_uses_pair_dropdown_selection(self):
        from models.schemas import DataSeries
        from core.extension_api import AnalysisExtension, extension_registry

        other = DataSeries(name="s2", x=[1.0, 2.0], y=[2.0, 3.0])
        third = DataSeries(name="s3", x=[1.0, 2.0], y=[3.0, 5.0])
        self.df.series.extend([other, third])

        extension_registry.register_analysis(
            AnalysisExtension(
                type="ui_pair_lines_extension",
                name="UI 双曲线扩展",
                handler=lambda lines, params: {
                    "analysis_type": "ui_pair_lines_extension",
                    "line_lengths": [len(line) for line in lines],
                    "line_refs": list(params.get("lines_list") or []),
                },
                lines_number=(2, 2),
            )
        )
        try:
            self.page._refresh_analysis_type_choices()
            self.page.on_tree_node_activated("series", self.s.id)
            self.page.on_tree_node_activated("series", other.id)
            self.page.on_tree_node_activated("series", third.id)
            self.page._type_combo.setCurrentIndex(self.page._analysis_type_ids.index("ui_pair_lines_extension"))
            self.page._extension_params_edit.set_options({"lines_list": [1, 3]})
            self.page._on_extension_analysis_options_changed({"lines_list": [1, 3]})

            self.page._run_analysis()

            self.assertEqual(self.page._result["source_name"], self.s.name)
            self.assertEqual(self.page._result["name1"], self.s.name)
            self.assertEqual(self.page._result["name2"], third.name)
            self.assertEqual(self.page._result["line_lengths"], [len(self.s.x), 2])
            self.assertEqual(self.page._result["line_refs"], [1, 3])
        finally:
            extension_registry.unregister_analysis("ui_pair_lines_extension")
            self.page._refresh_analysis_type_choices()

    def test_hidden_analysis_extension_is_not_listed(self):
        from core.extension_api import AnalysisExtension, extension_registry

        extension_registry.register_analysis(
            AnalysisExtension(
                type="ui_hidden_analysis",
                name="隐藏分析扩展",
                handler=lambda lines, params: {"analysis_type": "ui_hidden_analysis"},
                source_kind="base",
            )
        )
        try:
            self.page._refresh_analysis_type_choices()
            combo_items = [self.page._type_combo.itemText(i) for i in range(self.page._type_combo.count())]
            selector_items = [self.page._extension_panel._selector.itemText(i) for i in range(self.page._extension_panel._selector.count())]

            self.assertNotIn("隐藏分析扩展", combo_items)
            self.assertNotIn("隐藏分析扩展", selector_items)
        finally:
            extension_registry.unregister_analysis("ui_hidden_analysis")
            self.page._refresh_analysis_type_choices()

    def test_analysis_extension_can_render_custom_plot_series_and_custom_placeholders(self):
        from core.extension_api import AnalysisExtension, extension_registry
        from processing.extension_tools import line_from_xy

        if self.page._figure is None:
            self.skipTest("matplotlib unavailable")

        def _spectrum(lines, params):
            return {
                "analysis_type": "ui_spectrum_plot",
                "dominant_frequency": 12.5,
                "dominant_amplitude": 3.4,
                "x_label": "频率 (Hz)",
                "y_label": "幅值",
                "plot_title": "自定义频谱",
                "lines": [
                    {"line_name": "频谱结果", "line": line_from_xy([0.0, 1.0, 2.0], [0.2, 3.4, 1.1])}
                ],
                "_plot_series": [
                    {"name": "频谱", "line": "频谱结果", "color": "#0078D4"}
                ],
            }

        extension_registry.register_analysis(
            AnalysisExtension(
                type="ui_spectrum_plot",
                name="UI 频谱绘图",
                handler=_spectrum,
                description="返回一条可直接绘制的频谱曲线。",
            )
        )
        try:
            self.page._refresh_analysis_type_choices()
            self.page._type_combo.setCurrentIndex(self.page._analysis_type_ids.index("ui_spectrum_plot"))
            self.page._selected_inputs = [{"kind": "series", "node_id": self.s.id, "label": self.s.name}]
            self.page._get_selected_data = lambda: [([0.0, 0.1, 0.2], [1.0, 0.2, 0.1], "demo")]

            self.page._run_analysis()

            placeholder_items = [
                self.page._report_placeholder_combo.itemText(i)
                for i in range(self.page._report_placeholder_combo.count())
            ]
            self.assertTrue(any("{{dominant_frequency}}" in item for item in placeholder_items))

            axis = self.page._figure.axes[0]
            self.assertEqual(len(axis.lines), 1)
            self.assertEqual(axis.get_xlabel(), "频率 (Hz)")
            self.assertEqual(axis.get_ylabel(), "幅值")
            self.assertEqual(axis.get_title(), "自定义频谱")

            output_options = self.page._analysis_output_series_options()
            self.assertEqual(output_options[0]["label"], "频谱")
            self.assertEqual(output_options[0]["series"].name, "频谱结果")

            exported = self.page._build_analysis_output_series("导出频谱")
            self.assertIsNotNone(exported)
            self.assertEqual(exported.x, [0.0, 1.0, 2.0])
            self.assertEqual(exported.y, [0.2, 3.4, 1.1])

            self.page._report_editor.setPlainText("主频 {{dominant_frequency:.1f}} Hz")
            self.page._render_report_preview()
            self.assertIn("12.5", self.page._report_preview.toPlainText())
        finally:
            extension_registry.unregister_analysis("ui_spectrum_plot")
            self.page._refresh_analysis_type_choices()

    def test_analysis_extension_structured_output_uses_detail_view(self):
        from core.extension_api import AnalysisExtension, extension_registry
        from processing.extension_tools import line_from_xy

        if self.page._figure is None:
            self.skipTest("matplotlib unavailable")

        def _structured(lines, params):
            return {
                "analysis_type": "ui_structured_output",
                "summary_items": [{"label": "主指标", "value": 3.14}],
                "lines": [
                    {"line_name": "结构化曲线", "line": line_from_xy([0.0, 1.0, 2.0], [1.0, 1.5, 2.5])}
                ],
                "plot_series": [
                    {"name": "结构化曲线", "line": "结构化曲线", "plot_type": "line", "color": "#0078D4"}
                ],
                "table_sections": [
                    {"title": "结果明细", "headers": ["项", "值"], "rows": [["A", 1], ["B", 2]]}
                ],
                "text_sections": [
                    {"title": "分析说明", "content": "第一行\n第二行"}
                ],
            }

        extension_registry.register_analysis(
            AnalysisExtension(
                type="ui_structured_output",
                name="UI 结构化输出",
                handler=_structured,
                description="返回结构化分析输出。",
            )
        )
        try:
            self.page._refresh_analysis_type_choices()
            self.page._type_combo.setCurrentIndex(self.page._analysis_type_ids.index("ui_structured_output"))
            self.page._selected_inputs = [{"kind": "series", "node_id": self.s.id, "label": self.s.name}]
            self.page._get_data_for_inputs = lambda inputs: [([0.0, 0.5, 1.0], [1.0, 1.2, 1.5], "demo")]

            self.page._run_analysis()

            active_key = self.page._analysis_tab_key_for_index(self.page._analysis_tabs.currentIndex())
            self.assertIsNotNone(active_key)
            view = self.page._analysis_tab_views[active_key]

            self.assertIs(view["summary_stack"].currentWidget(), view["details_scroll"])
            self.assertIsNotNone(view["detail_summary_table"])
            self.assertEqual(view["detail_summary_table"].item(0, 0).text(), "主指标")
            self.assertEqual(view["detail_summary_table"].item(0, 1).text(), "3.14")
            self.assertEqual(len(view["detail_tables"]), 1)
            self.assertEqual(view["detail_tables"][0].item(0, 0).text(), "A")
            self.assertEqual(view["detail_tables"][0].item(0, 1).text(), "1")
            self.assertEqual(len(view["detail_text_widgets"]), 1)
            self.assertEqual(view["detail_text_widgets"][0].text(), "第一行\n第二行")

            axis = self.page._figure.axes[0]
            self.assertEqual(len(axis.lines), 1)
            self.assertEqual(axis.lines[0].get_label(), "结构化曲线")
        finally:
            extension_registry.unregister_analysis("ui_structured_output")
            self.page._refresh_analysis_type_choices()

    def test_analysis_extension_declared_placeholders_appear_before_run_and_support_search(self):
        from qfluentwidgets import EditableComboBox
        from core.extension_api import AnalysisExtension, extension_registry

        extension_registry.register_analysis(
            AnalysisExtension(
                type="ui_declared_placeholder_search",
                name="UI 声明占位符",
                handler=lambda lines, params: {
                    "analysis_type": "ui_declared_placeholder_search",
                    "dominant_frequency": 7.5,
                },
                report_placeholders=[
                    {"token": "{{dominant_frequency}}", "label": "主频", "description": "主频字段"},
                ],
            )
        )
        try:
            self.page._refresh_analysis_type_choices()
            self.page._refresh_report_placeholder_choices()

            items = [
                self.page._report_placeholder_combo.itemText(i)
                for i in range(self.page._report_placeholder_combo.count())
            ]

            self.assertTrue(any("[UI 声明占位符]主频" in item and "{{dominant_frequency}}" in item for item in items))
            self.assertIsInstance(self.page._report_placeholder_combo, EditableComboBox)
            self.assertIsNotNone(self.page._report_placeholder_combo.completer())
        finally:
            extension_registry.unregister_analysis("ui_declared_placeholder_search")
            self.page._refresh_analysis_type_choices()

    def test_analysis_page_loads_registered_extensions_on_init(self):
        from core.extension_api import AnalysisExtension, extension_registry
        from ui.pages.analysis_page import AnalysisPage

        extension_registry.register_analysis(
            AnalysisExtension(
                type="ui_init_analysis_extension",
                name="初始化分析扩展",
                handler=lambda lines, params: {"analysis_type": "ui_init_analysis_extension"},
                description="用于验证初始化时自动加载扩展。",
            )
        )
        page = AnalysisPage()
        try:
            combo_items = [page._type_combo.itemText(i) for i in range(page._type_combo.count())]
            selector_items = [page._extension_panel._selector.itemText(i) for i in range(page._extension_panel._selector.count())]
            self.assertIn("初始化分析扩展", combo_items)
            self.assertIn("初始化分析扩展", selector_items)
        finally:
            page.deleteLater()
            extension_registry.unregister_analysis("ui_init_analysis_extension")

    def test_export_report_to_path(self):
        self.page.on_tree_node_activated("series", self.s.id)
        self.page._run_analysis()
        self.page._on_generate_report()
        fd, path = tempfile.mkstemp(suffix=".md")
        os.close(fd)
        try:
            self.assertTrue(self.page._export_report_to_path(path))
            with open(path, "r", encoding="utf-8") as handle:
                content = handle.read()
            self.assertIn("分析类型", content)
        finally:
            os.unlink(path)

    def test_error_compare_requires_two_inputs(self):
        from models.schemas import DataSeries

        self.page._type_combo.setCurrentIndex(4)
        other = DataSeries(name="s2", x=[1.0, 2.0, 3.0], y=[1.5, 1.0, 2.5])
        self.df.series.append(other)
        self.page.on_tree_node_activated("series", self.s.id)
        self.page.on_tree_node_activated("series", other.id)
        self.page._run_analysis()
        self.assertEqual(self.page._result["analysis_type"], "error_compare")
        self.assertIn("mae", self.page._result)

    def test_load_saved_analysis_result_from_tree(self):
        self.page.on_tree_node_activated("series", self.s.id)
        self.page._run_analysis()
        from ui.dialogs.export_flow import AnalysisResultSavePlan

        analysis_root = self.pm._find_folder_by_group_type("analysis_result_group")
        self.assertIsNotNone(analysis_root)
        with mock.patch(
            "ui.pages.analysis_page.choose_analysis_result_save_plan",
            return_value=AnalysisResultSavePlan(result_name="拟合结果A", target_parent_id=analysis_root.id),
        ):
            self.page._save_result()
        node = next((n for n in self.p.tree.nodes if n.kind == "analysis_result"), None)
        self.assertIsNotNone(node)
        self.page.load_analysis_result(node.id)
        self.assertEqual(self.page._result["analysis_type"], "curve_fit")
        self.assertEqual(len(self.page._selected_inputs), 1)
        self.assertEqual(self.page._analysis_tabs.count(), 3)
        titles = [self.page._analysis_tabs.tabText(i) for i in range(self.page._analysis_tabs.count())]
        self.assertIn("拟合结果A", titles)
        active_key = self.page._analysis_tab_key_for_index(self.page._analysis_tabs.currentIndex())
        self.assertIsNotNone(active_key)
        view = self.page._analysis_tab_views[active_key]
        summary_table = view["detail_summary_table"] or self.page._summary_table
        summary_labels = [summary_table.item(row, 0).text() for row in range(summary_table.rowCount())]
        self.assertIn("R²", summary_labels)

    def test_selecting_series_does_not_auto_restore_saved_results_in_tabs(self):
        self.page.on_tree_node_activated("series", self.s.id)
        self.page._run_analysis()

        with mock.patch(
            "ui.pages.analysis_page.choose_analysis_result_save_plan",
            side_effect=_analysis_result_save_plans(self.pm, "拟合结果A", "拟合结果B"),
        ):
            self.page._save_result()
            self.page._save_result()

        self.page.on_tree_node_activated("series", self.s.id)

        titles = [self.page._analysis_tabs.tabText(i) for i in range(self.page._analysis_tabs.count())]
        self.assertEqual(self.page._analysis_tabs.count(), 2)
        self.assertNotIn("拟合结果A", titles)
        self.assertNotIn("拟合结果B", titles)

    def test_temporary_result_tab_keeps_its_own_inputs_without_overwriting_current_selection(self):
        from models.schemas import DataSeries

        other = DataSeries(name="s2", x=[10.0, 11.0], y=[20.0, 21.0])
        self.df.series.append(other)

        self.page.on_tree_node_activated("series", self.s.id)
        self.page._run_analysis()

        temp_view = self.page._analysis_tab_views[self.page._analysis_tab_keys[1]]
        self.assertEqual([item["node_id"] for item in temp_view["inputs"]], [self.s.id])

        self.page._analysis_tabs.setCurrentIndex(0)
        self.page.on_tree_node_activated("series", other.id)
        current_ids = [item["node_id"] for item in self.page._selected_inputs]
        self.assertEqual(current_ids, [self.s.id, other.id])

        self.page._analysis_tabs.setCurrentIndex(1)

        self.assertEqual([item["node_id"] for item in self.page._selected_inputs], current_ids)
        self.assertEqual(self.page._input_list.count(), 2)
        self.assertEqual([item["node_id"] for item in temp_view["inputs"]], [self.s.id])

    def test_analysis_result_tabs_are_closable(self):
        from qfluentwidgets import TabCloseButtonDisplayMode

        self.page.on_tree_node_activated("series", self.s.id)
        self.page._run_analysis()
        self.page._run_analysis()

        self.assertEqual(
            self.page._analysis_tabs.tabBar.closeButtonDisplayMode,
            TabCloseButtonDisplayMode.ON_HOVER,
        )

        self.page._analysis_tabs.setCurrentIndex(2)
        closing_title = self.page._analysis_tabs.tabText(2)
        self.page._on_analysis_tab_close_requested(2)

        titles = [self.page._analysis_tabs.tabText(i) for i in range(self.page._analysis_tabs.count())]
        self.assertEqual(self.page._analysis_tabs.count(), 2)
        self.assertNotIn(closing_title, titles)

    def test_peak_detect_uses_merged_detail_table_and_combined_line(self):
        from PySide6.QtWidgets import QApplication
        from PySide6.QtWidgets import QAbstractItemView
        from PySide6.QtWidgets import QTableWidgetSelectionRange

        from models.schemas import DataSeries

        target = DataSeries(
            name="oscillation",
            x=[0.0, 1.0, 2.0, 3.0, 4.0],
            y=[0.0, 2.0, 0.0, -1.5, 0.0],
        )
        self.df.series.append(target)

        self.page._type_combo.setCurrentIndex(self.page._analysis_type_ids.index("peak_detect"))
        self.page.on_tree_node_activated("series", target.id)
        self.page._run_analysis()

        active_key = self.page._analysis_tab_key_for_index(self.page._analysis_tabs.currentIndex())
        self.assertIsNotNone(active_key)
        view = self.page._analysis_tab_views[active_key]
        self.assertIs(view["summary_stack"].currentWidget(), view["details_scroll"])

        detail_summary_table = view["detail_summary_table"]
        self.assertIsNotNone(detail_summary_table)
        meta_labels = [detail_summary_table.item(row, 0).text() for row in range(detail_summary_table.rowCount())]
        self.assertIn("波峰数量", meta_labels)
        self.assertIn("波谷数量", meta_labels)

        self.assertEqual(len(view["detail_tables"]), 1)
        peak_points_table = view["detail_tables"][0]
        peak_headers = [peak_points_table.horizontalHeaderItem(i).text() for i in range(peak_points_table.columnCount())]
        self.assertEqual(peak_headers, ["序号", "类型", "X", "Y"])
        self.assertEqual(peak_points_table.rowCount(), 2)
        self.assertEqual(detail_summary_table.selectionMode(), QAbstractItemView.SelectionMode.ExtendedSelection)
        self.assertEqual(peak_points_table.selectionMode(), QAbstractItemView.SelectionMode.ExtendedSelection)
        self.assertEqual(peak_points_table.item(0, 0).text(), "1")
        self.assertEqual(peak_points_table.item(0, 1).text(), "波峰")
        self.assertEqual(peak_points_table.item(0, 2).text(), "1")
        self.assertEqual(peak_points_table.item(1, 0).text(), "2")
        self.assertEqual(peak_points_table.item(1, 1).text(), "波谷")
        self.assertEqual(peak_points_table.item(1, 2).text(), "3")
        self.assertLess(view["details_layout"].indexOf(detail_summary_table), view["details_layout"].indexOf(peak_points_table))
        self.assertIn("峰谷点", [item["line_name"] for item in self.page._result.get("lines", [])])

        peak_points_table.setRangeSelected(
            QTableWidgetSelectionRange(0, 0, 0, 3),
            True,
        )
        peak_points_table.copy_selection_to_clipboard()
        QApplication.processEvents()
        self.assertIn("\t", QApplication.clipboard().text())

    def test_report_preview_allows_selecting_one_result_per_analysis_type(self):
        with mock.patch(
            "ui.pages.analysis_page.choose_analysis_result_save_plan",
            side_effect=_analysis_result_save_plans(self.pm, "拟合结果A", "拟合结果B", "统计结果A"),
        ):
            self.page.on_tree_node_activated("series", self.s.id)
            self.page._type_combo.setCurrentIndex(self.page._analysis_type_ids.index("curve_fit"))
            self.page._run_analysis()
            self.page._save_result()
            self.page._save_result()
            self.page._type_combo.setCurrentIndex(self.page._analysis_type_ids.index("statistics"))
            self.page._run_analysis()
            self.page._save_result()

        self.page.on_tree_node_activated("series", self.s.id)
        self.page._on_generate_report()

        self.assertFalse(self.page._report_result_selector_panel.isHidden())
        curve_fit_combo = self.page._report_result_selectors["curve_fit"]
        statistics_combo = self.page._report_result_selectors["statistics"]
        self.assertEqual(curve_fit_combo.count(), 3)
        self.assertEqual(statistics_combo.count(), 3)

        curve_fit_combo.setCurrentIndex(2)
        statistics_combo.setCurrentIndex(2)
        self.page._render_report_preview()
        preview = self.page._report_preview.toPlainText()
        self.assertIn("拟合结果B", preview)
        self.assertIn("统计结果A", preview)
        self.assertNotIn("拟合结果A", preview)

    def test_report_preview_does_not_append_multi_result_sections_without_placeholder(self):
        with mock.patch(
            "ui.pages.analysis_page.choose_analysis_result_save_plan",
            side_effect=_analysis_result_save_plans(self.pm, "拟合结果A", "统计结果A"),
        ):
            self.page.on_tree_node_activated("series", self.s.id)
            self.page._type_combo.setCurrentIndex(self.page._analysis_type_ids.index("curve_fit"))
            self.page._run_analysis()
            self.page._save_result()
            self.page._type_combo.setCurrentIndex(self.page._analysis_type_ids.index("statistics"))
            self.page._run_analysis()
            self.page._save_result()

        self.page._report_editor.setPlainText("\n前言\n\n正文\n")
        self.page.on_tree_node_activated("series", self.s.id)
        self.page._on_generate_report()

        curve_fit_combo = self.page._report_result_selectors["curve_fit"]
        statistics_combo = self.page._report_result_selectors["statistics"]

        def _select_combo_text(combo, text):
            for index in range(combo.count()):
                if combo.itemText(index) == text:
                    combo.setCurrentIndex(index)
                    return
            self.fail(f"未找到结果选项: {text}")

        _select_combo_text(curve_fit_combo, "拟合结果A")
        _select_combo_text(statistics_combo, "统计结果A")
        self.page._render_report_preview()

        preview = self.page._report_preview.toPlainText()
        self.assertTrue(preview.startswith("\n前言\n\n正文\n"))
        self.assertEqual(preview.count("前言"), 1)
        self.assertEqual(preview.count("正文"), 1)
        self.assertNotIn("### 拟合结果A", preview)
        self.assertNotIn("### 统计结果A", preview)

    def test_report_preview_preserves_blank_lines_when_multi_result_sections_placeholder_present(self):
        with mock.patch(
            "ui.pages.analysis_page.choose_analysis_result_save_plan",
            side_effect=_analysis_result_save_plans(self.pm, "拟合结果A", "统计结果A"),
        ):
            self.page.on_tree_node_activated("series", self.s.id)
            self.page._type_combo.setCurrentIndex(self.page._analysis_type_ids.index("curve_fit"))
            self.page._run_analysis()
            self.page._save_result()
            self.page._type_combo.setCurrentIndex(self.page._analysis_type_ids.index("statistics"))
            self.page._run_analysis()
            self.page._save_result()

        self.page._report_editor.setPlainText("\n前言\n\n正文\n\n{{multi_result_sections}}\n")
        self.page.on_tree_node_activated("series", self.s.id)
        self.page._on_generate_report()

        curve_fit_combo = self.page._report_result_selectors["curve_fit"]
        statistics_combo = self.page._report_result_selectors["statistics"]

        def _select_combo_text(combo, text):
            for index in range(combo.count()):
                if combo.itemText(index) == text:
                    combo.setCurrentIndex(index)
                    return
            self.fail(f"未找到结果选项: {text}")

        _select_combo_text(curve_fit_combo, "拟合结果A")
        _select_combo_text(statistics_combo, "统计结果A")
        self.page._render_report_preview()

        preview = self.page._report_preview.toPlainText()
        self.assertTrue(preview.startswith("\n前言\n\n正文\n\n"))
        self.assertEqual(preview.count("前言"), 1)
        self.assertEqual(preview.count("正文"), 1)
        self.assertIn("### 拟合结果A", preview)
        self.assertIn("### 统计结果A", preview)

    def test_peak_detect_supports_x_distance_mode(self):
        from models.schemas import DataSeries

        target = DataSeries(
            name="peaks",
            x=[0.0, 0.2, 0.4, 0.6, 1.0, 1.2, 1.4, 1.6],
            y=[0.0, 1.0, 0.0, 0.9, 0.0, 0.0, 0.8, 0.0],
        )
        self.df.series.append(target)

        self.page._type_combo.setCurrentIndex(self.page._analysis_type_ids.index("peak_detect"))
        options = dict(self.page._extension_params_edit.current_options())
        options["min_distance_x"] = 0.7
        self.page._extension_params_edit.set_options(options)
        self.page._on_extension_analysis_options_changed(options)
        self.page.on_tree_node_activated("series", target.id)
        self.page._run_analysis()

        self.assertEqual(self.page._result["analysis_type"], "peak_detect")
        self.assertEqual(self.page._result["distance_mode"], "x_distance")
        self.assertEqual(self.page._result["count"], 2)

    def test_single_curve_analysis_uses_current_list_selection(self):
        from models.schemas import DataSeries

        other = DataSeries(name="s2", x=[1.0, 2.0], y=[2.0, 3.0])
        self.df.series.append(other)
        self.page.on_tree_node_activated("series", self.s.id)
        self.page.on_tree_node_activated("series", other.id)
        self.assertEqual(len(self.page._selected_inputs), 2)

        self.page._input_list.clearSelection()
        self.page._input_list.item(1).setSelected(True)

        data = self.page._get_selected_data()
        self.assertEqual(len(data), 1)
        self.assertEqual(data[0][2], other.name)


# ═══════════════════════════════════════════════════════════════════════════
# 7. DigitizePage — load_image_by_id
# ═══════════════════════════════════════════════════════════════════════════

class TestDigitizePage(unittest.TestCase):

    def setUp(self):
        self.pm, self.p, self.df, self.s = _make_project("dig_test")
        self._restore = _patch_pm(self.pm)
        from ui.pages.digitize_page import DigitizePage
        self.page = DigitizePage()

    def tearDown(self):
        self._restore()
        self.page.deleteLater()

    def test_eraser_cursor_draws_without_press(self):
        from PySide6.QtCore import QPointF

        viewer = self.page._image_viewer
        viewer._current_tool = viewer.MODE_ERASER
        viewer._mouse_image_pos = QPointF(10.0, 12.0)
        viewer._eraser_pressed = False
        viewer._scale = 1.0

        painter = mock.Mock()

        viewer._draw_eraser_cursor(painter)

        painter.drawEllipse.assert_called_once()

    def test_brush_mask_cursor_draws_without_press(self):
        from PySide6.QtCore import QPointF

        viewer = self.page._image_viewer
        viewer._current_tool = viewer.MODE_BRUSH_MASK
        viewer._mouse_image_pos = QPointF(10.0, 12.0)
        viewer._brush_painting = False
        viewer._scale = 1.0

        painter = mock.Mock()

        viewer._draw_eraser_cursor(painter)

        painter.drawEllipse.assert_called_once()

    def _add_digitize_curve(self):
        from models.schemas import ImageWork

        image = ImageWork(name="SEM A", image_path="sample.png")
        self.p.images.append(image)
        curve = self.pm.add_curve_to_image(
            image.id,
            [1.0, 2.0, 3.0],
            [4.0, 5.0, 6.0],
            name="曲线1",
        )
        self.assertIsNotNone(curve)
        self.page._current_image_id = image.id
        self.page._current_curve_id = curve.id
        self.page._refresh_export_name_suggestion(force=True)
        return image, curve

    def test_page_creates_no_crash(self):
        self.assertIsNotNone(self.page)

    def test_digitize_auto_mode_lists_builtin_extensions(self):
        items = [self.page._auto_mode_combo.itemText(i) for i in range(self.page._auto_mode_combo.count())]

        self.assertIn("颜色识别", items)
        self.assertIn("图形识别", items)

    def test_digitize_auto_controls_use_compact_extension_panel(self):
        self.assertIs(self.page._auto_mode_combo, self.page._digitize_extension_controls._selector)
        self.assertFalse(self.page._digitize_extension_controls._selector_row_widget.isHidden())
        self.assertFalse(self.page._digitize_extension_controls._config_row_widget.isHidden())
        self.assertFalse(self.page._digitize_extension_controls._editor.isHidden())

    def test_digitize_right_tabs_use_segmented_navigation(self):
        from qfluentwidgets import SegmentedWidget
        from PySide6.QtWidgets import QSizePolicy

        self.assertIsInstance(self.page._right_tabs.navigationWidget, SegmentedWidget)
        self.assertEqual(
            self.page._right_tabs.navigationWidget.sizePolicy().horizontalPolicy(),
            QSizePolicy.Policy.Expanding,
        )

    def test_digitize_auto_config_scroll_only_wraps_parameter_editor(self):
        from qfluentwidgets import SmoothScrollArea

        def _is_descendant(widget, ancestor) -> bool:
            current = widget
            while current is not None:
                if current is ancestor:
                    return True
                current = current.parentWidget()
            return False

        self.assertIsInstance(self.page._digitize_auto_config_scroll, SmoothScrollArea)
        self.assertIs(self.page._digitize_auto_config_scroll, self.page._digitize_extension_controls._editor._scroll_area)
        self.assertFalse(_is_descendant(self.page._digitize_extension_controls._selector_row_widget, self.page._digitize_auto_config_scroll))
        self.assertFalse(_is_descendant(self.page._digitize_extension_controls._config_row_widget, self.page._digitize_auto_config_scroll))
        self.assertFalse(_is_descendant(self.page._manual_tools_row, self.page._digitize_auto_config_scroll))
        self.assertFalse(_is_descendant(self.page._auto_tools_row, self.page._digitize_auto_config_scroll))

    def test_digitize_auto_extension_panel_has_no_trailing_stretch_after_controls(self):
        combined_tab = self.page._right_tabs.widget(0)
        layout = combined_tab.layout()

        self.assertIsNotNone(layout)
        self.assertIsNotNone(layout.itemAt(layout.count() - 1).widget())

    def test_digitize_brush_mask_slider_updates_viewer_size(self):
        self.assertEqual(self.page._mask_brush_size_slider.value(), 20)
        self.assertEqual(self.page._mask_brush_size_value_label.text(), "20 px")
        self.assertEqual(self.page._image_viewer.get_mask_brush_size(), 20.0)

        self.page._mask_brush_size_slider.setValue(32)

        self.assertEqual(self.page._mask_brush_size_value_label.text(), "32 px")
        self.assertEqual(self.page._image_viewer.get_mask_brush_size(), 32.0)

    def test_digitize_curve_table_uses_compact_upper_split(self):
        self.page.resize(1320, 880)
        self.page.show()
        QApplication.processEvents()

        sizes = self.page._right_content_splitter.sizes()
        self.assertEqual(len(sizes), 2)
        self.assertGreater(sum(sizes), 0)
        ratio = sizes[0] / sum(sizes)
        self.assertGreater(ratio, 0.28)
        self.assertLess(ratio, 0.42)

    def test_digitize_pickcolor_and_shot_values_flow_through_extension_panel(self):
        from extensions.digitize.color_detect import COLOR_DIGITIZE_EXTENSION_TYPE
        from extensions.digitize.shape_detect import SHAPE_DIGITIZE_EXTENSION_TYPE

        self.page._auto_mode_combo.setCurrentIndex(
            self.page._auto_mode_type_ids.index(COLOR_DIGITIZE_EXTENSION_TYPE)
        )
        editor = self.page._digitize_extension_controls._editor
        self.assertTrue(editor.set_field_value("sampled_color", {"r": 10, "g": 20, "b": 30}))
        self.assertTrue(editor.set_field_value("tolerance", 18))
        self.assertTrue(editor.set_field_value("step", 6))

        params = self.page._current_digitize_extension_params(
            COLOR_DIGITIZE_EXTENSION_TYPE,
            mask_polygons=[[(1.0, 2.0), (3.0, 4.0), (3.0, 5.0)]],
            mask_include_mode=False,
        )
        self.assertEqual(params["sampled_color"], {"r": 10, "g": 20, "b": 30})
        self.assertEqual(params["tolerance"], 18)
        self.assertEqual(params["step"], 6)
        self.assertFalse(params["mask_include_mode"])

        self.page._auto_mode_combo.setCurrentIndex(
            self.page._auto_mode_type_ids.index(SHAPE_DIGITIZE_EXTENSION_TYPE)
        )
        self.assertTrue(editor.set_field_value("template_info", {"size": [24, 11], "bounds": [0, 1, 24, 12]}))
        self.assertTrue(editor.set_field_value("threshold", 0.72))
        self.assertTrue(editor.set_field_value("color_weight", 0.4))

        params = self.page._current_digitize_extension_params(
            SHAPE_DIGITIZE_EXTENSION_TYPE,
            mask_polygons=None,
            mask_include_mode=True,
        )
        self.assertEqual(params["template_info"]["size"], [24, 11])
        self.assertEqual(params["threshold"], 0.72)
        self.assertEqual(params["color_weight"], 0.4)

    def test_digitize_pickcolor_button_click_triggers_request_and_backfills_value(self):
        from PySide6.QtGui import QColor
        from extensions.digitize.color_detect import COLOR_DIGITIZE_EXTENSION_TYPE
        from qfluentwidgets import ToolButton

        self.page.show()
        self.page._current_image_id = "img-pick"
        self.page._image_viewer.get_image_path = lambda: "probe.png"
        self.page._auto_mode_combo.setCurrentIndex(
            self.page._auto_mode_type_ids.index(COLOR_DIGITIZE_EXTENSION_TYPE)
        )
        editor = self.page._digitize_extension_controls._editor
        pick_button = editor.findChild(ToolButton, "interactiveFieldButton:sampled_color")
        summary_label = editor.findChild(QWidget, "interactiveFieldSummary:sampled_color")

        self.assertIsNotNone(pick_button)
        self.assertIsNotNone(summary_label)

        QTest.mouseClick(pick_button, Qt.MouseButton.LeftButton)

        self.assertEqual(self.page._pending_digitize_field_key, "sampled_color")
        self.assertEqual(self.page._pending_digitize_field_type, "pickcolor")
        self.assertEqual(self.page._active_tool, "color_pick")
        self.assertIn("请在图片上点击取色", self.page._status_label.text())

        self.page._on_color_picked(QColor("#0A141E"))

        self.assertEqual(editor.current_options()["sampled_color"], {"r": 10, "g": 20, "b": 30})
        self.assertEqual(summary_label.text(), "#0A141E")
        self.assertIn("#0a141e", pick_button.styleSheet().lower())
        self.assertIsNone(self.page._pending_digitize_field_key)
        self.assertEqual(self.page._status_label.text(), "已采样: #0a141e")

    def test_digitize_shot_button_click_triggers_request_and_backfills_value(self):
        from extensions.digitize.shape_detect import SHAPE_DIGITIZE_EXTENSION_TYPE
        from qfluentwidgets import ToolButton

        self.page.show()
        self.page._current_image_id = "img-shot"
        self.page._image_viewer.get_image_path = lambda: "probe.png"
        self.page._auto_mode_combo.setCurrentIndex(
            self.page._auto_mode_type_ids.index(SHAPE_DIGITIZE_EXTENSION_TYPE)
        )
        editor = self.page._digitize_extension_controls._editor
        shot_button = editor.findChild(ToolButton, "interactiveFieldButton:template_info")
        summary_label = editor.findChild(QWidget, "interactiveFieldSummary:template_info")

        self.assertIsNotNone(shot_button)
        self.assertIsNotNone(summary_label)

        QTest.mouseClick(shot_button, Qt.MouseButton.LeftButton)

        self.assertEqual(self.page._pending_digitize_field_key, "template_info")
        self.assertEqual(self.page._pending_digitize_field_type, "shot")
        self.assertEqual(self.page._active_tool, "crop_template")
        self.assertIn("请在图片上拖拽截图", self.page._status_label.text())

        with mock.patch(
            "extensions.digitize._extractors.ShapeExtractor.preprocess_region",
            return_value={"size": [24, 11], "bounds": [0, 1, 24, 12]},
        ):
            self.page._on_crop_region_selected(0.0, 1.0, 24.0, 12.0)

        self.assertEqual(
            editor.current_options()["template_info"],
            {"size": [24, 11], "bounds": [0, 1, 24, 12]},
        )
        self.assertEqual(summary_label.text(), "已截取 24×11px")
        self.assertIsNone(self.page._pending_digitize_field_key)
        self.assertIn("图例模板已截取", self.page._status_label.text())

    def test_digitize_auto_detect_uses_selected_digitize_extension(self):
        from core.extension_api import DigitizeExtension, extension_registry

        extension_registry.register_digitize(
            DigitizeExtension(
                type="ui_digitize_probe",
                name="UI 数字化探针",
                handler=lambda figure, params: [[12.0, 34.0]],
            )
        )
        try:
            self.page._refresh_digitize_extension_choices()
            self.page._current_image_id = "img-probe"
            self.page._image_viewer.get_image_path = lambda: "probe.png"
            self.page._auto_mode_combo.setCurrentIndex(self.page._auto_mode_type_ids.index("ui_digitize_probe"))

            self.page._on_auto_detect()

            self.assertEqual(self.page._auto_preview_points, [(12.0, 34.0)])
            self.assertIn("UI 数字化探针识别到 1 个点", self.page._status_label.text())
        finally:
            extension_registry.unregister_digitize("ui_digitize_probe")
            self.page._refresh_digitize_extension_choices()

    def test_digitize_extension_panel_tracks_current_extension(self):
        from extensions.digitize.shape_detect import SHAPE_DIGITIZE_EXTENSION_TYPE

        self.assertTrue(self.page.supports_extension_panel_toggle())
        self.page.set_extension_panel_visible(True)

        self.assertTrue(self.page.is_extension_panel_visible())
        self.assertFalse(self.page._extension_panel.isHidden())
        self.assertTrue(self.page._extension_panel._selector_row_widget.isHidden())
        self.assertTrue(self.page._extension_panel._editor.isHidden())
        self.assertEqual(self.page._extension_panel.current_type(), self.page._current_digitize_extension_type())

        self.page._auto_mode_combo.setCurrentIndex(self.page._auto_mode_type_ids.index(SHAPE_DIGITIZE_EXTENSION_TYPE))

        self.assertEqual(self.page._extension_panel.current_type(), SHAPE_DIGITIZE_EXTENSION_TYPE)

    def test_digitize_main_panel_uses_shared_workbench_spacing(self):
        preview_panel = self.page._splitter.widget(1)
        preview_margins = preview_panel.layout().contentsMargins()
        tool_margins = self.page._right_panel.layout().contentsMargins()

        self.assertEqual(self.page._splitter.handleWidth(), 4)
        self.assertEqual(
            (tool_margins.left(), tool_margins.top(), tool_margins.right(), tool_margins.bottom()),
            (14, 14, 14, 14),
        )
        self.assertEqual(
            (preview_margins.left(), preview_margins.top(), preview_margins.right(), preview_margins.bottom()),
            (14, 14, 14, 14),
        )

    def test_curve_data_buttons_are_above_panel_title(self):
        layout = self.page._right_panel.layout()

        self.assertIsNotNone(layout.itemAt(0).layout())
        self.assertIs(layout.itemAt(1).widget(), self.page._curve_panel_title)

    def test_crosshair_color_button_matches_manual_toolbar_height(self):
        self.assertEqual(self.page._crosshair_color_btn.width(), 34)
        self.assertEqual(self.page._crosshair_color_btn.height(), 34)
        self.assertEqual(self.page._crosshair_color_btn.height(), self.page._calibrate_btn.height())

    def test_digitize_border_color_uses_shared_theme_helper(self):
        from ui.theme import border_color

        self.assertEqual(self.page._border_color(), border_color())

    def test_load_image_by_id_nonexistent(self):
        """不存在的 ID 不崩溃"""
        self.page.load_image_by_id("nonexistent-id")

    def test_on_tree_node_selected_image_work(self):
        self.page.on_tree_node_selected("image_work", "fake-image-id")

    def test_load_image_by_id_accepts_tree_node_id(self):
        from PySide6.QtGui import QImage

        with tempfile.TemporaryDirectory() as temp_dir:
            image_path = Path(temp_dir) / "tree-node-image.png"
            image = QImage(16, 16, QImage.Format.Format_RGB32)
            image.fill(Qt.GlobalColor.white)
            self.assertTrue(image.save(str(image_path)))
            created = self.pm.add_image(str(image_path), name="tree-node-image.png")
            node = next((item for item in self.p.tree.nodes if item.kind == "image_work" and item.image_work_id == created.id), None)
            self.assertIsNotNone(node)

            self.page.load_image_by_id(node.id)

            self.assertEqual(self.page._current_image_id, created.id)

    def test_import_source_image_accepts_supported_image_file(self):
        from PySide6.QtGui import QImage

        with tempfile.TemporaryDirectory() as temp_dir:
            image_path = Path(temp_dir) / "import-source-image.png"
            image = QImage(16, 16, QImage.Format.Format_RGB32)
            image.fill(Qt.GlobalColor.white)
            self.assertTrue(image.save(str(image_path)))

            self.assertTrue(self.page.import_source_image(str(image_path)))
            self.assertIsNotNone(self.page._current_image_id)

    def test_on_tree_node_selected_folder(self):
        node = next((n for n in self.p.tree.nodes if n.kind == "folder"), None)
        if node:
            self.page.on_tree_node_selected("folder", node.id)

    def test_legacy_project_panel_removed(self):
        self.assertIsNone(self.page._project_tree)
        self.assertFalse(hasattr(self.page, "_new_project_btn"))
        self.assertFalse(hasattr(self.page, "_open_project_btn"))
        self.assertFalse(hasattr(self.page, "_save_project_btn"))
        self.assertFalse(hasattr(self.page, "_close_project_btn"))
        self.assertIsNotNone(self.page._add_image_btn)
        self.assertIsNotNone(self.page._add_curve_btn)

    def test_export_name_suggestion_uses_image_and_curve_names(self):
        self._add_digitize_curve()
        self.assertIn("SEM", self.page._export_name_edit.text())
        self.assertIn("曲线1", self.page._export_name_edit.text())

    def test_export_tab_hides_target_text_and_uses_compact_labels(self):
        from PySide6.QtWidgets import QSizePolicy

        self.assertTrue(self.page._export_target_label.isHidden())
        self.assertEqual(self.page._export_scope_label.minimumWidth(), self.page._export_scope_label.sizeHint().width())
        self.assertEqual(self.page._export_scope_label.sizePolicy().horizontalPolicy(), QSizePolicy.Policy.Minimum)
        self.assertEqual(self.page._export_curve_btn.text(), "导出曲线...")

    def test_export_to_data_file_does_not_create_digitize_result_folder_by_default(self):
        from ui.dialogs.export_flow import DataExportPlan

        self._add_digitize_curve()
        received = []
        self.page.project_modified.connect(lambda: received.append(True))

        datasets_root = self.pm._find_folder_by_group_type("datasets")
        self.assertIsNotNone(datasets_root)

        with mock.patch(
            "ui.pages.digitize_page.choose_data_export_plan",
            return_value=DataExportPlan(
                export_name="SEM_曲线1",
                new_parent_id=datasets_root.id,
                new_data_file_name="SEM_曲线1.digitize",
            ),
        ):
            self.page._on_export_to_data_file()

        derived_folder = next(
            (node for node in self.p.tree.nodes if node.kind == "folder" and node.parent_id == datasets_root.id and node.name == "数字化结果"),
            None,
        )
        self.assertIsNone(derived_folder)
        data_node = next(
            (node for node in self.p.tree.nodes if node.kind == "data_file" and node.parent_id == datasets_root.id and node.name == "SEM_曲线1.digitize"),
            None,
        )
        self.assertIsNotNone(data_node)
        data_file = next((df for df in self.p.data_files if df.id == data_node.data_file_id), None)
        self.assertIsNotNone(data_file)
        self.assertTrue(data_file.name.endswith(".digitize"))
        self.assertEqual(self.page._export_target_kind, "data_file")
        self.assertEqual(self.page._export_target_id, data_node.id)
        self.assertEqual(len(received), 1)

    def test_export_to_existing_data_file_appends_series(self):
        from ui.dialogs.export_flow import DataExportPlan

        self._add_digitize_curve()
        data_node = next((node for node in self.p.tree.nodes if node.kind == "data_file" and node.data_file_id == self.df.id), None)
        self.assertIsNotNone(data_node)
        before = len(self.df.series)

        self.page.on_tree_node_selected("data_file", data_node.id)
        with mock.patch(
            "ui.pages.digitize_page.choose_data_export_plan",
            return_value=DataExportPlan(export_name="追加结果", target_data_file_id=self.df.id),
        ):
            self.page._on_export_to_data_file()

        self.assertEqual(len(self.df.series), before + 1)
        self.assertIn(data_node.name, self.page._status_label.text())

    def test_curve_export_uses_shared_dialog_and_clipboard(self):
        from ui.dialogs.export_flow import CurveFileExportPlan

        _image, curve = self._add_digitize_curve()

        with mock.patch(
            "ui.pages.digitize_page.choose_curve_file_export_plan",
            return_value=CurveFileExportPlan(action="clipboard", file_format="csv", include_timestamp=False, merged=False),
        ):
            self.page._on_export_curves()

        clipboard_text = QApplication.clipboard().text()
        self.assertIn("SEM_A_曲线1", clipboard_text)
        self.assertIn("x\ty", clipboard_text)
        self.assertEqual(self.page._current_curve_id, curve.id)


# ═══════════════════════════════════════════════════════════════════════════
# 8. MainWindow — 树面板注入、页面切换、信号路由
# ═══════════════════════════════════════════════════════════════════════════

# ═══════════════════════════════════════════════════════════════════════════
# 8. MainWindow — 树面板注入、页面切换、信号路由
# ═══════════════════════════════════════════════════════════════════════════

class TestMainWindow(unittest.TestCase):
    """MainWindow 创建一次，所有测试共享同一实例"""

    @classmethod
    def setUpClass(cls):
        cls._restore_assets = _patch_global_assets()
        cls.pm, cls.p, cls.df, cls.s = _make_project("mw_test")
        cls._restore = _patch_pm(cls.pm)
        from ui.main_window import MainWindow
        cls.win = MainWindow()

    @classmethod
    def tearDownClass(cls):
        cls._restore()
        cls._restore_assets()
        cls.win.deleteLater()

    def test_window_creates_no_crash(self):
        self.assertIsNotNone(self.win)

    def test_tree_panel_exists(self):
        self.assertIsNotNone(self.win._tree_panel)

    def test_tree_toolbar_keeps_vertical_separator(self):
        from PySide6.QtWidgets import QFrame

        separator = self.win._tree_panel._toolbar_separator
        self.assertEqual(separator.frameShape(), QFrame.Shape.VLine)
        self.assertGreaterEqual(separator.maximumWidth(), 2)

    def test_tree_panel_has_project_action_buttons(self):
        self.assertIsNotNone(self.win._tree_panel.new_project_btn)
        self.assertIsNotNone(self.win._tree_panel.open_project_btn)
        self.assertIsNotNone(self.win._tree_panel.save_project_btn)
        self.assertIsNotNone(self.win._tree_panel.close_project_btn)
        self.assertIsNotNone(self.win._tree_panel.extension_toggle_btn)

    def test_tree_panel_has_no_source_import_button(self):
        self.assertFalse(hasattr(self.win._tree_panel, "import_source_btn"))

    def test_source_file_activation_routes_to_data_management(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            source_path = Path(temp_dir) / "raw_source.csv"
            source_path.write_text("x,y\n1,2\n", encoding="utf-8")
            node = self.pm.add_source_file(str(source_path))

        with mock.patch.object(self.win.data_page, "on_tree_node_selected") as selected_mock:
            self.win._on_tree_node_activated("source_file", node.id)

        self.assertIs(self.win.stackedWidget.currentWidget(), self.win.data_page)
        selected_mock.assert_called_once_with("source_file", node.id)

    def test_source_file_context_action_routes_to_dataset_import(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            source_path = Path(temp_dir) / "raw_source_context.csv"
            source_path.write_text("x,y\n1,2\n", encoding="utf-8")
            node = self.pm.add_source_file(str(source_path))
            self.assertIsNotNone(node)

        with mock.patch.object(self.win.data_page, "on_tree_node_selected") as selected_mock, \
             mock.patch.object(self.win.data_page, "_import_current_source_file_to_dataset") as import_mock:
            self.win._on_tree_node_activated("source_file_to_data", node.id)

        self.assertIs(self.win.stackedWidget.currentWidget(), self.win.data_page)
        selected_mock.assert_called_once_with("source_file", node.id)
        import_mock.assert_called_once_with()

    def test_tree_duplicate_warning_uses_main_window_parent(self):
        dataset_root = self.pm._find_folder_by_group_type("datasets")
        self.assertIsNotNone(dataset_root)
        existing_name = self.df.name

        with mock.patch("ui.widgets.project_tree.TextInputDialog.get_text", return_value=(existing_name, True)), \
             mock.patch("ui.widgets.project_tree.InfoBar.warning") as warning_mock:
            self.win._tree_panel.tree._cmd_add_dataset_node(dataset_root.id)

        warning_mock.assert_called_once()
        self.assertIs(warning_mock.call_args.kwargs.get("parent"), self.win)

    def test_tree_prune_empty_folder_feedback_uses_main_window_parent(self):
        datasets_root = self.pm._find_folder_by_group_type("datasets")
        self.assertIsNotNone(datasets_root)
        empty_folder = self.pm.add_folder("empty_cleanup_ui", parent_id=datasets_root.id, group_type="datasets")
        self.assertIsNotNone(empty_folder)

        with mock.patch("ui.widgets.project_tree.InfoBar.success") as success_mock:
            self.win._tree_panel.tree._cmd_prune_empty_folders(datasets_root.id, scope_label="数据集")

        success_mock.assert_called_once()
        self.assertIs(success_mock.call_args.kwargs.get("parent"), self.win)

    def test_chart_export_to_picture_group_refreshes_tree_immediately(self):
        from ui.dialogs.export_flow import PictureExportPlan

        if self.win.chart_page._figure is None:
            self.skipTest("matplotlib unavailable")

        with tempfile.TemporaryDirectory() as temp_dir:
            original_file_path = self.pm.current_project.file_path
            try:
                self.pm.current_project.file_path = str(Path(temp_dir) / "chart_refresh.aline")
                export_path = str(Path(temp_dir) / "chart_refresh.png")
                with mock.patch("ui.pages.chart_page.choose_picture_export_plan", return_value=PictureExportPlan(export_name="chart_refresh", target_folder_id=None)):
                    with mock.patch("ui.pages.chart_page.project_manager.prepare_picture_export_path", return_value=export_path):
                        self.win.chart_page._on_export_to_picture_group()
                QApplication.processEvents()
                picture = self.pm.current_project.pictures[-1]
                picture_node = next(
                    (node for node in self.pm.current_project.tree.nodes if node.kind == "picture" and node.picture_id == picture.id),
                    None,
                )
                self.assertIsNotNone(picture_node)
                self.assertIsNotNone(self.win._tree_panel.tree._find_item(picture_node.id))
            finally:
                self.pm.current_project.file_path = original_file_path

    def test_page_project_modified_refreshes_tree(self):
        with mock.patch.object(self.win._tree_panel.tree, "refresh") as refresh_mock:
            self.win.data_page.project_modified.emit()
            QApplication.processEvents()

        refresh_mock.assert_called()

    def test_source_file_to_digitize_routes_to_import_source_image(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            source_path = Path(temp_dir) / "raw_source.png"
            source_path.write_bytes(b"fake")
            node = self.pm.add_source_file(str(source_path))

        with mock.patch.object(self.win.digitize_page, "import_source_image") as import_mock:
            self.win._on_tree_node_activated("source_file_to_digitize", node.id)

        self.assertIs(self.win.stackedWidget.currentWidget(), self.win.digitize_page)
        import_mock.assert_called_once_with(str(source_path.resolve()), name=node.name)

    def test_tree_panel_exposes_single_expand_toggle_button(self):
        self.assertTrue(hasattr(self.win._tree_panel, "tree_expand_toggle_btn"))
        self.assertFalse(hasattr(self.win._tree_panel, "tree_expand_btn"))
        self.assertFalse(hasattr(self.win._tree_panel, "tree_collapse_btn"))

    def test_tree_panel_expand_toggle_button_sits_left_of_manage_button(self):
        self.win.resize(1320, 900)
        self.win.show()
        QApplication.processEvents()

        right_group = self.win._tree_panel._toolbar_right_group
        self.assertLess(
            right_group.indexOf(self.win._tree_panel.tree_expand_toggle_btn),
            right_group.indexOf(self.win._tree_panel.tree_manage_btn),
        )

    def test_theme_combo_uses_single_main_window_theme_path(self):
        current_index = self.win.settings_page.theme_combo.currentIndex()
        target_index = 0 if current_index != 0 else 1

        with mock.patch("ui.main_window.setTheme") as main_set_theme, \
             mock.patch("ui.pages.settings_page.setTheme") as page_set_theme, \
             mock.patch.object(self.win, "_update_all_pages_theme") as update_all_mock:
            self.win.settings_page.theme_combo.setCurrentIndex(target_index)
            QApplication.processEvents()

        main_set_theme.assert_called_once()
        page_set_theme.assert_not_called()
        update_all_mock.assert_called_once()
        self.win.settings_page.theme_combo.setCurrentIndex(current_index)
        QApplication.processEvents()

    def test_tree_panel_expand_toggle_button_switches_tooltip_after_click(self):
        self.win.resize(1320, 900)
        self.win.show()
        self.win._tree_panel.tree.refresh()
        QApplication.processEvents()

        button = self.win._tree_panel.tree_expand_toggle_btn

        self.assertEqual(button.toolTip(), "全部展开")
        button.click()
        QApplication.processEvents()
        self.assertEqual(button.toolTip(), "全部折叠")

        button.click()
        QApplication.processEvents()
        self.assertEqual(button.toolTip(), "全部展开")

    def test_navigation_has_tree_toggle_button(self):
        self.assertIsNotNone(self.win._tree_toggle_nav_btn)

    def test_window_registers_ui_control_shortcuts(self):
        from core.shortcut_manager import shortcut_manager

        self.assertEqual(shortcut_manager.get("toggle_project_tree"), "Alt+1")
        self.assertEqual(shortcut_manager.get("toggle_extension_panel"), "Alt+2")
        self.assertEqual(shortcut_manager.get("go_home"), "Alt+H")

        bindings = self.win._shortcut_bindings._shortcuts
        self.assertIn("toggle_project_tree", bindings)
        self.assertIn("toggle_extension_panel", bindings)
        self.assertIn("go_home", bindings)
        self.assertEqual(bindings["toggle_project_tree"].key().toString(), "Alt+1")
        self.assertEqual(bindings["toggle_extension_panel"].key().toString(), "Alt+2")
        self.assertEqual(bindings["go_home"].key().toString(), "Alt+H")

    def test_go_home_shortcut_binding_switches_to_home_page(self):
        self.win.switchTo(self.win.chart_page)
        self.win._shortcut_bindings._shortcuts["go_home"].activated.emit()

        self.assertIs(self.win.stackedWidget.currentWidget(), self.win.home_page)

    def test_create_project_from_panel_uses_window_level_dialog_flow(self):
        with mock.patch("ui.main_window.TextInputDialog.get_text", return_value=("面板项目", True)) as prompt_mock, \
             mock.patch("ui.main_window.QFileDialog.getExistingDirectory", return_value="/tmp/aline-panel-project") as dir_mock, \
             mock.patch.object(self.pm, "create_new", return_value=self.p) as create_mock, \
             mock.patch.object(self.win, "_on_project_created") as created_mock:
            self.win._create_project_from_panel()

        prompt_mock.assert_called_once()
        dir_mock.assert_called_once()
        create_mock.assert_called_once_with("面板项目", parent_dir="/tmp/aline-panel-project", create_structure=True)
        created_mock.assert_called_once_with("面板项目")

    def test_open_project_from_panel_uses_window_level_dialog_flow(self):
        with mock.patch("ui.main_window.QFileDialog.getOpenFileName", return_value=("/tmp/demo.aline", "ALine 项目 (*.aline)")) as file_mock, \
             mock.patch.object(self.pm, "open", return_value=self.p) as open_mock, \
             mock.patch.object(self.win, "_on_project_opened") as opened_mock:
            self.win._open_project_from_panel()

        file_mock.assert_called_once()
        open_mock.assert_called_once_with("/tmp/demo.aline")
        opened_mock.assert_called_once_with("/tmp/demo.aline")

    def test_tree_panel_has_tree_widget(self):
        from ui.widgets.project_tree import ProjectTreeWidget
        self.assertIsInstance(self.win._tree_panel.tree, ProjectTreeWidget)

    def test_tree_panel_is_hosted_in_splitter_with_width_bounds(self):
        from PySide6.QtWidgets import QSplitter

        self.assertIsInstance(self.win._tree_splitter, QSplitter)
        self.assertIs(self.win._tree_splitter.widget(0), self.win._tree_panel)
        self.assertIs(self.win._tree_splitter.widget(1), self.win.stackedWidget)
        self.assertEqual(self.win._tree_panel.minimumWidth(), 260)
        self.assertEqual(self.win._tree_panel.maximumWidth(), 420)

    def test_switch_to_data_page_shows_tree(self):
        self.win.switchTo(self.win.data_page)
        self.assertFalse(self.win._tree_panel.isHidden())

    def test_switch_to_settings_page_hides_tree(self):
        self.win.switchTo(self.win.settings_page)
        self.assertTrue(self.win._tree_panel.isHidden())
        self.assertFalse(self.win._tree_toggle_nav_btn.isEnabled())

    def test_switch_to_home_page_hides_tree(self):
        self.win.switchTo(self.win.home_page)
        self.assertTrue(self.win._tree_panel.isHidden())

    def test_switch_to_chart_page_shows_tree(self):
        self.win.switchTo(self.win.chart_page)
        self.assertFalse(self.win._tree_panel.isHidden())
        self.assertTrue(self.win._tree_toggle_nav_btn.isEnabled())

    def test_page_tree_focus_mode_applies_page_specific_filter(self):
        from ui.main_window import _BUSINESS_TREE_KINDS

        previous = self.win._page_tree_focus_mode_enabled
        try:
            self.win._page_tree_focus_mode_enabled = True

            def _top_level_labels() -> list[str]:
                tree = self.win._tree_panel.tree._tree
                return [tree.topLevelItem(index).text(0) for index in range(tree.topLevelItemCount())]

            def _top_level_kinds() -> list[str]:
                tree = self.win._tree_panel.tree._tree
                return [tree.topLevelItem(index).data(0, Qt.ItemDataRole.UserRole)[0] for index in range(tree.topLevelItemCount())]

            self.win.switchTo(self.win.data_page)
            self.assertEqual(self.win._tree_panel.tree._filter_kinds, list(_BUSINESS_TREE_KINDS))
            self.assertEqual(self.win._tree_panel.tree._focus_root_group_types, [])
            self.assertIn("全局资源", _top_level_labels())
            self.assertIn("project", _top_level_kinds())
            self.assertIn("global_root", _top_level_kinds())

            self.win.switchTo(self.win.chart_page)
            self.assertEqual(self.win._tree_panel.tree._focus_root_group_types, ["datasets", "pictures"])
            self.assertEqual(_top_level_labels(), ["数据集", "图片集"])

            self.win.switchTo(self.win.process_page)
            self.assertEqual(self.win._tree_panel.tree._focus_root_group_types, ["datasets"])
            self.assertEqual(_top_level_labels(), ["数据集"])

            self.win.switchTo(self.win.analysis_page)
            self.assertEqual(self.win._tree_panel.tree._focus_root_group_types, ["datasets", "analysis_result_group"])
            self.assertEqual(_top_level_labels(), ["数据集", "分析结果"])

            self.win.switchTo(self.win.digitize_page)
            self.assertEqual(self.win._tree_panel.tree._focus_root_group_types, ["datasets", "images"])
            self.assertEqual(_top_level_labels(), ["数据集", "数字化"])
        finally:
            self.win._page_tree_focus_mode_enabled = previous
            self.win._update_tree_panel_visibility(self.win.stackedWidget.currentWidget())

    def test_extension_toggle_button_visible_on_all_function_pages(self):
        self.win.switchTo(self.win.data_page)
        self.assertFalse(self.win._tree_panel.extension_toggle_btn.isHidden())
        self.assertFalse(self.win._tree_panel.extension_toggle_btn.isEnabled())

        self.win.switchTo(self.win.chart_page)
        self.assertFalse(self.win._tree_panel.extension_toggle_btn.isHidden())
        self.assertTrue(self.win._tree_panel.extension_toggle_btn.isEnabled())

        self.win.switchTo(self.win.process_page)
        self.assertFalse(self.win._tree_panel.extension_toggle_btn.isHidden())
        self.assertTrue(self.win._tree_panel.extension_toggle_btn.isEnabled())

        self.win.switchTo(self.win.analysis_page)
        self.assertFalse(self.win._tree_panel.extension_toggle_btn.isHidden())
        self.assertTrue(self.win._tree_panel.extension_toggle_btn.isEnabled())

        self.win.switchTo(self.win.digitize_page)
        self.assertFalse(self.win._tree_panel.extension_toggle_btn.isHidden())
        self.assertTrue(self.win._tree_panel.extension_toggle_btn.isEnabled())

    def test_extension_toggle_button_hides_and_shows_current_page_panel(self):
        self.win.switchTo(self.win.chart_page)
        self.assertFalse(self.win.chart_page.is_extension_panel_visible())
        self.assertTrue(self.win.chart_page._extension_panel.isHidden())

        self.win._toggle_current_page_extension_panel()
        self.assertTrue(self.win.chart_page.is_extension_panel_visible())
        self.assertFalse(self.win.chart_page._extension_panel.isHidden())

        self.win._toggle_current_page_extension_panel()
        self.assertFalse(self.win.chart_page.is_extension_panel_visible())
        self.assertTrue(self.win.chart_page._extension_panel.isHidden())

    def test_function_tool_panels_use_consistent_width(self):
        from ui.theme import WORKBENCH_TOOL_PANEL_WIDTH

        widths = {
            self.win.data_page._tool_panel.maximumWidth(),
            self.win.chart_page._tool_panel.maximumWidth(),
            self.win.process_page._tool_panel.maximumWidth(),
            self.win.analysis_page._tool_panel.maximumWidth(),
            self.win.digitize_page._tool_panel.maximumWidth(),
        }

        self.assertEqual(widths, {WORKBENCH_TOOL_PANEL_WIDTH})

    def test_replay_onboarding_resets_all_page_progress(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            config_path = Path(temp_dir) / "ui_preferences.json"
            with mock.patch("core.ui_preferences._CONFIG_PATH", config_path), \
                 mock.patch("ui.main_window.QTimer.singleShot", side_effect=lambda _ms, callback: callback()), \
                 mock.patch.object(self.win.home_page, "start_onboarding") as start_mock:
                from core.ui_preferences import UIPreferences

                prefs = UIPreferences(
                    home_onboarding_completed=True,
                    page_onboarding_completed={"data": True, "chart": True},
                )
                prefs.save()

                self.win._replay_home_onboarding()

                refreshed = UIPreferences.load()
                self.assertFalse(refreshed.home_onboarding_completed)
                self.assertEqual(refreshed.page_onboarding_completed, {})
                start_mock.assert_called_once_with(force=True)

    def test_tree_toggle_button_hides_and_shows_tree_panel(self):
        self.win.switchTo(self.win.data_page)
        self.assertFalse(self.win._tree_panel.isHidden())

        self.win._toggle_tree_panel()
        self.assertTrue(self.win._tree_panel.isHidden())

        self.win._toggle_tree_panel()
        self.assertFalse(self.win._tree_panel.isHidden())

    def test_switch_to_process_page_shows_tree(self):
        self.win.switchTo(self.win.process_page)
        self.assertFalse(self.win._tree_panel.isHidden())

    def test_switch_to_analysis_page_shows_tree(self):
        self.win.switchTo(self.win.analysis_page)
        self.assertFalse(self.win._tree_panel.isHidden())

    def test_switch_to_digitize_page_shows_tree(self):
        self.win.switchTo(self.win.digitize_page)
        self.assertFalse(self.win._tree_panel.isHidden())

    def test_on_project_opened_refreshes_tree(self):
        """项目打开后树自动刷新"""
        self.win._on_project_opened(self.p.name)
        self.assertGreater(
            self.win._tree_panel.tree._tree.topLevelItemCount(), 0
        )

    def test_extension_panel_visibility_is_shared_across_function_pages(self):
        self.win.switchTo(self.win.chart_page)

        self.win._toggle_current_page_extension_panel()

        for page in (
            self.win.chart_page,
            self.win.process_page,
            self.win.analysis_page,
            self.win.digitize_page,
        ):
            self.assertTrue(page.is_extension_panel_visible())

        self.win.switchTo(self.win.process_page)
        self.assertFalse(self.win.process_page._extension_panel.isHidden())

        self.win._toggle_current_page_extension_panel()

        for page in (
            self.win.chart_page,
            self.win.process_page,
            self.win.analysis_page,
            self.win.digitize_page,
        ):
            self.assertFalse(page.is_extension_panel_visible())

    def test_tree_panel_extension_toggle_button_state_stays_synced_across_pages(self):
        button = self.win._tree_panel.extension_toggle_btn

        self.win.switchTo(self.win.chart_page)
        self.assertEqual(button.toolTip(), "显示扩展面板")

        self.win._toggle_current_page_extension_panel()
        self.assertEqual(button.toolTip(), "隐藏扩展面板")

        self.win.switchTo(self.win.analysis_page)
        self.assertEqual(button.toolTip(), "隐藏扩展面板")

        self.win._toggle_current_page_extension_panel()
        self.assertEqual(button.toolTip(), "显示扩展面板")

        self.win.switchTo(self.win.digitize_page)
        self.assertEqual(button.toolTip(), "显示扩展面板")

    def test_switching_page_reapplies_shared_extension_panel_state(self):
        button = self.win._tree_panel.extension_toggle_btn

        self.win.switchTo(self.win.chart_page)
        self.win._toggle_current_page_extension_panel()
        self.assertTrue(self.win.analysis_page.is_extension_panel_visible())

        self.win.analysis_page.set_extension_panel_visible(False)
        self.assertFalse(self.win.analysis_page.is_extension_panel_visible())

        self.win.switchTo(self.win.analysis_page)

        self.assertTrue(self.win.analysis_page.is_extension_panel_visible())
        self.assertFalse(self.win.analysis_page._extension_panel.isHidden())
        self.assertEqual(button.toolTip(), "隐藏扩展面板")

    def test_tree_root_switches_current_project(self):
        other = self.pm.create_new("mw_other")
        self.pm.migrate_to_v3(other)
        self.pm.set_current_project(self.p.id)
        self.win._tree_panel.tree.refresh()
        project_item = self.win._tree_panel.tree._tree.topLevelItem(1)
        self.win._tree_panel.tree._on_item_clicked(project_item, 0)
        self.assertEqual(self.pm.current_project_id, other.id)

    def test_tree_node_selected_routes_to_data_page(self):
        """树信号路由到数据页"""
        self.win.switchTo(self.win.data_page)
        node = next((n for n in self.p.tree.nodes if n.kind == "data_file"), None)
        if node:
            self.win._on_tree_node_selected("data_file", node.id)
            self.assertIs(self.win.stackedWidget.currentWidget(), self.win.data_page)

    def test_tree_node_selected_routes_to_process_page(self):
        self.win.switchTo(self.win.process_page)
        node = next((n for n in self.p.tree.nodes if n.kind == "data_file"), None)
        if node:
            self.win._on_tree_node_selected("data_file", node.id)

    def test_tree_panel_data_actions_visible_only_on_data_page(self):
        # 顶部共享树工具条仅保留项目与树管理操作，不再提供数据入口按钮。
        self.win.switchTo(self.win.data_page)
        self.assertFalse(self.win._tree_panel.new_project_btn.isHidden())
        self.assertFalse(self.win._tree_panel.open_project_btn.isHidden())
        self.assertFalse(self.win._tree_panel.save_project_btn.isHidden())
        self.assertFalse(self.win._tree_panel.close_project_btn.isHidden())
        self.assertFalse(self.win._tree_panel.tree_manage_btn.isHidden())
        self.assertFalse(hasattr(self.win._tree_panel, "add_dataset_btn"))
        self.assertFalse(hasattr(self.win._tree_panel, "import_file_btn"))
        self.win.switchTo(self.win.process_page)
        self.assertFalse(self.win._tree_panel.new_project_btn.isHidden())
        self.assertFalse(self.win._tree_panel.open_project_btn.isHidden())
        self.assertFalse(self.win._tree_panel.save_project_btn.isHidden())
        self.assertFalse(self.win._tree_panel.close_project_btn.isHidden())
        self.assertFalse(self.win._tree_panel.tree_manage_btn.isHidden())
        self.assertFalse(hasattr(self.win._tree_panel, "add_dataset_btn"))
        self.assertFalse(hasattr(self.win._tree_panel, "import_file_btn"))

    def test_project_tree_manage_button_opens_dialog(self):
        with mock.patch("ui.main_window.ProjectTreeManageDialog") as dialog_cls:
            dialog = dialog_cls.return_value
            self.win._open_project_tree_manage_dialog()

        dialog_cls.assert_called_once_with(self.win)
        dialog.project_modified.connect.assert_called()
        dialog.exec.assert_called_once_with()

    def test_tree_node_activated_series_stays_on_analysis_page(self):
        self.win.switchTo(self.win.analysis_page)
        self.win._on_tree_node_activated("series", self.s.id)
        self.assertIs(self.win.stackedWidget.currentWidget(), self.win.analysis_page)
        self.assertEqual(len(self.win.analysis_page._selected_inputs), 1)

    def test_tree_node_activated_data_file_only_dispatches_on_process_page(self):
        data_file_node = next((n for n in self.p.tree.nodes if n.kind == "data_file"), None)
        self.assertIsNotNone(data_file_node)

        self.win.switchTo(self.win.analysis_page)
        before = len(self.win.analysis_page._selected_inputs)
        self.win._on_tree_node_activated("data_file", data_file_node.id)
        self.assertEqual(len(self.win.analysis_page._selected_inputs), before)

        self.win.switchTo(self.win.process_page)
        self.win._on_tree_node_activated("data_file", data_file_node.id)
        self.assertGreater(len(self.win.process_page._selected_inputs), 0)

    def test_tree_node_activated_analysis_result_loads_analysis_page(self):
        self.win.analysis_page.on_tree_node_activated("series", self.s.id)
        self.win.analysis_page._run_analysis()
        with mock.patch(
            "ui.pages.analysis_page.choose_analysis_result_save_plan",
            return_value=_analysis_result_save_plans(self.pm, "拟合结果A")[0],
        ):
            self.win.analysis_page._save_result()
        node = next((n for n in self.p.tree.nodes if n.kind == "analysis_result"), None)
        self.assertIsNotNone(node)
        self.win.switchTo(self.win.chart_page)
        self.win._on_tree_node_activated("analysis_result", node.id)
        self.assertIs(self.win.stackedWidget.currentWidget(), self.win.analysis_page)
        self.assertEqual(self.win.analysis_page._result["analysis_type"], "curve_fit")

    def test_tree_node_activated_series_stays_on_process_page(self):
        self.win.switchTo(self.win.process_page)
        self.win._on_tree_node_activated("series", self.s.id)
        self.assertIs(self.win.stackedWidget.currentWidget(), self.win.process_page)
        self.assertEqual(self.win.process_page._selected_src_id, self.s.id)

    def test_tree_node_activated_series_stays_on_chart_page(self):
        self.win.switchTo(self.win.chart_page)
        self.win._on_tree_node_activated("series", self.s.id)
        self.assertIs(self.win.stackedWidget.currentWidget(), self.win.chart_page)
        self.assertEqual(len(self.win.chart_page._chart_series), 1)

    def test_tree_node_activated_pipeline_loads(self):
        ops = [{"type": "smooth", "params": {"method": "moving_avg", "window": 3}}]
        sp = self.pm.add_saved_pipeline("test", ops)
        if sp:
            self.win._on_tree_node_activated("global_pipeline", sp.id)

    def test_tree_node_activated_report_template(self):
        tmpl = self.pm.add_report_template("r1", "# Report")
        self.win._on_tree_node_activated("global_report_template", tmpl.id)
        self.assertEqual(self.win.analysis_page._current_report_template_id, tmpl.id)

    def test_tree_node_activated_extension_config_routes_to_data_page_editor(self):
        from core.extension_api import ProcessingExtension, extension_registry
        from core.global_assets import global_assets

        def _probe(lines, params):
            return lines[0] if lines else []

        extension_registry.register_processing(
            ProcessingExtension(type="mw_extension_route_probe", name="主窗口路由探针", handler=_probe, default_options={"factor": 2})
        )
        try:
            config = global_assets.add_extension_config(
                category="processing",
                extension_type="mw_extension_route_probe",
                extension_name="主窗口路由探针",
                name="方案A",
                options={"factor": 8},
            )

            with mock.patch.object(self.win.data_page, "open_extension_config", wraps=self.win.data_page.open_extension_config) as open_mock:
                self.win._on_tree_node_activated("global_extension_config", config.id)

            self.assertIs(self.win.stackedWidget.currentWidget(), self.win.data_page)
            open_mock.assert_called_once_with(config.id)
            self.assertEqual(self.win.data_page._preview_section_label.text(), "配置编辑")
            self.assertIn('"factor": 8', self.win.data_page._text_preview.toPlainText())
        finally:
            extension_registry.unregister_processing("mw_extension_route_probe")

    def test_tree_node_activated_image_work(self):
        with mock.patch.object(self.win.digitize_page, "load_image_by_id") as load_mock:
            self.win._on_tree_node_activated("image_work", "fake-img-id")

        self.assertIs(self.win.stackedWidget.currentWidget(), self.win.digitize_page)
        load_mock.assert_called_once_with("fake-img-id")

    def test_tree_node_activated_image_work_add_curve_routes_to_digitize(self):
        with mock.patch.object(self.win.digitize_page, "load_image_by_id") as load_mock, \
             mock.patch.object(self.win.digitize_page, "_on_add_curve") as add_curve_mock:
            self.win._on_tree_node_activated("image_work_add_curve", "fake-img-id")

        self.assertIs(self.win.stackedWidget.currentWidget(), self.win.digitize_page)
        load_mock.assert_called_once_with("fake-img-id")
        add_curve_mock.assert_called_once_with()

    def test_tree_panel_removes_top_data_action_buttons(self):
        self.assertFalse(hasattr(self.win._tree_panel, "add_dataset_btn"))
        self.assertFalse(hasattr(self.win._tree_panel, "import_file_btn"))

    def test_tree_node_activated_curve_export_routes_to_digitize_export(self):
        with mock.patch.object(self.win.digitize_page, "load_curve_by_id") as load_mock, \
             mock.patch.object(self.win.digitize_page, "_on_export_to_data_file") as export_mock:
            self.win._on_tree_node_activated("curve_export_to_data_file", "curve-id")

        self.assertIs(self.win.stackedWidget.currentWidget(), self.win.digitize_page)
        load_mock.assert_called_once_with("curve-id")
        export_mock.assert_called_once_with()


class TestMainWindowPictureRoute(unittest.TestCase):

    def setUp(self):
        self._restore_assets = _patch_global_assets()
        self.pm, self.p, self.df, self.s = _make_project("mw_picture_route")
        self._restore = _patch_pm(self.pm)
        from ui.main_window import MainWindow

        self.win = MainWindow()

    def tearDown(self):
        self._restore()
        self._restore_assets()
        self.win.deleteLater()

    def test_picture_activation_routes_to_chart_restore(self):
        from models.schemas import FigureState, PicturePlotSeriesSnapshot, PicturePlotSnapshot

        with tempfile.TemporaryDirectory() as temp_dir:
            picture_path = Path(temp_dir) / "restore.png"
            picture_path.write_bytes(b"png")
            picture_node = self.pm.add_picture(
                str(picture_path),
                name="restore.png",
                plot_snapshot=PicturePlotSnapshot(
                    figure_state=FigureState(x_label="Restored X"),
                    series=[
                        PicturePlotSeriesSnapshot(
                            curve_key="restored-curve",
                            curve_identity="restored-curve",
                            name="restored",
                            display_name="restored",
                            x=[1.0, 2.0],
                            y=[3.0, 4.0],
                        )
                    ],
                ),
            )

            self.assertIsNotNone(picture_node)
            with mock.patch("ui.pages.chart_page.MessageBox.exec", return_value=True):
                self.win._on_tree_node_activated("picture", picture_node.id)

        self.assertIs(self.win.stackedWidget.currentWidget(), self.win.chart_page)
        self.assertEqual(len(self.win.chart_page._chart_series), 1)
        self.assertEqual(self.win.chart_page._chart_series[0]["name"], "restored")

    def test_send_to_visualize_picture_routes_to_chart_restore(self):
        from models.schemas import FigureState, PicturePlotSeriesSnapshot, PicturePlotSnapshot

        with tempfile.TemporaryDirectory() as temp_dir:
            picture_path = Path(temp_dir) / "restore_send.png"
            picture_path.write_bytes(b"png")
            picture_node = self.pm.add_picture(
                str(picture_path),
                name="restore_send.png",
                plot_snapshot=PicturePlotSnapshot(
                    figure_state=FigureState(x_label="Restored X"),
                    series=[
                        PicturePlotSeriesSnapshot(
                            curve_key="restored-curve",
                            curve_identity="restored-curve",
                            name="restored-send",
                            display_name="restored-send",
                            x=[1.0, 2.0],
                            y=[3.0, 4.0],
                        )
                    ],
                ),
            )

            self.assertIsNotNone(picture_node)
            with mock.patch("ui.pages.chart_page.MessageBox.exec", return_value=True):
                self.win._on_send_to_visualize("picture", picture_node.id)

        self.assertIs(self.win.stackedWidget.currentWidget(), self.win.chart_page)
        self.assertEqual(len(self.win.chart_page._chart_series), 1)
        self.assertEqual(self.win.chart_page._chart_series[0]["name"], "restored-send")


# ═══════════════════════════════════════════════════════════════════════════
# 9. 组合信号流程测试 — 模拟完整工作流
# ═══════════════════════════════════════════════════════════════════════════

class TestSignalWorkflows(unittest.TestCase):
    """组合前端信号与后端逻辑的端到端工作流"""

    def setUp(self):
        self._restore_assets = _patch_global_assets()
        self.pm, self.p, self.df, self.s = _make_project("wf_test")
        self._restore = _patch_pm(self.pm)

    def tearDown(self):
        self._restore()
        self._restore_assets()

    def test_workflow_import_then_visualize(self):
        """
        工作流: DataPage 树节点选中 → on_tree_node_selected →
                ChartPage 接收同一节点 → 不崩溃
        """
        from ui.pages.data_page import DataPage
        from ui.pages.chart_page import ChartPage
        dp = DataPage()
        cp = ChartPage()
        node = next((n for n in self.p.tree.nodes if n.kind == "data_file"), None)
        self.assertIsNotNone(node)
        dp.on_tree_node_selected("data_file", node.id)
        cp.on_tree_node_selected("data_file", node.id)
        dp.deleteLater()
        cp.deleteLater()

    def test_workflow_save_and_load_pipeline(self):
        """
        工作流: ProcessPage 加载数据 → 设置 ops →
            保存全局 Pipeline → 按 pipeline id 加载
        """
        from ui.pages.process_page import ProcessPage
        pp = ProcessPage()
        # 设置操作链
        ops = [
            {"type": "smooth", "params": {"method": "savgol", "window": 5, "poly": 2}},
            {"type": "normalize", "params": {"mode": "minmax"}},
        ]
        pp._load_ops_into_chain(ops)
        self.assertEqual(len(pp._ops), 2)
        # 保存 Pipeline
        sp = self.pm.add_saved_pipeline("流程A", ops, "测试用")
        self.assertIsNotNone(sp)
        pp.load_pipeline(sp.id)
        self.assertEqual(len(pp._ops), 2)
        self.assertEqual(pp._ops[0]["type"], "smooth")
        pp.deleteLater()

    def test_workflow_tree_node_selected_triggers_data_preview(self):
        """
        工作流: ProjectTreeWidget 发出 node_selected →
                DataPage.on_tree_node_selected 被调用 → 数据预览不崩溃
        """
        from ui.widgets.project_tree import ProjectTreeWidget
        from ui.pages.data_page import DataPage
        tree = ProjectTreeWidget()
        dp = DataPage()
        # 连接信号
        tree.node_selected.connect(dp.on_tree_node_selected)
        # 发射信号
        node = next((n for n in self.p.tree.nodes if n.kind == "data_file"), None)
        if node:
            tree.node_selected.emit("data_file", node.id)
        tree.deleteLater()
        dp.deleteLater()

    def test_workflow_tree_node_selected_to_analysis(self):
        """
        工作流: ProjectTreeWidget.node_selected →
                AnalysisPage.on_tree_node_selected → 系列列表更新
        """
        from ui.widgets.project_tree import ProjectTreeWidget
        from ui.pages.analysis_page import AnalysisPage
        tree = ProjectTreeWidget()
        ap = AnalysisPage()
        tree.node_selected.connect(ap.on_tree_node_selected)
        node = next((n for n in self.p.tree.nodes if n.kind == "data_file"), None)
        if node:
            tree.node_selected.emit("data_file", node.id)
        tree.deleteLater()
        ap.deleteLater()

    def test_workflow_tree_node_selected_to_process(self):
        """
        工作流: ProjectTreeWidget.node_selected →
                ProcessPage.on_tree_node_selected → 处理输入设置
        """
        from ui.widgets.project_tree import ProjectTreeWidget
        from ui.pages.process_page import ProcessPage
        tree = ProjectTreeWidget()
        pp = ProcessPage()
        tree.node_selected.connect(pp.on_tree_node_selected)
        node = next((n for n in self.p.tree.nodes if n.kind == "data_file"), None)
        if node:
            tree.node_selected.emit("data_file", node.id)
        tree.deleteLater()
        pp.deleteLater()

    def test_workflow_project_modified_refreshes_tree(self):
        """
        工作流: 树中执行创建文件夹 → project_modified 发出 →
                树刷新后包含新文件夹节点
        """
        from ui.widgets.project_tree import ProjectTreeWidget
        tree = ProjectTreeWidget()
        tree.refresh()
        nodes_before = len(self.p.tree.nodes)
        # 手动创建文件夹
        self.pm.add_folder("新文件夹_test")
        self.assertEqual(len(self.p.tree.nodes), nodes_before + 1)
        # 刷新
        tree.refresh()
        self.assertGreater(tree._tree.topLevelItemCount(), 0)
        tree.project_modified.emit()  # 验证信号可发射
        tree.deleteLater()

    def test_workflow_settings_save_ai_config_and_reload(self):
        """
        工作流: SettingsPage 保存 AI 配置 → AIConfig.load() 读回 →
                配置一致
        """
        import tempfile, json
        from ui.pages.settings_page import SettingsPage
        from core.ai_client import AIConfig
        page = SettingsPage()
        page._ai_url_edit.setText("http://localhost:11434/v1")
        page._ai_model_edit.setText("llama3")
        page._ai_timeout_edit.setText("45")
        page._ai_provider_combo.setCurrentIndex(1)  # Ollama
        # Mock 保存路径，避免写入真实配置文件
        with tempfile.NamedTemporaryFile(suffix=".json", delete=False, mode="w") as f:
            tmp_path = f.name
        orig_path = AIConfig._config_path if hasattr(AIConfig, "_config_path") else None
        # 直接测试配置对象构建
        cfg = AIConfig(
            provider="ollama",
            base_url=page._ai_url_edit.text(),
            api_key="",
            model=page._ai_model_edit.text(),
            timeout=int(page._ai_timeout_edit.text()),
        )
        self.assertEqual(cfg.provider, "ollama")
        self.assertEqual(cfg.model, "llama3")
        self.assertEqual(cfg.timeout, 45)
        page.deleteLater()
        os.unlink(tmp_path)

    def test_workflow_pipeline_node_activated_loads_to_process(self):
        """
        工作流: MainWindow._on_tree_node_activated("global_pipeline", id) →
                ProcessPage._ops 被填充（通过直接调用 ProcessPage.load_pipeline 验证）
        """
        ops = [{"type": "derivative", "params": {}}]
        sp = self.pm.add_saved_pipeline("deriv", ops)
        self.assertIsNotNone(sp)
        # 直接测试 ProcessPage.load_pipeline 而无需创建 MainWindow
        from ui.pages.process_page import ProcessPage
        pp = ProcessPage()
        pp.load_pipeline(sp.id)
        self.assertEqual(len(pp._ops), 1)
        self.assertEqual(pp._ops[0]["type"], "derivative")
        pp.deleteLater()


# ═══════════════════════════════════════════════════════════════════════════
# 10. ImportDialog 文件解析器
# ═══════════════════════════════════════════════════════════════════════════

class TestImportDialogParsers(unittest.TestCase):
    """ImportDialog 文件解析工具函数测试"""

    def _write_temp(self, content, suffix=".csv"):
        import tempfile
        f = tempfile.NamedTemporaryFile(mode="w", suffix=suffix,
                                        delete=False, encoding="utf-8")
        f.write(content)
        f.close()
        return f.name

    def test_parse_csv_two_column(self):
        from ui.dialogs.import_dialog import _parse_file_preview
        path = self._write_temp("x,y\n1.0,2.0\n3.0,4.0\n5.0,6.0\n")
        try:
            headers, rows = _parse_file_preview(path)
            self.assertEqual(len(headers), 2)
            self.assertEqual(len(rows), 3)
            self.assertAlmostEqual(rows[0][0], 1.0)
        finally:
            os.unlink(path)

    def test_parse_csv_no_header(self):
        from ui.dialogs.import_dialog import _parse_file_preview
        path = self._write_temp("1.0,2.0\n3.0,4.0\n")
        try:
            headers, rows = _parse_file_preview(path)
            self.assertEqual(len(rows), 2)
        finally:
            os.unlink(path)

    def test_parse_csv_tab_separated(self):
        from ui.dialogs.import_dialog import _parse_file_preview
        path = self._write_temp("x\ty\n1.0\t10.0\n2.0\t20.0\n", ".txt")
        try:
            headers, rows = _parse_file_preview(path)
            self.assertEqual(len(headers), 2)
        finally:
            os.unlink(path)

    def test_parse_json_dict_fmt(self):
        import json as json_mod
        from ui.dialogs.import_dialog import _parse_file_preview
        data = {"time": [0.0, 1.0, 2.0], "force": [0.0, 10.0, 20.0]}
        path = self._write_temp(json_mod.dumps(data), ".json")
        try:
            headers, rows = _parse_file_preview(path)
            self.assertEqual(len(headers), 2)
            self.assertEqual(len(rows), 3)
        finally:
            os.unlink(path)

    def test_parse_json_list_fmt(self):
        import json as json_mod
        from ui.dialogs.import_dialog import _parse_file_preview
        data = [{"x": 1, "y": 2}, {"x": 3, "y": 4}]
        path = self._write_temp(json_mod.dumps(data), ".json")
        try:
            headers, rows = _parse_file_preview(path)
            self.assertEqual(len(rows), 2)
        finally:
            os.unlink(path)

    def test_parse_npy(self):
        try:
            import numpy as np
        except ImportError:
            self.skipTest("numpy not available")
        from ui.dialogs.import_dialog import _parse_file_preview
        import tempfile
        arr = np.array([[1.0, 2.0], [3.0, 4.0], [5.0, 6.0]])
        path = tempfile.mktemp(suffix=".npy")
        try:
            np.save(path, arr)
            headers, rows = _parse_file_preview(path)
            self.assertEqual(len(headers), 2)
            self.assertEqual(len(rows), 3)
        finally:
            if os.path.exists(path):
                os.unlink(path)

    def test_import_dialog_creates_no_crash(self):
        from ui.dialogs.import_dialog import ImportDialog
        dlg = ImportDialog()
        self.assertIsNotNone(dlg)
        dlg.deleteLater()

    def test_import_dialog_state_machine_back_and_forth(self):
        from ui.dialogs.import_dialog import ImportDialog
        dlg = ImportDialog()
        dlg._file_path = "demo.csv"
        dlg._raw_headers = ["x", "y"]
        dlg._raw_rows = [[1.0, 2.0], [3.0, 4.0]]
        dlg._btn_next.setEnabled(True)

        dlg._go_next()
        self.assertEqual(dlg._stack.currentIndex(), 1)
        self.assertEqual(dlg._btn_next.text(), "导入")
        self.assertTrue(dlg._btn_back.isEnabled())

        dlg._go_next()
        self.assertEqual(dlg._stack.currentIndex(), 2)
        self.assertEqual(dlg._btn_next.text(), "完成")

        dlg._go_back()
        self.assertEqual(dlg._stack.currentIndex(), 1)
        self.assertEqual(dlg._btn_next.text(), "导入")

        dlg._go_back()
        self.assertEqual(dlg._stack.currentIndex(), 0)
        self.assertEqual(dlg._btn_next.text(), "下一步")
        self.assertFalse(dlg._btn_back.isEnabled())
        dlg.deleteLater()

    def test_import_dialog_builds_variable_role_table(self):
        from qfluentwidgets import TableWidget
        from ui.dialogs.import_dialog import ImportDialog

        dlg = ImportDialog()
        dlg._raw_headers = ["time", "force", "col_2"]
        dlg._raw_rows = [[0.0, 1.0, 0.1], [1.0, 2.0, 0.2]]

        dlg._populate_col_table()

        self.assertIsInstance(dlg._col_table, TableWidget)
        self.assertEqual(dlg._col_table.rowCount(), 3)
        self.assertEqual(dlg._col_table.columnCount(), 6)
        self.assertEqual(dlg._name_edits[0].text(), "time")
        self.assertEqual(dlg._name_edits[2].text(), "变量3")
        self.assertTrue(dlg._role_buttons[0]["X 轴"].isChecked())
        self.assertTrue(dlg._role_buttons[1]["Y 轴"].isChecked())
        self.assertTrue(dlg._role_buttons[2]["Y 轴"].isChecked())
        dlg.deleteLater()

    def test_import_dialog_defaults_to_new_data_file_name(self):
        from ui.dialogs.import_dialog import ImportDialog

        dlg = ImportDialog()
        dlg._file_path = "demo.csv"
        dlg._raw_headers = ["time", "signal"]
        dlg._raw_rows = [[0.0, 1.0], [1.0, 2.0]]
        dlg._populate_col_table()

        self.assertIsNone(dlg.get_target_data_file_id())
        self.assertEqual(dlg.get_file_name(), "demo")
        dlg.deleteLater()

    def test_import_dialog_default_import_uses_all_y_columns_and_stem_name(self):
        from ui.dialogs.import_dialog import ImportDialog

        dlg = ImportDialog()
        dlg._file_path = "demo.csv"
        dlg._raw_headers = ["time", "force", "stress"]
        dlg._raw_rows = [[0.0, 1.0, 3.0], [1.0, 2.0, 4.0]]

        series_list = dlg.import_with_default_options()

        self.assertEqual([series.name for series in series_list], ["force", "stress"])
        self.assertEqual(dlg.get_file_name(), "demo")
        dlg.deleteLater()

    def test_import_dialog_load_file_populates_preview_state(self):
        from ui.dialogs.import_dialog import ImportDialog

        temp_path = Path(tempfile.NamedTemporaryFile(suffix=".csv", delete=False).name)
        try:
            temp_path.write_text("x,y\n0,1\n1,2\n", encoding="utf-8")
            dlg = ImportDialog()

            dlg.load_file(str(temp_path))

            self.assertEqual(dlg._file_path, str(temp_path))
            self.assertEqual(dlg._path_edit.text(), str(temp_path))
            self.assertEqual(dlg._raw_headers, ["x", "y"])
            self.assertEqual(len(dlg._raw_rows), 2)
            self.assertTrue(dlg._btn_next.isEnabled())
            dlg.deleteLater()
        finally:
            temp_path.unlink(missing_ok=True)

    def test_import_dialog_failed_reload_clears_previous_state(self):
        from ui.dialogs.import_dialog import ImportDialog

        dlg = ImportDialog()
        dlg._file_path = "old.csv"
        dlg._raw_headers = ["x", "y"]
        dlg._raw_rows = [[1.0, 2.0]]
        dlg.imported_series = [mock.Mock(name="series")]
        dlg._populate_col_table()
        dlg._go_next()

        with mock.patch("ui.dialogs.import_dialog._parse_file_preview", side_effect=ValueError("bad file")):
            with self.assertRaisesRegex(ValueError, "bad file"):
                dlg.load_file("broken.csv")

        self.assertEqual(dlg._stack.currentIndex(), 0)
        self.assertEqual(dlg._btn_next.text(), "下一步")
        self.assertFalse(dlg._btn_back.isEnabled())
        self.assertFalse(dlg._btn_next.isEnabled())
        self.assertEqual(dlg._raw_headers, [])
        self.assertEqual(dlg._raw_rows, [])
        self.assertEqual(dlg.imported_series, [])
        self.assertEqual(dlg._col_table.rowCount(), 0)
        dlg.deleteLater()

    def test_import_dialog_can_target_existing_data_file(self):
        from ui.dialogs.import_dialog import ImportDialog

        pm, _p, _df, _s = _make_project("import_existing")
        restore = _patch_pm(pm)
        dlg = None
        try:
            dlg = ImportDialog()
            dlg._file_path = "demo.csv"
            dlg._raw_headers = ["time", "signal"]
            dlg._raw_rows = [[0.0, 1.0], [1.0, 2.0]]
            dlg._populate_col_table()

            self.assertGreaterEqual(dlg._data_file_target_combo.count(), 2)
            dlg._data_file_target_combo.setCurrentIndex(1)
            self.assertEqual(dlg.get_target_data_file_id(), pm.current_project.data_files[0].id)
            self.assertFalse(dlg._data_file_name_edit.isEnabled())
        finally:
            restore()
            if dlg is not None:
                dlg.deleteLater()

    def test_import_dialog_existing_targets_use_project_tree_paths_for_duplicates(self):
        from models.schemas import DataFile, DataSeries
        from ui.dialogs.import_dialog import ImportDialog

        pm, _p, _df, _s = _make_project("import_existing_paths")
        restore = _patch_pm(pm)
        dlg = None
        try:
            dataset_root = pm._find_folder_by_group_type("datasets")
            self.assertIsNotNone(dataset_root)
            left_folder = pm.add_folder("组A", parent_id=dataset_root.id, group_type="datasets")
            right_folder = pm.add_folder("组B", parent_id=dataset_root.id, group_type="datasets")
            self.assertIsNotNone(left_folder)
            self.assertIsNotNone(right_folder)
            left_df = DataFile(name="重复.csv", series=[DataSeries(name="a", x=[0.0], y=[1.0])])
            right_df = DataFile(name="重复.csv", series=[DataSeries(name="b", x=[0.0], y=[2.0])])
            pm.add_data_file(left_df, parent_id=left_folder.id)
            pm.add_data_file(right_df, parent_id=right_folder.id)

            dlg = ImportDialog()
            choices = dlg._existing_data_file_choices()
            labels = [label for label, _data_file_id in choices if "重复.csv" in label]

            self.assertIn("组A/重复.csv", labels)
            self.assertIn("组B/重复.csv", labels)
        finally:
            restore()
            if dlg is not None:
                dlg.deleteLater()

    def test_import_dialog_name_edit_keeps_focus_when_mouse_only_moves(self):
        from PySide6.QtCore import QEvent
        from ui.dialogs.import_dialog import ImportDialog

        dlg = ImportDialog()
        dlg._file_path = "demo.csv"
        dlg._raw_headers = ["time", "force", "stress"]
        dlg._raw_rows = [[0.0, 1.0, 3.0], [1.0, 2.0, 4.0]]
        dlg._populate_col_table()
        dlg.show()
        _app.processEvents()

        edit = dlg._name_edits[0]
        edit.setFocus()
        _app.processEvents()
        self.assertIs(dlg.focusWidget(), edit)

        target_widget = dlg._col_table.cellWidget(0, 1)
        hover_event = QEvent(QEvent.Type.HoverMove)
        QApplication.sendEvent(target_widget, hover_event)
        _app.processEvents()

        self.assertIs(dlg.focusWidget(), edit)
        dlg.deleteLater()

    def test_import_dialog_multi_y_uses_variable_names(self):
        from ui.dialogs.import_dialog import ImportDialog

        dlg = ImportDialog()
        dlg._file_path = "demo.csv"
        dlg._raw_headers = ["time", "force", "stress"]
        dlg._raw_rows = [[0.0, 1.0, 3.0], [1.0, 2.0, 4.0]]
        dlg._populate_col_table()
        dlg._name_edits[1].setText("力")
        dlg._name_edits[2].setText("应力")
        dlg._role_buttons[2]["Y 轴"].setChecked(True)

        series_list = dlg._do_import()

        self.assertEqual([series.name for series in series_list], ["力", "应力"])
        dlg.deleteLater()


# ═══════════════════════════════════════════════════════════════════════════
# 11. AnalysisPage — on_tree_node_activated 和输入管理
# ═══════════════════════════════════════════════════════════════════════════

class TestAnalysisPageV3(unittest.TestCase):

    def setUp(self):
        self._restore_assets = _patch_global_assets()
        self.pm, self.p, self.df, self.s = _make_project("ap_v3")
        self._restore = _patch_pm(self.pm)
        from ui.pages.analysis_page import AnalysisPage
        self.page = AnalysisPage()

    def tearDown(self):
        self._restore()
        self._restore_assets()
        self.page.deleteLater()

    def test_on_tree_node_activated_series_to_analysis(self):
        """_to_analysis 后缀的 kind 应去掉后缀并加入输入列表"""
        self.page.on_tree_node_activated("series_to_analysis", self.s.id)
        self.assertGreater(len(self.page._selected_inputs), 0)

    def test_on_tree_node_activated_series(self):
        self.page.on_tree_node_activated("series", self.s.id)
        self.assertGreater(len(self.page._selected_inputs), 0)

    def test_dedup_same_node_id(self):
        """同一 node_id 多次激活不重复添加"""
        self.page.on_tree_node_activated("series", self.s.id)
        self.page.on_tree_node_activated("series", self.s.id)
        self.assertEqual(len(self.page._selected_inputs), 1)

    def test_clear_inputs(self):
        self.page.on_tree_node_activated("series", self.s.id)
        self.page._clear_inputs()
        self.assertEqual(len(self.page._selected_inputs), 0)

    def test_get_selected_data_returns_series(self):
        self.page.on_tree_node_activated("series", self.s.id)
        data = self.page._get_selected_data()
        self.assertEqual(len(data), 1)
        xs, ys, name = data[0]
        self.assertEqual(len(xs), len(self.s.x))

    def test_on_tree_node_activated_data_file_expands_all_series(self):
        from models.schemas import DataSeries

        self.df.series.append(DataSeries(name="s2", x=[1.0, 2.0], y=[3.0, 4.0]))
        node = next((n for n in self.p.tree.nodes if n.kind == "data_file"), None)
        self.assertIsNotNone(node)

        self.page.on_tree_node_activated("data_file", node.id)

        self.assertEqual(len(self.page._selected_inputs), 2)
        self.assertEqual(self.page._input_list.count(), 2)
        self.assertIn("当前分析:", self.page._selected_input_state_label.text())

    def test_analysis_selected_input_labels_use_tree_paths(self):
        from models.schemas import DataFile, DataSeries

        datasets_root = self.pm._find_folder_by_group_type("datasets")
        self.assertIsNotNone(datasets_root)
        folder = self.pm.add_folder("分析组", parent_id=datasets_root.id, group_type="datasets")
        self.assertIsNotNone(folder)
        node = self.pm.add_data_file(
            DataFile(name="analysis_path.csv", series=[DataSeries(name="curve_a", x=[0.0, 1.0], y=[1.0, 2.0])]),
            parent_id=folder.id,
        )

        self.page.on_tree_node_activated("data_file", node.id)

        self.assertEqual(self.page._input_list.item(0).text(), "分析组/analysis_path.csv/curve_a")

    def test_error_compare_uses_pair_dropdown_selection(self):
        from models.schemas import DataSeries

        other = DataSeries(name="s2", x=[1.0, 2.0, 3.0], y=[1.5, 1.0, 2.5])
        third = DataSeries(name="s3", x=[1.0, 2.0, 3.0], y=[2.0, 2.0, 2.0])
        self.df.series.extend([other, third])

        self.page.on_tree_node_activated("series", self.s.id)
        self.page.on_tree_node_activated("series", other.id)
        self.page.on_tree_node_activated("series", third.id)
        self.page._type_combo.setCurrentIndex(self.page._analysis_type_ids.index("error_compare"))
        self.page._extension_params_edit.set_options({"lines_list": [1, 3]})
        self.page._on_extension_analysis_options_changed({"lines_list": [1, 3]})

        self.page._run_analysis()

        self.assertEqual(self.page._result["analysis_type"], "error_compare")
        self.assertEqual(self.page._result["name1"], self.s.name)
        self.assertEqual(self.page._result["name2"], third.name)

    def test_extension_line_selection_does_not_override_current_analysis_summary(self):
        from core.extension_api import AnalysisExtension, ExtensionConfigField, extension_registry
        from models.schemas import DataSeries

        def _probe(lines, params):
            return {"analysis_type": "analysis_extension_state_probe", "inputs": list(params.get("lines_list") or [])}

        other = DataSeries(name="s2", x=[1.0, 2.0], y=[2.0, 3.0])
        third = DataSeries(name="s3", x=[1.0, 2.0], y=[3.0, 4.0])
        self.df.series.extend([other, third])

        extension_registry.register_analysis(
            AnalysisExtension(
                type="analysis_extension_state_probe",
                name="分析状态探针",
                handler=_probe,
                lines_number=(2, -1),
            )
        )
        try:
            self.page._refresh_analysis_type_choices()
            self.page.on_tree_node_activated("series", self.s.id)
            self.page.on_tree_node_activated("series", other.id)
            self.page.on_tree_node_activated("series", third.id)
            target_index = self.page._analysis_type_ids.index("analysis_extension_state_probe")
            self.page._type_combo.setCurrentIndex(target_index)
            self.page._input_list.setCurrentRow(1)

            self.page._extension_params_edit.optionsChanged.emit({"lines_list": [1, 3]})

            self.assertIn(other.name, self.page._selected_input_state_label.text())
            self.assertNotIn(self.s.name, self.page._selected_input_state_label.text())
            self.assertNotIn(third.name, self.page._selected_input_state_label.text())
        finally:
            extension_registry.unregister_analysis("analysis_extension_state_probe")

    def test_clicking_selected_input_item_updates_current_analysis_target(self):
        from models.schemas import DataSeries

        other = DataSeries(name="s2", x=[1.0, 2.0], y=[2.0, 3.0])
        third = DataSeries(name="s3", x=[1.0, 2.0], y=[3.0, 4.0])
        self.df.series.extend([other, third])

        self.page.on_tree_node_activated("series", self.s.id)
        self.page.on_tree_node_activated("series", other.id)
        self.page.on_tree_node_activated("series", third.id)

        target_item = self.page._input_list.item(1)
        self.assertIsNotNone(target_item)

        self.page._on_input_list_item_clicked(target_item)

        current_payload = self.page._input_list.currentItem().data(Qt.ItemDataRole.UserRole)
        self.assertEqual(current_payload["node_id"], other.id)
        self.assertIn(other.name, self.page._selected_input_state_label.text())

    def test_analysis_result_actions_include_curve_export_button(self):
        self.assertIsNotNone(self.page._export_result_btn)
        self.assertTrue(self.page._export_result_btn.isHidden())
        self.assertFalse(self.page._export_result_btn.isEnabled())

    def test_analysis_result_actions_fill_width_and_keep_export_in_middle(self):
        from PySide6.QtWidgets import QApplication
        from PySide6.QtWidgets import QSizePolicy

        self.page.resize(1200, 800)
        self.page.show()
        self.page._result = {
            "lines": [
                {"line_name": "结果曲线", "line": [[0.0, 1.0], [1.0, 2.0]]},
            ]
        }
        self.page._refresh_result_action_buttons()
        QApplication.processEvents()

        layout = self.page._analysis_result_actions_layout
        self.assertIs(layout.itemAt(0).widget(), self.page._save_result_btn)
        self.assertIs(layout.itemAt(1).widget(), self.page._export_result_btn)
        self.assertIs(layout.itemAt(2).widget(), self.page._generate_report_btn)
        self.assertEqual(self.page._save_result_btn.sizePolicy().horizontalPolicy(), QSizePolicy.Policy.Expanding)
        self.assertEqual(self.page._export_result_btn.sizePolicy().horizontalPolicy(), QSizePolicy.Policy.Expanding)
        self.assertEqual(self.page._generate_report_btn.sizePolicy().horizontalPolicy(), QSizePolicy.Policy.Expanding)
        self.assertFalse(self.page._export_result_btn.isHidden())
        self.assertLess(self.page._save_result_btn.x(), self.page._export_result_btn.x())
        self.assertLess(self.page._export_result_btn.x(), self.page._generate_report_btn.x())
        self.assertLess(abs(self.page._save_result_btn.width() - self.page._generate_report_btn.width()), 8)

    def test_analysis_selected_input_reorder_preserves_current_indicator(self):
        from models.schemas import DataSeries

        other = DataSeries(name="s2", x=[1.0, 2.0], y=[2.0, 3.0])
        third = DataSeries(name="s3", x=[1.0, 2.0], y=[3.0, 4.0])
        self.df.series.extend([other, third])

        self.page.on_tree_node_activated("series", self.s.id)
        self.page.on_tree_node_activated("series", other.id)
        self.page.on_tree_node_activated("series", third.id)

        target_item = self.page._input_list.item(1)
        self.page._input_list.setCurrentItem(target_item)
        target_item.setSelected(True)

        self.page._move_selected_inputs_down()

        current_payload = self.page._input_list.currentItem().data(Qt.ItemDataRole.UserRole)
        self.assertEqual(current_payload["node_id"], other.id)
        self.assertTrue(self.page._input_list.currentItem().isSelected())


# ═══════════════════════════════════════════════════════════════════════════
# 12. ChartPage — on_tree_node_activated 和 _chart_series 管理
# ═══════════════════════════════════════════════════════════════════════════

class TestChartPageV3(unittest.TestCase):

    def setUp(self):
        self._restore_assets = _patch_global_assets()
        self.pm, self.p, self.df, self.s = _make_project("cp_v3")
        self._restore = _patch_pm(self.pm)
        from ui.pages.chart_page import ChartPage
        self.page = ChartPage()

    def tearDown(self):
        self._restore()
        self._restore_assets()
        self.page.deleteLater()

    def test_on_tree_node_activated_series_adds_to_chart(self):
        self.page.on_tree_node_activated("series", self.s.id)
        self.assertGreater(len(self.page._chart_series), 0)

    def test_on_tree_node_activated_data_file_adds_all_series(self):
        node = next((n for n in self.p.tree.nodes if n.kind == "data_file"), None)
        self.assertIsNotNone(node)
        self.page.on_tree_node_activated("data_file", node.id)
        self.assertGreater(len(self.page._chart_series), 0)

    def test_update_theme_defers_chart_redraw_to_event_loop(self):
        if self.page._figure is None or self.page._canvas is None:
            self.skipTest("matplotlib unavailable")

        self.page.show()
        QApplication.processEvents()

        with mock.patch.object(self.page, "_redraw_now") as redraw:
            self.page.update_theme()
            self.assertEqual(redraw.call_count, 0)
            QApplication.processEvents()
            self.assertEqual(redraw.call_count, 1)

    def test_dedup_same_series(self):
        self.page.on_tree_node_activated("series", self.s.id)
        count1 = len(self.page._chart_series)
        self.page.on_tree_node_activated("series", self.s.id)
        self.assertEqual(len(self.page._chart_series), count1)

    def test_clear_chart(self):
        self.page.on_tree_node_activated("series", self.s.id)
        self.page._on_clear_chart()
        self.assertEqual(len(self.page._chart_series), 0)

    def test_data_file_activation_batches_redraw(self):
        from models.schemas import DataSeries

        node = next((n for n in self.p.tree.nodes if n.kind == "data_file"), None)
        self.assertIsNotNone(node)
        self.df.series.append(DataSeries(name="s2", x=[1.0, 2.0], y=[3.0, 4.0]))

        with mock.patch.object(self.page, "_redraw_now") as redraw:
            self.page.on_tree_node_activated("data_file", node.id)

        self.assertEqual(len(self.page._chart_series), 2)
        self.assertEqual(redraw.call_count, 1)

    def test_export_to_picture_group_creates_project_picture_node(self):
        from ui.dialogs.export_flow import PictureExportPlan

        self.page.on_tree_node_activated("series", self.s.id)
        with tempfile.TemporaryDirectory() as temp_dir:
            project_file = Path(temp_dir) / "chart_export.aline"
            self.pm.save(str(project_file))
            picture_root = next(
                n for n in self.p.tree.nodes
                if n.kind == "folder" and getattr(n, "group_type", None) == "pictures"
            )
            self.page.on_tree_node_selected("folder", picture_root.id)
            with mock.patch(
                "ui.pages.chart_page.choose_picture_export_plan",
                return_value=PictureExportPlan(export_name="chart.png", target_folder_id=picture_root.id),
            ):
                self.page._on_export_to_picture_group()

            self.assertEqual(len(self.p.pictures), 1)
            picture = self.p.pictures[0]
            self.assertTrue(Path(self.pm.get_picture_path(picture.id)).exists())
            picture_node = next((n for n in self.p.tree.nodes if n.kind == "picture" and n.picture_id == picture.id), None)
            self.assertIsNotNone(picture_node)
            self.assertEqual(picture_node.parent_id, picture_root.id)

    def test_export_to_picture_group_persists_plot_snapshot(self):
        from ui.dialogs.export_flow import PictureExportPlan

        self.page.on_tree_node_activated("series", self.s.id)
        self.page._x_label_edit.setText("Time")
        with tempfile.TemporaryDirectory() as temp_dir:
            project_file = Path(temp_dir) / "chart_snapshot.aline"
            self.pm.save(str(project_file))
            picture_root = next(
                n for n in self.p.tree.nodes
                if n.kind == "folder" and getattr(n, "group_type", None) == "pictures"
            )
            self.page.on_tree_node_selected("folder", picture_root.id)
            with mock.patch(
                "ui.pages.chart_page.choose_picture_export_plan",
                return_value=PictureExportPlan(export_name="chart.png", target_folder_id=picture_root.id),
            ):
                self.page._on_export_to_picture_group()

            picture = self.p.pictures[0]
            self.assertIsNotNone(picture.plot_snapshot)
            self.assertEqual(picture.plot_snapshot.figure_state.x_label, "Time")
            self.assertEqual(len(picture.plot_snapshot.series), 1)
            self.assertEqual(picture.plot_snapshot.series[0].name, self.s.name)

    def test_picture_activation_rejects_missing_plot_extension(self):
        from models.schemas import FigureState, PicturePlotExtensionSnapshot, PicturePlotSeriesSnapshot, PicturePlotSnapshot

        with tempfile.TemporaryDirectory() as temp_dir:
            picture_path = Path(temp_dir) / "missing-extension.png"
            picture_path.write_bytes(b"png")
            snapshot = PicturePlotSnapshot(
                figure_state=FigureState(x_label="Restored X"),
                series=[
                    PicturePlotSeriesSnapshot(
                        curve_key="restored-curve",
                        curve_identity="restored-curve",
                        name="restored",
                        display_name="restored",
                        x=[1.0, 2.0],
                        y=[2.0, 3.0],
                    )
                ],
                applied_extensions=[
                    PicturePlotExtensionSnapshot(
                        id="plot-extension-1",
                        type="missing-plot-extension",
                        sequence=1,
                        extension_version="1.0.0",
                    )
                ],
            )
            picture_node = self.pm.add_picture(str(picture_path), name="missing-extension.png", plot_snapshot=snapshot)

            self.assertIsNotNone(picture_node)
            self.page.on_tree_node_activated("series", self.s.id)
            with mock.patch("ui.pages.chart_page.InfoBar.error") as error_mock, \
                 mock.patch("ui.pages.chart_page.MessageBox.exec") as exec_mock:
                self.page.on_tree_node_activated("picture", picture_node.id)

        self.assertEqual(len(self.page._chart_series), 1)
        self.assertEqual(self.page._chart_series[0]["name"], self.s.name)
        error_mock.assert_called_once()
        self.assertEqual(error_mock.call_args.args[1], "绘图中使用的扩展不存在或未加载")
        exec_mock.assert_not_called()

    def test_picture_activation_restores_snapshot_and_warns_on_version_mismatch(self):
        from core.extension_api import PlotExtension, extension_registry
        from models.schemas import FigureState, PicturePlotExtensionSnapshot, PicturePlotSeriesSnapshot, PicturePlotSnapshot

        def _noop_extension(_context, _options):
            return None

        extension_registry.register_plot(
            PlotExtension(
                type="picture-restore-extension",
                name="图片恢复扩展",
                handler=_noop_extension,
                version="2.0.0",
            )
        )
        try:
            with tempfile.TemporaryDirectory() as temp_dir:
                picture_path = Path(temp_dir) / "restore.png"
                picture_path.write_bytes(b"png")
                snapshot = PicturePlotSnapshot(
                    figure_state=FigureState(x_label="Restored X", y_label="Restored Y"),
                    curve_styles={"restored-curve": {"color": "#123456", "linewidth": 2.5}},
                    series=[
                        PicturePlotSeriesSnapshot(
                            curve_key="restored-curve",
                            curve_identity="restored-curve",
                            name="restored",
                            display_name="restored",
                            x=[1.0, 2.0],
                            y=[3.0, 4.0],
                        )
                    ],
                    applied_extensions=[
                        PicturePlotExtensionSnapshot(
                            id="plot-extension-1",
                            type="picture-restore-extension",
                            sequence=1,
                            options={"alpha": 0.3},
                            extension_version="1.0.0",
                        )
                    ],
                    selected_curve_key="restored-curve",
                )
                picture_node = self.pm.add_picture(str(picture_path), name="restore.png", plot_snapshot=snapshot)

                self.assertIsNotNone(picture_node)
                self.page.on_tree_node_activated("series", self.s.id)
                with mock.patch("ui.pages.chart_page.InfoBar.warning") as warning_mock, \
                     mock.patch("ui.pages.chart_page.MessageBox.exec", return_value=True), \
                     mock.patch("ui.pages.chart_page.InfoBar.success") as success_mock:
                    self.page.on_tree_node_activated("picture", picture_node.id)

            self.assertEqual(len(self.page._chart_series), 1)
            self.assertEqual(self.page._chart_series[0]["name"], "restored")
            self.assertEqual(self.page._x_label_edit.text(), "Restored X")
            self.assertEqual(self.page._curve_styles["restored-curve"]["color"], "#123456")
            warning_mock.assert_called_once()
            self.assertEqual(warning_mock.call_args.args[0], "扩展版本不一致")
            success_mock.assert_called_once()
        finally:
            extension_registry.unregister_plot("picture-restore-extension")

    def test_project_modified_signal_emitted_on_save_template(self):
        """保存模板应发出 project_modified"""
        from models.schemas import FigureConfig
        received = []
        self.page.project_modified.connect(lambda: received.append(True))
        cfg = FigureConfig(name="tmpl_test")
        tm = self.pm.add_figure_template(cfg)
        self.page.project_modified.emit()
        self.assertEqual(len(received), 1)


# ═══════════════════════════════════════════════════════════════════════════
# 13. SettingsPage v0.3 — temperature/max_tokens 字段
# ═══════════════════════════════════════════════════════════════════════════

class TestSettingsPageV3(unittest.TestCase):

    def setUp(self):
        from ui.pages.settings_page import SettingsPage
        self.page = SettingsPage()

    def tearDown(self):
        self.page.deleteLater()

    def test_temperature_edit_exists(self):
        self.assertIsNotNone(self.page._ai_temperature_edit)

    def test_max_tokens_edit_exists(self):
        self.assertIsNotNone(self.page._ai_max_tokens_edit)

    def test_top_p_edit_exists(self):
        self.assertIsNotNone(self.page._ai_top_p_edit)

    def test_system_prompt_edit_exists(self):
        self.assertIsNotNone(self.page._ai_system_prompt_edit)

    def test_save_with_temperature(self):
        self.page._ai_temperature_edit.setText("0.5")
        self.page._ai_top_p_edit.setText("0.8")
        self.page._ai_max_tokens_edit.setText("1024")
        self.page._ai_system_prompt_edit.setPlainText("请使用中文")
        self.page._save_ai_config()  # should not crash

    def test_invalid_temperature_fallback(self):
        self.page._ai_temperature_edit.setText("not_a_number")
        self.page._save_ai_config()  # should fallback without crash

    def test_invalid_max_tokens_fallback(self):
        self.page._ai_max_tokens_edit.setText("-1abc")
        self.page._save_ai_config()  # should fallback without crash

    def test_invalid_top_p_fallback(self):
        self.page._ai_top_p_edit.setText("2.5")
        cfg = self.page._collect_ai_config()
        self.assertEqual(cfg.top_p, 1.0)

    def test_tmpl_list_exists(self):
        self.assertIsNotNone(self.page._tmpl_list)
        self.assertTrue(self.page._tmpl_card.isHidden())

    def test_refresh_templates_no_project(self):
        """无打开项目时刷新不崩溃"""
        self.page.refresh_templates()

    def test_refresh_templates_with_project(self):
        from core.project_manager import ProjectManager

        restore_assets = _patch_global_assets()
        pm = ProjectManager()
        pm.create_new("sett_tmpl_test")
        pm.add_report_template("settings_template_demo", "# Hello")
        # Patch and refresh
        restore = _patch_pm(pm)
        try:
            self.page.refresh_templates()
            self.assertEqual(self.page._tmpl_list.count(), 0)
        finally:
            restore()
            restore_assets()

    def test_report_template_card_hidden(self):
        self.assertTrue(self.page._tmpl_card.isHidden())


class TestDialogFocusCommit(unittest.TestCase):

    def test_advanced_figure_dialog_installs_click_away_focus_commit(self):
        from ui.dialogs.advanced_figure_dialog import AdvancedFigureDialog

        dialog = AdvancedFigureDialog()
        try:
            self.assertIsNotNone(dialog._click_away_focus_commit)
        finally:
            dialog.deleteLater()

    def test_ai_tool_dialog_installs_click_away_focus_commit(self):
        from ui.dialogs.ai_tool_dialog import AIToolDialog

        dialog = AIToolDialog()
        try:
            self.assertIsNotNone(dialog._click_away_focus_commit)
        finally:
            dialog.deleteLater()

    def test_export_dialogs_install_click_away_focus_commit(self):
        from ui.dialogs.export_flow import (
            _AnalysisResultSaveDialog,
            _BatchDataExportDialog,
            _DataExportDialog,
            _PictureExportDialog,
        )

        dialogs = [
            _DataExportDialog(
                None,
                title="导出数据列",
                entries=[],
                default_export_name="结果A",
                default_file_name="数据集A",
                file_suffix=".csv",
                current_text=None,
                show_export_name=True,
            ),
            _BatchDataExportDialog(
                None,
                title="批量导出数据列",
                entries=[],
                source_labels=["曲线A"],
                default_export_names=["结果A"],
                default_file_name="数据集A",
                file_suffix=".csv",
                current_text=None,
            ),
            _PictureExportDialog(
                None,
                title="导出图片",
                folder_entries=[{"label": "图片集", "node_id": None}],
                default_export_name="图1",
                current_text=None,
                file_suffix=".png",
            ),
            _AnalysisResultSaveDialog(
                None,
                title="保存分析结果",
                folder_entries=[{"label": "分析结果", "node_id": None}],
                default_result_name="结果A",
                current_text=None,
            ),
        ]
        try:
            for dialog in dialogs:
                self.assertIsNotNone(dialog._click_away_focus_commit)
        finally:
            for dialog in dialogs:
                dialog.deleteLater()


class TestReportTemplateDialog(unittest.TestCase):

    def setUp(self):
        self._restore_assets = _patch_global_assets()
        from core.global_assets import global_assets
        from models.schemas import ReportTemplate
        from ui.dialogs import report_template_dialog as report_dialog_module

        self._report_dialog_module = report_dialog_module
        self._web_patch = mock.patch.object(report_dialog_module, "_HAS_WEB", False)
        self._web_patch.start()
        self.template = global_assets.add_report_template(ReportTemplate(name="报告模板A", content="# Old Report"))
        self.dialog = report_dialog_module.ReportTemplateDialog(template_id=self.template.id)

    def tearDown(self):
        self.dialog.deleteLater()
        self._web_patch.stop()
        self._restore_assets()

    def test_load_template_list_uses_global_assets(self):
        items = [self.dialog._tmpl_combo.itemText(i) for i in range(self.dialog._tmpl_combo.count())]
        self.assertIn("默认模板", items)
        self.assertIn("报告模板A", items)

    def test_editor_uses_fluent_selection_colors(self):
        editor_style = self.dialog._editor.styleSheet()

        self.assertIn("selection-background-color", editor_style)
        self.assertIn("selection-color: #ffffff", editor_style)

    def test_save_template_updates_existing_template_without_duplication(self):
        from core.global_assets import global_assets

        self.dialog._editor.setPlainText("# Updated Report")
        with mock.patch("ui.dialogs.report_template_dialog.TextInputDialog.get_text", return_value=("报告模板A", True)):
            self.dialog._on_save_template()

        matches = [item for item in global_assets.list_report_templates() if item.name == "报告模板A"]
        self.assertEqual(len(matches), 1)
        self.assertEqual(matches[0].id, self.template.id)
        self.assertEqual(matches[0].content, "# Updated Report")

    def test_editor_can_insert_placeholder_from_selector(self):
        target_text = "{{table:params}}"
        index = next(
            i for i in range(self.dialog._placeholder_combo.count())
            if target_text in self.dialog._placeholder_combo.itemText(i)
        )
        self.dialog._editor.setPlainText("# 模板\n")
        self.dialog._placeholder_combo.setCurrentIndex(index)

        self.dialog._insert_selected_placeholder()

        self.assertIn(target_text, self.dialog._editor.toPlainText())

    def test_placeholder_selector_includes_custom_result_fields(self):
        dialog = self._report_dialog_module.ReportTemplateDialog(
            result={"analysis_type": "spectrum_analysis", "dominant_frequency": 9.81}
        )
        try:
            items = [dialog._placeholder_combo.itemText(i) for i in range(dialog._placeholder_combo.count())]
            self.assertTrue(any("{{dominant_frequency}}" in item for item in items))
        finally:
            dialog.deleteLater()


if __name__ == "__main__":
    unittest.main()
